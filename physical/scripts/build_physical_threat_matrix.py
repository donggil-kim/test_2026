# -*- coding: utf-8 -*-
"""
물리·인적 보안위협 매트릭스 빌더
------------------------------
분류 : 물리·인적 공격면 7개 도메인 → 위협분류(Lv2) → 세부위협(Lv3) — data/taxonomy.yaml
교차 : ATT&CK v19.2(물리·인적 접점) · MITRE CTID 내부자 TTP 지식베이스(ITKB) · 다른 매트릭스 ID(IDT·SCT·OTC·CL·AI·OT)
근거 : 물리·인적 보안사고 DB(data/incidents.yaml) · ATT&CK 절차(물리·인적 맥락) · ITKB 내부자 관측 기법 · 정부 경보 · 연구·시연
대응 : NIST SP 800-53 Rev.5 · ISO/IEC 27001:2022 부속서 A 5~7 · ISMS-P · CERT 내부자 가이드 7판 실천과제
참고 : NPSA·CISA 지침 등 참고 문헌(data/frameworks.yaml references) — 행 refs와 문구 '참고:' 인용
논리 : AI v3.2 · 클라우드 v5 · OT · 공급망 · 아이덴티티 매트릭스와 같은 발생가능성 수식(상 = 실제 사고 2건 이상),
       심각도는 물리·인적 결과 기준, 정부 경보·위협인텔·ITKB·ATT&CK은 '중'까지만

문구 : data/text/*.yaml(핵심 요약·요약설명·참조·탐지·대응) — 빌드 전에 scripts/validate.py로 사례·참조 ID를 대조
시나리오 : data/scenarios.yaml(실제 사고 기반 공격 체인)

실행 : python3 scripts/build_physical_threat_matrix.py [--data-only] [--allow-missing-text] [--skip-validate]
                                                      [--version v2] [--worksheet <폴더>]
       → output/통합_물리인적보안위협_매트릭스_<버전>.xlsx · .csv
"""
import collections
import csv
import datetime
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import taxonomy_rules as R  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
REF = ROOT / "reference"
VERSION = "v2"
NAME = "통합_물리인적보안위협_매트릭스"


def out_paths(version):
    return ROOT / "output" / f"{NAME}_{version}.xlsx", ROOT / "output" / f"{NAME}_{version}.csv"


def load_yaml(p, default=None):
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else default


# ===========================================================================
# 1) 원천 데이터
# ===========================================================================
TAX = load_yaml(DATA / "taxonomy.yaml")
FW = load_yaml(DATA / "frameworks.yaml")
INCIDENTS = load_yaml(DATA / "incidents.yaml", []) or []
CHANGELOG = load_yaml(DATA / "changelog.yaml", {})
SCENARIOS = load_yaml(DATA / "scenarios.yaml", {}) or {}
ATK = json.loads((REF / "attack" / "enterprise-attack-v19.2-subset.json").read_text(encoding="utf-8"))
ITKB_DOC = json.loads((REF / "itkb" / "insider-threat-ttp-kb.json").read_text(encoding="utf-8"))
ITKB = ITKB_DOC["techniques"]
NIST = json.loads((REF / "nist" / "sp800-53r5-names.json").read_text(encoding="utf-8"))["names"]
LINKS = json.loads((REF / "links" / "other_matrices.json").read_text(encoding="utf-8"))["ids"]
REFS = FW["references"]
REF_SHORT = {v["short"]: k for k, v in REFS.items()}

TEXT = {}
for f in sorted((DATA / "text").glob("*.yaml")):
    TEXT.update(load_yaml(f, {}) or {})

DOMAINS, ROWS = [], []
for d in TAX["domains"]:
    DOMAINS.append(d)
    for l2 in d["lv2"]:
        for x in l2["lv3"]:
            ROWS.append(dict(x, code=d["code"], domain=f"[{d['code']}] {d['ko']}", lv2_id=l2["id"], lv2=l2["name"],
                             single=(x["id"] == l2["id"])))
ROW = {o["id"]: o for o in ROWS}
DOMAIN = {d["code"]: d for d in DOMAINS}

for e in INCIDENTS:
    e["date"] = str(e["date"])
    e["status"] = R.incident_status(e)
    e["recent"] = e["date"] >= R.RECENT_FROM
    e["domains"] = sorted({k.split("-")[1] for k in e["pht"]}, key=R.DOMAIN_ORDER.index)
INC = {e["id"]: e for e in INCIDENTS}

INSIDER_TYPES = {"내부자(악의)", "내부자(비의도)", "공모", "제3자"}

# ===========================================================================
# 2) 근거 집계
# ===========================================================================
EV = {o["id"]: dict(real=set(), recent=set(), kr=set(), advisory=set(), intel=set(), research=set(), excluded=set(),
                    subj=set(), itkb=set(), sector=collections.Counter(), impact=collections.Counter())
      for o in ROWS}

for e in INCIDENTS:
    for k in e["pht"]:
        x = EV[k]
        if e["status"] == "실제 사고":
            x["real"].add(e["id"])
            if e["recent"]:
                x["recent"].add(e["id"])
            if e.get("region") == "한국":
                x["kr"].add(e["id"])
            x["sector"].update(e["sector"])
            x["impact"].update(i for i in e["impact"] if i not in ("영향 없음(사전 차단)", "영향 미상"))
        elif e["status"] == "정부 경보":
            x["advisory"].add(e["id"])
        elif e["status"] == "위협인텔":
            x["intel"].add(e["id"])
        elif e["status"] == "실증·연구":
            x["research"].add(e["id"])
        else:
            x["excluded"].add(e["id"])


def atk_subjects(tid):
    """기법을 쓰는 ATT&CK 주체 — 물리 기법은 전체, 사람 대상 정찰·위장 기법은 인적 단서가 있는 절차만, 나머지는 세지 않음"""
    return {p["subject"] for p in ATK["procedures"].get(tid, []) if R.attack_counts(tid, p["text"])}


for o in ROWS:
    for t in o["attack"]:
        EV[o["id"]]["subj"] |= atk_subjects(t)
    # ITKB — 위협 주체에 내부자·공모·제3자가 있는 행만 '내부자 관측 기법'으로 표시
    if set(o["actors"]) & INSIDER_TYPES:
        EV[o["id"]]["itkb"] = {t for t in o["attack"] if t in ITKB}


# ===========================================================================
# 3) 위험 평가
# ===========================================================================
def likelihood(x):
    if len(x["real"]) >= R.LIKELY_HIGH_REAL:
        return "상"
    if x["real"] or x["subj"] or x["itkb"] or x["advisory"] or x["intel"] or x["research"]:
        return "중"
    return "하"


def evidence_level(x):
    if x["real"]:
        return "실제 사고 확인"
    if x["subj"] or x["itkb"] or x["advisory"] or x["intel"]:
        return "실사용 기법 포함"
    if x["research"]:
        return "실증·공개 취약점"
    return "이론·시나리오"


for o in ROWS:
    x = EV[o["id"]]
    o["likelihood"] = likelihood(x)
    o["risk"] = R.RISK_MATRIX[(o["likelihood"], o["severity"])]
    o["level"] = evidence_level(x)
    o["n_real"], o["n_recent"], o["n_kr"] = len(x["real"]), len(x["recent"]), len(x["kr"])
    o["n_adv"], o["n_intel"], o["n_research"] = len(x["advisory"]), len(x["intel"]), len(x["research"])
    o["n_atk"], o["n_itkb"] = len(x["subj"]), len(x["itkb"])
    o["lk_why"] = (f"실제 사고 {o['n_real']}건(최근 {o['n_recent']} · 국내 {o['n_kr']}) · 정부 경보 {o['n_adv']}건 · "
                   f"위협인텔 {o['n_intel']}건 · ATT&CK 사례 {o['n_atk']}건 · ITKB 관측 기법 {o['n_itkb']}개 · 실증·연구 {o['n_research']}건")
    txt = TEXT.get(o["id"]) or {}
    o["oneline"], o["summary"], o["reference"], o["detect"] = (txt.get(k) or "" for k in ("oneline", "summary", "reference",
                                                                                         "detect"))
    o["text_src"] = "분석가 작성" if txt else "문구 미작성"

rank_key = {r: i for i, r in enumerate(R.RISK_ORDER)}
for i, o in enumerate(sorted(ROWS, key=lambda o: (rank_key[o["risk"]], -o["n_real"], -o["n_recent"], o["id"])), 1):
    o["priority"] = i


def fw_pair(cat, key):
    v = FW[cat][key] if key in FW[cat] else FW[cat][int(key)] if str(key).isdigit() and int(key) in FW[cat] else [""]
    return v if isinstance(v, list) else [v]


def fw_label(cat, key):
    v = fw_pair(cat, key)
    return f"{key} {v[-1]}"


def nist_label(k):
    return f"{k} {NIST.get(k, '')}"


def atk_label(t):
    return f"{t} {ATK['techniques'][t]['name']}"


def link_label(k):
    return f"{k} {LINKS[k]['name']}" if k in LINKS else k


def ref_label(k):
    return REFS[k]["short"] if k in REFS else k


def inc_label(i):
    e = INC[i]
    return f"{i} {e['title']}({e['date']})"


def _bullets(block, section):
    """'■ 단락' 아래 '- ' 줄 목록"""
    out, on = [], False
    for ln in (block or "").splitlines():
        if ln.startswith("■ "):
            on = ln[2:].strip() == section
        elif on and ln.startswith("- "):
            out.append(ln[2:].strip())
    return out


def first_case(o):
    """대표 사례 = '■ 실제 사례' 첫 줄 → '[라벨] 사례명(시점) (사고 ID)' (보고서용 간략 표기)
    문구 미작성 행은 매핑된 실제 사고 중 최신 1건(사고 DB 제목·시점)"""
    if not o["reference"]:
        real = sorted(EV[o["id"]]["real"], key=lambda i: (INC[i]["date"], i), reverse=True)
        if real:
            return f"[실제 사고] {INC[real[0]]['title']}({INC[real[0]]['date']}) ({real[0]})"
        other = sorted(EV[o["id"]]["advisory"] | EV[o["id"]]["intel"] | EV[o["id"]]["research"],
                       key=lambda i: (INC[i]["date"], i), reverse=True)
        if other:
            lab = {"정부 경보": "[정부 경보]", "위협인텔": "[위협인텔]", "실증·연구": "[실증]"}[INC[other[0]]["status"]]
            return f"{lab} {INC[other[0]]['title']}({INC[other[0]]['date']}) ({other[0]})"
        return "공개 사고 미확인"
    for b in _bullets(o["reference"], "실제 사례"):
        if b.startswith("공개 사고 미확인"):
            return "공개 사고 미확인"
        ids = re.findall(r"PHI-\d{3}", b)
        return b.split(": ", 1)[0] + (f" ({', '.join(ids)})" if ids else "")
    return ""


def first_control(o):
    """핵심 대응 = '■ 대응' 첫 줄, 문구 미작성 행은 매핑된 대응 기준 ID"""
    b = [ln for ln in _bullets(o["detect"], "대응") if not ln.startswith("참고:")]
    if b:
        return b[0]
    ids = [f"NIST {'·'.join(o['nist'])}" if o["nist"] else "",
           f"ISO 27001 {'·'.join(str(k) for k in o['iso'])}" if o["iso"] else "",
           f"ISMS-P {'·'.join(str(k) for k in o['ismsp'])}" if o["ismsp"] else ""]
    return "대응 기준: " + ", ".join(i for i in ids if i) if any(ids) else ""


def run_validate(mode):
    """scripts/validate.py 실행 — 오류가 있으면 빌드 중단, 결과 줄을 개요 시트에 기록"""
    cmd = [sys.executable, str(Path(__file__).resolve().parent / "validate.py")]
    if mode:
        cmd.append(mode)
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    lines = [ln for ln in r.stdout.splitlines() if ln.strip()]
    if r.returncode != 0:
        print("\n".join(lines))
        raise SystemExit("검증 오류 — 문구·데이터를 고친 뒤 다시 빌드(--skip-validate로 무시 가능)")
    return next((ln for ln in reversed(lines) if ln.startswith("검증 결과")), "")


VALIDATION = "검증 생략(--skip-validate)"

# ===========================================================================
# 4) xlsx 출력 — 클라우드 v5 · OT · 공급망 · 아이덴티티 서식 준용
# ===========================================================================
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "실사용 기법 포함": "2E86C1",
               "실증·공개 취약점": "B9770E", "이론·시나리오": "7F8C8D"}
HDR, NAVY, GRAY = "34495E", "1F3864", "8C8C8C"


def _w(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hdr(ws, row, cols, start=1, color=HDR):
    for j, c in enumerate(cols, start):
        cell = ws.cell(row, j, c)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BORDER


def _risk(cell, risk):
    bg, fg = RISK_COLOR[risk]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=risk in ("매우 높음", "높음"), color=fg)
    cell.alignment = CENTER


def _level(cell, level):
    cell.font = Font(bold=True, color=LEVEL_COLOR[level])
    cell.alignment = CENTER


def _row(ws, r, vals, start=1):
    for j, v in enumerate(vals, start):
        c = ws.cell(r, j, v)
        c.alignment = WRAP
        c.border = BORDER


def _title(ws, text, sub=None):
    ws.cell(1, 1, text).font = Font(size=14, bold=True, color=NAVY)
    if sub:
        ws.cell(2, 1, sub).font = Font(size=9, color="595959")


def _section(ws, r, text):
    ws.cell(r, 1, text).font = Font(bold=True, size=11, color=NAVY)
    return r + 1


def summary_cell(o):
    return o["summary"] or f"(문구 미작성) {o['en']}"


def _kv_sheet(ws, rows, widths=(22, 140)):
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    _w(ws, list(widths))


def write_xlsx(path):
    wb = openpyxl.Workbook()
    today = datetime.date.today().isoformat()
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    sc = collections.Counter(o["severity"] for o in ROWS)
    st = collections.Counter(e["status"] for e in INCIDENTS)
    real = [e for e in INCIDENTS if e["status"] == "실제 사고"]
    kr = [e for e in real if e.get("region") == "한국"]
    written = sum(1 for o in ROWS if o["text_src"] == "분석가 작성")
    top = sorted(ROWS, key=lambda o: o["priority"])[:10]
    top_stage = [o for o in sorted(ROWS, key=lambda o: o["priority"]) if o["code"] != "IM"][:10]
    n_lv2 = sum(len(d["lv2"]) for d in DOMAINS)
    years = sorted({e["date"][:4] for e in INCIDENTS})
    regions = collections.Counter(e.get("region") for e in INCIDENTS)
    reused = [e for e in INCIDENTS if e.get("origin")]
    atk_rows = [o for o in ROWS if o["attack"]]
    itkb_rows = [o for o in ROWS if EV[o["id"]]["itkb"]]
    scn = SCENARIOS.get("scenarios") or []

    # ---------------- 개요 ----------------
    ws = wb.active
    ws.title = "개요"
    meta = [
        (f"물리·인적 보안위협 매트릭스 {VERSION}", ""),
        ("기준", "물리·인적 공격면 7개 도메인 → 위협분류 → 세부위협 분류에 ATT&CK v19.2(물리·인적 접점)·MITRE 내부자 TTP 지식베이스를 교차 매핑하고, "
                 "AI v3.2·클라우드 v5·OT·공급망·아이덴티티 매트릭스의 위험평가·근거 로직을 이식. 드론·물리 침입·인적 포섭·위장취업·내부자·물리 반출을 다룸"),
        ("생성일", today), ("", ""),
        ("■ 기준 데이터", ""),
        ("물리·인적 보안사고 DB", f"{len(INCIDENTS)}건({years[0]}~{years[-1]}, 공개 출처) — 실제 사고 {st['실제 사고']}건 · 정부 경보 "
                            f"{st.get('정부 경보', 0)}건 · 위협인텔 {st.get('위협인텔', 0)}건 · 실증·연구 {st.get('실증·연구', 0)}건 · 집계 제외 "
                            f"{st.get('제외(검증)', 0)}건, 국내 {regions.get('한국', 0)}건(실제 사고 {len(kr)}건), 최근(2025~) 실제 사고 "
                            f"{sum(1 for e in real if e['recent'])}건, 아이덴티티 사고 DB 재인용 {len(reused)}건(origin 표기)"),
        ("MITRE ATT&CK", f"Enterprise v19.2 — 물리·인적 접점 기법 {len(ATK['techniques'])}개를 {len(atk_rows)}개 행에 연계, 나머지 "
                         f"{len(ROWS) - len(atk_rows)}개 행은 미연계 사유 기록. 사례 수는 물리·인적 맥락 절차만 집계"),
        ("내부자 TTP 지식베이스", f"MITRE CTID ITKB {ITKB_DOC['version']} — 관측 기법 {ITKB_DOC['counts']['techniques']}개 중 행 ATT&CK과 겹치는 "
                             f"{len(itkb_rows)}개 행(위협 주체에 내부자·공모·제3자가 있는 행)을 '내부자 관측 기법'으로 표시(Apache-2.0)"),
        ("대응 기준", "NIST SP 800-53 Rev.5(PE·PS·MP 등) · ISO/IEC 27001:2022 부속서 A 5~7 · ISMS-P 2.2~2.4(+관련 항목) · CERT 내부자 가이드 7판 실천과제 22개"),
        ("참고 문헌", f"{len(REFS)}종 — " + " · ".join(v["short"] for v in REFS.values())
         + " (분류 근거·위협 내용·대응 문구의 인용 자료, 발생가능성 산정에는 넣지 않음)"),
        ("다른 매트릭스", "아이덴티티 v2(IDT-) · 공급망 v2(SCT-) · OT v5(OTC-) · 통합 AI·클라우드·OT v2 요약 ID(AI-·CL-·OT-) 연계 열과 "
                       "'다른 매트릭스 연계' 시트 — 계정·디지털 경로는 원 매트릭스 ID로 넘김"),
        ("", ""),
        ("■ 시트 구성", ""),
        ("보고서용 간략 매트릭스", "본문 삽입용 — 우선순위 순 세부위협 · 핵심 요약 · 위험도 · 실제 사고 수 · 대표 사례 1건 · 핵심 대응 1줄"),
        ("매트릭스 뷰", "도메인(열)별 세부위협을 위험도 색으로 배치한 한눈 보기"),
        ("통합 매트릭스", "분류 체계 · 교차 매핑 · 위험 평가 · 실제 근거 · 대응 기준 · 참고 문헌 전체 열(문구 포함)"),
        ("통합매트릭스_LITE", "핵심 열 발췌(필터·보고용)"),
        ("도메인 요약", "도메인별 경계(넣지 않는 것) · 위협 수 · 위험도/근거 수준 분포 · 실제 사고 · 최고위험 항목 · 맥락 통계(정부 통계)"),
        ("주체·업종 요약", "위협 주체 유형(외부자·내부자·공모·제3자)·적용 프로파일·업종·지역별 세부위협과 실제 사고"),
        ("공격 체인 시나리오", (f"실제 사고 기반 공격 흐름 {len(scn)}개 — 단계별 세부위협(PHT-ID)·행위·탐지 포인트·초크 포인트" if scn
                             else "시트 양식만 수록 — 시나리오(data/scenarios.yaml)는 v2에서 작성")),
        ("역매핑_사고사례", f"사고 DB {len(INCIDENTS)}건과 매핑 세부위협·법적 처리·출처(근거 추적용)"),
        ("교차 매핑 연계", "ATT&CK 기법별 연계 행·물리·인적 맥락 사례 주체 수, ITKB 관측 기법 70개의 연계 여부, 관측 가능한 인적 지표(OHI)"),
        ("대응 기준 연계", "NIST 800-53 · ISO 27001 · ISMS-P · CERT 실천과제별 연계 세부위협과 위험도 분포(통제 우선순위 근거)"),
        ("참고 문헌", "NPSA·CISA 지침 등 참고 문헌 카탈로그 — 쓰는 자리·주 도메인·인용 세부위협·URL·이용 조건·확인 방법"),
        ("다른 매트릭스 연계", "아이덴티티·공급망·OT·클라우드·AI 매트릭스 항목별로 연결된 물리·인적 세부위협"),
        ("평가 기준", "발생가능성·심각도·위험도·근거수준 산정 규칙, 사고 집계·ATT&CK 맥락·ITKB 기준, 경계 규칙, 문구 작성 규칙"),
        ("변경이력", "버전별 변경 내역"), ("", ""),
        ("■ 분류 체계", ""),
        ("Lv1 도메인", "물리·인적 공격면 7개 — " + " → ".join(f"[{d['code']}] {d['ko']}" for d in DOMAINS)),
        ("Lv2 위협분류", f"{n_lv2}개(PHT-<도메인>-NN). 하위 세부위협이 하나뿐이면 Lv2 = Lv3"),
        ("Lv3 세부위협", f"{len(ROWS)}개(PHT-<도메인>-NN.m) — 요약설명 · 참조(물리 관점·실제 사례) · 탐지·대응. ID는 고정(재사용 안 함)"),
        ("설계 원칙", "공격면·주체 7개 도메인 + 공격 단계 열(정찰·접근·실행·지속·내부자화·반출·영향). 한 행 = 한 행위이며, 여러 단계에 걸친 흐름은 "
                    "공격 체인 시나리오로 이음. 국가배후 여부는 행 속성이 아니라 사고 DB 행위자 필드로만 다룸"),
        ("경계", "RC = 사람이 하는 정보 수집 활동 · SV = 장비 기반 수집 · HU = 외부 행위자가 사람을 움직이는 행위 · IN = 내부 신분을 얻거나 가진 주체 · "
               "EX = 물리 반출 경로 · PA = 시설·장비에 대한 물리 접근 · IM = 결과. 계정·디지털 경로는 아이덴티티·클라우드·공급망 매트릭스 ID로 연계"),
        ("", ""),
        ("■ 결과 요약", ""),
        ("세부위협(Lv3)", f"{len(ROWS)}개 / 위협분류 {n_lv2}개 / 도메인 {len(DOMAINS)}개"),
        ("위험도", " · ".join(f"{k} {rc.get(k, 0)}" for k in R.RISK_ORDER)),
        ("근거 수준", " · ".join(f"{k} {lc.get(k, 0)}" for k in R.LEVELS)),
        ("심각도", " · ".join(f"{k} {sc.get(k, 0)}" for k in ("상", "중", "하"))
         + " — '상'은 핵심기술 통째 유출·기반시설 운영·안전·장기 잠복 내부 접근일 때만(평가 기준 시트)"),
        ("물리 관점 문구", f"{written}/{len(ROWS)}개 작성(핵심 요약·요약설명·참조·탐지·대응)"
                       + (" — 전 항목 완료" if written == len(ROWS) else
                          f" — 나머지 {len(ROWS) - written}개는 '문구 미작성'(요약설명 자리에 영문명 회색 표시, 대표 사례는 사고 DB 최신 근거, "
                          "핵심 대응은 매핑된 대응 기준 ID)")
                       + f" | scripts/validate.py {VALIDATION}"),
        ("우선 위협 Top 10", "\n".join(f"{o['priority']}. {o['id']} {o['name']} — {o['risk']}, 실제 사고 {o['n_real']}건"
                                    f"(최근 {o['n_recent']} · 국내 {o['n_kr']})" for o in top)),
        ("우선 위협 Top 10(IM 제외)", "\n".join(f"{o['priority']}. {o['id']} {o['name']} — {o['risk']}, 실제 사고 {o['n_real']}건"
                                           for o in top_stage)),
        ("", ""), ("■ 주의", ""),
        ("발생가능성", "공개 사고 기준이라 기소·판결로 드러나는 기술 유출·내부자 사건(IN·HU·EX)과 공항 운영 방해 드론 사건이 높게 나오고, "
                     "물리 침입·감시·정찰(PA·RC·SV 일부)처럼 공개되지 않거나 귀속되지 않는 위협은 과소평가될 수 있음. 수식은 다른 매트릭스와 같게 "
                     f"유지(상 = 실제 사고 {R.LIKELY_HIGH_REAL}건 이상)"),
        ("근거 상한", "정부 경보·위협인텔·ATT&CK 사례·ITKB 관측 기법·연구 실증은 발생가능성 '중'까지만 반영. 정부 통계(검거·적발 건수)는 사고로 세지 않고 "
                    "'도메인 요약' 맥락 통계로만 표시"),
        ("드론", "공항·기반시설 운영 방해와 군사시설 촬영 사례는 많지만 기업 기술 스파이로 귀속된 공개 사례는 확인하지 못함. 드론 무력화·전파 교란은 "
               "법적 권한이 있는 기관의 영역이라 대응 문구에서 권고하지 않음(탐지·식별·신고·차폐 중심)"),
        ("영향·피해 도메인", "[IM] 행은 결과 유형이라 여러 사고가 몰려 우선순위 상위에 오름 — 통제 우선순위는 앞 단계 도메인과 '대응 기준 연계' 시트로 판단"),
        ("근거 중복", "한 사고가 여러 단계(포섭 → 내부자 반출 → 기술 유출)에 동시에 매핑되므로 세부위협별 사고 수 합은 사고 DB 건수보다 큼. "
                    "아이덴티티 사고 DB와 겹치는 사고는 origin 열로 표시(매트릭스 간 합산 금지)"),
        ("다른 매트릭스와 비교", f"사고 DB 규모·출처가 매트릭스마다 달라(클라우드 680건·OT 85건·공급망 167건·신원 126건·물리·인적 {len(INCIDENTS)}건) "
                           "위험도·사고 수는 같은 매트릭스 안에서 비교. 같은 위협도 매트릭스별 심각도 기준이 달라 등급이 다를 수 있음"),
        ("방어 자료", "행·문구는 위험 자산·탐지 지점·통제 중심으로 상위 수준에서 서술 — 침입·도청·촬영·포섭의 실행 방법은 다루지 않음. 인적 지표는 객관적 "
                    "행동·직무 사실만 쓰고 출신·종교·성별 등 보호 특성은 쓰지 않음"),
        ("출처 표기", "MITRE ATT&CK® © The MITRE Corporation · ITKB © MITRE CTID(Apache-2.0) · NIST SP 800-53(공공 영역) · ISO/IEC 27001(유료 표준, "
                    "번호·제목만) · ISMS-P · CERT Common Sense Guide © CMU(번호·제목만) · NPSA(Crown copyright·GOV.UK OGL v3.0, 요지 의역) · "
                    "CISA(공공 영역). 사고별 출처는 '역매핑_사고사례' 시트"),
    ]
    _kv_sheet(ws, meta, (24, 140))
    ws.cell(1, 1).font = Font(size=15, bold=True, color=NAVY)

    # ---------------- 보고서용 간략 매트릭스 ----------------
    ws = wb.create_sheet("보고서용 간략 매트릭스")
    bcols = ["순위", "도메인(Lv1)", "PHT-ID", "세부위협(Lv3)", "핵심 요약", "위험도", "발생가능성", "심각도", "실제 사고 수",
             "국내", "최근(2025~)", "대표 사례", "핵심 대응"]
    ws.cell(1, 1, f"보고서용 간략 매트릭스 {VERSION} — 본문 삽입용(상세 문구는 '통합 매트릭스')").font = Font(size=13, bold=True, color=NAVY)
    ws.cell(2, 1, "순위 = 위험도 → 실제 사고 수 → 최근 사고(2025~) → PHT-ID | 대표 사례 = 참조 '■ 실제 사례' 첫 줄 | "
                  f"핵심 대응 = '■ 대응' 첫 줄 | {today}").font = Font(size=9, color="595959")
    _hdr(ws, 4, bcols)
    for r, o in enumerate(sorted(ROWS, key=lambda o: o["priority"]), 5):
        _row(ws, r, [o["priority"], o["domain"], o["id"], o["name"], o["oneline"] or summary_cell(o), o["risk"], o["likelihood"],
                     o["severity"], o["n_real"], o["n_kr"], o["n_recent"], first_case(o), first_control(o)])
        ws.cell(r, 3).font = Font(bold=True)
        ws.cell(r, 4).font = Font(bold=True)
        if not o["oneline"]:
            ws.cell(r, 5).font = Font(color=GRAY)
        _risk(ws.cell(r, 6), o["risk"])
        for j in (1, 7, 8, 9, 10, 11):
            ws.cell(r, j).alignment = CENTER
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(bcols))}{len(ROWS) + 4}"
    _w(ws, [6, 16, 12, 28, 44, 9, 7, 7, 7, 6, 7, 46, 54])
    ws.print_title_rows = "4:4"

    # ---------------- 매트릭스 뷰 ----------------
    ws = wb.create_sheet("매트릭스 뷰")
    by_c = collections.defaultdict(list)
    for o in ROWS:
        by_c[o["code"]].append(o)
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도 (빨강 매우 높음 · 주황 높음 · 노랑 보통 · 초록 낮음), [n] = 실제 사고 수, "
                  "괄호 안 숫자 = 세부위협 수").font = Font(bold=True, color=NAVY)
    for j, d in enumerate(DOMAINS, 1):
        c = ws.cell(2, j, f"[{d['code']}] {d['ko']}\n{d['en']}\n({len(by_c[d['code']])})")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        for i, o in enumerate(sorted(by_c[d["code"]], key=lambda o: (rank_key[o["risk"]], -o["n_real"], o["id"])), 3):
            cell = ws.cell(i, j, f"{o['id']} {o['name']}" + (f" [{o['n_real']}]" if o["n_real"] else ""))
            _risk(cell, o["risk"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9, bold=o["risk"] == "매우 높음", color=RISK_COLOR[o["risk"]][1])
            cell.border = BORDER
    ws.row_dimensions[2].height = 48
    ws.freeze_panes = "A3"
    _w(ws, [26] * len(DOMAINS))

    # ---------------- 통합 매트릭스 ----------------
    ws = wb.create_sheet("통합 매트릭스")
    g_class = ["도메인(Lv1)", "PHT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "참조", "영문명", "공격 단계", "설명 출처"]
    g_cross = ([f"주체:{a}" for a in R.ACTOR_TYPES] + ["적용 프로파일", "자산·통제 지점", "ATT&CK ID", "ATT&CK 기법명", "ATT&CK 미연계 사유",
                                                    "ITKB 내부자 관측 기법", "아이덴티티 연계", "공급망 연계", "OT 연계", "클라우드 연계",
                                                    "AI 연계"])
    g_risk = ["발생가능성", "심각도", "위험도", "우선순위", "발생가능성 근거 (자동 산정)", "심각도 근거"]
    g_ev = ["실제 사고 수", "최근 사고(2025~)", "국내 실제 사고", "정부 경보", "위협인텔", "실증·연구 수", "ATT&CK 사례 수(물리·인적 맥락)",
            "근거 수준", "관련 사례 ID", "주요 업종(실제 사고)", "ATT&CK 사례 주체(일부)"]
    g_ctl = ["NIST SP 800-53", "ISO/IEC 27001:2022", "ISMS-P", "CERT 실천과제", "탐지·대응 포인트", "참고 문헌"]
    groups = [("분류 체계", len(g_class), "2E5496"), ("교차 매핑", len(g_cross), "1F7A8C"),
              ("위험 평가", len(g_risk), "A04000"), ("실제 근거", len(g_ev), "1E8449"), ("대응 기준·참고 문헌", len(g_ctl), "6C3483")]
    cols = g_class + g_cross + g_risk + g_ev + g_ctl
    t = ws.cell(1, 1, f"물리·인적 보안위협 매트릭스 {VERSION} — 분류체계 · 위험평가 · 실제근거 · 대응기준")
    t.font = Font(size=13, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, f"물리·인적 보안사고 DB {len(INCIDENTS)}건 + ATT&CK v19.2 물리·인적 접점 + MITRE 내부자 TTP 지식베이스 | {today}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    c0 = 1
    for g, span, color in groups:
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + span - 1)
        c = ws.cell(3, c0, g)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        c0 += span
    _hdr(ws, 4, cols)
    ci = {c: i + 1 for i, c in enumerate(cols)}
    for r, o in enumerate(ROWS, 5):
        x = EV[o["id"]]
        subj = sorted(x["subj"])
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summary_cell(o), o["reference"], o["en"],
                     ", ".join(o["stage"]), o["text_src"]]
             + ["●" if a in o["actors"] else "" for a in R.ACTOR_TYPES]
             + [" · ".join(o["profiles"]), " · ".join(o["assets"]), ", ".join(o["attack"]),
                "\n".join(atk_label(t) for t in o["attack"]), o.get("attack_none", ""),
                "\n".join(atk_label(t) for t in sorted(x["itkb"])),
                "\n".join(link_label(k) for k in o["idt"]), "\n".join(link_label(k) for k in o["sct"]),
                "\n".join(link_label(k) for k in o["ot"]), "\n".join(link_label(k) for k in o["cloud"]),
                "\n".join(link_label(k) for k in o["ai"]),
                o["likelihood"], o["severity"], o["risk"], o["priority"], o["lk_why"], o["severity_why"],
                o["n_real"], o["n_recent"], o["n_kr"], o["n_adv"], o["n_intel"], o["n_research"], o["n_atk"], o["level"],
                ", ".join(sorted(x["real"]))
                + (("\n정부 경보: " + ", ".join(sorted(x["advisory"]))) if x["advisory"] else "")
                + (("\n위협인텔: " + ", ".join(sorted(x["intel"]))) if x["intel"] else "")
                + (("\n실증: " + ", ".join(sorted(x["research"]))) if x["research"] else "")
                + (("\n집계 제외: " + ", ".join(sorted(x["excluded"]))) if x["excluded"] else ""),
                " · ".join(f"{k} {v}" for k, v in x["sector"].most_common(4)),
                ", ".join(f"{s} {ATK['subjects'][s]['name']}" for s in subj[:12]) + (f" 외 {len(subj) - 12}" if len(subj) > 12 else ""),
                "\n".join(nist_label(k) for k in o["nist"]), "\n".join(fw_label("iso27001", str(k)) for k in o["iso"]),
                "\n".join(fw_label("ismsp", str(k)) for k in o["ismsp"]), "\n".join(fw_label("cert", k) for k in o["cert"]),
                o["detect"], "\n".join(ref_label(k) for k in o["refs"])])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, ci["요약설명"]).font = Font(color=GRAY)
        for c in [f"주체:{a}" for a in R.ACTOR_TYPES] + ["발생가능성", "심각도", "우선순위"]:
            ws.cell(r, ci[c]).alignment = CENTER
        _risk(ws.cell(r, ci["위험도"]), o["risk"])
        _level(ws.cell(r, ci["근거 수준"]), o["level"])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(ROWS) + 4}"
    widths = {"도메인(Lv1)": 16, "PHT-ID": 12, "위협분류(Lv2)": 18, "세부위협(Lv3)": 26, "핵심 요약": 34, "요약설명": 60, "참조": 80,
              "영문명": 26, "공격 단계": 12, "설명 출처": 9, "적용 프로파일": 18, "자산·통제 지점": 22, "ATT&CK ID": 14, "ATT&CK 기법명": 30,
              "ATT&CK 미연계 사유": 30, "ITKB 내부자 관측 기법": 26, "아이덴티티 연계": 26, "공급망 연계": 26, "OT 연계": 22,
              "클라우드 연계": 20, "AI 연계": 20, "발생가능성": 7, "심각도": 7, "위험도": 8, "우선순위": 7,
              "발생가능성 근거 (자동 산정)": 30, "심각도 근거": 44, "실제 사고 수": 7, "최근 사고(2025~)": 7, "국내 실제 사고": 7,
              "정부 경보": 7, "위협인텔": 7, "실증·연구 수": 7, "ATT&CK 사례 수(물리·인적 맥락)": 8, "근거 수준": 11, "관련 사례 ID": 24,
              "주요 업종(실제 사고)": 22, "ATT&CK 사례 주체(일부)": 30, "NIST SP 800-53": 26, "ISO/IEC 27001:2022": 26,
              "ISMS-P": 18, "CERT 실천과제": 30, "탐지·대응 포인트": 70, "참고 문헌": 22}
    _w(ws, [widths.get(c, 6) for c in cols])

    # ---------------- 통합매트릭스_LITE ----------------
    ws = wb.create_sheet("통합매트릭스_LITE")
    ws.cell(1, 1, f"물리·인적 보안위협 매트릭스 {VERSION} — LITE").font = Font(size=13, bold=True, color=NAVY)
    lcols = ["도메인(Lv1)", "PHT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "공격 단계", "위협 주체", "적용 프로파일",
             "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수", "국내", "최근 사고(2025~)", "탐지·대응 포인트"]
    lc_ = {c: i + 1 for i, c in enumerate(lcols)}
    _hdr(ws, 2, lcols)
    for r, o in enumerate(ROWS, 3):
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summary_cell(o), ", ".join(o["stage"]),
                     " · ".join(o["actors"]), " · ".join(o["profiles"]), o["likelihood"], o["severity"], o["risk"], o["priority"],
                     o["level"], o["n_real"], o["n_kr"], o["n_recent"], o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, lc_["요약설명"]).font = Font(color=GRAY)
        _risk(ws.cell(r, lc_["위험도"]), o["risk"])
        _level(ws.cell(r, lc_["근거 수준"]), o["level"])
        for c in ("발생가능성", "심각도", "우선순위", "실제 사고 수", "국내", "최근 사고(2025~)"):
            ws.cell(r, lc_[c]).alignment = CENTER
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(lcols))}{len(ROWS) + 2}"
    _w(ws, [16, 12, 18, 26, 34, 60, 12, 18, 22, 7, 7, 8, 7, 11, 7, 6, 7, 70])

    # ---------------- 도메인 요약 ----------------
    ws = wb.create_sheet("도메인 요약")
    dcols = ["도메인", "설명", "넣지 않는 것(경계)", "위협분류(Lv2)", "세부위협(Lv3)"] + R.RISK_ORDER + R.LEVELS + \
            ["매핑 실제 사고(중복 제거)", "국내", "최근 사고(2025~)", "최고위험 세부위협(상위 3)", "맥락 통계(사고로 세지 않음)"]
    _hdr(ws, 1, dcols)
    for i, d in enumerate(DOMAINS, 2):
        os_ = by_c[d["code"]]
        rr = collections.Counter(o["risk"] for o in os_)
        ll = collections.Counter(o["level"] for o in os_)
        incs = set().union(*(EV[o["id"]]["real"] for o in os_))
        top3 = sorted(os_, key=lambda o: o["priority"])[:3]
        ctx = "\n".join(f"· {c['text']} ({c['source'].split(' | ')[0]})" for c in d.get("context") or [])
        _row(ws, i, [f"[{d['code']}] {d['ko']}", d["desc"], d["excludes"], len(d["lv2"]), len(os_)]
             + [rr.get(k, 0) for k in R.RISK_ORDER] + [ll.get(k, 0) for k in R.LEVELS]
             + [len(incs), sum(1 for e in incs if INC[e].get("region") == "한국"), sum(1 for e in incs if INC[e]["recent"]),
                "\n".join(f"{o['id']} {o['name']}({o['risk']}, {o['n_real']}건)" for o in top3), ctx])
    ws.freeze_panes = "B2"
    _w(ws, [20, 50, 40, 8, 8, 7, 7, 7, 7, 9, 9, 9, 9, 10, 7, 9, 50, 60])

    # ---------------- 주체·업종 요약 ----------------
    ws = wb.create_sheet("주체·업종 요약")
    _title(ws, "위협 주체 유형·적용 프로파일·업종·지역별 요약 — 실제 사고 기준(한 사고가 여러 유형·업종에 걸치면 각각 집계)",
           "세부위협 수 = 그 주체·프로파일을 대상으로 하는 세부위협(taxonomy actors·profiles) · 업종·지역은 사고 DB 기준")
    ecols = ["구분", "항목", "세부위협 수", "매우 높음·높음", "실제 사고", "국내", "최근(2025~)", "주요 영향", "대표 세부위협(사고 수 상위 3)",
             "최근 사례(최대 3)"]
    r = 4

    def summary_block(label, items, row_pred, inc_pred):
        nonlocal r
        _hdr(ws, r, ecols)
        r += 1
        for it in items:
            rows_ = [o for o in ROWS if row_pred(o, it)] if row_pred else []
            er = [e for e in INCIDENTS if e["status"] == "실제 사고" and inc_pred(e, it)] if inc_pred else []
            imp = collections.Counter(x for e in er for x in e["impact"] if x not in ("영향 없음(사전 차단)", "영향 미상"))
            rows_c = collections.Counter(k for e in er for k in e["pht"])
            recent = sorted(er, key=lambda e: e["date"], reverse=True)[:3]
            if not rows_c and rows_:
                rows_c = collections.Counter({o["id"]: o["n_real"] for o in sorted(rows_, key=lambda o: o["priority"])[:3]})
            _row(ws, r, [label, it, len(rows_) if row_pred else "", sum(1 for o in rows_ if o["risk"] in ("매우 높음", "높음"))
                         if row_pred else "", len(er) if inc_pred else "", sum(1 for e in er if e.get("region") == "한국") if inc_pred else "",
                         sum(1 for e in er if e["recent"]) if inc_pred else "", " · ".join(f"{k} {v}" for k, v in imp.most_common(3)),
                         "\n".join(f"{k} {ROW[k]['name']}({v})" for k, v in rows_c.most_common(3)),
                         "\n".join(inc_label(e["id"]) for e in recent)])
            r += 1
        r += 1

    summary_block("위협 주체", R.ACTOR_TYPES, lambda o, it: it in o["actors"], lambda e, it: it in e["actor_type"])
    summary_block("적용 프로파일", R.PROFILES, lambda o, it: it in o["profiles"], None)
    summary_block("업종(사고 DB)", [s for s in R.SECTORS if any(s in e["sector"] for e in INCIDENTS)], None,
                  lambda e, it: it in e["sector"])
    summary_block("지역(사고 DB)", [g for g in R.REGIONS if any(e.get("region") == g for e in INCIDENTS)], None,
                  lambda e, it: e.get("region") == it)
    ws.freeze_panes = "C5"
    _w(ws, [12, 18, 8, 9, 8, 6, 8, 34, 46, 60])

    # ---------------- 공격 체인 시나리오 ----------------
    ws = wb.create_sheet("공격 체인 시나리오")
    ws.cell(1, 1, "공격 체인 시나리오 — 실제 사고를 물리·인적 공격 흐름으로 재구성. 각 단계는 세부위협(PHT-ID, 위험도 색)에 연결하고, "
                  "흐름을 가장 싸게 끊는 통제를 초크 포인트로 표시").font = Font(bold=True, color=NAVY)
    ws.cell(2, 1, "행위는 사고 DB 요약에 있는 사실만 적음(data/scenarios.yaml, scripts/validate.py가 ID 대조)").font = Font(size=9, color="595959")
    r = 4
    if not scn:
        ws.cell(r, 1, "작성 예정 — data/scenarios.yaml은 v2에서 작성").font = Font(color=GRAY)
    for s in scn:
        c = ws.cell(r, 1, f"{s['id']}. {s['title']}")
        c.font = Font(size=12, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
        info = [("개요", s["summary"]), ("근거 사고", " · ".join(inc_label(i) for i in s["incidents"]))]
        if s.get("ref"):
            info.append(("참조", s["ref"]))
        for k, v in info:
            ws.cell(r, 1, k).font = Font(bold=True, size=9, color="595959")
            ws.cell(r, 2, v).alignment = WRAP
            ws.cell(r, 2).font = Font(size=9)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            r += 1
        _hdr(ws, r, ["#", "공격 단계", "PHT-ID", "세부위협", "행위", "탐지 포인트"])
        r += 1
        for i, st_ in enumerate(s["steps"], 1):
            rows_ = [ROW[k] for k in st_["pht"]]
            _row(ws, r, [i, st_["단계"], "\n".join(o["id"] for o in rows_), "\n".join(o["name"] for o in rows_), st_["행위"],
                         st_["탐지"]])
            _risk(ws.cell(r, 3), min((o["risk"] for o in rows_), key=lambda x: rank_key[x]))
            ws.cell(r, 1).alignment = CENTER
            r += 1
        ws.cell(r, 1, "초크").font = Font(bold=True, color="A04000")
        ws.cell(r, 2, "\n".join(f"· {c}" for c in s["chokepoints"])).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 15 * len(s["chokepoints"]) + 4
        r += 2
    _w(ws, [6, 12, 14, 30, 56, 50])

    # ---------------- 역매핑_사고사례 ----------------
    ws = wb.create_sheet("역매핑_사고사례")
    icols = ["사례 ID", "시점", "사례명", "사례유형", "검증 수준", "집계 상태", "법적 처리", "지역", "업종", "행위자", "주체 유형", "영향",
             "요약", "도메인", "매핑 세부위협", "ATT&CK 참조", "다른 DB 사고 ID", "출처"]
    _hdr(ws, 1, icols)
    for i, e in enumerate(sorted(INCIDENTS, key=lambda e: (e["date"], e["id"])), 2):
        _row(ws, i, [e["id"], e["date"], e["title"], e["kind"], e["verification"], e["status"], e.get("legal", ""), e.get("region", ""),
                     ", ".join(e["sector"]), e.get("actor", ""), ", ".join(e["actor_type"]), ", ".join(e["impact"]), e["summary"],
                     " · ".join(e["domains"]), "\n".join(f"{k} {ROW[k]['name']}" for k in e["pht"]),
                     ", ".join(e.get("attack_ref") or []), ", ".join(e.get("origin") or []), "\n".join(e["sources"])])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(icols))}{len(INCIDENTS) + 1}"
    _w(ws, [9, 8, 34, 9, 13, 10, 10, 7, 16, 18, 14, 18, 70, 10, 40, 12, 10, 70])

    # ---------------- 교차 매핑 연계 ----------------
    ws = wb.create_sheet("교차 매핑 연계")
    _title(ws, "교차 매핑 항목별 연계 세부위협 — ATT&CK 물리·인적 접점 · MITRE 내부자 TTP 지식베이스(ITKB) · 관측 가능한 인적 지표(OHI)",
           "ATT&CK 사례 수 집계: 전체 = 물리 기법(장치 추가·물리 매체·근접 Wi-Fi·하드웨어 공급망·이동식 매체) · 인적 단서 = 위장 프로필·인물 정찰 "
           "기법 중 LinkedIn·채용·위장 프로필·직원 식별 단서가 있는 절차 · 미집계 = 절차가 원격 사이버 수법뿐(교차 매핑만)")
    r = 4
    r = _section(ws, r, f"■ MITRE ATT&CK v19.2 — 행에 연계한 기법 {len(ATK['techniques'])}개")
    _hdr(ws, r, ["ID", "기법명", "전술", "연계 세부위협", "집계 규칙", "물리·인적 맥락 사례 주체 수", "ITKB 관측", "연계 세부위협 위험도"])
    r += 1
    for t in sorted(ATK["techniques"]):
        rows = [o for o in ROWS if t in o["attack"]]
        rule = "전체" if t in R.ATTACK_COUNT_ALL else "인적 단서" if t in R.ATTACK_COUNT_KEYWORD else "미집계"
        rcnt = collections.Counter(o["risk"] for o in rows)
        _row(ws, r, [t, ATK["techniques"][t]["name"], ATK["techniques"][t]["tactics"], "\n".join(f"{o['id']} {o['name']}" for o in rows),
                     rule, len(atk_subjects(t)), "●" if t in ITKB else "", " · ".join(f"{x} {rcnt[x]}" for x in R.RISK_ORDER if rcnt[x])])
        r += 1
    r += 1
    r = _section(ws, r, f"■ ATT&CK 미연계 세부위협 {len(ROWS) - len(atk_rows)}개와 사유")
    _hdr(ws, r, ["PHT-ID", "세부위협", "미연계 사유", "", "", "", "", ""])
    r += 1
    for o in [o for o in ROWS if not o["attack"]]:
        _row(ws, r, [o["id"], o["name"], o.get("attack_none", "")])
        r += 1
    r += 1
    r = _section(ws, r, f"■ ITKB 관측 기법 {ITKB_DOC['counts']['techniques']}개(기법 {ITKB_DOC['counts']['parent']} · 하위 "
                        f"{ITKB_DOC['counts']['sub']}) — {ITKB_DOC['version']}")
    _hdr(ws, r, ["ID(v14)", "기법명", "전술", "연계 세부위협", "v19.2 ID", "완화책 수", "데이터 소스 수", "비고"])
    r += 1
    for t, v in ITKB.items():
        rows = [o for o in ROWS if t in EV[o["id"]]["itkb"]]
        note = v.get("note", "") or ("" if rows else "IT 시스템 위 내부자 행위 — 디지털 경로는 아이덴티티·클라우드 매트릭스 범위(물리·인적 행 미연계)")
        _row(ws, r, [t, v["name"], ", ".join(v["tactics"]), "\n".join(f"{o['id']} {o['name']}" for o in rows), v.get("v19", ""),
                     len(v["mitigations"]), len(v["datasources"]), note])
        r += 1
    r += 1
    r = _section(ws, r, "■ 관측 가능한 인적 지표(OHI) 9종 — IN 행 탐지 문구 참고(객관적 행동·직무 사실만, 보호 특성 제외)")
    for k in ITKB_DOC["ohi"]:
        _row(ws, r, ["OHI", k])
        r += 1
    _w(ws, [12, 36, 30, 46, 14, 12, 12, 44])

    # ---------------- 대응 기준 연계 ----------------
    ws = wb.create_sheet("대응 기준 연계")
    _title(ws, "대응 기준별 연계 세부위협 — 연계 위협의 위험도 분포로 통제 우선순위를 가늠",
           "NIST 800-53 = 조직 통제 · ISO 27001 = 국제 인증 통제 · ISMS-P = 국내 인증기준 · CERT = 내부자 위험 실천과제")
    r = 4
    used_nist = sorted({k for o in ROWS for k in o["nist"]})
    ctl_sets = [("NIST SP 800-53 Rev.5", "nist", used_nist, lambda k: NIST.get(k, ""), lambda k: ""),
                ("ISO/IEC 27001:2022 부속서 A", "iso", [str(k) for k in FW["iso27001"]], lambda k: fw_pair("iso27001", k)[0],
                 lambda k: fw_pair("iso27001", k)[1]),
                ("ISMS-P 인증기준", "ismsp", [str(k) for k in FW["ismsp"]], lambda k: "", lambda k: fw_pair("ismsp", k)[0]),
                ("CERT Common Sense Guide 7판 실천과제", "cert", [str(k) for k in FW["cert"]], lambda k: fw_pair("cert", int(k))[0],
                 lambda k: fw_pair("cert", int(k))[1])]
    for title, fld, keys, en, ko in ctl_sets:
        r = _section(ws, r, f"■ {title}")
        _hdr(ws, r, ["ID", "명칭(영문)", "명칭(국문)", "연계 세부위협 수", "매우 높음·높음 위협 수", "연계 세부위협", "미연계 사유"])
        r += 1
        for k in keys:
            rows = [o for o in ROWS if k in [str(v) for v in o[fld]]]
            hi = sum(1 for o in rows if o["risk"] in ("매우 높음", "높음"))
            why = R.FRAMEWORK_NOT_MAPPED.get((fld, k)) or R.FRAMEWORK_NOT_MAPPED.get((fld, int(k)) if k.isdigit() else None, "")
            _row(ws, r, [k, en(k), ko(k), len(rows), hi,
                         "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in sorted(rows, key=lambda o: o["priority"])), why])
            r += 1
        r += 1
    _w(ws, [10, 50, 32, 8, 10, 56, 40])

    # ---------------- 참고 문헌 ----------------
    ws = wb.create_sheet("참고 문헌")
    _title(ws, f"참고 문헌 카탈로그 {len(REFS)}종 — 분류 근거·위협 내용·대응 문구의 인용 자료(발생가능성 산정에는 넣지 않음)",
           "행 refs → '통합 매트릭스' 참고 문헌 열, 문구는 ■ 물리 관점·■ 대응 끝에 '참고: 출처 〈자료〉'로만 인용(작업방향 6.5.4). 원문·PDF 비수록")
    rcols = ["키", "인용 표기", "자료", "발행", "연도", "쓰는 자리", "주 도메인", "인용 세부위협", "URL", "아카이브", "이용 조건", "확인 방법", "확인일"]
    _hdr(ws, 4, rcols)
    for r, (k, v) in enumerate(REFS.items(), 5):
        rows = [o for o in ROWS if k in o["refs"]]
        _row(ws, r, [k, v["short"], v["title"], v["publisher"], str(v["year"]), " · ".join(v["use"]), " · ".join(v["domains"]),
                     "\n".join(f"{o['id']} {o['name']}" for o in rows), "\n".join(v["urls"]), v.get("archive", ""), v["terms"],
                     v["verify"], str(v["checked"])])
    ws.freeze_panes = "C5"
    _w(ws, [12, 24, 44, 18, 10, 18, 14, 40, 50, 30, 40, 40, 11])

    # ---------------- 다른 매트릭스 연계 ----------------
    ws = wb.create_sheet("다른 매트릭스 연계")
    _title(ws, "다른 매트릭스 항목 ↔ 물리·인적 세부위협 — 계정·디지털 경로는 원 매트릭스에서, 사람·시설·물리 경로 통제는 이 매트릭스에서",
           "등급은 각 매트릭스 기준(사고 DB·심각도 기준이 달라 직접 비교하지 않음). 원 매트릭스 위험도는 각 매트릭스 산출물 기준")
    _hdr(ws, 4, ["매트릭스", "연계 ID", "원 매트릭스 항목명", "원 매트릭스 위험도", "연계 물리·인적 세부위협", "물리·인적 위험도(최고)"])
    r = 5
    for fld, mx in (("idt", "아이덴티티"), ("sct", "공급망"), ("ot", "OT"), ("cloud", "클라우드"), ("ai", "AI")):
        for k in sorted({k for o in ROWS for k in o[fld]}):
            rows = sorted([o for o in ROWS if k in o[fld]], key=lambda o: o["priority"])
            best = min((o["risk"] for o in rows), key=lambda x: rank_key[x])
            _row(ws, r, [mx, k, LINKS.get(k, {}).get("name", ""), LINKS.get(k, {}).get("risk", ""),
                         "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in rows), best])
            _risk(ws.cell(r, 6), best)
            r += 1
    ws.freeze_panes = "C5"
    _w(ws, [10, 14, 40, 10, 60, 12])

    # ---------------- 평가 기준 ----------------
    ws = wb.create_sheet("평가 기준")
    A = TAX["assessment"]
    crit = [
        ("■ 위험도", ""),
        ("위험도", "발생가능성 × 심각도 3×3 표(AI·클라우드·OT·공급망·아이덴티티 매트릭스와 동일) — 상×상 매우 높음, 상×중·중×상 높음, "
                "상×하·중×중·하×상 보통, 나머지 낮음"),
        ("발생가능성", A["likelihood_rule"]),
        ("심각도 상", A["severity_rule"]["상"]),
        ("심각도 중", A["severity_rule"]["중"]),
        ("심각도 하", A["severity_rule"]["하"]),
        ("심각도 점검", A["severity_check"]),
        ("우선순위", "위험도 → 실제 사고 수 → 최근 사고(2025~) → PHT-ID 순(1 = 최우선)"),
        ("근거 수준", "실제 사고 확인(사고 DB 실제 사고) > 실사용 기법 포함(ATT&CK 물리·인적 맥락 사례·ITKB 내부자 관측 기법·정부 경보·위협인텔) > "
                    "실증·공개 취약점(사고 DB 연구·시연) > 이론·시나리오"),
        ("", ""), ("■ 근거 집계", ""),
        ("실제 사고", f"사례유형 ∈ {{{'·'.join(sorted(R.REAL_KINDS))}}}이고 검증 수준 ∉ {{{'·'.join(sorted(R.EXCLUDED_VERIFICATION))}}}. "
                    "정부 경보는 실제 사고로 세지 않고 실사용 근거('중' 상한)로만 반영(작업방향 8장)"),
        ("최근 사고", f"사고 인지·공개 시점 {R.RECENT_FROM} 이후"),
        ("국내 사고", "피해 지역(region)이 한국인 실제 사고 — 국내 보고서 독자용 보조 열"),
        ("사고 매핑", "분석가가 사고 경위를 읽고 단계별 세부위협에 매핑(한 사고가 포섭·반출·기술 유출 등 여러 행에 매핑될 수 있음). "
                    "ATT&CK 그룹·소프트웨어·캠페인 ID는 대조·표시용이며 자동 매핑에 쓰지 않음. 다른 매트릭스 사고 DB와 같은 사고는 origin으로 표시"),
        ("ATT&CK 사례 수", "행의 ATT&CK 기법을 쓰는 주체(그룹·소프트웨어·캠페인) 수. 장치 추가·물리 매체 반출·근접 Wi-Fi·하드웨어 공급망·이동식 매체 "
                         "기법은 절차 전체를, 위장 프로필·인물 정찰 기법은 절차 설명에 인적 단서(LinkedIn·채용·위장 프로필·직원 식별 등)가 있는 것만 셈. "
                         "유효 계정·정보 저장소 수집·원격 접속 도구·사칭 등은 절차가 원격 사이버 수법뿐이라 세지 않음('교차 매핑 연계' 시트)"),
        ("ITKB 관측 기법", "MITRE CTID 내부자 TTP 지식베이스에서 실제 내부자 사례로 관측된 기법이 행의 ATT&CK과 겹치고, 행의 위협 주체에 내부자·공모·"
                        "제3자가 있으면 '내부자 관측 기법'으로 표시하고 실사용 근거('중' 상한)로 반영"),
        ("맥락 통계", "경찰청 검거 건수·국정감사 드론 적발 건수 같은 정부 통계는 사고로 세지 않고 '도메인 요약'에만 표시"),
        ("", ""), ("■ 경계 규칙(한 행 = 한 행위)", ""),
    ] + [(f"[{d['code']}] {d['ko']}", f"{d['desc']} | 넣지 않는 것: {d['excludes']}") for d in DOMAINS] + [
        ("", ""), ("■ 다른 매트릭스와의 근거 대응", ""),
        ("아이덴티티 매트릭스", "실제 사고 ↔ 신원 사고 DB / 실사용 ↔ ATT&CK 신원 맥락·KEV(신원)·위협인텔 / 실증 ↔ 연구·시연·SAT 시연"),
        ("공급망 매트릭스", "실제 사고 ↔ 공급망 사고 DB / 실사용 ↔ ATT&CK 공급망 맥락·KEV(공급망) / 실증 ↔ 연구·시연"),
        ("물리·인적 매트릭스", "실제 사고 ↔ 물리·인적 사고 DB / 실사용 ↔ ATT&CK 물리·인적 맥락·ITKB 관측 기법·정부 경보·위협인텔 / 실증 ↔ 연구·시연"),
        ("", ""), ("■ 사례 라벨(문구 작성용)", ""),
        ("[실제 사고]", "사고 DB에서 '실제 사고'로 집계되는 사례"),
        ("[정부 경보]", "사고 DB에서 '정부 경보'로 집계되는 정부 기관 경보·조사 결과"),
        ("[위협인텔]", "사고 DB에서 '위협인텔'로 집계되는 보안업체·언론의 위협 보고"),
        ("[실증]", "연구자·보안업체가 재현한 공격(사고 DB 연구·시연)"),
        ("[시나리오]", "공개 사례가 없는 위협의 가상 흐름(사고 ID 없이)"),
        ("참고:", "참고 문헌 인용 — ■ 물리 관점·■ 대응 끝에 '참고: 출처 〈자료〉'로만(사례 라벨 아님)"),
        ("", ""), ("■ 물리 관점 문구 작성 규칙(data/text/*.yaml)", ""),
        ("핵심 요약", "한 줄 50자 이내, '수단 + 대상 + 결과'를 명사형으로 — 보고서 본문·슬라이드용"),
        ("요약설명", "2줄 개조식 — ① 공격자(내부자 등)는 [수단·경로]로 [행위]할 수 있음 ② 피해·확산 특성 또는 탐지·방어가 어려운 이유"),
        ("참조", "■ 물리 관점(대상·통제 지점·수법(식별 수준)·연계 PHT-ID·다른 매트릭스 ID, 끝에 '참고:') / ■ 실제 사례('[라벨] 사례명(YYYY-MM): "
               "경위·결과 (사고 ID)', 공개 사례가 없으면 '공개 사고 미확인 — 사유')"),
        ("탐지·대응", "■ 탐지 / ■ 대응(예방 → 차단 → 탐지 순, 합법적 대응만, 매핑된 대응 기준 ID 병기, 끝에 '참고:')"),
        ("사실 근거", "수치·사례는 사고 DB 요약·출처에 있는 것만 씀. 용어는 자격증명(인증정보·크리덴셜 쓰지 않음)·반출(행위)/유출(결과)"),
        ("방어 관점", "수법은 식별에 필요한 수준까지만 적고 단계별 실행 방법은 쓰지 않음. 드론 무력화·전파 교란은 권고하지 않음(탐지·식별·신고·차폐). "
                    "임직원 모니터링은 개인정보보호법·노사 협의 범위 안에서만, 인적 지표는 객관적 행동·직무 사실만"),
        ("자동 검증", "scripts/validate.py — 인용 사고가 그 행에 매핑됐는지, 라벨이 집계 상태와 맞는지, 시점이 사고 DB와 같은지, 설명 속 수치가 사고 요약에 "
                    "있는지, PHT·ATT&CK·NIST·ISO·ISMS-P·CERT·다른 매트릭스 ID가 존재하고 행 매핑과 맞는지, '참고:'가 카탈로그 이름과 맞는지, 실제 사고가 "
                    "매핑된 행이 [실제 사고] 사례를 인용하는지 대조(빌드 전에 자동 실행)"),
    ]
    _kv_sheet(ws, crit, (22, 140))

    # ---------------- 변경이력 ----------------
    ws = wb.create_sheet("변경이력")
    _hdr(ws, 1, ["버전", "날짜", "구분", "대상", "변경 내용", "사유"])
    for i, ent in enumerate(CHANGELOG.get("entries", []), 2):
        _row(ws, i, ent)
    _w(ws, [6, 11, 10, 22, 80, 50])

    for s in wb.worksheets:
        s.sheet_view.zoomScale = 90
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToWidth = 1
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.fitToHeight = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path):
    cols = ["도메인", "PHT-ID", "위협분류", "세부위협", "핵심 요약", "요약설명", "공격 단계", "위협 주체", "적용 프로파일", "ATT&CK",
            "ITKB 관측 기법", "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수", "국내 실제 사고", "최근 사고", "정부 경보",
            "ATT&CK 사례 수", "관련 사례 ID", "NIST", "ISO 27001", "ISMS-P", "CERT", "참고 문헌", "아이덴티티 연계", "공급망 연계", "OT 연계",
            "클라우드 연계", "AI 연계"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for o in ROWS:
            w.writerow([o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], o["summary"].replace("\n", " "),
                        ", ".join(o["stage"]), " · ".join(o["actors"]), " · ".join(o["profiles"]), ", ".join(o["attack"]),
                        ", ".join(sorted(EV[o["id"]]["itkb"])), o["likelihood"], o["severity"], o["risk"], o["priority"], o["level"],
                        o["n_real"], o["n_kr"], o["n_recent"], o["n_adv"], o["n_atk"], ", ".join(sorted(EV[o["id"]]["real"])),
                        ", ".join(o["nist"]), ", ".join(str(k) for k in o["iso"]), ", ".join(str(k) for k in o["ismsp"]),
                        ", ".join(str(k) for k in o["cert"]), ", ".join(o["refs"]), ", ".join(o["idt"]), ", ".join(o["sct"]),
                        ", ".join(o["ot"]), ", ".join(o["cloud"]), ", ".join(o["ai"])])


def write_worksheet(outdir):
    """문구 작성용 근거 정리본 — 도메인별 마크다운(행별 매핑 사고·ATT&CK 주체·ITKB·대응 기준·참고 문헌)"""
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    for d in DOMAINS:
        lines = [f"# [{d['code']}] {d['ko']} — 문구 작성용 근거 정리본", "", f"- 범위: {d['desc']}", f"- 넣지 않는 것: {d['excludes']}", ""]
        for o in [o for o in ROWS if o["code"] == d["code"]]:
            x = EV[o["id"]]
            lines += [f"## {o['id']} {o['name']} ({o['en']})", "",
                      f"- 위험: {o['likelihood']}×{o['severity']}={o['risk']} · {o['lk_why']}",
                      f"- 단계: {', '.join(o['stage'])} · 주체: {', '.join(o['actors'])} · 자산: {', '.join(o['assets'])}",
                      f"- ATT&CK: {', '.join(atk_label(t) for t in o['attack']) or '미연계 — ' + o.get('attack_none', '')}",
                      f"- ITKB 관측: {', '.join(sorted(x['itkb'])) or '-'}",
                      f"- 연계: IDT {o['idt']} · SCT {o['sct']} · OT {o['ot']} · CL {o['cloud']} · AI {o['ai']}",
                      f"- 대응: NIST {o['nist']} · ISO 27001 {o['iso']} · ISMS-P {o['ismsp']} · CERT {o['cert']}",
                      f"- 참고 문헌: {', '.join(ref_label(k) for k in o['refs']) or '-'}",
                      f"- 심각도 근거: {o['severity_why']}", "", "### 사고 DB"]
            for i in sorted(x["real"] | x["advisory"] | x["intel"] | x["research"] | x["excluded"], key=lambda i: INC[i]["date"]):
                e = INC[i]
                lines.append(f"- {i} [{e['status']}] {e['title']}({e['date']}) — {' '.join(e['summary'].split())}"
                             + (f" ATT&CK {', '.join(e['attack_ref'])}" if e.get("attack_ref") else ""))
            if x["subj"]:
                lines += ["", "### ATT&CK 주체(물리·인적 맥락)", ", ".join(f"{s} {ATK['subjects'][s]['name']}" for s in sorted(x["subj"]))]
            lines.append("")
        (outdir / f"{d['code']}.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"근거 정리본 → {outdir}")


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--worksheet", help="문구 작성용 근거 정리본(도메인별 .md)을 만들 폴더")
    ap.add_argument("--skip-validate", action="store_true", help="scripts/validate.py 검증을 건너뜀")
    ap.add_argument("--version", help="출력 파일 버전 표기(기본 v2)")
    ap.add_argument("--data-only", action="store_true", help="문구 검사 없이 원천 데이터만 검증하고 빌드(v1 — 문구 작성 전 산출물)")
    ap.add_argument("--allow-missing-text", action="store_true",
                    help="문구가 없는 세부위협을 오류 대신 '문구 미작성'으로 두고 빌드(작성 중간 산출물용)")
    args = ap.parse_args()
    if args.version:
        VERSION = args.version
    if args.data_only:  # v1 재현 — 문구·시나리오를 빼고 분류·근거·평가만 출력
        for o in ROWS:
            o["oneline"] = o["summary"] = o["reference"] = o["detect"] = ""
            o["text_src"] = "문구 미작성"
        SCENARIOS.clear()
    OUT, CSV_OUT = out_paths(VERSION)
    if not args.skip_validate:
        VALIDATION = run_validate("--data-only" if args.data_only else "--allow-missing-text" if args.allow_missing_text else "")
        print(VALIDATION)
    if args.worksheet:
        write_worksheet(args.worksheet)
    write_xlsx(OUT)
    write_csv(CSV_OUT)
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    print(f"세부위협 {len(ROWS)}개 · 사고 {len(INCIDENTS)}건 → {OUT.name}")
    print("위험도", dict(rc), "| 근거 수준", dict(lc))
