# -*- coding: utf-8 -*-
"""
통합 소프트웨어 공급망 보안위협 매트릭스 빌더
-------------------------------------------
분류 : 공급망 단계(공격면) 8개 도메인 → 위협분류(Lv2) → 세부위협(Lv3) — data/taxonomy.yaml
교차 : ATT&CK v19.2 · SLSA v1.2 · OWASP CI/CD Top 10 · OWASP OSS Top 10 · SAP Risk Explorer(AV) · CNCF 침해 유형
근거 : 공급망 보안사고 DB(data/incidents.yaml) · ATT&CK 절차(공급망 맥락) · CISA KEV(공급망 판정) · SAP 문헌 · OSV 악성 패키지
논리 : 통합 AI 매트릭스 v3.2 · 클라우드 v5 · OT와 같은 발생가능성 수식(상 = 실제 사고 2건 이상), 심각도는 공급망 결과 기준

실행 : python3 scripts/build_supplychain_threat_matrix.py
       → output/통합_공급망보안위협_매트릭스_<버전>.xlsx · .csv
"""
import collections
import csv
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import taxonomy_rules as R  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
REF = ROOT / "reference"
VERSION = "v1"
OUT = ROOT / "output" / f"통합_공급망보안위협_매트릭스_{VERSION}.xlsx"
CSV_OUT = ROOT / "output" / f"통합_공급망보안위협_매트릭스_{VERSION}.csv"


def load_yaml(p, default=None):
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else default


# ===========================================================================
# 1) 원천 데이터
# ===========================================================================
TAX = load_yaml(DATA / "taxonomy.yaml")
FW = load_yaml(DATA / "frameworks.yaml")
INCIDENTS = load_yaml(DATA / "incidents.yaml")
CHANGELOG = load_yaml(DATA / "changelog.yaml", {})
SCENARIOS = load_yaml(DATA / "scenarios.yaml", {})
ATK = json.loads((REF / "attack" / "enterprise-attack-v19.2-subset.json").read_text(encoding="utf-8"))
SAP_AV = {x["avId"]: x for x in json.loads((REF / "sap-risk-explorer" / "attackvectors.json").read_text(encoding="utf-8"))}
SAP_SG = {x["sgId"]: x["sgName"] for x in json.loads((REF / "sap-risk-explorer" / "safeguards.json").read_text(encoding="utf-8"))}
SAP_REFS = json.loads((REF / "sap-risk-explorer" / "references.json").read_text(encoding="utf-8"))
KEV = {v["cveID"]: v for v in json.loads((REF / "advisories" / "known_exploited_vulnerabilities.json")
                                          .read_text(encoding="utf-8"))["vulnerabilities"]}
KEV_META = json.loads((REF / "advisories" / "known_exploited_vulnerabilities.json").read_text(encoding="utf-8"))

TEXT = {}
for f in sorted((DATA / "text").glob("*.yaml")):
    TEXT.update(load_yaml(f, {}) or {})

with open(DATA / "osv_malicious_counts.csv", encoding="utf-8") as f:
    OSV_NOTE = f.readline().lstrip("# ").strip()
    OSV = list(csv.reader(f))

DOMAINS, ROWS = [], []
for d in TAX["domains"]:
    DOMAINS.append(d)
    for l2 in d["lv2"]:
        for x in l2["lv3"]:
            ROWS.append(dict(x, code=d["code"], domain=f"[{d['code']}] {d['ko']}", lv2_id=l2["id"], lv2=l2["name"],
                             single=(x["id"] == l2["id"])))
ROW = {o["id"]: o for o in ROWS}


def incident_status(e):
    if e["kind"] in R.RESEARCH_KINDS:
        return "실증·연구"
    if e["kind"] in R.INTEL_KINDS:
        return "위협인텔"
    if e["verification"] in R.EXCLUDED_VERIFICATION:
        return "제외(검증)"
    return "실제 사고"


for e in INCIDENTS:
    e["date"] = str(e["date"])
    e["status"] = incident_status(e)
    e["recent"] = e["date"] >= R.RECENT_FROM

INC = {e["id"]: e for e in INCIDENTS}

# ===========================================================================
# 2) 근거 집계
# ===========================================================================
EV = {o["id"]: dict(real=set(), recent=set(), research=set(), intel=set(), excluded=set(), subj=set(),
                    kev=set(), sap=set(), eco=collections.Counter(), impact=collections.Counter())
      for o in ROWS}

for e in INCIDENTS:
    for k in e["sct"]:
        x = EV[k]
        if e["status"] == "실제 사고":
            x["real"].add(e["id"])
            if e["recent"]:
                x["recent"].add(e["id"])
            x["eco"].update(e["ecosystem"])
            x["impact"].update(i for i in e["impact"] if i != "영향 없음(사전 차단)")
        elif e["status"] == "실증·연구":
            x["research"].add(e["id"])
        elif e["status"] == "위협인텔":
            x["intel"].add(e["id"])
        else:
            x["excluded"].add(e["id"])


def atk_subjects(tid):
    """기법을 쓰는 ATT&CK 주체 — 공급망 고유 기법은 전체, 범용 기법은 절차 설명에 공급망 단서가 있는 것만"""
    out = set()
    for p in ATK["procedures"].get(tid, []):
        if tid in R.ATTACK_SC_SPECIFIC or R.ATTACK_SC_KEYWORDS.search(p["text"]):
            out.add(p["subject"])
    return out


for o in ROWS:
    for t in o["attack"]:
        EV[o["id"]]["subj"] |= atk_subjects(t)

KEV_ROWS = {}
for cve, (cat, rows, note) in R.KEV_SUPPLYCHAIN.items():
    if cve not in KEV:
        raise SystemExit(f"KEV 카탈로그에 없는 CVE: {cve}")
    for k in rows:
        EV[k]["kev"].add(cve)
    KEV_ROWS[cve] = rows

AV_ROWS = collections.defaultdict(set)
for o in ROWS:
    for av in o["sap"]:
        AV_ROWS[av].add(o["id"])
SAP_DOCS = []   # (index, title, year, tags, link, [AV])
for i, r in enumerate(SAP_REFS):
    tags = set(r.get("tags", {}).get("contents", []))
    if "attack" in tags or not (tags & R.SAP_RESEARCH_TAGS):
        continue
    avs = [v["avId"] for v in r.get("vectors", []) if v["avId"] in AV_ROWS]
    if not avs:
        continue
    SAP_DOCS.append((i, r["title"], r.get("tags", {}).get("year"), sorted(tags), r["link"], avs))
    for av in avs:
        for k in AV_ROWS[av]:
            EV[k]["sap"].add(i)


def row_safeguards(o):
    sg = set()
    for av in o["sap"]:
        for info in SAP_AV.get(av, {}).get("info", []):
            sg.update(m["sgId"] for m in info.get("Mapped Safeguard", []))
    return sorted(sg)


# ===========================================================================
# 3) 위험 평가
# ===========================================================================
def likelihood(x):
    if len(x["real"]) >= R.LIKELY_HIGH_REAL:
        return "상"
    if x["real"] or x["subj"] or x["kev"] or x["intel"] or x["research"] or x["sap"]:
        return "중"
    return "하"


def evidence_level(x):
    if x["real"]:
        return "실제 사고 확인"
    if x["subj"] or x["kev"] or x["intel"]:
        return "실사용 기법 포함"
    if x["research"] or x["sap"]:
        return "실증·공개 취약점"
    return "이론·시나리오"


TACTIC_KO = FW["attack_tactics_ko"]


def tactics_ko(tid):
    t = ATK["techniques"][tid]["tactics"] or ""
    return [TACTIC_KO.get(s.strip(), s.strip()) for s in t.split(",") if s.strip()]


for o in ROWS:
    x = EV[o["id"]]
    o["likelihood"] = likelihood(x)
    o["risk"] = R.RISK_MATRIX[(o["likelihood"], o["severity"])]
    o["level"] = evidence_level(x)
    o["n_real"], o["n_recent"], o["n_research"] = len(x["real"]), len(x["recent"]), len(x["research"])
    o["n_atk"], o["n_kev"], o["n_sap"] = len(x["subj"]), len(x["kev"]), len(x["sap"])
    o["lk_why"] = (f"실제 사고 {o['n_real']}건(최근 {o['n_recent']}) · ATT&CK 사례 {o['n_atk']}건 · KEV {o['n_kev']}건 · "
                   f"실증·연구 {o['n_research']}건 · 공격 분류 문헌 {o['n_sap']}건")
    o["safeguards"] = row_safeguards(o)
    txt = TEXT.get(o["id"]) or {}
    o["summary"], o["reference"], o["detect"] = (txt.get(k) or "" for k in ("summary", "reference", "detect"))
    o["text_src"] = "분석가 작성" if txt else "작성 예정"

rank_key = {r: i for i, r in enumerate(R.RISK_ORDER)}
for i, o in enumerate(sorted(ROWS, key=lambda o: (rank_key[o["risk"]], -o["n_real"], -o["n_recent"], o["id"])), 1):
    o["priority"] = i


def fw_label(cat, key, with_ko=True):
    v = FW[cat][key]
    if isinstance(v, list):
        return f"{key} {v[1] if with_ko else v[0]}" if len(v) > 1 else f"{key} {v[0]}"
    return f"{key} {v}"


def s2c2f_label(k):
    p, lv, en, ko = FW["s2c2f"][k]
    return f"{k} {ko}"


def nist_label(k):
    return f"{k} {FW['nist80053'][k]}"


def atk_label(t):
    return f"{t} {ATK['techniques'][t]['name']}"


def inc_label(i):
    e = INC[i]
    return f"{i} {e['title']}({e['date']})"


# ===========================================================================
# 4) xlsx 출력 — 클라우드 v5 · OT 서식 준용
# ===========================================================================
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "실사용 기법 포함": "2E86C1",
               "실증·공개 취약점": "B9770E", "이론·시나리오": "7F8C8D"}
HDR, NAVY, GRAY = "34495E", "1F3864", "8C8C8C"


def _w(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hdr(ws, row, cols, start=1, color=HDR):
    for j, c in enumerate(cols, start):
        cell = ws.cell(row, j, c)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BORDER


def _risk(cell, risk):
    bg, fg = RISK_COLOR[risk]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=risk in ("매우 높음", "높음"), color=fg)
    cell.alignment = CENTER


def _level(cell, level):
    cell.font = Font(bold=True, color=LEVEL_COLOR[level])
    cell.alignment = CENTER


def _row(ws, r, vals, start=1):
    for j, v in enumerate(vals, start):
        c = ws.cell(r, j, v)
        c.alignment = WRAP
        c.border = BORDER


def _title(ws, text, sub=None):
    ws.cell(1, 1, text).font = Font(size=14, bold=True, color=NAVY)
    if sub:
        ws.cell(2, 1, sub).font = Font(size=9, color="595959")


def _section(ws, r, text):
    ws.cell(r, 1, text).font = Font(bold=True, size=11, color=NAVY)
    return r + 1


def summary_cell(o):
    return o["summary"] or f"(작성 예정) {o['en']}"


def osv_total(keys):
    hdr = OSV[0]
    tot = 0
    for r in OSV[1:]:
        if r[0] in keys:
            tot += int(r[hdr.index("합계")])
    return tot


OSV_KEYS = {"npm": ["npm"], "PyPI": ["pypi"], "RubyGems": ["rubygems"], "PHP(Packagist·PEAR)": ["packagist"],
            "기타 패키지 생태계": ["nuget", "go", "crates.io", "maven", "git"],
            "IDE·브라우저 확장": ["vscode", "vscode:open-vsx.org"]}


def write_xlsx(path):
    wb = openpyxl.Workbook()
    today = datetime.date.today().isoformat()
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    sc = collections.Counter(o["severity"] for o in ROWS)
    st = collections.Counter(e["status"] for e in INCIDENTS)
    real = [e for e in INCIDENTS if e["status"] == "실제 사고"]
    kr = [e for e in real if e.get("region") == "한국"]
    written = sum(1 for o in ROWS if o["text_src"] == "분석가 작성")
    top = sorted(ROWS, key=lambda o: o["priority"])[:10]
    n_lv2 = sum(len(d["lv2"]) for d in DOMAINS)
    osv_all = osv_total([r[0] for r in OSV[1:]])
    kev_by_cat = collections.Counter(R.KEV_SUPPLYCHAIN[c][0] for c in R.KEV_SUPPLYCHAIN)
    years = sorted({e["date"][:4] for e in INCIDENTS})

    # ---------------- 개요 ----------------
    ws = wb.active
    ws.title = "개요"
    meta = [
        (f"통합 소프트웨어 공급망 보안위협 매트릭스 {VERSION}", ""),
        ("기준", "공급망 단계(공격면) 8개 도메인 → 위협분류 → 세부위협 분류에 ATT&CK v19.2·SLSA v1.2·OWASP CI/CD·OSS Top 10·"
                 "SAP 공격 분류를 교차 매핑하고, 통합 AI 매트릭스 v3.2·클라우드 v5·OT 매트릭스의 위험평가·근거 로직을 이식"),
        ("생성일", today), ("", ""),
        ("■ 기준 데이터", ""),
        ("공급망 보안사고 DB", f"{len(INCIDENTS)}건({years[0]}~{years[-1]}, 공개 출처) — 실제 사고 {st['실제 사고']}건 · "
                         f"실증·연구 {st['실증·연구']}건 · 위협인텔 {st.get('위협인텔', 0)}건 · 집계 제외 {st.get('제외(검증)', 0)}건, "
                         f"국내 실제 사고 {len(kr)}건, 최근(2025~) 실제 사고 {sum(1 for e in real if e['recent'])}건"),
        ("MITRE ATT&CK", f"Enterprise v19.2 — 매핑 기법 {len(ATK['techniques'])}개, 그 기법을 쓰는 주체 {len(ATK['subjects'])}개의 절차 "
                         "중 공급망 맥락 절차만 'ATT&CK 사례 수'로 집계(공급망 고유 기법은 전체)"),
        ("SLSA v1.2", "위협 모델 (A)~(I)·의존성·가용성·검증 위협 — 세부위협별 'SLSA 위협' 열"),
        ("OWASP", "Top 10 CI/CD Security Risks(2022) · Top 10 Open Source Software Risks(2023) — ID·명칭만 연계"),
        ("SAP Risk Explorer", f"Ladisa et al. 공급망 공격 분류 공격 벡터 {len(SAP_AV)}개 전체를 세부위협에 매핑 · 연구·실증 문헌 "
                              f"{len(SAP_DOCS)}건을 '공격 분류 문헌' 근거로 집계(공격 사례 문헌은 사고 DB로 이관)"),
        ("CISA KEV", f"{KEV_META.get('catalogVersion')}판 {KEV_META.get('count')}건 중 공급망 판정 {len(R.KEV_SUPPLYCHAIN)}건 — "
                     + " · ".join(f"{k} {v}" for k, v in kev_by_cat.items())),
        ("OSV 악성 패키지", f"OpenSSF Malicious Packages {osv_all:,}건({OSV_NOTE}) — 생태계·연도별 현황('악성 패키지 현황' 시트)"),
        ("대응 기준", "OpenSSF S2C2F · NIST SSDF(SP 800-218) · OpenSSF Scorecard · NIST SP 800-53 Rev.5 · SAP 대응책(SG)"),
        ("다른 매트릭스", "통합 AI·클라우드 매트릭스 v1(AI-·CL- 요약 ID) · 통합 OT 매트릭스(OTC-) 연계 열"), ("", ""),
        ("■ 시트 구성", ""),
        ("매트릭스 뷰", "도메인(열)별 세부위협을 위험도 색으로 배치한 한눈 보기"),
        ("통합 매트릭스", "분류 체계 · 교차 매핑 · 위험 평가 · 실제 근거 · 대응 기준 전체 열"),
        ("통합매트릭스_LITE", "핵심 열 발췌(필터·보고용)"),
        ("도메인 요약", "도메인별 위협 수 · 위험도/근거수준 분포 · 매핑 사고 수 · 최고위험 항목"),
        ("생태계 요약", "생태계·경로(npm·PyPI·Actions·확장·상용 SW 업데이트 등)별 실제 사고·영향·대표 세부위협·OSV 악성 패키지 수"),
        ("역매핑_사고사례", f"사고 DB {len(INCIDENTS)}건과 매핑 세부위협·출처(근거 추적용)"),
        ("KEV 근거", "CISA KEV 공급망 판정 내역(구분·매핑 세부위협·관련 사고)"),
        ("악성 패키지 현황", "OSV 형식 악성 패키지 보고의 생태계·연도별 건수와 해석 유의사항"),
        ("프레임워크 연계", "SLSA · OWASP CI/CD · OWASP OSS · CNCF 침해 유형 · SAP 공격 벡터 항목별 연계 세부위협(커버리지)"),
        ("대응 기준 연계", "S2C2F · SSDF · Scorecard · NIST 800-53 · SAP 대응책별 연계 세부위협과 위험도 분포(통제 우선순위 근거)"),
        ("평가 기준", "발생가능성·심각도·위험도·근거수준 산정 규칙, 사고 집계·ATT&CK 맥락·KEV 판정 기준"),
        ("변경이력", "버전별 변경 내역"), ("", ""),
        ("■ 분류 체계", ""),
        ("Lv1 도메인", "공급망 단계(공격면) 8개 — " + " → ".join(f"[{d['code']}] {d['ko']}" for d in DOMAINS)),
        ("Lv2 위협분류", f"{n_lv2}개(SCT-<도메인>-NN). 하위 세부위협이 하나뿐이면 Lv2 = Lv3"),
        ("Lv3 세부위협", f"{len(ROWS)}개(SCT-<도메인>-NN.m) — 요약설명 · 참조(공급망 관점·실제 사례) · 탐지·대응. ID는 고정(재사용 안 함)"),
        ("설계 원칙", "AI 매트릭스처럼 공격면(공급망 단계) 도메인을 쓰고, 클라우드·OT처럼 ATT&CK 전술을 '공격 단계' 열로 병기해 "
                    "다른 매트릭스와 나란히 비교할 수 있게 함. 'SLSA 위협' 열은 생산(A~G)·소비(H·I) 단계 위치를 함께 보여 줌"),
        ("적용 프로파일", " · ".join(f"{k}: {v}" for k, v in TAX["profiles"].items())), ("", ""),
        ("■ 결과 요약", ""),
        ("세부위협(Lv3)", f"{len(ROWS)}개 / 위협분류 {n_lv2}개 / 도메인 {len(DOMAINS)}개"),
        ("위험도", " · ".join(f"{k} {rc.get(k, 0)}" for k in R.RISK_ORDER)),
        ("근거 수준", " · ".join(f"{k} {lc.get(k, 0)}" for k in R.LEVELS)),
        ("심각도", " · ".join(f"{k} {sc.get(k, 0)}" for k in ("상", "중", "하"))
         + " — '하'에 해당하는 공급망 위협(라이선스·품질 위험 등)은 범위에서 제외"),
        ("공급망 관점 문구", f"{written}/{len(ROWS)}개 작성" + (" — 전 항목 완료" if written == len(ROWS) else
                                                       " — 나머지는 영문명만 표시(회색, '작성 예정')")),
        ("우선 위협 Top 10", "\n".join(f"{o['priority']}. {o['id']} {o['name']} — {o['risk']}, 실제 사고 {o['n_real']}건"
                                    f"(최근 {o['n_recent']})" for o in top)),
        ("", ""), ("■ 주의", ""),
        ("발생가능성", f"공개 사고 기준이라 보고가 많은 오픈소스(npm·PyPI)·대형 사건 유형이 높게 나오고, 내부 빌드·서명 침해처럼 "
                     f"드러나기 어려운 위협은 과소평가될 수 있음. 수식은 AI·클라우드·OT 매트릭스와 같게 유지(상 = 실제 사고 "
                     f"{R.LIKELY_HIGH_REAL}건 이상)"),
        ("다른 매트릭스와 비교", "사고 DB 규모·출처가 매트릭스마다 달라(클라우드 680건·OT 85건·공급망 "
                           f"{len(INCIDENTS)}건) 위험도·사고 수는 같은 매트릭스 안에서 비교. 같은 위협(예: 악성 모델 유입)도 "
                           "매트릭스별 심각도 기준이 달라 등급이 다를 수 있음"),
        ("영향·확산 도메인", "[IM] 행은 공격 결과 유형이라 여러 사고가 몰려 우선순위 상위에 오름 — 통제 우선순위는 앞 단계 도메인과 "
                         "'대응 기준 연계' 시트(연계 위협의 위험도 분포)로 판단"),
        ("근거 중복", "한 사고가 여러 단계(계정 탈취 → 악성 게시 → 웜 확산)에 동시에 매핑되므로 세부위협별 사고 수를 합하면 "
                    "사고 DB 건수보다 큼. KEV 등재 사례 일부는 사고 DB에도 있어 근거로 이중 표시됨(발생가능성은 실제 사고 수로만 '상' 판정)"),
        ("OSV 악성 패키지", "보고 ID 연도 기준 건수이며, 2025년 npm 급증의 대부분은 tea.xyz 보상 토큰을 노린 자동 대량 게시(약 15만 건, "
                        "2025-11 OpenSSF 등재)로 자격증명 탈취 코드가 없는 스팸성 패키지 — 건수를 위험 규모로 읽지 말 것"),
        ("최근 사고", "2025~2026년 사고는 1차 공개 자료·보안업체 분석 기준으로 정리했으며 추가 조사로 수치·귀속이 바뀔 수 있음"),
        ("출처 표기", "MITRE ATT&CK® © The MITRE Corporation · SLSA(Community Specification License 1.0, v1.2) · OWASP(CC BY-SA 4.0, "
                    "ID·명칭만) · SAP Risk Explorer(Apache-2.0) · CNCF TAG Security 카탈로그(CC BY 4.0) · OpenSSF S2C2F·Scorecard · "
                    "NIST SP 800-218·800-53(공공 영역) · CISA KEV(공공 영역) · OpenSSF Malicious Packages(Apache-2.0). "
                    "사고별 출처는 '역매핑_사고사례' 시트"),
    ]
    for i, (a, b) in enumerate(meta, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    ws.cell(1, 1).font = Font(size=15, bold=True, color=NAVY)
    _w(ws, [22, 140])

    # ---------------- 매트릭스 뷰 ----------------
    ws = wb.create_sheet("매트릭스 뷰")
    by_c = collections.defaultdict(list)
    for o in ROWS:
        by_c[o["code"]].append(o)
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도 (빨강 매우 높음 · 주황 높음 · 노랑 보통 · 초록 낮음), [n] = 실제 사고 수, "
                  "괄호 안 숫자 = 세부위협 수").font = Font(bold=True, color=NAVY)
    for j, d in enumerate(DOMAINS, 1):
        c = ws.cell(2, j, f"[{d['code']}] {d['ko']}\n{d['en']}\n({len(by_c[d['code']])})")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        for i, o in enumerate(sorted(by_c[d["code"]], key=lambda o: (rank_key[o["risk"]], -o["n_real"], o["id"])), 3):
            cell = ws.cell(i, j, f"{o['id']} {o['name']}" + (f" [{o['n_real']}]" if o["n_real"] else ""))
            _risk(cell, o["risk"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9, bold=o["risk"] == "매우 높음", color=RISK_COLOR[o["risk"]][1])
            cell.border = BORDER
    ws.row_dimensions[2].height = 48
    ws.freeze_panes = "A3"
    _w(ws, [24] * len(DOMAINS))

    # ---------------- 통합 매트릭스 ----------------
    ws = wb.create_sheet("통합 매트릭스")
    g_class = ["도메인(Lv1)", "SCT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "요약설명", "참조", "영문명", "공격 단계(ATT&CK 전술)",
               "설명 출처"]
    g_cross = (["대상 자산"] + [f"프로파일:{p}" for p in R.PROFILES]
               + ["ATT&CK ID", "SLSA v1.2 위협", "OWASP CI/CD Top10", "OWASP OSS Top10", "SAP 공격 벡터(AV)", "CNCF 침해 유형",
                  "클라우드 매트릭스 연계", "AI 매트릭스 연계", "OT 매트릭스 연계", "ATT&CK 기법명", "적용 프로파일 수"])
    g_risk = ["발생가능성", "심각도", "위험도", "우선순위", "발생가능성 근거 (자동 산정)", "심각도 근거"]
    g_ev = ["실제 사고 수", "최근 사고(2025~)", "ATT&CK 사례 수(공급망 맥락)", "KEV(공급망)", "실증·연구 수", "공격 분류 문헌(SAP)",
            "근거 수준", "관련 사례 ID", "주요 생태계(실제 사고)", "ATT&CK 사례 주체(일부)"]
    g_ctl = ["S2C2F 요구사항", "NIST SSDF 실천과제", "OpenSSF Scorecard", "NIST SP 800-53", "SAP 대응책(SG)", "탐지·대응 포인트"]
    groups = [("분류 체계", len(g_class), "2E5496"), ("교차 매핑", len(g_cross), "1F7A8C"),
              ("위험 평가", len(g_risk), "A04000"), ("실제 근거", len(g_ev), "1E8449"), ("대응 기준", len(g_ctl), "6C3483")]
    cols = g_class + g_cross + g_risk + g_ev + g_ctl
    t = ws.cell(1, 1, f"통합 소프트웨어 공급망 보안위협 매트릭스 {VERSION} — 분류체계 · 위험평가 · 실제근거 · 대응기준")
    t.font = Font(size=13, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, f"공급망 보안사고 DB {len(INCIDENTS)}건 + ATT&CK v19.2 + CISA KEV {len(R.KEV_SUPPLYCHAIN)}건 + SAP 공격 분류 "
                  f"{len(SAP_AV)}개 AV | {today}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    c0 = 1
    for g, span, color in groups:
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + span - 1)
        c = ws.cell(3, c0, g)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        c0 += span
    _hdr(ws, 4, cols)
    ci = {c: i + 1 for i, c in enumerate(cols)}
    for r, o in enumerate(ROWS, 5):
        x = EV[o["id"]]
        prof = ["●" if p in o["profiles"] else "" for p in R.PROFILES]
        subj = sorted(x["subj"])
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], summary_cell(o), o["reference"], o["en"], ", ".join(o["stage"]),
                     o["text_src"], ", ".join(o["assets"])] + prof
             + [", ".join(o["attack"]), "\n".join(fw_label("slsa", k) for k in o["slsa"]),
                "\n".join(fw_label("owasp_cicd", k) for k in o["cicd"]), "\n".join(fw_label("owasp_oss", k) for k in o["oss"]),
                "\n".join(f"{k} {FW['sap_ko'][k]}" for k in o["sap"]), ", ".join(o["cncf"]),
                ", ".join(o["cloud"]), ", ".join(o["ai"]), ", ".join(o["ot"]),
                "\n".join(atk_label(t) for t in o["attack"]), len(o["profiles"]),
                o["likelihood"], o["severity"], o["risk"], o["priority"], o["lk_why"], o["severity_why"],
                o["n_real"], o["n_recent"], o["n_atk"], o["n_kev"], o["n_research"], o["n_sap"], o["level"],
                ", ".join(sorted(x["real"])) + (("\n실증: " + ", ".join(sorted(x["research"]))) if x["research"] else ""),
                " · ".join(f"{k} {v}" for k, v in x["eco"].most_common(4)),
                ", ".join(f"{s} {ATK['subjects'][s]['name']}" for s in subj[:12]) + (f" 외 {len(subj) - 12}" if len(subj) > 12 else ""),
                "\n".join(s2c2f_label(k) for k in o["s2c2f"]), "\n".join(fw_label("ssdf", k) for k in o["ssdf"]),
                "\n".join(f"{k}" for k in o["scorecard"]), "\n".join(nist_label(k) for k in o["nist"]),
                "\n".join(f"{k} {SAP_SG.get(k, '')}" for k in o["safeguards"]), o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, ci["요약설명"]).font = Font(color=GRAY)
        for c in [f"프로파일:{p}" for p in R.PROFILES] + ["발생가능성", "심각도", "우선순위", "적용 프로파일 수"]:
            ws.cell(r, ci[c]).alignment = CENTER
        _risk(ws.cell(r, ci["위험도"]), o["risk"])
        _level(ws.cell(r, ci["근거 수준"]), o["level"])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(ROWS) + 4}"
    widths = {"도메인(Lv1)": 14, "SCT-ID": 12, "위협분류(Lv2)": 18, "세부위협(Lv3)": 24, "요약설명": 60, "참조": 80, "영문명": 26,
              "공격 단계(ATT&CK 전술)": 14, "설명 출처": 9, "대상 자산": 18, "ATT&CK ID": 14, "SLSA v1.2 위협": 22,
              "OWASP CI/CD Top10": 22, "OWASP OSS Top10": 20, "SAP 공격 벡터(AV)": 24, "CNCF 침해 유형": 14,
              "클라우드 매트릭스 연계": 12, "AI 매트릭스 연계": 12, "OT 매트릭스 연계": 10, "ATT&CK 기법명": 30,
              "적용 프로파일 수": 7, "발생가능성": 7, "심각도": 7, "위험도": 8, "우선순위": 7, "발생가능성 근거 (자동 산정)": 30,
              "심각도 근거": 40, "실제 사고 수": 7, "최근 사고(2025~)": 7, "ATT&CK 사례 수(공급망 맥락)": 8, "KEV(공급망)": 7,
              "실증·연구 수": 7, "공격 분류 문헌(SAP)": 8, "근거 수준": 11, "관련 사례 ID": 22, "주요 생태계(실제 사고)": 22,
              "ATT&CK 사례 주체(일부)": 30, "S2C2F 요구사항": 26, "NIST SSDF 실천과제": 26, "OpenSSF Scorecard": 16,
              "NIST SP 800-53": 26, "SAP 대응책(SG)": 26, "탐지·대응 포인트": 70}
    _w(ws, [widths.get(c, 6) for c in cols])

    # ---------------- 통합매트릭스_LITE ----------------
    ws = wb.create_sheet("통합매트릭스_LITE")
    ws.cell(1, 1, f"통합 소프트웨어 공급망 보안위협 매트릭스 {VERSION} — LITE").font = Font(size=13, bold=True, color=NAVY)
    lcols = ["도메인(Lv1)", "SCT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "요약설명", "공격 단계", "적용 프로파일", "발생가능성", "심각도",
             "위험도", "우선순위", "근거 수준", "실제 사고 수", "최근 사고(2025~)", "탐지·대응 포인트"]
    _hdr(ws, 2, lcols)
    for r, o in enumerate(ROWS, 3):
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], summary_cell(o), ", ".join(o["stage"]), " · ".join(o["profiles"]),
                     o["likelihood"], o["severity"], o["risk"], o["priority"], o["level"], o["n_real"], o["n_recent"], o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, 5).font = Font(color=GRAY)
        _risk(ws.cell(r, 10), o["risk"])
        _level(ws.cell(r, 12), o["level"])
        for j in (8, 9, 11, 13, 14):
            ws.cell(r, j).alignment = CENTER
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(lcols))}{len(ROWS) + 2}"
    _w(ws, [14, 12, 18, 24, 60, 14, 18, 7, 7, 8, 7, 11, 7, 7, 70])

    # ---------------- 도메인 요약 ----------------
    ws = wb.create_sheet("도메인 요약")
    dcols = ["도메인", "설명", "위협분류(Lv2)", "세부위협(Lv3)"] + R.RISK_ORDER + R.LEVELS + \
            ["매핑 실제 사고(중복 제거)", "최근 사고(2025~)", "최고위험 세부위협(상위 3)"]
    _hdr(ws, 1, dcols)
    for i, d in enumerate(DOMAINS, 2):
        os_ = by_c[d["code"]]
        rr = collections.Counter(o["risk"] for o in os_)
        ll = collections.Counter(o["level"] for o in os_)
        incs = set().union(*(EV[o["id"]]["real"] for o in os_))
        top3 = sorted(os_, key=lambda o: o["priority"])[:3]
        _row(ws, i, [f"[{d['code']}] {d['ko']}", d["desc"], len(d["lv2"]), len(os_)] + [rr.get(k, 0) for k in R.RISK_ORDER]
             + [ll.get(k, 0) for k in R.LEVELS] + [len(incs), sum(1 for e in incs if INC[e]["recent"]),
                                                  "\n".join(f"{o['id']} {o['name']}({o['risk']}, {o['n_real']}건)" for o in top3)])
    _w(ws, [16, 60, 8, 8, 7, 7, 7, 7, 9, 9, 9, 9, 10, 9, 50])

    # ---------------- 생태계 요약 ----------------
    ws = wb.create_sheet("생태계 요약")
    _title(ws, "생태계·경로별 요약 — 실제 사고 기준(한 사고가 여러 생태계에 걸치면 각각 집계)",
           "OSV 악성 패키지 수는 OpenSSF Malicious Packages 보고 건수(생태계 합계, '악성 패키지 현황' 시트 참고)")
    ecols = ["구분", "생태계·경로", "실제 사고", "최근(2025~)", "실증·연구", "주요 영향", "대표 세부위협(사고 수 상위 3)",
             "최근 사례(최대 3)", "OSV 악성 패키지 보고"]
    _hdr(ws, 4, ecols)
    for i, eco in enumerate(R.ECOSYSTEMS, 5):
        es = [e for e in INCIDENTS if eco in e["ecosystem"]]
        er = [e for e in es if e["status"] == "실제 사고"]
        imp = collections.Counter(x for e in er for x in e["impact"] if x != "영향 없음(사전 차단)")
        rows_c = collections.Counter(k for e in er for k in e["sct"])
        recent = sorted(er, key=lambda e: e["date"], reverse=True)[:3]
        _row(ws, i, [R.ECOSYSTEM_GROUP[eco], eco, len(er), sum(1 for e in er if e["recent"]),
                     sum(1 for e in es if e["status"] == "실증·연구"),
                     " · ".join(f"{k} {v}" for k, v in imp.most_common(3)),
                     "\n".join(f"{k} {ROW[k]['name']}({v})" for k, v in rows_c.most_common(3)),
                     "\n".join(inc_label(e["id"]) for e in recent),
                     f"{osv_total(OSV_KEYS[eco]):,}" if eco in OSV_KEYS else ""])
    ws.freeze_panes = "C5"
    _w(ws, [14, 22, 8, 8, 8, 30, 46, 52, 12])

    # ---------------- 역매핑_사고사례 ----------------
    ws = wb.create_sheet("역매핑_사고사례")
    icols = ["사례 ID", "시점", "사례명", "사례유형", "검증 수준", "집계 상태", "생태계·경로", "지역", "행위자", "영향", "요약",
             "매핑 세부위협", "ATT&CK 참조", "KEV", "출처"]
    _hdr(ws, 1, icols)
    for i, e in enumerate(sorted(INCIDENTS, key=lambda e: (e["date"], e["id"])), 2):
        _row(ws, i, [e["id"], e["date"], e["title"], e["kind"], e["verification"], e["status"], ", ".join(e["ecosystem"]),
                     e.get("region", ""), e.get("actor", ""), ", ".join(e["impact"]), e["summary"],
                     "\n".join(f"{k} {ROW[k]['name']}" for k in e["sct"]), ", ".join(e.get("attack_ref") or []),
                     ", ".join(e.get("kev") or []), "\n".join(e["sources"])])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(icols))}{len(INCIDENTS) + 1}"
    _w(ws, [9, 8, 34, 9, 13, 10, 18, 7, 18, 18, 70, 40, 14, 16, 70])

    # ---------------- KEV 근거 ----------------
    ws = wb.create_sheet("KEV 근거")
    _title(ws, f"CISA KEV 공급망 판정 — {KEV_META.get('catalogVersion')}판 {KEV_META.get('count')}건 중 {len(R.KEV_SUPPLYCHAIN)}건",
           "판정 기준: 침해된 소프트웨어(CWE-506 등)·업데이트 무결성 결함(CWE-494)·공급망 인프라(CI/CD·SCM·아티팩트 저장소·RMM·배포 도구)·"
           "제품에 내장되어 대량 악용된 오픈소스 구성요소. 실사용 근거('중' 상한)로만 반영")
    r = 4
    for cat in R.KEV_CATEGORIES:
        r = _section(ws, r, f"■ {cat} ({kev_by_cat[cat]})")
        kc = ["CVE", "벤더", "제품", "취약점명", "KEV 등재일", "랜섬웨어 악용", "CWE", "매핑 세부위협", "비고"]
        _hdr(ws, r, kc)
        r += 1
        for cve, (c, rows, note) in sorted(R.KEV_SUPPLYCHAIN.items(), key=lambda kv: KEV[kv[0]]["dateAdded"], reverse=True):
            if c != cat:
                continue
            v = KEV[cve]
            _row(ws, r, [cve, v["vendorProject"], v["product"], v["vulnerabilityName"], v["dateAdded"],
                         v.get("knownRansomwareCampaignUse", ""), ", ".join(v.get("cwes") or []),
                         "\n".join(f"{k} {ROW[k]['name']}" for k in rows), note])
            r += 1
        r += 1
    _w(ws, [16, 16, 24, 46, 11, 9, 12, 40, 34])

    # ---------------- 악성 패키지 현황 ----------------
    ws = wb.create_sheet("악성 패키지 현황")
    _title(ws, "OSV 형식 악성 패키지 보고 현황(OpenSSF Malicious Packages)", OSV_NOTE)
    _hdr(ws, 4, OSV[0])
    for i, rr in enumerate(OSV[1:], 5):
        _row(ws, i, [rr[0]] + [int(v) for v in rr[1:]])
    n = len(OSV) + 4
    notes = [
        "■ 해석 유의",
        "· 연도는 보고 ID(MAL-YYYY-NNNN) 연도로, 패키지가 게시된 시점과 다를 수 있음",
        "· 2025년 npm 급증의 대부분은 tea.xyz 보상 토큰을 노린 자동 대량 게시(약 15만 건, 2025-11 Amazon·OpenSSF가 일괄 등재)로, "
        "자격증명 탈취 코드가 없는 스팸성 패키지 — 건수가 곧 피해 규모는 아님",
        "· 보고는 자동 분석·보안업체 피드에 좌우되어 npm·PyPI 비중이 실제 위험보다 크게 나타날 수 있음(감시가 약한 생태계는 과소보고)",
        "· 개별 패키지는 세부위협 행에 직접 집계하지 않고 SCT-DP-01·DP-03(이름 혼동·독자 악성 패키지)의 배경 근거로만 사용",
        "· 출처: https://github.com/ossf/malicious-packages (Apache-2.0), 2025 급증 배경: "
        "https://www.darkreading.com/application-security/150000-packages-flood-npm-registry-token-farming",
    ]
    for j, t in enumerate(notes, n + 1):
        ws.cell(j, 1, t).font = Font(bold=t.startswith("■"), color=NAVY if t.startswith("■") else "000000")
    _w(ws, [22] + [10] * (len(OSV[0]) - 1))

    # ---------------- 프레임워크 연계 ----------------
    ws = wb.create_sheet("프레임워크 연계")
    _title(ws, "프레임워크 항목별 연계 세부위협 — 커버리지 점검(미연계 항목은 사유 표기)")
    r = 3
    fw_sets = [("SLSA v1.2 위협", "slsa", lambda k: FW["slsa"][k][0], lambda k: FW["slsa"][k][1]),
               ("OWASP Top 10 CI/CD Security Risks", "cicd", lambda k: FW["owasp_cicd"][k][0], lambda k: FW["owasp_cicd"][k][1]),
               ("OWASP Top 10 OSS Risks", "oss", lambda k: FW["owasp_oss"][k][0], lambda k: FW["owasp_oss"][k][1]),
               ("CNCF 침해 유형", "cncf", lambda k: k, lambda k: FW["cncf"][k]),
               ("SAP Risk Explorer 공격 벡터(AV)", "sap", lambda k: SAP_AV[k]["avName"], lambda k: FW["sap_ko"][k])]
    catkey = {"slsa": "slsa", "cicd": "owasp_cicd", "oss": "owasp_oss", "cncf": "cncf", "sap": "sap_ko"}
    for title, fld, en, ko in fw_sets:
        r = _section(ws, r, f"■ {title}")
        _hdr(ws, r, ["ID", "명칭(영문)", "명칭(국문)", "연계 세부위협 수", "연계 세부위협", "위험도 분포", "실제 사고(중복 제거)", "미연계 사유"])
        r += 1
        keys = list(FW[catkey[fld]].keys()) if fld != "sap" else list(SAP_AV.keys())
        for k in keys:
            rows = [o for o in ROWS if k in o[fld]]
            incs = set().union(*(EV[o["id"]]["real"] for o in rows)) if rows else set()
            rcnt = collections.Counter(o["risk"] for o in rows)
            why = R.FRAMEWORK_NOT_MAPPED.get((catkey[fld] if fld != "oss" else "owasp_oss", k), "")
            if not rows and not why and fld == "sap" and k in ("AV-000", "AV-001", "AV-200", "AV-300", "AV-400", "AV-500"):
                why = "상위(분기) 노드 — 하위 AV로 매핑"
            _row(ws, r, [k, en(k), ko(k), len(rows), "\n".join(f"{o['id']} {o['name']}" for o in rows),
                         " · ".join(f"{x} {rcnt[x]}" for x in R.RISK_ORDER if rcnt[x]), len(incs), why])
            r += 1
        r += 1
    _w(ws, [12, 40, 30, 8, 50, 26, 9, 40])

    # ---------------- 대응 기준 연계 ----------------
    ws = wb.create_sheet("대응 기준 연계")
    _title(ws, "대응 기준별 연계 세부위협 — 연계 위협의 위험도 분포로 통제 우선순위를 가늠",
           "S2C2F = 소비자(오픈소스 사용 조직) 요구사항 · SSDF = 생산자(개발 조직) 실천과제 · Scorecard = 오픈소스 프로젝트 점검 지표 · "
           "NIST 800-53 = 조직 통제 · SAP SG = 공격 벡터에 매핑된 대응책")
    r = 4
    ctl_sets = [("OpenSSF S2C2F 요구사항", "s2c2f", lambda k: FW["s2c2f"][k][2], lambda k: FW["s2c2f"][k][3],
                 lambda k: f"{FW['s2c2f'][k][0]} · {FW['s2c2f'][k][1]}"),
                ("NIST SSDF(SP 800-218) 실천과제", "ssdf", lambda k: FW["ssdf"][k][0], lambda k: FW["ssdf"][k][1], lambda k: ""),
                ("OpenSSF Scorecard 점검 항목", "scorecard", lambda k: k, lambda k: FW["scorecard"][k], lambda k: ""),
                ("NIST SP 800-53 Rev.5", "nist", lambda k: FW["nist80053"][k], lambda k: "", lambda k: ""),
                ("SAP 대응책(SG)", "safeguards", lambda k: SAP_SG[k], lambda k: "", lambda k: "")]
    ckey = {"s2c2f": "s2c2f", "ssdf": "ssdf", "scorecard": "scorecard", "nist": "nist80053"}
    for title, fld, en, ko, extra in ctl_sets:
        r = _section(ws, r, f"■ {title}")
        _hdr(ws, r, ["ID", "명칭(영문)", "명칭(국문)", "실천·성숙도", "연계 세부위협 수", "매우 높음·높음 위협 수", "연계 세부위협", "미연계 사유"])
        r += 1
        keys = list(FW[ckey[fld]].keys()) if fld in ckey else sorted(SAP_SG)
        for k in keys:
            rows = [o for o in ROWS if k in o[fld]]
            hi = sum(1 for o in rows if o["risk"] in ("매우 높음", "높음"))
            why = R.FRAMEWORK_NOT_MAPPED.get((ckey.get(fld, fld), k), "")
            _row(ws, r, [k, en(k), ko(k), extra(k), len(rows), hi,
                         "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in sorted(rows, key=lambda o: o["priority"])),
                         why])
            r += 1
        r += 1
    _w(ws, [12, 46, 30, 16, 8, 10, 56, 36])

    # ---------------- 평가 기준 ----------------
    ws = wb.create_sheet("평가 기준")
    crit = [
        ("■ 위험도", ""),
        ("위험도", "발생가능성 × 심각도 3×3 표(AI·클라우드·OT 매트릭스와 동일) — 상×상 매우 높음, 상×중·중×상 높음, 상×하·중×중·하×상 보통, 나머지 낮음"),
        ("발생가능성", TAX["assessment"]["likelihood_rule"]),
        ("심각도 상", TAX["assessment"]["severity_rule"]["상"]),
        ("심각도 중", TAX["assessment"]["severity_rule"]["중"]),
        ("심각도 하", TAX["assessment"]["severity_rule"]["하"]),
        ("우선순위", "위험도 → 실제 사고 수 → 최근 사고(2025~) → SCT-ID 순(1 = 최우선)"),
        ("근거 수준", "실제 사고 확인(사고 DB 실제 사고) > 실사용 기법 포함(ATT&CK 공급망 맥락 사례·KEV·위협인텔) > "
                    "실증·공개 취약점(사고 DB 연구·시연·SAP 연구 문헌) > 이론·시나리오"),
        ("", ""), ("■ 근거 집계", ""),
        ("실제 사고", f"사례유형 ∈ {{{'·'.join(sorted(R.REAL_KINDS))}}}이고 검증 수준 ∉ {{{'·'.join(sorted(R.EXCLUDED_VERIFICATION))}}}. "
                    "연구자의 실증·노출 발견은 '연구·시연'(실증 근거)으로 분리 — 클라우드 v5의 '실제 사고/연구·노출' 구분과 같음"),
        ("최근 사고", f"사고 시점 {R.RECENT_FROM} 이후(통합 매트릭스 v1과 같은 기준)"),
        ("사고 매핑", "분석가가 사고 경위를 읽고 단계별 세부위협에 매핑(한 사고가 계정 탈취·악성 게시·확산 등 여러 행에 매핑될 수 있음). "
                    "ATT&CK 캠페인·소프트웨어 ID는 대조·표시용이며 자동 매핑에 쓰지 않음"),
        ("ATT&CK 사례 수", "행의 ATT&CK 기법을 쓰는 주체(그룹·소프트웨어·캠페인) 수. 공급망 고유 기법(T1195.001·.002, T1677, T1176.002, "
                         "T1072, T1199, T1553.002 등)은 절차 전체를, 범용 기법은 절차 설명에 공급망 단서(npm·저장소·CI/CD·빌드·코드 서명·"
                         "업데이트 등)가 있는 것만 셈"),
        ("KEV(공급망)", "CISA KEV 중 침해된 소프트웨어·업데이트 무결성 결함·공급망 인프라(CI/CD·SCM·아티팩트 저장소·RMM·배포 도구)·"
                       "대량 악용된 오픈소스 구성요소만 판정('KEV 근거' 시트). 실사용 근거로 '중'까지만 반영"),
        ("공격 분류 문헌", "SAP Risk Explorer 문헌 중 연구·실증 성격(peer-reviewed·proof-of-concept·vulnerability·thesis·presentation)을 "
                       "세부 AV 매핑으로 집계. 공격 사례 문헌은 사고 DB로 옮겨 중복 집계하지 않음"),
        ("", ""), ("■ 다른 매트릭스와의 근거 대응", ""),
        ("AI 매트릭스 v3.2", "실제 사고 ↔ ATLAS Incident·OWASP 인용 사고 / 실사용 ↔ Realized 기법 / 실증 ↔ 실증·CVE"),
        ("클라우드 매트릭스 v5", "실제 사고 ↔ 클라우드 사고 DB 실제 사고 / 실사용 ↔ ATT&CK 클라우드 사례 / 실증 ↔ 연구·노출·벤더 매트릭스"),
        ("OT 매트릭스", "실제 사고 ↔ OT 사고 DB / 실사용 ↔ ATT&CK ICS 절차·KEV(OT) / 실증 ↔ 연구·시연·ICS 권고 CVE"),
        ("공급망 매트릭스", "실제 사고 ↔ 공급망 사고 DB / 실사용 ↔ ATT&CK 공급망 맥락 절차·KEV(공급망) / 실증 ↔ 연구·시연·SAP 연구 문헌"),
        ("", ""), ("■ 사례 라벨(문구 작성용)", ""),
        ("[실제 사고]", "사고 DB에서 '실제 사고'로 집계되는 사례"),
        ("[공개 취약점]", "CVE·KEV·공식 보안 권고로 확인된 결함"),
        ("[실증]", "연구자·레드팀·보안업체가 실제 제품·서비스에서 재현한 공격(사고 DB 연구·시연)"),
        ("[시나리오]", "기준 문서(SLSA·OWASP 등)의 가정 사례"),
    ]
    for i, (a, b) in enumerate(crit, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    _w(ws, [20, 140])

    # ---------------- 변경이력 ----------------
    ws = wb.create_sheet("변경이력")
    _hdr(ws, 1, ["버전", "날짜", "구분", "대상", "변경 내용", "사유"])
    for i, ent in enumerate(CHANGELOG.get("entries", []), 2):
        _row(ws, i, ent)
    _w(ws, [6, 11, 10, 22, 80, 50])

    for s in wb.worksheets:
        s.sheet_view.zoomScale = 90
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToWidth = 1
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.fitToHeight = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path):
    cols = ["도메인", "SCT-ID", "위협분류", "세부위협", "요약설명", "공격 단계", "적용 프로파일", "ATT&CK", "SLSA", "OWASP CI/CD",
            "OWASP OSS", "SAP AV", "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수", "최근 사고",
            "ATT&CK 사례 수", "KEV", "관련 사례 ID"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for o in ROWS:
            w.writerow([o["domain"], o["id"], o["lv2"], o["name"], o["summary"].replace("\n", " "), ", ".join(o["stage"]),
                        " · ".join(o["profiles"]), ", ".join(o["attack"]), ", ".join(o["slsa"]), ", ".join(o["cicd"]),
                        ", ".join(o["oss"]), ", ".join(o["sap"]), o["likelihood"], o["severity"], o["risk"], o["priority"],
                        o["level"], o["n_real"], o["n_recent"], o["n_atk"], o["n_kev"], ", ".join(sorted(EV[o["id"]]["real"]))])


if __name__ == "__main__":
    write_xlsx(OUT)
    write_csv(CSV_OUT)
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    print(f"세부위협 {len(ROWS)}개 · 사고 {len(INCIDENTS)}건 → {OUT.name}")
    print("위험도", dict(rc), "| 근거 수준", dict(lc))
