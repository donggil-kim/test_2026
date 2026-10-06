# -*- coding: utf-8 -*-
"""
통합 클라우드 보안위협 매트릭스 빌더 (v2)
-----------------------------------------
뼈대 : MITRE ATT&CK Cloud 킬체인 (전술 → 기법 → 하위기법/벤더 항목)
교차 : AWS TTC · Azure ATRM · K8s 매트릭스 · CSA Top Threats 2026 · 통합 AI 매트릭스(UT)
근거 : ATT&CK 클라우드 실제 사례 + 클라우드 보안사고 DB(680건) + CSA 2026 사례
논리 : 통합 AI 보안위협 매트릭스 v3.2의 위험평가·근거수준 로직 이식

실행 : python scripts/build_cloud_threat_matrix.py
"""
import openpyxl, re, collections, datetime, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from taxonomy_rules import (TACTIC_CODE, CSA_NAMES, CSA_MAP, CSA_NOTE, AI_UT_MAP, SEVERITY,
                            TACTIC_SEVERITY, DEPRECATED, DROP, MERGE, NAME_OVERRIDE, KEYWORD_RULES)
import csv

VERSION = "v5"
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_MATRIX = os.path.join(HERE, "data", "Cloud_ATTACK_Matrix_Integrated_v19.2.xlsx")
SRC_INC = os.path.join(HERE, "data", "Cloud_IncidentDB_v2_LITE.xlsx")
OUT = os.path.join(HERE, "output", f"통합_클라우드보안위협_매트릭스_{VERSION}.xlsx")
SRC_ATTACK = os.path.join(HERE, "data", "enterprise-attack-v19.2-techniques.xlsx")
OVERRIDE = os.path.join(HERE, "data", "threat_text_override.csv")
TEMPLATE = os.path.join(HERE, "output", "threat_text_template.csv")
OV_COLS = ["ATT&CK ID", "ATT&CK 기법명", "Cloud 범위", "현재 요약설명(참고용)",
           "세부위협명(입력)", "요약설명(입력)", "참조(입력)", "탐지·대응(입력)"]

RISK_MATRIX = {
    ("상", "상"): "매우 높음", ("상", "중"): "높음", ("상", "하"): "보통",
    ("중", "상"): "높음", ("중", "중"): "보통", ("중", "하"): "낮음",
    ("하", "상"): "보통", ("하", "중"): "낮음", ("하", "하"): "낮음",
}
REAL_KINDS = {"사고", "캠페인", "사례연구(익명)", "규제공시(8-K)", "위협인텔 보고서"}
RECENT_FROM = "2025"


def parent_tid(tid):
    return tid.split(".")[0] if tid and tid.startswith("T") else tid


def norm_tid(t):
    return DEPRECATED.get(t, t)


# ---------------------------------------------------------------------------
# 1) ATT&CK 통합 매트릭스 → 리프(Lv3) 골격
# ---------------------------------------------------------------------------
def vendor_items(cell):
    if not cell:
        return []
    return re.findall(r"•\s*([A-Za-z0-9.\-]+)\s", cell)


def load_skeleton():
    wb = openpyxl.load_workbook(SRC_MATRIX, read_only=True)
    rows = list(wb["통합 매트릭스"].iter_rows(min_row=2, values_only=True))
    has_child = set()
    for r in rows:
        if "하위기법" in (r[1] or ""):
            has_child.add(parent_tid(r[2]))

    domains, leaves = [], []
    tactic = None
    seq = collections.Counter()
    cur = None  # 현재 Lv2
    for r in rows:
        gubun, tid, name = (r[1] or ""), (r[2] or ""), (r[3] or "")
        if gubun == "전술":
            tactic = r[0]
            code, ko = TACTIC_CODE[tactic]
            domains.append(dict(tactic=tactic, code=code, ko=ko, tid=tid,
                                desc=(r[4] or "").strip(), label=f"[{code}] {ko}"))
            continue
        code, ko = TACTIC_CODE[tactic]
        disp = re.sub(r"^" + re.escape(tid) + r"\s*", "", name).strip()
        base = dict(tactic=tactic, domain=f"[{code}] {ko}", tid=tid, gubun=gubun,
                    scope=r[5] or "", plat_cloud=r[6] or "", plat_other=r[7] or "",
                    aws=vendor_items(r[8]), azure=vendor_items(r[9]), k8s=vendor_items(r[10]),
                    desc=(r[4] or "").strip(), atk_cases=int(r[12] or 0),
                    atk_summary=r[11] or "", link=r[13] or "")
        is_sub = "하위기법" in gubun
        if is_sub and cur and cur["tid"] == parent_tid(tid):
            cur["n"] += 1
            leaves.append(dict(base, ctc=f"{cur['code']}.{cur['n']}", lv2=cur["name"],
                               lv2code=cur["code"], lv3=disp))
            continue
        seq[tactic] += 1
        lv2code = f"CTC-{code}-{seq[tactic]:02d}"
        cur = dict(code=lv2code, name=disp, tid=tid, n=0)
        if tid in has_child and "연결용" not in gubun:
            # Lv2 그룹 : 상위기법 자체 근거가 있으면 '.0 (일반)' 리프로 보존
            leaves.append(dict(base, ctc=f"{lv2code}.0", lv2=disp, lv2code=lv2code,
                               lv3=f"{disp} (일반·상위기법)", is_general=True))
        else:
            leaves.append(dict(base, ctc=lv2code, lv2=disp, lv2code=lv2code, lv3=disp))
    return domains, leaves


# ---------------------------------------------------------------------------
# 2) 사고 DB → 기법 매핑 (초안 ID + 키워드/근본원인 규칙)
# ---------------------------------------------------------------------------
def load_incidents(valid_tids):
    wb = openpyxl.load_workbook(SRC_INC, read_only=True)
    incs = []
    for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        e = dict(id=r[0], title=r[1] or "", summary=r[2] or "", kind=r[3] or "",
                 src=r[4] or "", org=r[5] or "", actor=r[6] or "", date=str(r[7] or ""),
                 rel=r[8] or "", layer=r[9] or "", root=r[11] or r[10] or "",
                 impact=r[13] or r[12] or "", draft=r[21] or "",
                 # 키워드 매핑용: 제목·요약·원문 필드(초기침투·확산·영향·사용 기법·공격 대상·도구) / 구조 필드(전체값)
                 free=" ".join(str(x) for x in (r[1], r[2], *r[15:21]) if x),
                 cond=dict(root=f"{r[10] or ''} {r[11] or ''}", layer=r[9] or "",
                           impact=f"{r[12] or ''} {r[13] or ''}", kind=r[3] or ""))
        incs.append(e)

    for e in incs:
        links = {}  # tid -> 근거
        draft = [norm_tid(t) for t in re.findall(r"T\d{4}(?:\.\d{3})?", e["draft"])]
        for t in draft:
            links[t] = "초안ID"
        # 매트릭스에 없는 기법은 상위기법으로 올림
        clean = {}
        for t, why in links.items():
            if t in valid_tids:
                clean[t] = why
            elif parent_tid(t) in valid_tids:
                clean.setdefault(parent_tid(t), why)
        e["links"] = keyword_links(e, clean, valid_tids)
        e["counted"] = e["rel"] != "비클라우드 의심"
        e["real"] = e["kind"] in REAL_KINDS
        e["recent"] = e["date"][:4] >= RECENT_FROM
        e["template"] = e["summary"].startswith("Wiz가 ")
    return incs


def _kw_match(rule, e):
    tid, kw, cond = rule
    c = e["cond"]
    for k in ("root", "layer", "impact"):
        if k in cond and not any(v in c[k] for v in cond[k]):
            return False
    if c["kind"] in cond.get("skip_kind", ()):
        return False
    if cond.get("neg") and re.search(cond["neg"], e["free"], re.I):
        return False
    return re.search(kw, e["free"], re.I) is not None


def keyword_links(e, links, valid_tids):
    """초안 ID 매핑에 키워드 매핑(KEYWORD_RULES)을 더한다.
    - 같은 기법이나 그 하위기법이 이미 있으면 추가하지 않음
    - 상위기법 태그만 있고 키워드가 하위기법을 가리키면 상위 태그를 하위기법으로 세분"""
    links = dict(links)
    refined = set()
    for rule in KEYWORD_RULES:
        t = rule[0]
        if t not in valid_tids or t in links or any(parent_tid(x) == t and x != t for x in links):
            continue
        if not _kw_match(rule, e):
            continue
        p = parent_tid(t)
        if p != t and (p in links or p in refined):
            if p in links:
                refined.add(p)
                del links[p]
            links[t] = "초안ID→키워드 세분"
        else:
            links[t] = "키워드"
    return links


# ---------------------------------------------------------------------------
# 3) 평가 로직 (AI 매트릭스 v3.2 수식과 통일)
#   AI v3.2 : 상 = 실제 사고(ATLAS Incident + OWASP 인용) ≥ 2
#             중 = 실제 사고 1 또는 실증 사례·Realized 기법·공개 취약점 존재
#   대응    : 실제 사고 ↔ 사고 DB 실제 사고 / Realized 기법 ↔ ATT&CK 클라우드 사례
#             실증·공개 취약점 ↔ 사고 DB 연구·노출 + 벤더 매트릭스 문서화
# ---------------------------------------------------------------------------
LIKELY_HIGH_REAL = 2
def evidence_level(real, research, atk, vendor_n):
    if real >= 1:
        return "실제 사고 확인"
    if atk >= 1:
        return "실사용 기법 포함"
    if research >= 1 or vendor_n >= 1:
        return "실증·공개 취약점"
    return "이론·시나리오"


def likelihood(real, research, atk, vendor_n):
    if real >= LIKELY_HIGH_REAL:
        return "상"
    if real >= 1 or atk >= 1 or research >= 1 or vendor_n >= 1:
        return "중"
    return "하"


def severity(tid, tactic):
    for key in (tid, parent_tid(tid)):
        if key in SEVERITY:
            return SEVERITY[key]
    g = TACTIC_SEVERITY[tactic]
    return g, f"전술 기준값({TACTIC_CODE[tactic][1]})"


# ---------------------------------------------------------------------------
# 4) 설명 보조 : ATT&CK 원문 목록 복원 · 사용자 입력 문구
# ---------------------------------------------------------------------------
def _clean_md(t):
    t = re.sub(r"\(Citation:[^)]*\)", "", t)
    t = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", t)
    t = re.sub(r"</?[a-z]+>", "", t).replace("`", "")
    return re.sub(r"\s+", " ", t).strip()


def load_attack_lists():
    """ATT&CK 원문에서 첫 문단 뒤 글머리표 목록을 추출 (첫 문단이 ':'로 끝나는 경우 복원용)"""
    if not os.path.exists(SRC_ATTACK):
        return {}
    wb = openpyxl.load_workbook(SRC_ATTACK, read_only=True)
    lists = {}
    for r in wb["techniques"].iter_rows(min_row=2, values_only=True):
        items = [_clean_md(x[2:]) for x in (r[3] or "").split("\n") if x.startswith("* ")]
        if items:
            lists[r[0]] = [i if len(i) <= 220 else i[:217] + "…" for i in items]
    return lists


def load_override():
    if not os.path.exists(OVERRIDE):
        return {}
    with open(OVERRIDE, encoding="utf-8-sig") as f:
        return {r["ATT&CK ID"].strip(): r for r in csv.DictReader(f) if r.get("ATT&CK ID")}


def final_tid(t):
    while t in MERGE:
        t = MERGE[t]
    return t


# ---------------------------------------------------------------------------
# 5) 조립
# ---------------------------------------------------------------------------
def build():
    domains, leaves = load_skeleton()
    valid = {l["tid"] for l in leaves}
    incs = load_incidents(valid)
    lists = load_attack_lists()
    override = load_override()

    by_tid = collections.defaultdict(list)
    for e in incs:
        for t, why in e["links"].items():
            by_tid[t].append((e, why))

    # 기법(고유 ID) 단위 근거 집계 → 통합 대상으로 합산
    first = {}
    for l in leaves:
        first.setdefault(l["tid"], l)
    agg = collections.defaultdict(lambda: dict(pairs=[], atk=0, aws=[], azure=[], k8s=[],
                                                merged=[], atk_lines=[]))
    lost = []
    for tid, l in first.items():
        if tid in DROP:
            lost += [e["id"] for e, _ in by_tid.get(tid, []) if e["counted"]]
            continue
        a = agg[final_tid(tid)]
        a["pairs"] += by_tid.get(tid, [])
        a["atk"] += l["atk_cases"]
        for k in ("aws", "azure", "k8s"):
            a[k] += [x for x in l[k] if x not in a[k]]
        a["atk_lines"] += [x for x in l["atk_summary"].split("\n") if x.strip()]
        if final_tid(tid) != tid:
            a["merged"].append(f"{tid} {re.sub(r' [(]일반·상위기법[)]$', '', l['lv3'])}")

    rows = []
    for l in leaves:
        tid = l["tid"]
        if tid in DROP or tid in MERGE:
            continue
        a = agg[tid]
        pairs = [(e, w) for e, w in a["pairs"] if e["counted"]]
        vendor_n = len(a["aws"]) + len(a["azure"]) + len(a["k8s"])
        if l.get("is_general") and not (pairs or a["atk"] or vendor_n):
            continue
        ids = {e["id"] for e, _ in pairs}
        real = len({e["id"] for e, _ in pairs if e["real"]})
        research = len(ids) - real
        recent = len({e["id"] for e, _ in pairs if e["real"] and e["recent"]})
        atk = a["atk"]
        lk = likelihood(real, research, atk, vendor_n)
        sv, sv_why = severity(tid, l["tactic"])
        csa_ids = sorted(set(CSA_MAP.get(tid, CSA_MAP.get(parent_tid(tid), []))))

        # 요약설명 : 기존 번역 설명(첫 문단). ':'로 끝나면 목록을 참조로 넘김
        summary = l["desc"]
        restored = []
        if summary.rstrip().endswith(":"):
            restored = lists.get(tid, [])
            summary = summary.rstrip().rstrip(":").rstrip() + " (세부 항목은 참조 열 참조)"
        base_summary = summary  # override 적용 전의 기존 번역 요약(템플릿 참고용)

        # 참조 : 원문 보충 · 실제 사고 · ATT&CK 사례 · 통합 기법 (기존 데이터 재배치)
        ref = []
        if restored:
            ref.append("■ ATT&CK 원문 목록\n" + "\n".join(f"- {x}" for x in restored))
        uniq = {}
        for e, _ in pairs:
            uniq.setdefault(e["id"], e)
        evs = sorted(uniq.values(), key=lambda e: (e["real"], not e["template"], e["date"]), reverse=True)
        if evs:
            lines = []
            for e in evs[:3]:
                tag = "실제 사고" if e["real"] else "실증·연구"
                body = "" if e["template"] else f": {e['summary'][:140]}"
                lines.append(f"- [{tag}] {e['title']}({e['date'][:7]}, {e['id']}){body}")
            more = f"\n- 외 {len(evs) - 3}건(관련 사례 ID 열)" if len(evs) > 3 else ""
            ref.append("■ 실제 사례(사고 DB)\n" + "\n".join(lines) + more)
        if a["atk_lines"]:
            al = [x.lstrip("• ").strip() for x in a["atk_lines"] if x.startswith("•")][:4]
            if al:
                ref.append(f"■ ATT&CK 사례(행위자, 총 {atk}건)\n" + "\n".join(f"- {x}" for x in al))
        if a["merged"]:
            ref.append("■ 통합된 기법\n" + "\n".join(f"- {x}" for x in a["merged"]))
        if not ref:
            ref.append("■ 실제 사례 미확인 — ATT&CK·벤더 매트릭스상 가능 기법")

        lv3 = NAME_OVERRIDE.get(tid, l["lv3"])
        ov = override.get(tid, {})
        lv3 = (ov.get("세부위협명(입력)") or "").strip() or lv3
        summary = (ov.get("요약설명(입력)") or "").strip() or summary
        reference = (ov.get("참조(입력)") or "").strip() or "\n\n".join(ref)
        detect = (ov.get("탐지·대응(입력)") or "").strip()

        rows.append(dict(
            l, lv3=lv3, aws=a["aws"], azure=a["azure"], k8s=a["k8s"], vendor_n=vendor_n,
            summary=summary, reference=reference, detect=detect, base_summary=base_summary,
            text_src="사용자 입력" if (ov.get("요약설명(입력)") or "").strip() else "기존 번역",
            merged=", ".join(m.split(" ")[0] for m in a["merged"]),
            csa=" · ".join(f"SI-{i:02d} {CSA_NAMES[i]}" for i in csa_ids), csa_ids=csa_ids,
            ut=AI_UT_MAP.get(tid, AI_UT_MAP.get(parent_tid(tid), "")),
            likelihood=lk, severity=sv, sev_why=sv_why, risk=RISK_MATRIX[(lk, sv)],
            lk_why=f"실제 사고 {real}건(최근 {recent}) · 실증·연구 {research}건 · ATT&CK 사례 {atk}건 · 벤더 기법 {vendor_n}개",
            real=real, research=research, recent=recent, atk=atk,
            level=evidence_level(real, research, atk, vendor_n),
            inc_ids=", ".join(sorted(ids)),
            basis=", ".join(sorted({w for _, w in pairs})),
        ))

    # ID 재부여 : 전술별 Lv2 순번, Lv2 내 '.0(일반)' → .0, 하위기법 → .1..
    out = []
    groups = collections.OrderedDict()
    for r in rows:
        groups.setdefault((r["tactic"], r["lv2code"]), []).append(r)
    seq = collections.Counter()
    for (tactic, _), grp in groups.items():
        seq[tactic] += 1
        code = f"CTC-{TACTIC_CODE[tactic][0]}-{seq[tactic]:02d}"
        if len(grp) == 1 and (grp[0].get("is_general") or "." not in grp[0]["ctc"]):
            g = grp[0]
            g.update(ctc=code, lv2code=code)
            if g.get("is_general") and g["lv3"].endswith("(일반·상위기법)"):
                g["lv3"] = NAME_OVERRIDE.get(g["tid"], g["lv2"])
            out.append(g)
            continue
        n = 0
        for g in grp:
            if g.get("is_general"):
                g.update(ctc=f"{code}.0", lv2code=code)
            elif "." in g["ctc"]:
                n += 1
                g.update(ctc=f"{code}.{n}", lv2code=code)
            else:
                g.update(ctc=code, lv2code=code)
            out.append(g)

    bad = [o["tid"] for o in out if o["summary"].rstrip().endswith(":")]
    assert not bad, f"요약설명이 ':'로 끝나는 행: {bad}"
    stats = dict(dropped=sorted(DROP & set(first)), merged=sorted(set(MERGE) & set(first)),
                 lost_incidents=sorted(set(lost)))
    return domains, out, incs, stats


def write_template(out):
    """사용자 문구 입력용 CSV (고유 기법 단위, 현재 요약을 참고용으로 포함)"""
    seen, rows = set(), []
    for o in out:
        if o["tid"] in seen:
            continue
        seen.add(o["tid"])
        rows.append({OV_COLS[0]: o["tid"], OV_COLS[1]: o["lv3"], OV_COLS[2]: o["scope"],
                     OV_COLS[3]: o.get("base_summary", o["summary"]),
                     OV_COLS[4]: "", OV_COLS[5]: "", OV_COLS[6]: "", OV_COLS[7]: ""})
    with open(TEMPLATE, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=OV_COLS)
        w.writeheader()
        w.writerows(rows)
    return TEMPLATE


# ===========================================================================
#  xlsx 출력
# ===========================================================================
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
SCOPE_COLOR = {"Cloud 전용": "D6EAF8", "혼합": "FEF9E7", "Cloud 외 연관": "EAECEE",
               "연결용": "EAECEE", "ATT&CK 미대응": "FADBD8"}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "실사용 기법 포함": "2E86C1",
               "실증·공개 취약점": "B9770E", "이론·시나리오": "7F8C8D"}
HDR = "34495E"


def _w(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hdr(ws, row, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c)
        cell.fill = PatternFill("solid", fgColor=HDR)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BORDER


def _risk(cell, risk):
    bg, fg = RISK_COLOR[risk]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=risk in ("매우 높음", "높음"), color=fg)
    cell.alignment = CENTER


def _row(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.alignment = WRAP
        c.border = BORDER


def write_xlsx(domains, out, incs, stats):
    wb = openpyxl.Workbook()
    today = datetime.date.today().isoformat()
    counted = [e for e in incs if e["counted"]]
    mapped = [e for e in counted if e["links"]]
    kw_only = [e for e in mapped if all(w == "키워드" for w in e["links"].values())]

    # ---------------- 개요 ----------------
    ws = wb.active
    ws.title = "개요"
    rc = collections.Counter(o["risk"] for o in out)
    lc = collections.Counter(o["level"] for o in out)
    unique_tids = len({o["tid"] for o in out})
    rewritten = len({o["tid"] for o in out if o["text_src"] == "사용자 입력"})
    meta = [
        (f"통합 클라우드 보안위협 매트릭스 {VERSION}", ""),
        ("기준", "MITRE ATT&CK Cloud 킬체인(전술→기법→하위기법) 뼈대 · 통합 AI 보안위협 매트릭스 v3.2의 위험평가·근거 로직 이식"),
        ("생성일", today), ("", ""),
        ("■ 기준 데이터", ""),
        ("MITRE ATT&CK", "Enterprise v19.2 Cloud(IaaS·SaaS·Office Suite·Identity Provider) — 전술 14 · 기법/하위기법 및 클라우드 실제 사례"),
        ("벤더 매트릭스", "AWS Threat Technique Catalog · Azure Threat Research Matrix · Threat Matrix for Kubernetes"),
        ("CSA Top Threats 2026", "11대 위협 연계(기법 단위, 하위기법 지정 시 우선) — 원문 미수록, 연계 ID만"),
        ("클라우드 보안사고 DB", f"{len(incs)}건(2010~2026, Wiz·ramimac·SEC·GTI·MS) — 집계 대상 {len(counted)}건(비클라우드 의심 제외), 기법 매핑 {len(mapped)}건(키워드 규칙으로만 매핑 {len(kw_only)}건 포함)"),
        ("통합 AI 매트릭스 v3.2", "AI 관련 클라우드 기법에 UT-ID 교차 표시"), ("", ""),
        ("■ 시트 구성", ""),
        ("매트릭스 뷰", "전술(열)별 세부위협을 위험도 색으로 배치한 한눈 보기"),
        ("통합 매트릭스", "분류체계 · 교차매핑 · 위험평가 · 실제근거 · 탐지·대응 전체 열"),
        ("통합매트릭스_LITE", "핵심 열 발췌(필터·보고용) + 탐지·대응 포인트"),
        ("도메인 요약", "전술별 위협 수 · 위험도/근거수준 분포 · 매핑 사고 수 · 최고위험 항목"),
        ("CSA 2026 연계", "11대 위협별 연계 세부위협·사고 수와 연계 기준(SI-02·SI-08 교차 위협, SI-06은 AI 매트릭스 UT 연계)"),
        ("역매핑_사고사례", "사고 680건과 매핑된 기법·매핑 근거(근거 추적용)"),
        ("평가 기준", "발생가능성·심각도·위험도·근거수준 산정 규칙"), ("", ""),
        ("■ 분류 체계", ""),
        ("Lv1 도메인", "ATT&CK 전술. 예: [IA] 초기 침투, [IM] 영향"),
        ("Lv2 위협분류", "ATT&CK 기법. 예: CTC-IA-02 = T1190"),
        ("Lv3 세부위협", "하위기법/벤더 항목 — 요약설명 · 참조(클라우드 관점 · 실제 사례 등) · 탐지·대응 포인트. '.0 (일반·상위기법)'은 하위기법으로 특정되지 않은 상위기법 근거를 보존한 행"),
        ("클라우드 특화 검토", f"일반 엔터프라이즈 기법 {len(stats['dropped'])}개 삭제({', '.join(stats['dropped'])}), {len(stats['merged'])}개 통합(근거·벤더 항목은 대상 행에 합산, '통합된 기법' 열)"),
        ("설명 문구(v4)", "요약설명·참조·탐지·대응은 클라우드 관점으로 재작성한 문구(data/threat_text_override.csv)를 우선 적용하며, 미작성 기법은 기존 번역 설명과 사례 데이터 재배치를 사용('설명 출처' 열). 요약설명은 2줄 개조식, 참조는 '클라우드 관점·공격 시나리오·실제 사례' 구성"),
        ("", ""), ("■ 결과 요약", ""),
        ("세부위협(Lv3)", f"{len(out)}건 / 도메인 {len(domains)}개"),
        ("위험도", " · ".join(f"{k} {rc.get(k, 0)}" for k in RISK_COLOR)),
        ("근거 수준", " · ".join(f"{k} {lc.get(k, 0)}" for k in LEVEL_COLOR)),
        ("클라우드 관점 재작성", f"고유 기법 {rewritten}/{unique_tids}개 적용(요약설명·참조·탐지·대응). " + ("전 기법 완료" if rewritten >= unique_tids else "나머지는 기존 번역 설명 유지 — 후속 작성 예정")),
        ("", ""), ("■ 주의", ""),
        ("사고 매핑", f"사고 DB의 ATT&CK 초안 ID(폐기 ID는 v19.2로 변환) + 키워드 규칙 {len(KEYWORD_RULES)}개(v5, 근거 '키워드'·'초안ID→키워드 세분'). 세분되지 않은 상위기법 태그는 '.0 (일반·상위기법)' 행에 보존. 매핑 근거는 '역매핑_사고사례'에서 사고별로 확인"),
        ("위험평가", f"발생가능성은 근거에서 자동 산정(AI 매트릭스 v3.2와 동일: 상=실제 사고 {LIKELY_HIGH_REAL}건 이상), 심각도는 기법별 기준값 → 조직 맥락에 맞게 검토 권장"),
        ("중복 표시", "한 기법이 여러 전술에 속하면 전술마다 반복 표시(ATT&CK 원칙). 근거 수는 동일 기법 기준"),
        ("출처 표기", "MITRE ATT&CK © The MITRE Corporation / CSA·AWS·Microsoft 각 원저작권자. 배포 전 각 출처 이용약관 확인"),
    ]
    for i, (a, b) in enumerate(meta, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color="1F3864" if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    ws.cell(1, 1).font = Font(size=15, bold=True, color="1F3864")
    _w(ws, [24, 120])

    # ---------------- 매트릭스 뷰 ----------------
    ws = wb.create_sheet("매트릭스 뷰")
    order = [d["tactic"] for d in domains]
    by_t = collections.defaultdict(list)
    for o in out:
        by_t[o["tactic"]].append(o)
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도 (빨강 매우 높음 · 주황 높음 · 노랑 보통 · 초록 낮음), [n] = 실제 사고 수")
    ws.cell(1, 1).font = Font(bold=True, color="1F3864")
    for j, t in enumerate(order, 1):
        d = domains[j - 1]
        c = ws.cell(2, j, f"{d['label']}\n{t}\n({len(by_t[t])})")
        c.fill = PatternFill("solid", fgColor="1F3864"); c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        rank = {"매우 높음": 0, "높음": 1, "보통": 2, "낮음": 3}
        for i, o in enumerate(sorted(by_t[t], key=lambda o: (rank[o["risk"]], -o["real"], o["ctc"])), 3):
            cell = ws.cell(i, j, f"{o['ctc']} {o['lv3']}" + (f" [{o['real']}]" if o["real"] else ""))
            _risk(cell, o["risk"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9, bold=o["risk"] == "매우 높음", color=RISK_COLOR[o["risk"]][1])
            cell.border = BORDER
    ws.row_dimensions[2].height = 48
    ws.freeze_panes = "A3"
    _w(ws, [22] * len(order))

    # ---------------- 통합 매트릭스 ----------------
    ws = wb.create_sheet("통합 매트릭스")
    groups = [("분류 체계", 13, "2E5496"), ("교차 매핑", 5, "1F7A8C"),
              ("위험 평가", 5, "A04000"), ("실제 근거", 8, "1E8449"), ("탐지·대응", 1, "6C3483")]
    cols = ["도메인(Lv1)", "CTC-ID", "위협분류(Lv2)", "세부위협(Lv3)", "요약설명", "참조",
            "ATT&CK ID", "통합된 기법", "구분", "Cloud 범위", "Cloud 플랫폼", "기타 플랫폼", "설명 출처",
            "CSA Top Threats 2026", "AI 매트릭스 연계(UT)", "AWS (TTC)", "Azure (ATRM)", "Kubernetes",
            "발생가능성", "심각도", "위험도", "발생가능성 근거 (자동 산정)", "심각도 근거",
            "실제 사고 수", "최근 사고(2025~)", "실증·연구 수", "ATT&CK 사례 수", "근거 수준",
            "사고 매핑 근거", "관련 사례 ID", "ATT&CK 링크", "탐지·대응 포인트"]
    t = ws.cell(1, 1, f"통합 클라우드 보안위협 매트릭스 {VERSION} — 분류체계 · 위험평가 · 실제근거")
    t.font = Font(size=13, bold=True, color="FFFFFF"); t.fill = PatternFill("solid", fgColor="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, f"ATT&CK Cloud v19.2 + AWS TTC · Azure ATRM · K8s + CSA Top Threats 2026 + 클라우드 보안사고 DB {len(incs)}건 | {today}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    c0 = 1
    for g, span, color in groups:
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + span - 1)
        c = ws.cell(3, c0, g); c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF"); c.alignment = CENTER
        c0 += span
    _hdr(ws, 4, cols)
    for r, o in enumerate(out, 5):
        _row(ws, r, [o["domain"], o["ctc"], o["lv2"], o["lv3"], o["summary"], o["reference"],
                     o["tid"], o["merged"], o["gubun"], o["scope"], o["plat_cloud"], o["plat_other"], o["text_src"],
                     o["csa"], o["ut"], "\n".join(o["aws"]), "\n".join(o["azure"]), "\n".join(o["k8s"]),
                     o["likelihood"], o["severity"], o["risk"], o["lk_why"], o["sev_why"],
                     o["real"], o["recent"], o["research"], o["atk"], o["level"],
                     o["basis"], o["inc_ids"], o["link"], o.get("detect", "")])
        ws.cell(r, 10).fill = PatternFill("solid", fgColor=SCOPE_COLOR.get(o["scope"], "FFFFFF"))
        _risk(ws.cell(r, 21), o["risk"])
        ws.cell(r, 28).font = Font(bold=True, color=LEVEL_COLOR[o["level"]])
        for cc in (19, 20, 24, 25, 26, 27):
            ws.cell(r, cc).alignment = CENTER
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(out) + 4}"
    _w(ws, [15, 13, 22, 26, 48, 60, 11, 14, 13, 12, 15, 15, 10,
            26, 16, 22, 22, 18, 8, 7, 9, 30, 28, 7, 8, 7, 7, 12, 16, 22, 28, 52])

    # ---------------- LITE ----------------
    ws = wb.create_sheet("통합매트릭스_LITE")
    lcols = ["도메인(Lv1)", "CTC-ID", "위협분류(Lv2)", "세부위협(Lv3)", "요약설명", "ATT&CK ID", "Cloud 범위",
             "CSA Top Threats 2026", "발생가능성", "심각도", "위험도", "근거 수준",
             "실제 사고 수", "ATT&CK 사례 수", "탐지·대응 포인트"]
    ws.cell(1, 1, f"통합 클라우드 보안위협 매트릭스 {VERSION} — LITE").font = Font(size=12, bold=True, color="1F3864")
    _hdr(ws, 2, lcols)
    for r, o in enumerate(out, 3):
        _row(ws, r, [o["domain"], o["ctc"], o["lv2"], o["lv3"], o["summary"], o["tid"], o["scope"], o["csa"],
                     o["likelihood"], o["severity"], o["risk"], o["level"], o["real"], o["atk"], o.get("detect", "")])
        ws.cell(r, 7).fill = PatternFill("solid", fgColor=SCOPE_COLOR.get(o["scope"], "FFFFFF"))
        _risk(ws.cell(r, 11), o["risk"])
        ws.cell(r, 12).font = Font(bold=True, color=LEVEL_COLOR[o["level"]])
        for cc in (9, 10, 13, 14):
            ws.cell(r, cc).alignment = CENTER
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(lcols))}{len(out) + 2}"
    _w(ws, [15, 13, 22, 26, 48, 11, 12, 28, 8, 7, 9, 13, 8, 8, 52])

    # ---------------- 도메인 요약 ----------------
    ws = wb.create_sheet("도메인 요약")
    dcols = ["도메인", "ATT&CK 전술", "전술 설명", "위협분류(Lv2)", "세부위협(Lv3)",
             "매우 높음", "높음", "보통", "낮음", "실제 사고 확인", "실사용 기법", "실증·공개", "이론",
             "매핑 사고 수(중복 제거)", "최고위험 세부위협(상위 3)"]
    _hdr(ws, 1, dcols)
    for r, d in enumerate(domains, 2):
        os_ = by_t[d["tactic"]]
        rc_ = collections.Counter(o["risk"] for o in os_)
        lc_ = collections.Counter(o["level"] for o in os_)
        ids = set()
        for o in os_:
            ids |= {x for x in o["inc_ids"].split(", ") if x.startswith("INC")}
        top = sorted(os_, key=lambda o: ({"매우 높음": 0, "높음": 1, "보통": 2, "낮음": 3}[o["risk"]], -o["real"], -o["atk"]))[:3]
        _row(ws, r, [d["label"], f"{d['tid']} {d['tactic']}", d["desc"], len({o["lv2code"] for o in os_}), len(os_),
                     rc_["매우 높음"], rc_["높음"], rc_["보통"], rc_["낮음"],
                     lc_["실제 사고 확인"], lc_["실사용 기법 포함"], lc_["실증·공개 취약점"], lc_["이론·시나리오"],
                     len(ids), "\n".join(f"{o['ctc']} {o['lv3']} ({o['risk']})" for o in top)])
        for cc in range(4, 15):
            ws.cell(r, cc).alignment = CENTER
    ws.freeze_panes = "B2"
    _w(ws, [16, 26, 40, 9, 9, 8, 8, 8, 8, 9, 9, 9, 8, 11, 48])

    # ---------------- CSA 2026 연계 ----------------
    ws = wb.create_sheet("CSA 2026 연계")
    _hdr(ws, 1, ["CSA 이슈", "위협명", "연계 세부위협 수", "그중 위험 높음 이상", "연계 사고 수",
                 "그중 실제 사고", "연계 기준", "연계 세부위협(위험 높음 이상, 없으면 전체)"])
    real_ids = {e["id"] for e in incs if e["counted"] and e["real"]}
    r = 2
    for i in range(1, 12):
        rel = [o for o in out if i in o["csa_ids"]]
        hi = [o for o in rel if o["risk"] in ("매우 높음", "높음")]
        ids = set()
        for o in rel:
            ids |= {x for x in o["inc_ids"].split(", ") if x}
        note = CSA_NOTE.get(i, "기법 주제 기준 1차 매핑(scripts/taxonomy_rules.py의 CSA_MAP)")
        _row(ws, r, [f"SI-{i:02d}", CSA_NAMES[i], len(rel), len(hi), len(ids), len(ids & real_ids), note,
                     "\n".join(sorted({f"{o['ctc']} {o['lv3']} ({o['risk']})" for o in (hi or rel)}))[:3000]])
        for cc in range(3, 7):
            ws.cell(r, cc).alignment = CENTER
        r += 1
    ws.freeze_panes = "C2"
    _w(ws, [9, 24, 10, 11, 9, 9, 44, 70])

    # ---------------- 역매핑_사고사례 ----------------
    ws = wb.create_sheet("역매핑_사고사례")
    icols = ["사건ID", "대표 제목", "기준일", "사례유형", "출처", "클라우드 관련도", "집계 포함",
             "서비스 계층", "근본원인", "영향유형", "DB 초안 ATT&CK", "매핑 기법(근거)", "사건 요약"]
    _hdr(ws, 1, icols)
    for r, e in enumerate(sorted(incs, key=lambda x: x["date"]), 2):
        _row(ws, r, [e["id"], e["title"], e["date"], e["kind"], e["src"], e["rel"],
                     "Y" if e["counted"] else "N(비클라우드 의심)", e["layer"], e["root"], e["impact"],
                     e["draft"], "\n".join(f"{t} ({w})" for t, w in sorted(e["links"].items())), e["summary"]])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(icols))}{len(incs) + 1}"
    _w(ws, [10, 34, 10, 12, 9, 13, 10, 20, 20, 16, 14, 30, 70])

    # ---------------- 평가 기준 ----------------
    ws = wb.create_sheet("평가 기준")
    txt = [
        ("■ 발생가능성 (근거에서 자동 산정, AI 매트릭스 v3.2 수식과 동일)", ""),
        ("상", f"실제 사고 ≥{LIKELY_HIGH_REAL}건"),
        ("중", "실제 사고 1건, 또는 ATT&CK 클라우드 사례·실증 연구·벤더 특화기법 중 하나 이상 존재"),
        ("하", "근거 없음(이론·시나리오)"),
        ("근거 대응", "AI 매트릭스 '실제 사고(ATLAS Incident+OWASP 인용)' ↔ 사고 DB 실제 사고 / 'Realized 기법' ↔ ATT&CK 클라우드 사례 / "
                      "'실증 사례·공개 취약점' ↔ 사고 DB 연구·노출 + 벤더 매트릭스 문서화. ATT&CK 사례는 건수와 무관하게 '중'까지만 반영"), ("", ""),
        ("■ 심각도 (기법별 기준값, 근거 열 참조)", ""),
        ("상", "클라우드 계정·테넌트 장악, 대규모 데이터 유출·파괴·암호화, 핵심 자격증명·서명키 탈취, 과금·자원 대량 손실"),
        ("중", "단일 워크로드·서비스 범위 침해, 지속성·은닉 확보 등 후속 공격 기반"),
        ("하", "열람·탐색 위주로 직접 피해 없음(후속 공격 준비)"), ("", ""),
        ("■ 위험도 (발생가능성 × 심각도)", "심각도 상 / 중 / 하"),
        ("발생가능성 상", "매우 높음 / 높음 / 보통"),
        ("발생가능성 중", "높음 / 보통 / 낮음"),
        ("발생가능성 하", "보통 / 낮음 / 낮음"), ("", ""),
        ("■ 근거 수준 (높은 순)", ""),
        ("실제 사고 확인", "사고 DB의 실제 사고(사고·캠페인·사례연구·공시·위협인텔)와 연결"),
        ("실사용 기법 포함", "실제 사고 연결은 없으나 ATT&CK 그룹·캠페인·악성코드의 실제 사용 사례 존재"),
        ("실증·공개 취약점", "연구·취약점 공개 사례 또는 벤더(AWS/Azure/K8s) 매트릭스 문서화만 존재"),
        ("이론·시나리오", "매트릭스상 가능 기법이나 사례·문서 근거 없음"), ("", ""),
        ("■ 사고 → 기법 매핑", ""),
        ("초안ID", "사고 DB의 ATT&CK 초안 ID 그대로(폐기 ID는 v19.2로 변환: T1562→T1685 등)"),
        ("키워드", f"초안 ID를 보완하는 2차 매핑(v5): 사고 제목·요약·원문 필드의 키워드와 근본원인·서비스 계층·영향유형 조건으로 기법 연결 — scripts/taxonomy_rules.py KEYWORD_RULES({len(KEYWORD_RULES)}개)"),
        ("초안ID→키워드 세분", "초안에 상위기법만 있고 키워드가 하위기법을 가리키면 하위기법으로 옮김(예: T1496 → T1496.001 컴퓨트 채굴, T1566 → T1566.004 비싱)"),
        ("상위기법 태그", "하위기법으로 세분되지 않은 상위기법 태그는 해당 기법의 '.0 (일반·상위기법)' 행으로 집계"),
        ("집계 제외", "클라우드 관련도 '비클라우드 의심' 사고"),
        ("최근 사고", f"기준일 {RECENT_FROM}년 이후 실제 사고"),
    ]
    for i, (a, b) in enumerate(txt, 1):
        ws.cell(i, 1, a).font = Font(bold=True, color="1F3864" if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    _w(ws, [22, 110])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    domains, out, incs, stats = build()
    path = write_xlsx(domains, out, incs, stats)
    print("문구 입력 템플릿:", write_template(out))
    print("삭제", len(stats["dropped"]), "/ 통합", len(stats["merged"]),
          "/ 삭제 기법에만 연결돼 빠진 사고", stats["lost_incidents"])
    C = collections.Counter
    counted = [e for e in incs if e["counted"]]
    print("저장:", path)
    print("도메인", len(domains), "/ 세부위협", len(out))
    print("사고: 집계대상", len(counted), "/ 매핑", sum(1 for e in counted if e["links"]))
    print("위험도", C(o["risk"] for o in out))
    print("근거수준", C(o["level"] for o in out))
