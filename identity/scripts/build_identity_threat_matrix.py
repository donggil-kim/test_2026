# -*- coding: utf-8 -*-
"""
통합 신원(아이덴티티)·계정 보안위협 매트릭스 빌더
-----------------------------------------------
분류 : 신원 공격면 10개 도메인 → 위협분류(Lv2) → 세부위협(Lv3) — data/taxonomy.yaml
교차 : ATT&CK v19.2 · Browser & Identity Attacks Matrix(SAT) · OWASP NHI Top 10 · NIST SP 800-63-4 · OWASP ASVS 5.0 · OWASP OAT
근거 : 신원 보안사고 DB(data/incidents.yaml) · ATT&CK 절차(신원 맥락) · CISA KEV(신원 판정) · SAT 공개 시연
대응 : CSA CCM v4.1 · CIS Controls v8.1 · NIST SP 800-53 Rev.5 · ISMS-P
논리 : AI v3.2 · 클라우드 v5 · OT · 공급망 매트릭스와 같은 발생가능성 수식(상 = 실제 사고 2건 이상), 심각도는 신원 결과 기준

문구 : data/text/*.yaml(핵심 요약·요약설명·참조·탐지·대응) — 빌드 전에 scripts/validate.py로 사례·참조 ID를 대조
시나리오 : data/scenarios.yaml(실제 사고 기반 공격 체인)

실행 : python3 scripts/build_identity_threat_matrix.py [--allow-missing-text] [--skip-validate] [--worksheet <폴더>]
       → output/통합_신원보안위협_매트릭스_<버전>.xlsx · .csv
"""
import collections
import csv
import datetime
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import taxonomy_rules as R  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
REF = ROOT / "reference"
VERSION = "v2"
OUT = ROOT / "output" / f"통합_신원보안위협_매트릭스_{VERSION}.xlsx"
CSV_OUT = ROOT / "output" / f"통합_신원보안위협_매트릭스_{VERSION}.csv"


def load_yaml(p, default=None):
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else default


# ===========================================================================
# 1) 원천 데이터
# ===========================================================================
TAX = load_yaml(DATA / "taxonomy.yaml")
FW = load_yaml(DATA / "frameworks.yaml")
INCIDENTS = load_yaml(DATA / "incidents.yaml")
CHANGELOG = load_yaml(DATA / "changelog.yaml", {})
SCENARIOS = load_yaml(DATA / "scenarios.yaml", {}) or {}
ATK = json.loads((REF / "attack" / "enterprise-attack-v19.2-subset.json").read_text(encoding="utf-8"))
SAT_DOC = json.loads((REF / "push-bia" / "techniques.json").read_text(encoding="utf-8"))
SAT = SAT_DOC["techniques"]
NIST = json.loads((REF / "nist" / "sp800-53r5-names.json").read_text(encoding="utf-8"))["names"]
LINKS = json.loads((REF / "links" / "other_matrices.json").read_text(encoding="utf-8"))["ids"]
KEV_META = json.loads((REF / "advisories" / "known_exploited_vulnerabilities.json").read_text(encoding="utf-8"))
KEV = {v["cveID"]: v for v in KEV_META["vulnerabilities"]}

TEXT = {}
for f in sorted((DATA / "text").glob("*.yaml")):
    TEXT.update(load_yaml(f, {}) or {})

DOMAINS, ROWS = [], []
for d in TAX["domains"]:
    DOMAINS.append(d)
    for l2 in d["lv2"]:
        for x in l2["lv3"]:
            ROWS.append(dict(x, code=d["code"], domain=f"[{d['code']}] {d['ko']}", lv2_id=l2["id"], lv2=l2["name"],
                             single=(x["id"] == l2["id"])))
ROW = {o["id"]: o for o in ROWS}

for e in INCIDENTS:
    e["date"] = str(e["date"])
    e["status"] = R.incident_status(e)
    e["recent"] = e["date"] >= R.RECENT_FROM
INC = {e["id"]: e for e in INCIDENTS}

# ===========================================================================
# 2) 근거 집계
# ===========================================================================
EV = {o["id"]: dict(real=set(), recent=set(), research=set(), intel=set(), excluded=set(), subj=set(),
                    kev=set(), sat=set(), env=collections.Counter(), impact=collections.Counter())
      for o in ROWS}

for e in INCIDENTS:
    for k in e["idt"]:
        x = EV[k]
        if e["status"] == "실제 사고":
            x["real"].add(e["id"])
            if e["recent"]:
                x["recent"].add(e["id"])
            x["env"].update(e["id_env"])
            x["impact"].update(i for i in e["impact"] if i != "영향 없음(사전 차단)")
        elif e["status"] == "실증·연구":
            x["research"].add(e["id"])
        elif e["status"] == "위협인텔":
            x["intel"].add(e["id"])
        else:
            x["excluded"].add(e["id"])


def atk_subjects(tid):
    """기법을 쓰는 ATT&CK 주체 — 신원 기법은 전체, 범용 기법은 절차 설명에 신원 단서가 있는 것만"""
    out = set()
    for p in ATK["procedures"].get(tid, []):
        if R.attack_counts_all(tid) or R.ATTACK_ID_KEYWORDS.search(p["text"]):
            out.add(p["subject"])
    return out


for o in ROWS:
    for t in o["attack"]:
        EV[o["id"]]["subj"] |= atk_subjects(t)
    # SAT 공개 시연 — 앱별 예시(PoC)가 있는 기법만 실증 근거로 센다
    EV[o["id"]]["sat"] = {s for s in o["sat"] if SAT[s]["n_examples"] > 0}

for cve, (cat, rows, note) in R.KEV_IDENTITY.items():
    if cve not in KEV:
        raise SystemExit(f"KEV 카탈로그에 없는 CVE: {cve}")
    for k in rows:
        EV[k]["kev"].add(cve)


# ===========================================================================
# 3) 위험 평가
# ===========================================================================
def likelihood(x):
    if len(x["real"]) >= R.LIKELY_HIGH_REAL:
        return "상"
    if x["real"] or x["subj"] or x["kev"] or x["intel"] or x["research"] or x["sat"]:
        return "중"
    return "하"


def evidence_level(x):
    if x["real"]:
        return "실제 사고 확인"
    if x["subj"] or x["kev"] or x["intel"]:
        return "실사용 기법 포함"
    if x["research"] or x["sat"]:
        return "실증·공개 취약점"
    return "이론·시나리오"


TACTIC_KO = FW["attack_tactics_ko"]

for o in ROWS:
    x = EV[o["id"]]
    o["likelihood"] = likelihood(x)
    o["risk"] = R.RISK_MATRIX[(o["likelihood"], o["severity"])]
    o["level"] = evidence_level(x)
    o["n_real"], o["n_recent"], o["n_research"] = len(x["real"]), len(x["recent"]), len(x["research"])
    o["n_atk"], o["n_kev"], o["n_sat"], o["n_intel"] = len(x["subj"]), len(x["kev"]), len(x["sat"]), len(x["intel"])
    o["lk_why"] = (f"실제 사고 {o['n_real']}건(최근 {o['n_recent']}) · ATT&CK 사례 {o['n_atk']}건 · KEV {o['n_kev']}건 · "
                   f"위협인텔 {o['n_intel']}건 · 실증·연구 {o['n_research']}건 · SAT 시연 {o['n_sat']}건")
    txt = TEXT.get(o["id"]) or {}
    o["oneline"], o["summary"], o["reference"], o["detect"] = (txt.get(k) or "" for k in ("oneline", "summary", "reference",
                                                                                         "detect"))
    o["text_src"] = "분석가 작성" if txt else "문구 미작성"

rank_key = {r: i for i, r in enumerate(R.RISK_ORDER)}
for i, o in enumerate(sorted(ROWS, key=lambda o: (rank_key[o["risk"]], -o["n_real"], -o["n_recent"], o["id"])), 1):
    o["priority"] = i


def fw_pair(cat, key):
    v = FW[cat][key]
    return v if isinstance(v, list) else [v]


def fw_label(cat, key):
    v = fw_pair(cat, key)
    return f"{key} {v[-1]}"


def nist_label(k):
    return f"{k} {NIST.get(k, '')}"


def atk_label(t):
    return f"{t} {ATK['techniques'][t]['name']}"


def sat_label(s):
    return f"{s} {SAT[s]['name']}"


def link_label(k):
    return f"{k} {LINKS[k]['name']}" if k in LINKS else k


def inc_label(i):
    e = INC[i]
    return f"{i} {e['title']}({e['date']})"


def _bullets(block, section):
    """'■ 단락' 아래 '- ' 줄 목록"""
    out, on = [], False
    for ln in (block or "").splitlines():
        if ln.startswith("■ "):
            on = ln[2:].strip() == section
        elif on and ln.startswith("- "):
            out.append(ln[2:].strip())
    return out


def first_case(o):
    """대표 사례 = '■ 실제 사례' 첫 줄 → '[라벨] 사례명(시점) (사고 ID)' (보고서용 간략 표기)
    문구 미작성 행은 매핑된 실제 사고 중 최신 1건(사고 DB 제목·시점)"""
    if not o["reference"]:
        real = sorted(EV[o["id"]]["real"], key=lambda i: (INC[i]["date"], i), reverse=True)
        return f"[실제 사고] {INC[real[0]]['title']}({INC[real[0]]['date']}) ({real[0]})" if real else "공개 사고 미확인"
    for b in _bullets(o["reference"], "실제 사례"):
        if b.startswith("공개 사고 미확인"):
            return "공개 사고 미확인"
        if b.startswith("참고:"):
            continue
        ids = re.findall(r"IDI-\d{3}", b)
        return b.split(": ", 1)[0] + (f" ({', '.join(ids)})" if ids else "")
    return ""


def first_control(o):
    """핵심 대응 = '■ 대응' 첫 줄, 문구 미작성 행은 매핑된 대응 기준 ID"""
    b = _bullets(o["detect"], "대응")
    if b:
        return b[0]
    ids = [f"CIS {'·'.join(str(k) for k in o['cis'])}" if o["cis"] else "",
           f"NIST {'·'.join(o['nist'])}" if o["nist"] else "",
           f"ISMS-P {'·'.join(str(k) for k in o['ismsp'])}" if o["ismsp"] else ""]
    return "대응 기준: " + ", ".join(i for i in ids if i) if any(ids) else ""


def run_validate(allow_missing_text=False):
    """scripts/validate.py 실행 — 오류가 있으면 빌드 중단, 결과 줄을 개요 시트에 기록"""
    cmd = [sys.executable, str(Path(__file__).resolve().parent / "validate.py")]
    if allow_missing_text:
        cmd.append("--allow-missing-text")
    r = subprocess.run(cmd,
                       capture_output=True, text=True, encoding="utf-8")
    lines = [ln for ln in r.stdout.splitlines() if ln.strip()]
    if r.returncode != 0:
        print("\n".join(lines))
        raise SystemExit("검증 오류 — 문구·데이터를 고친 뒤 다시 빌드(--skip-validate로 무시 가능)")
    return next((ln for ln in reversed(lines) if ln.startswith("검증 결과")), "")


VALIDATION = "검증 생략(--skip-validate)"

# ===========================================================================
# 4) xlsx 출력 — 클라우드 v5 · OT · 공급망 서식 준용
# ===========================================================================
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "실사용 기법 포함": "2E86C1",
               "실증·공개 취약점": "B9770E", "이론·시나리오": "7F8C8D"}
HDR, NAVY, GRAY = "34495E", "1F3864", "8C8C8C"


def _w(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hdr(ws, row, cols, start=1, color=HDR):
    for j, c in enumerate(cols, start):
        cell = ws.cell(row, j, c)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BORDER


def _risk(cell, risk):
    bg, fg = RISK_COLOR[risk]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=risk in ("매우 높음", "높음"), color=fg)
    cell.alignment = CENTER


def _level(cell, level):
    cell.font = Font(bold=True, color=LEVEL_COLOR[level])
    cell.alignment = CENTER


def _row(ws, r, vals, start=1):
    for j, v in enumerate(vals, start):
        c = ws.cell(r, j, v)
        c.alignment = WRAP
        c.border = BORDER


def _title(ws, text, sub=None):
    ws.cell(1, 1, text).font = Font(size=14, bold=True, color=NAVY)
    if sub:
        ws.cell(2, 1, sub).font = Font(size=9, color="595959")


def _section(ws, r, text):
    ws.cell(r, 1, text).font = Font(bold=True, size=11, color=NAVY)
    return r + 1


def summary_cell(o):
    return o["summary"] or f"(문구 미작성) {o['en']}"


def write_xlsx(path):
    wb = openpyxl.Workbook()
    today = datetime.date.today().isoformat()
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    sc = collections.Counter(o["severity"] for o in ROWS)
    st = collections.Counter(e["status"] for e in INCIDENTS)
    real = [e for e in INCIDENTS if e["status"] == "실제 사고"]
    kr = [e for e in real if e.get("region") == "한국"]
    written = sum(1 for o in ROWS if o["text_src"] == "분석가 작성")
    top = sorted(ROWS, key=lambda o: o["priority"])[:10]
    top_stage = [o for o in sorted(ROWS, key=lambda o: o["priority"]) if o["code"] != "IM"][:10]
    n_lv2 = sum(len(d["lv2"]) for d in DOMAINS)
    kev_by_cat = collections.Counter(v[0] for v in R.KEV_IDENTITY.values())
    years = sorted({e["date"][:4] for e in INCIDENTS})
    sat_mapped = sorted({s for o in ROWS for s in o["sat"]})
    reused = [e for e in INCIDENTS if e.get("origin")]

    # ---------------- 개요 ----------------
    ws = wb.active
    ws.title = "개요"
    meta = [
        (f"통합 신원(아이덴티티)·계정 보안위협 매트릭스 {VERSION}", ""),
        ("기준", "신원 공격면 10개 도메인 → 위협분류 → 세부위협 분류에 ATT&CK v19.2·Browser & Identity Attacks Matrix·OWASP NHI Top 10·"
                 "NIST SP 800-63-4·OWASP ASVS 5.0을 교차 매핑하고, AI v3.2·클라우드 v5·OT·공급망 매트릭스의 위험평가·근거 로직을 이식"),
        ("생성일", today), ("", ""),
        ("■ 기준 데이터", ""),
        ("신원 보안사고 DB", f"{len(INCIDENTS)}건({years[0]}~{years[-1]}, 공개 출처) — 실제 사고 {st['실제 사고']}건 · "
                        f"실증·연구 {st['실증·연구']}건 · 위협인텔 {st.get('위협인텔', 0)}건 · 집계 제외 {st.get('제외(검증)', 0)}건, "
                        f"국내 실제 사고 {len(kr)}건, 최근(2025~) 실제 사고 {sum(1 for e in real if e['recent'])}건, "
                        f"공급망 사고 DB 재인용 {len(reused)}건(origin 표기)"),
        ("MITRE ATT&CK", f"Enterprise v19.2 — 매핑 기법 {len(ATK['techniques'])}개, 그 기법을 쓰는 주체 {len(ATK['subjects'])}개의 절차 "
                         "중 신원 맥락 절차만 'ATT&CK 사례 수'로 집계(신원 기법은 전체)"),
        ("Browser & Identity Attacks Matrix", f"Push Security(구 SaaS Attacks Matrix) 커밋 {SAT_DOC['commit'][:7]} — 기법 {len(SAT)}개 중 "
                                              f"{len(sat_mapped)}개 매핑, 앱별 시연(예시)이 있는 기법은 '실증' 근거로만 집계(CC BY 4.0)"),
        ("OWASP · NIST", "OWASP Non-Human Identities Top 10(2025) · ASVS 5.0 · Automated Threats · NIST SP 800-63-4 — ID·명칭(주제)만 연계"),
        ("CISA KEV", f"{KEV_META.get('catalogVersion')}판 {KEV_META.get('count')}건 중 신원 판정 {len(R.KEV_IDENTITY)}건 — "
                     + " · ".join(f"{k} {v}" for k, v in kev_by_cat.items())),
        ("대응 기준", "CSA CCM v4.1(IAM 등) · CIS Controls v8.1(5·6) · NIST SP 800-53 Rev.5 · ISMS-P 인증기준"),
        ("다른 매트릭스", "통합 AI·클라우드·OT 매트릭스 v2 요약 ID(AI-·CL-·OT-) · 공급망 매트릭스 v2(SCT-) 연계 열과 '다른 매트릭스 연계' 시트"),
        ("", ""),
        ("■ 시트 구성", ""),
        ("보고서용 간략 매트릭스", "본문 삽입용 — 우선순위 순 세부위협 · 핵심 요약 · 위험도 · 실제 사고 수 · 대표 사례 1건 · 핵심 대응 1줄"),
        ("매트릭스 뷰", "도메인(열)별 세부위협을 위험도 색으로 배치한 한눈 보기"),
        ("통합 매트릭스", "분류 체계 · 교차 매핑 · 위험 평가 · 실제 근거 · 대응 기준 전체 열(문구 포함)"),
        ("통합매트릭스_LITE", "핵심 열 발췌(필터·보고용)"),
        ("도메인 요약", "도메인별 위협 수 · 위험도/근거수준 분포 · 매핑 사고 수 · 최고위험 항목"),
        ("신원 유형·ID 환경 요약", "신원 유형(인력·특권·NHI·고객)과 ID 환경(온프레미스 AD·클라우드 IdP·SaaS 등)별 실제 사고·영향·대표 세부위협"),
        ("공격 체인 시나리오", (f"실제 사고 기반 공격 흐름 {len(SCENARIOS['scenarios'])}개 — 단계별 세부위협(IDT-ID)·행위·탐지 포인트·초크 포인트"
                             if SCENARIOS.get("scenarios") else "시트 양식만 수록 — 시나리오(data/scenarios.yaml)는 다음 단계에서 작성")),
        ("역매핑_사고사례", f"사고 DB {len(INCIDENTS)}건과 매핑 세부위협·출처(근거 추적용)"),
        ("KEV 근거", "CISA KEV 신원 판정 내역(구분·매핑 세부위협·비고)"),
        ("프레임워크 연계", "SAT · OWASP NHI · NIST SP 800-63-4 · ASVS · OAT 항목별 연계 세부위협(커버리지)"),
        ("대응 기준 연계", "CSA CCM · CIS Controls · NIST 800-53 · ISMS-P별 연계 세부위협과 위험도 분포(통제 우선순위 근거)"),
        ("다른 매트릭스 연계", "AI·클라우드·OT·공급망 매트릭스 항목별로 연결된 신원 세부위협(같은 위협을 오가며 보기)"),
        ("평가 기준", "발생가능성·심각도·위험도·근거수준 산정 규칙, 사고 집계·ATT&CK 맥락·KEV 판정 기준, 문구 작성 규칙"),
        ("변경이력", "버전별 변경 내역"), ("", ""),
        ("■ 분류 체계", ""),
        ("Lv1 도메인", "신원 공격면 10개 — " + " → ".join(f"[{d['code']}] {d['ko']}" for d in DOMAINS)),
        ("Lv2 위협분류", f"{n_lv2}개(IDT-<도메인>-NN). 하위 세부위협이 하나뿐이면 Lv2 = Lv3"),
        ("Lv3 세부위협", f"{len(ROWS)}개(IDT-<도메인>-NN.m) — 요약설명 · 참조(신원 관점·실제 사례) · 탐지·대응. ID는 고정(재사용 안 함)"),
        ("설계 원칙", "공급망 매트릭스처럼 공격면 도메인을 쓰고 ATT&CK 전술을 '공격 단계' 열로 병기 — 같은 기법(예: 토큰 탈취)이 "
                    "인력·특권·비인간 신원·고객마다 다른 통제로 막히므로 공격면 축이 통제 지점을 잘 드러냄"),
        ("신원 유형", " · ".join(f"{k}: {v}" for k, v in TAX["id_types"].items())), ("", ""),
        ("■ 결과 요약", ""),
        ("세부위협(Lv3)", f"{len(ROWS)}개 / 위협분류 {n_lv2}개 / 도메인 {len(DOMAINS)}개"),
        ("위험도", " · ".join(f"{k} {rc.get(k, 0)}" for k in R.RISK_ORDER)),
        ("근거 수준", " · ".join(f"{k} {lc.get(k, 0)}" for k in R.LEVELS)),
        ("심각도", " · ".join(f"{k} {sc.get(k, 0)}" for k in ("상", "중", "하"))
         + " — '상'은 신원 체계 장악·다수 계정 동시 장악·정상 권한 대량 피해일 때만(평가 기준 시트)"),
        ("신원 관점 문구", f"{written}/{len(ROWS)}개 작성(핵심 요약·요약설명·참조·탐지·대응)"
                       + (" — 전 항목 완료" if written == len(ROWS) else
                          f" — 나머지 {len(ROWS) - written}개는 '문구 미작성'(요약설명 자리에 영문명 회색 표시, 대표 사례는 사고 DB 최신 "
                          "실제 사고, 핵심 대응은 매핑된 대응 기준 ID)")
                       + f" | scripts/validate.py {VALIDATION}"),
        ("우선 위협 Top 10", "\n".join(f"{o['priority']}. {o['id']} {o['name']} — {o['risk']}, 실제 사고 {o['n_real']}건"
                                    f"(최근 {o['n_recent']})" for o in top)),
        ("우선 위협 Top 10(IM 제외)", "\n".join(f"{o['priority']}. {o['id']} {o['name']} — {o['risk']}, 실제 사고 {o['n_real']}건"
                                           for o in top_stage)),
        ("", ""), ("■ 주의", ""),
        ("발생가능성", f"공개 사고 기준이라 공시·소송·정부 경보로 드러난 대형 사건(헬프데스크 사칭·MFA 미적용·크리덴셜 스터핑)이 높게 나오고, "
                     f"디렉터리 내부 공격(Golden Ticket 등)처럼 경위가 공개되지 않는 위협은 과소평가될 수 있음. 수식은 다른 매트릭스와 같게 "
                     f"유지(상 = 실제 사고 {R.LIKELY_HIGH_REAL}건 이상)"),
        ("다른 매트릭스와 비교", "사고 DB 규모·출처가 매트릭스마다 달라(클라우드 680건·OT 85건·공급망 167건·신원 "
                           f"{len(INCIDENTS)}건) 위험도·사고 수는 같은 매트릭스 안에서 비교. 같은 위협도 매트릭스별 심각도 기준이 달라 "
                           "등급이 다를 수 있음('다른 매트릭스 연계' 시트로 오가며 보기)"),
        ("영향·확산 도메인", "[IM] 행은 결과 유형이라 여러 사고가 몰려 우선순위 상위에 오름 — 통제 우선순위는 앞 단계 도메인과 '대응 기준 연계' 시트로 판단"),
        ("근거 중복", "한 사고가 여러 단계(헬프데스크 사칭 → IdP 장악 → 랜섬웨어)에 동시에 매핑되므로 세부위협별 사고 수 합은 사고 DB "
                    "건수보다 큼. 공급망 사고 DB와 겹치는 사고는 origin 열로 표시(매트릭스 간 합산 금지)"),
        ("SAT 시연", "Browser & Identity Attacks Matrix의 예시는 앱별 시연(PoC)이며 관측 기반만이 아니라 탐색적 기법을 포함 — '실증' 근거로만 "
                   "쓰고 실제 사고로 세지 않음"),
        ("최근 사고", "2025~2026년 사고는 1차 공개 자료·정부 발표·보안업체 분석 기준이며 추가 조사로 수치·귀속이 바뀔 수 있음"),
        ("출처 표기", "MITRE ATT&CK® © The MITRE Corporation · Browser & Identity Attacks Matrix © Push Security(CC BY 4.0) · "
                    "OWASP NHI Top 10·ASVS(CC BY-SA 4.0)·Automated Threats(CC BY-SA 3.0) — ID·명칭만 · NIST SP 800-63-4·800-53(공공 영역) · "
                    "CSA CCM · CIS Controls(CC BY-NC-ND 4.0) — ID·명칭만 · ISMS-P · CISA KEV(공공 영역). 사고별 출처는 '역매핑_사고사례' 시트"),
    ]
    for i, (a, b) in enumerate(meta, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    ws.cell(1, 1).font = Font(size=15, bold=True, color=NAVY)
    _w(ws, [24, 140])

    # ---------------- 보고서용 간략 매트릭스 ----------------
    ws = wb.create_sheet("보고서용 간략 매트릭스")
    bcols = ["순위", "도메인(Lv1)", "IDT-ID", "세부위협(Lv3)", "핵심 요약", "위험도", "발생가능성", "심각도", "실제 사고 수",
             "최근(2025~)", "대표 사례", "핵심 대응"]
    t = ws.cell(1, 1, f"보고서용 간략 매트릭스 {VERSION} — 본문 삽입용(상세 문구는 '통합 매트릭스')")
    t.font = Font(size=13, bold=True, color=NAVY)
    ws.cell(2, 1, "순위 = 위험도 → 실제 사고 수 → 최근 사고(2025~) → IDT-ID | 대표 사례 = 참조 '■ 실제 사례' 첫 줄 | "
                  f"핵심 대응 = '■ 대응' 첫 줄 | {today}").font = Font(size=9, color="595959")
    _hdr(ws, 4, bcols)
    for r, o in enumerate(sorted(ROWS, key=lambda o: o["priority"]), 5):
        _row(ws, r, [o["priority"], o["domain"], o["id"], o["name"], o["oneline"] or summary_cell(o), o["risk"], o["likelihood"],
                     o["severity"], o["n_real"], o["n_recent"], first_case(o), first_control(o)])
        ws.cell(r, 3).font = Font(bold=True)
        ws.cell(r, 4).font = Font(bold=True)
        _risk(ws.cell(r, 6), o["risk"])
        for j in (1, 7, 8, 9, 10):
            ws.cell(r, j).alignment = CENTER
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(bcols))}{len(ROWS) + 4}"
    _w(ws, [6, 16, 12, 28, 44, 9, 7, 7, 7, 7, 46, 54])
    ws.print_title_rows = "4:4"

    # ---------------- 매트릭스 뷰 ----------------
    ws = wb.create_sheet("매트릭스 뷰")
    by_c = collections.defaultdict(list)
    for o in ROWS:
        by_c[o["code"]].append(o)
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도 (빨강 매우 높음 · 주황 높음 · 노랑 보통 · 초록 낮음), [n] = 실제 사고 수, "
                  "괄호 안 숫자 = 세부위협 수").font = Font(bold=True, color=NAVY)
    for j, d in enumerate(DOMAINS, 1):
        c = ws.cell(2, j, f"[{d['code']}] {d['ko']}\n{d['en']}\n({len(by_c[d['code']])})")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        for i, o in enumerate(sorted(by_c[d["code"]], key=lambda o: (rank_key[o["risk"]], -o["n_real"], o["id"])), 3):
            cell = ws.cell(i, j, f"{o['id']} {o['name']}" + (f" [{o['n_real']}]" if o["n_real"] else ""))
            _risk(cell, o["risk"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9, bold=o["risk"] == "매우 높음", color=RISK_COLOR[o["risk"]][1])
            cell.border = BORDER
    ws.row_dimensions[2].height = 48
    ws.freeze_panes = "A3"
    _w(ws, [24] * len(DOMAINS))

    # ---------------- 통합 매트릭스 ----------------
    ws = wb.create_sheet("통합 매트릭스")
    g_class = ["도메인(Lv1)", "IDT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "참조", "영문명",
               "공격 단계(ATT&CK 전술)", "설명 출처"]
    g_cross = ([f"신원:{p}" for p in R.ID_TYPES] + ["ID 환경", "ATT&CK ID", "Browser & Identity Attacks(SAT)", "OWASP NHI Top10",
                                                   "NIST SP 800-63-4", "OWASP ASVS 5.0", "OWASP 자동화 위협",
                                                   "클라우드 매트릭스 연계", "AI 매트릭스 연계", "OT 매트릭스 연계", "공급망 매트릭스 연계",
                                                   "ATT&CK 기법명"])
    g_risk = ["발생가능성", "심각도", "위험도", "우선순위", "발생가능성 근거 (자동 산정)", "심각도 근거"]
    g_ev = ["실제 사고 수", "최근 사고(2025~)", "ATT&CK 사례 수(신원 맥락)", "KEV(신원)", "위협인텔", "실증·연구 수", "SAT 시연",
            "근거 수준", "관련 사례 ID", "주요 ID 환경(실제 사고)", "ATT&CK 사례 주체(일부)"]
    g_ctl = ["CSA CCM v4.1", "CIS Controls v8.1", "NIST SP 800-53", "ISMS-P", "탐지·대응 포인트"]
    groups = [("분류 체계", len(g_class), "2E5496"), ("교차 매핑", len(g_cross), "1F7A8C"),
              ("위험 평가", len(g_risk), "A04000"), ("실제 근거", len(g_ev), "1E8449"), ("대응 기준", len(g_ctl), "6C3483")]
    cols = g_class + g_cross + g_risk + g_ev + g_ctl
    t = ws.cell(1, 1, f"통합 신원·계정 보안위협 매트릭스 {VERSION} — 분류체계 · 위험평가 · 실제근거 · 대응기준")
    t.font = Font(size=13, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, f"신원 보안사고 DB {len(INCIDENTS)}건 + ATT&CK v19.2 + CISA KEV {len(R.KEV_IDENTITY)}건 + Browser & Identity "
                  f"Attacks Matrix {len(SAT)}개 기법 | {today}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    c0 = 1
    for g, span, color in groups:
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + span - 1)
        c = ws.cell(3, c0, g)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        c0 += span
    _hdr(ws, 4, cols)
    ci = {c: i + 1 for i, c in enumerate(cols)}
    for r, o in enumerate(ROWS, 5):
        x = EV[o["id"]]
        prof = ["●" if p in o["id_types"] else "" for p in R.ID_TYPES]
        subj = sorted(x["subj"])
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summary_cell(o), o["reference"], o["en"],
                     ", ".join(o["stage"]), o["text_src"]] + prof
             + [", ".join(o["id_envs"]), ", ".join(o["attack"]), "\n".join(sat_label(s) for s in o["sat"]),
                "\n".join(fw_label("owasp_nhi", k) for k in o["nhi"]), "\n".join(fw_label("sp80063", k) for k in o["sp80063"]),
                "\n".join(fw_label("asvs", k) for k in o["asvs"]), "\n".join(fw_label("oat", k) for k in o["oat"]),
                "\n".join(link_label(k) for k in o["cloud"]), "\n".join(link_label(k) for k in o["ai"]),
                "\n".join(link_label(k) for k in o["ot"]), "\n".join(link_label(k) for k in o["sct"]),
                "\n".join(atk_label(t) for t in o["attack"]),
                o["likelihood"], o["severity"], o["risk"], o["priority"], o["lk_why"], o["severity_why"],
                o["n_real"], o["n_recent"], o["n_atk"], o["n_kev"], o["n_intel"], o["n_research"], o["n_sat"], o["level"],
                ", ".join(sorted(x["real"])) + (("\n실증: " + ", ".join(sorted(x["research"]))) if x["research"] else "")
                + (("\n위협인텔: " + ", ".join(sorted(x["intel"]))) if x["intel"] else ""),
                " · ".join(f"{k} {v}" for k, v in x["env"].most_common(4)),
                ", ".join(f"{s} {ATK['subjects'][s]['name']}" for s in subj[:12]) + (f" 외 {len(subj) - 12}" if len(subj) > 12 else ""),
                "\n".join(fw_label("ccm", k) for k in o["ccm"]), "\n".join(fw_label("cis", str(k)) for k in o["cis"]),
                "\n".join(nist_label(k) for k in o["nist"]), "\n".join(fw_label("ismsp", str(k)) for k in o["ismsp"]),
                o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, ci["요약설명"]).font = Font(color=GRAY)
        for c in [f"신원:{p}" for p in R.ID_TYPES] + ["발생가능성", "심각도", "우선순위"]:
            ws.cell(r, ci[c]).alignment = CENTER
        _risk(ws.cell(r, ci["위험도"]), o["risk"])
        _level(ws.cell(r, ci["근거 수준"]), o["level"])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(ROWS) + 4}"
    widths = {"도메인(Lv1)": 16, "IDT-ID": 12, "위협분류(Lv2)": 18, "세부위협(Lv3)": 26, "핵심 요약": 34, "요약설명": 60, "참조": 80,
              "영문명": 26, "공격 단계(ATT&CK 전술)": 14, "설명 출처": 9, "ID 환경": 18, "ATT&CK ID": 14,
              "Browser & Identity Attacks(SAT)": 24, "OWASP NHI Top10": 22, "NIST SP 800-63-4": 24, "OWASP ASVS 5.0": 22,
              "OWASP 자동화 위협": 18, "클라우드 매트릭스 연계": 22, "AI 매트릭스 연계": 22, "OT 매트릭스 연계": 20,
              "공급망 매트릭스 연계": 24, "ATT&CK 기법명": 30, "발생가능성": 7, "심각도": 7, "위험도": 8, "우선순위": 7,
              "발생가능성 근거 (자동 산정)": 30, "심각도 근거": 40, "실제 사고 수": 7, "최근 사고(2025~)": 7,
              "ATT&CK 사례 수(신원 맥락)": 8, "KEV(신원)": 7, "위협인텔": 7, "실증·연구 수": 7, "SAT 시연": 7, "근거 수준": 11,
              "관련 사례 ID": 24, "주요 ID 환경(실제 사고)": 22, "ATT&CK 사례 주체(일부)": 30, "CSA CCM v4.1": 26,
              "CIS Controls v8.1": 26, "NIST SP 800-53": 26, "ISMS-P": 18, "탐지·대응 포인트": 70}
    _w(ws, [widths.get(c, 6) for c in cols])

    # ---------------- 통합매트릭스_LITE ----------------
    ws = wb.create_sheet("통합매트릭스_LITE")
    ws.cell(1, 1, f"통합 신원·계정 보안위협 매트릭스 {VERSION} — LITE").font = Font(size=13, bold=True, color=NAVY)
    lcols = ["도메인(Lv1)", "IDT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "공격 단계", "신원 유형", "ID 환경",
             "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수", "최근 사고(2025~)", "탐지·대응 포인트"]
    lc_ = {c: i + 1 for i, c in enumerate(lcols)}
    _hdr(ws, 2, lcols)
    for r, o in enumerate(ROWS, 3):
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summary_cell(o), ", ".join(o["stage"]),
                     " · ".join(o["id_types"]), " · ".join(o["id_envs"]), o["likelihood"], o["severity"], o["risk"], o["priority"],
                     o["level"], o["n_real"], o["n_recent"], o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, lc_["요약설명"]).font = Font(color=GRAY)
        _risk(ws.cell(r, lc_["위험도"]), o["risk"])
        _level(ws.cell(r, lc_["근거 수준"]), o["level"])
        for c in ("발생가능성", "심각도", "우선순위", "실제 사고 수", "최근 사고(2025~)"):
            ws.cell(r, lc_[c]).alignment = CENTER
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(lcols))}{len(ROWS) + 2}"
    _w(ws, [16, 12, 18, 26, 34, 60, 14, 16, 22, 7, 7, 8, 7, 11, 7, 7, 70])

    # ---------------- 도메인 요약 ----------------
    ws = wb.create_sheet("도메인 요약")
    dcols = ["도메인", "설명", "위협분류(Lv2)", "세부위협(Lv3)"] + R.RISK_ORDER + R.LEVELS + \
            ["매핑 실제 사고(중복 제거)", "최근 사고(2025~)", "최고위험 세부위협(상위 3)"]
    _hdr(ws, 1, dcols)
    for i, d in enumerate(DOMAINS, 2):
        os_ = by_c[d["code"]]
        rr = collections.Counter(o["risk"] for o in os_)
        ll = collections.Counter(o["level"] for o in os_)
        incs = set().union(*(EV[o["id"]]["real"] for o in os_))
        top3 = sorted(os_, key=lambda o: o["priority"])[:3]
        _row(ws, i, [f"[{d['code']}] {d['ko']}", d["desc"], len(d["lv2"]), len(os_)] + [rr.get(k, 0) for k in R.RISK_ORDER]
             + [ll.get(k, 0) for k in R.LEVELS] + [len(incs), sum(1 for e in incs if INC[e]["recent"]),
                                                  "\n".join(f"{o['id']} {o['name']}({o['risk']}, {o['n_real']}건)" for o in top3)])
    _w(ws, [20, 60, 8, 8, 7, 7, 7, 7, 9, 9, 9, 9, 10, 9, 50])

    # ---------------- 신원 유형·ID 환경 요약 ----------------
    ws = wb.create_sheet("신원 유형·ID 환경 요약")
    _title(ws, "신원 유형·ID 환경별 요약 — 실제 사고 기준(한 사고가 여러 유형·환경에 걸치면 각각 집계)",
           "세부위협 수 = 그 유형·환경을 대상으로 하는 세부위협(taxonomy id_types·id_envs) · 위험도 분포는 해당 세부위협 기준")
    ecols = ["구분", "항목", "세부위협 수", "매우 높음·높음", "실제 사고", "최근(2025~)", "실증·연구", "주요 영향", "대표 세부위협(사고 수 상위 3)",
             "최근 사례(최대 3)"]
    r = 4
    for label, fld_row, fld_inc, items in (("신원 유형", "id_types", "id_type", R.ID_TYPES), ("ID 환경", "id_envs", "id_env", R.ID_ENVS)):
        _hdr(ws, r, ecols)
        r += 1
        for it in items:
            rows_ = [o for o in ROWS if it in o[fld_row]]
            es = [e for e in INCIDENTS if it in e[fld_inc]]
            er = [e for e in es if e["status"] == "실제 사고"]
            imp = collections.Counter(x for e in er for x in e["impact"] if x != "영향 없음(사전 차단)")
            rows_c = collections.Counter(k for e in er for k in e["idt"])
            recent = sorted(er, key=lambda e: e["date"], reverse=True)[:3]
            _row(ws, r, [label, it, len(rows_), sum(1 for o in rows_ if o["risk"] in ("매우 높음", "높음")), len(er),
                         sum(1 for e in er if e["recent"]), sum(1 for e in es if e["status"] == "실증·연구"),
                         " · ".join(f"{k} {v}" for k, v in imp.most_common(3)),
                         "\n".join(f"{k} {ROW[k]['name']}({v})" for k, v in rows_c.most_common(3)),
                         "\n".join(inc_label(e["id"]) for e in recent)])
            r += 1
        r += 1
    ws.freeze_panes = "C5"
    _w(ws, [10, 18, 8, 9, 8, 8, 8, 30, 46, 56])

    # ---------------- 공격 체인 시나리오 ----------------
    ws = wb.create_sheet("공격 체인 시나리오")
    ws.cell(1, 1, "공격 체인 시나리오 — 실제 사고를 신원 공격 흐름으로 재구성. 각 단계는 세부위협(IDT-ID, 위험도 색)에 연결하고, "
                  "흐름을 가장 싸게 끊는 통제를 초크 포인트로 표시").font = Font(bold=True, color=NAVY)
    ws.cell(2, 1, "행위는 사고 DB 요약에 있는 사실만 적음(data/scenarios.yaml, scripts/validate.py가 ID 대조)").font = Font(size=9, color="595959")
    r = 4
    if not SCENARIOS.get("scenarios"):
        ws.cell(r, 1, "작성 예정 — data/scenarios.yaml이 아직 없음").font = Font(color=GRAY)
    for scn in SCENARIOS.get("scenarios", []):
        c = ws.cell(r, 1, f"{scn['id']}. {scn['title']}")
        c.font = Font(size=12, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
        info = [("개요", scn["summary"]), ("근거 사고", " · ".join(inc_label(i) for i in scn["incidents"]))]
        if scn.get("ref"):
            info.append(("참조", scn["ref"]))
        for k, v in info:
            ws.cell(r, 1, k).font = Font(bold=True, size=9, color="595959")
            ws.cell(r, 2, v).alignment = WRAP
            ws.cell(r, 2).font = Font(size=9)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            r += 1
        _hdr(ws, r, ["#", "공격 단계", "IDT-ID", "세부위협", "행위", "탐지 포인트"])
        r += 1
        for i, st_ in enumerate(scn["steps"], 1):
            rows_ = [ROW[k] for k in st_["idt"]]
            _row(ws, r, [i, st_["단계"], "\n".join(o["id"] for o in rows_), "\n".join(o["name"] for o in rows_), st_["행위"],
                         st_["탐지"]])
            _risk(ws.cell(r, 3), min((o["risk"] for o in rows_), key=lambda x: rank_key[x]))
            ws.cell(r, 1).alignment = CENTER
            r += 1
        ws.cell(r, 1, "초크").font = Font(bold=True, color="A04000")
        ws.cell(r, 2, "\n".join(f"· {c}" for c in scn["chokepoints"])).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 15 * len(scn["chokepoints"]) + 4
        r += 2
    _w(ws, [6, 14, 14, 30, 56, 50])

    # ---------------- 역매핑_사고사례 ----------------
    ws = wb.create_sheet("역매핑_사고사례")
    icols = ["사례 ID", "시점", "사례명", "사례유형", "검증 수준", "집계 상태", "ID 환경", "신원 유형", "지역", "행위자", "영향", "요약",
             "매핑 세부위협", "ATT&CK 참조", "KEV", "다른 DB 사고 ID", "출처"]
    _hdr(ws, 1, icols)
    for i, e in enumerate(sorted(INCIDENTS, key=lambda e: (e["date"], e["id"])), 2):
        _row(ws, i, [e["id"], e["date"], e["title"], e["kind"], e["verification"], e["status"], ", ".join(e["id_env"]),
                     ", ".join(e["id_type"]), e.get("region", ""), e.get("actor", ""), ", ".join(e["impact"]), e["summary"],
                     "\n".join(f"{k} {ROW[k]['name']}" for k in e["idt"]), ", ".join(e.get("attack_ref") or []),
                     ", ".join(e.get("kev") or []), ", ".join(e.get("origin") or []), "\n".join(e["sources"])])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(icols))}{len(INCIDENTS) + 1}"
    _w(ws, [9, 8, 34, 9, 13, 10, 18, 12, 8, 18, 18, 70, 40, 14, 16, 10, 70])

    # ---------------- KEV 근거 ----------------
    ws = wb.create_sheet("KEV 근거")
    _title(ws, f"CISA KEV 신원 판정 — {KEV_META.get('catalogVersion')}판 {KEV_META.get('count')}건 중 {len(R.KEV_IDENTITY)}건",
           "판정: 인증 계열 CWE 후보 + 신원·접근 인프라 제품 후보 → 수동 판정. 신원 체계를 직접 노리는 것만(일반 원격 코드 실행·"
           "소비자 기기·OT 장비 제외). 실사용 근거('중' 상한)로만 반영")
    r = 4
    for cat in R.KEV_CATEGORIES:
        r = _section(ws, r, f"■ {cat} ({kev_by_cat[cat]})")
        kc = ["CVE", "벤더", "제품", "취약점명", "KEV 등재일", "랜섬웨어 악용", "CWE", "매핑 세부위협", "비고"]
        _hdr(ws, r, kc)
        r += 1
        for cve, (c, rows, note) in sorted(R.KEV_IDENTITY.items(), key=lambda kv: KEV[kv[0]]["dateAdded"], reverse=True):
            if c != cat:
                continue
            v = KEV[cve]
            _row(ws, r, [cve, v["vendorProject"], v["product"], v["vulnerabilityName"], v["dateAdded"],
                         v.get("knownRansomwareCampaignUse", ""), ", ".join(v.get("cwes") or []),
                         "\n".join(f"{k} {ROW[k]['name']}" for k in rows), note])
            r += 1
        r += 1
    _w(ws, [16, 16, 24, 46, 11, 9, 12, 40, 40])

    # ---------------- 프레임워크 연계 ----------------
    ws = wb.create_sheet("프레임워크 연계")
    _title(ws, "교차 매핑 항목별 연계 세부위협 — 커버리지 점검(미연계 항목은 사유 표기)")
    r = 3
    fw_sets = [("Browser & Identity Attacks Matrix (Push Security)", "sat", list(SAT.keys()),
                lambda k: SAT[k]["name"], lambda k: ", ".join(SAT[k]["tactics"])),
               ("OWASP Non-Human Identities Top 10 (2025)", "nhi", list(FW["owasp_nhi"]),
                lambda k: fw_pair("owasp_nhi", k)[0], lambda k: fw_pair("owasp_nhi", k)[1]),
               ("NIST SP 800-63-4 (주제 코드)", "sp80063", list(FW["sp80063"]),
                lambda k: fw_pair("sp80063", k)[0], lambda k: fw_pair("sp80063", k)[1]),
               ("OWASP ASVS 5.0", "asvs", list(FW["asvs"]), lambda k: fw_pair("asvs", k)[0], lambda k: fw_pair("asvs", k)[1]),
               ("OWASP Automated Threats", "oat", list(FW["oat"]), lambda k: fw_pair("oat", k)[0], lambda k: fw_pair("oat", k)[1])]
    for title, fld, keys, en, ko in fw_sets:
        r = _section(ws, r, f"■ {title}")
        _hdr(ws, r, ["ID", "명칭(영문)", "명칭(국문)·전술", "연계 세부위협 수", "연계 세부위협", "위험도 분포", "실제 사고(중복 제거)", "미연계 사유"])
        r += 1
        for k in keys:
            rows = [o for o in ROWS if k in o[fld]]
            incs = set().union(*(EV[o["id"]]["real"] for o in rows)) if rows else set()
            rcnt = collections.Counter(o["risk"] for o in rows)
            why = R.FRAMEWORK_NOT_MAPPED.get((fld, k), "")
            _row(ws, r, [k, en(k), ko(k), len(rows), "\n".join(f"{o['id']} {o['name']}" for o in rows),
                         " · ".join(f"{x} {rcnt[x]}" for x in R.RISK_ORDER if rcnt[x]), len(incs), why])
            r += 1
        r += 1
    _w(ws, [12, 40, 30, 8, 50, 26, 9, 40])

    # ---------------- 대응 기준 연계 ----------------
    ws = wb.create_sheet("대응 기준 연계")
    _title(ws, "대응 기준별 연계 세부위협 — 연계 위협의 위험도 분포로 통제 우선순위를 가늠",
           "CCM = 클라우드 통제 · CIS = 우선 실행 세이프가드 · NIST 800-53 = 조직 통제 · ISMS-P = 국내 인증기준")
    r = 4
    used_nist = sorted({k for o in ROWS for k in o["nist"]})
    ctl_sets = [("CSA CCM v4.1", "ccm", list(FW["ccm"]), lambda k: fw_pair("ccm", k)[0], lambda k: fw_pair("ccm", k)[1]),
                ("CIS Controls v8.1", "cis", [str(k) for k in FW["cis"]], lambda k: fw_pair("cis", k)[0],
                 lambda k: fw_pair("cis", k)[1]),
                ("NIST SP 800-53 Rev.5", "nist", used_nist, lambda k: NIST.get(k, ""), lambda k: ""),
                ("ISMS-P 인증기준", "ismsp", [str(k) for k in FW["ismsp"]], lambda k: "", lambda k: fw_pair("ismsp", k)[0])]
    for title, fld, keys, en, ko in ctl_sets:
        r = _section(ws, r, f"■ {title}")
        _hdr(ws, r, ["ID", "명칭(영문)", "명칭(국문)", "연계 세부위협 수", "매우 높음·높음 위협 수", "연계 세부위협", "미연계 사유"])
        r += 1
        for k in keys:
            rows = [o for o in ROWS if k in [str(v) for v in o[fld]]]
            hi = sum(1 for o in rows if o["risk"] in ("매우 높음", "높음"))
            why = R.FRAMEWORK_NOT_MAPPED.get((fld, k), "")
            _row(ws, r, [k, en(k), ko(k), len(rows), hi,
                         "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in sorted(rows, key=lambda o: o["priority"])), why])
            r += 1
        r += 1
    _w(ws, [12, 46, 30, 8, 10, 56, 36])

    # ---------------- 다른 매트릭스 연계 ----------------
    ws = wb.create_sheet("다른 매트릭스 연계")
    _title(ws, "다른 매트릭스 항목 ↔ 신원 세부위협 — 같은 위협을 영역별 매트릭스와 오가며 보기",
           "등급은 각 매트릭스 기준(사고 DB·심각도 기준이 달라 직접 비교하지 않음). 원 매트릭스 위험도는 각 매트릭스 산출물 기준")
    _hdr(ws, 4, ["매트릭스", "연계 ID", "원 매트릭스 항목명", "원 매트릭스 위험도", "연계 신원 세부위협", "신원 매트릭스 위험도(최고)"])
    r = 5
    for fld, mx in (("cloud", "클라우드"), ("ai", "AI"), ("ot", "OT"), ("sct", "공급망")):
        keys = sorted({k for o in ROWS for k in o[fld]})
        for k in keys:
            rows = sorted([o for o in ROWS if k in o[fld]], key=lambda o: o["priority"])
            best = min((o["risk"] for o in rows), key=lambda x: rank_key[x])
            _row(ws, r, [mx, k, LINKS.get(k, {}).get("name", ""), LINKS.get(k, {}).get("risk", ""),
                         "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in rows), best])
            _risk(ws.cell(r, 6), best)
            r += 1
    ws.freeze_panes = "C5"
    _w(ws, [10, 14, 36, 10, 60, 12])

    # ---------------- 평가 기준 ----------------
    ws = wb.create_sheet("평가 기준")
    crit = [
        ("■ 위험도", ""),
        ("위험도", "발생가능성 × 심각도 3×3 표(AI·클라우드·OT·공급망 매트릭스와 동일) — 상×상 매우 높음, 상×중·중×상 높음, 상×하·중×중·하×상 보통, 나머지 낮음"),
        ("발생가능성", TAX["assessment"]["likelihood_rule"]),
        ("심각도 상", TAX["assessment"]["severity_rule"]["상"]),
        ("심각도 중", TAX["assessment"]["severity_rule"]["중"]),
        ("심각도 하", TAX["assessment"]["severity_rule"]["하"]),
        ("심각도 점검", TAX["assessment"]["severity_check"]),
        ("우선순위", "위험도 → 실제 사고 수 → 최근 사고(2025~) → IDT-ID 순(1 = 최우선)"),
        ("근거 수준", "실제 사고 확인(사고 DB 실제 사고) > 실사용 기법 포함(ATT&CK 신원 맥락 사례·KEV(신원)·위협인텔) > "
                    "실증·공개 취약점(사고 DB 연구·시연·SAT 시연) > 이론·시나리오"),
        ("", ""), ("■ 근거 집계", ""),
        ("실제 사고", f"사례유형 ∈ {{{'·'.join(sorted(R.REAL_KINDS))}}}이고 검증 수준 ∉ {{{'·'.join(sorted(R.EXCLUDED_VERIFICATION))}}}. "
                    "연구자의 실증·노출 발견은 '연구·시연'(실증 근거)으로 분리"),
        ("최근 사고", f"사고 인지·공개 시점 {R.RECENT_FROM} 이후"),
        ("사고 매핑", "분석가가 사고 경위를 읽고 단계별 세부위협에 매핑(한 사고가 헬프데스크 사칭·IdP 장악·랜섬웨어 등 여러 행에 매핑될 수 있음). "
                    "ATT&CK 그룹·소프트웨어·캠페인 ID는 대조·표시용이며 자동 매핑에 쓰지 않음. 다른 매트릭스 사고 DB와 같은 사고는 origin으로 표시"),
        ("ATT&CK 사례 수", "행의 ATT&CK 기법을 쓰는 주체(그룹·소프트웨어·캠페인) 수. 신원 기법(무차별 대입·MFA 우회·인증 변조·Kerberos 티켓·"
                         "대체 인증 자료·유효 계정·계정 조작 등)은 절차 전체를, 범용 기법(공개 앱 악용·피싱·신뢰 관계·영향 등)은 절차 설명에 "
                         "신원 단서(자격증명·계정·MFA·토큰·세션·SSO·Kerberos·헬프데스크 등)가 있는 것만 셈"),
        ("KEV(신원)", "CISA KEV 중 원격 접속·경계 장비 인증 우회·세션 노출·하드코딩 자격증명·디렉터리/Kerberos/NTLM·서명 검증·IdP 제품·"
                     "자격증명 노출만 판정('KEV 근거' 시트). 실사용 근거로 '중'까지만 반영"),
        ("SAT 시연", "Browser & Identity Attacks Matrix 기법 중 앱별 시연 예시가 있는 것을 '실증' 근거로 집계(실제 사고 아님)"),
        ("", ""), ("■ 다른 매트릭스와의 근거 대응", ""),
        ("AI 매트릭스 v3.2", "실제 사고 ↔ ATLAS Incident·OWASP 인용 사고 / 실사용 ↔ Realized 기법 / 실증 ↔ 실증·CVE"),
        ("클라우드 매트릭스 v5", "실제 사고 ↔ 클라우드 사고 DB 실제 사고 / 실사용 ↔ ATT&CK 클라우드 사례 / 실증 ↔ 연구·노출·벤더 매트릭스"),
        ("OT 매트릭스", "실제 사고 ↔ OT 사고 DB / 실사용 ↔ ATT&CK ICS 절차·KEV(OT) / 실증 ↔ 연구·시연·ICS 권고 CVE"),
        ("공급망 매트릭스", "실제 사고 ↔ 공급망 사고 DB / 실사용 ↔ ATT&CK 공급망 맥락 절차·KEV(공급망) / 실증 ↔ 연구·시연·SAP 연구 문헌"),
        ("신원 매트릭스", "실제 사고 ↔ 신원 사고 DB / 실사용 ↔ ATT&CK 신원 맥락 절차·KEV(신원)·위협인텔 / 실증 ↔ 연구·시연·SAT 시연"),
        ("", ""), ("■ 사례 라벨(문구 작성용)", ""),
        ("[실제 사고]", "사고 DB에서 '실제 사고'로 집계되는 사례"),
        ("[위협인텔]", "사고 DB에서 '위협인텔'로 집계되는 보안업체 위협 보고"),
        ("[공개 취약점]", "CVE·KEV·공식 보안 권고로 확인된 결함"),
        ("[실증]", "연구자·레드팀·보안업체가 실제 제품·서비스에서 재현한 공격(사고 DB 연구·시연)"),
        ("[시연]", "Browser & Identity Attacks Matrix의 앱별 시연 예시(SAT-ID 병기)"),
        ("", ""), ("■ 신원 관점 문구 작성 규칙(data/text/*.yaml)", ""),
        ("핵심 요약", "한 줄 50자 이내, '수단 + 대상 + 결과'를 명사형으로 — 보고서 본문·슬라이드용"),
        ("요약설명", "2줄 개조식 — ① 공격자는 [수단·경로]로 [행위]할 수 있음 ② 피해·확산 특성 또는 탐지·방어가 어려운 이유"),
        ("참조", "■ 신원 관점(대상·경로·수법·변형·연계 IDT-ID·다른 매트릭스 ID) / ■ 실제 사례('[라벨] 사례명(YYYY-MM): 경위·결과 "
               "(사고 ID·CVE·ATT&CK ID)', 공개 사례가 없으면 '공개 사고 미확인 — 사유')"),
        ("탐지·대응", "■ 탐지 / ■ 대응(예방 → 차단 → 탐지 순, 운영 제약 반영, 매핑된 대응 기준 ID 병기)"),
        ("사실 근거", "수치·사례는 사고 DB 요약·출처에 있는 것만 씀. 용어는 자격증명(인증정보·크리덴셜 쓰지 않음)·반출(행위)/유출(결과)"),
        ("자동 검증", "scripts/validate.py — 인용 사고가 그 행에 매핑됐는지, 라벨이 집계 상태와 맞는지, 시점·CVE·ATT&CK ID가 사고 DB와 "
                    "같은지, 설명 속 수치가 사고 요약에 있는지, IDT·SAT·NHI·CCM·CIS·NIST·ISMS-P·다른 매트릭스 ID가 존재하고 행 매핑과 "
                    "맞는지, '실제 사고 확인' 행이 [실제 사고] 사례를 인용하는지 대조(빌드 전에 자동 실행)"),
    ]
    for i, (a, b) in enumerate(crit, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    _w(ws, [20, 140])

    # ---------------- 변경이력 ----------------
    ws = wb.create_sheet("변경이력")
    _hdr(ws, 1, ["버전", "날짜", "구분", "대상", "변경 내용", "사유"])
    for i, ent in enumerate(CHANGELOG.get("entries", []), 2):
        _row(ws, i, ent)
    _w(ws, [6, 11, 10, 22, 80, 50])

    for s in wb.worksheets:
        s.sheet_view.zoomScale = 90
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToWidth = 1
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.fitToHeight = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path):
    cols = ["도메인", "IDT-ID", "위협분류", "세부위협", "핵심 요약", "요약설명", "공격 단계", "신원 유형", "ID 환경", "ATT&CK", "SAT",
            "OWASP NHI", "NIST SP 800-63-4", "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수", "최근 사고",
            "ATT&CK 사례 수", "KEV", "관련 사례 ID", "클라우드 연계", "AI 연계", "OT 연계", "공급망 연계"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for o in ROWS:
            w.writerow([o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], o["summary"].replace("\n", " "),
                        ", ".join(o["stage"]), " · ".join(o["id_types"]), " · ".join(o["id_envs"]), ", ".join(o["attack"]),
                        ", ".join(o["sat"]), ", ".join(o["nhi"]), ", ".join(o["sp80063"]), o["likelihood"], o["severity"],
                        o["risk"], o["priority"], o["level"], o["n_real"], o["n_recent"], o["n_atk"], o["n_kev"],
                        ", ".join(sorted(EV[o["id"]]["real"])), ", ".join(o["cloud"]), ", ".join(o["ai"]), ", ".join(o["ot"]),
                        ", ".join(o["sct"])])


def write_worksheet(outdir):
    """문구 작성용 근거 정리본 — 도메인별 마크다운(행별 매핑 사고·KEV·ATT&CK 주체·SAT)"""
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    for d in DOMAINS:
        lines = [f"# [{d['code']}] {d['ko']} — 문구 작성용 근거 정리본", ""]
        for o in [o for o in ROWS if o["code"] == d["code"]]:
            x = EV[o["id"]]
            lines += [f"## {o['id']} {o['name']} ({o['en']})", "",
                      f"- 위험: {o['likelihood']}×{o['severity']}={o['risk']} · {o['lk_why']}",
                      f"- ATT&CK: {', '.join(atk_label(t) for t in o['attack'])}",
                      "- SAT: " + ", ".join(f"{sat_label(s)}(예시 {SAT[s]['n_examples']})" for s in o["sat"]),
                      f"- 연계: 클라우드 {o['cloud']} · AI {o['ai']} · OT {o['ot']} · 공급망 {o['sct']}",
                      f"- 대응: CCM {o['ccm']} · CIS {o['cis']} · NIST {o['nist']} · ISMS-P {o['ismsp']}", "", "### 사고 DB"]
            for i in sorted(x["real"] | x["research"] | x["intel"], key=lambda i: INC[i]["date"]):
                e = INC[i]
                lines.append(f"- {i} [{e['status']}] {e['title']}({e['date']}) — {' '.join(e['summary'].split())}"
                             + (f" KEV {', '.join(e['kev'])}" if e.get("kev") else "")
                             + (f" ATT&CK {', '.join(e['attack_ref'])}" if e.get("attack_ref") else ""))
            if x["kev"]:
                lines += ["", "### KEV(신원)"] + [f"- {c} {KEV[c]['vendorProject']} {KEV[c]['product']} — "
                                                   f"{R.KEV_IDENTITY[c][2]}" for c in sorted(x["kev"])]
            if x["subj"]:
                lines += ["", "### ATT&CK 주체(신원 맥락)", ", ".join(f"{s} {ATK['subjects'][s]['name']}" for s in sorted(x["subj"]))]
            lines.append("")
        (outdir / f"{d['code']}.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"근거 정리본 → {outdir}")


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--worksheet", help="문구 작성용 근거 정리본(도메인별 .md)을 만들 폴더")
    ap.add_argument("--skip-validate", action="store_true", help="scripts/validate.py 검증을 건너뜀")
    ap.add_argument("--version", help="출력 파일 버전 표기(기본 v2)")
    ap.add_argument("--allow-missing-text", action="store_true",
                    help="문구가 없는 세부위협을 오류 대신 '문구 미작성'으로 두고 빌드(작성 중간 산출물용)")
    args = ap.parse_args()
    if args.version:
        VERSION = args.version
        OUT = ROOT / "output" / f"통합_신원보안위협_매트릭스_{VERSION}.xlsx"
        CSV_OUT = ROOT / "output" / f"통합_신원보안위협_매트릭스_{VERSION}.csv"
    if not args.skip_validate:
        VALIDATION = run_validate(args.allow_missing_text)
        print(VALIDATION)
    if args.worksheet:
        write_worksheet(args.worksheet)
    write_xlsx(OUT)
    write_csv(CSV_OUT)
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    print(f"세부위협 {len(ROWS)}개 · 사고 {len(INCIDENTS)}건 → {OUT.name}")
    print("위험도", dict(rc), "| 근거 수준", dict(lc))
