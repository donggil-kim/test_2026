# -*- coding: utf-8 -*-
"""
고도화 엔터프라이즈 침해 보안위협 매트릭스 빌더
------------------------------------------------
분류 : ATT&CK Enterprise 침투~영향 12개 전술 → 위협분류(Lv2) → 세부위협(Lv3) — data/taxonomy.yaml
고유 : 레거시 방어 공백 유형 · 요구 방어 수준(레거시/EDR·XDR/행위·네트워크·아이덴티티)
교차 : ATT&CK v19.2(1차·주체 절차 수=빈발 신호) · CISA KEV · LOLBAS·LOLDrivers · 통합 v4.1 연계 ID
근거 : 사고 DB + ATT&CK 절차 사례 수 + KEV 악용 · 대응 NIST SP 800-53 Rev.5
논리 : 발생가능성 상 = 실제 사고 2건 이상. 선정은 고도화·최근빈발·레거시난제로 가리고, 우선순위에 레거시 공백·방어 난이도 보조(D5).
실행 : python3 scripts/build_enterprise_threat_matrix.py [--version v1] [--skip-validate]
"""
import argparse
import collections
import csv
import datetime
import json
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import taxonomy_rules as R  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
REF = ROOT / "reference"
NAME = "통합_고도화엔터프라이즈_보안위협_매트릭스"


def ld(p, default=None):
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else default


TAX = ld(DATA / "taxonomy.yaml")
FW = ld(DATA / "frameworks.yaml")
INCIDENTS = ld(DATA / "incidents.yaml", []) or []
CHANGELOG = ld(DATA / "changelog.yaml", {}) or {}
SCENARIOS = ld(DATA / "scenarios.yaml", {}) or {}
ATK = json.loads((REF / "attack" / "enterprise-attack-v19.2-subset.json").read_text(encoding="utf-8"))
KEV = json.loads((REF / "kev" / "kev-subset.json").read_text(encoding="utf-8"))
NIST = json.loads((REF / "nist" / "sp800-53r5-names.json").read_text(encoding="utf-8"))["names"]
LINKS = json.loads((REF / "links" / "other_matrices.json").read_text(encoding="utf-8"))["ids"]
REFS = FW["references"]

TEXT = {}
for f in sorted((DATA / "text").glob("*.yaml")) if (DATA / "text").exists() else []:
    TEXT.update(ld(f, {}) or {})

DOMAINS, ROWS = [], []
for d in TAX["domains"]:
    DOMAINS.append(d)
    for l2 in d["lv2"]:
        for x in l2["lv3"]:
            ROWS.append(dict(x, code=d["code"], domain=f"[{d['code']}] {d['ko']}", lv2_id=l2["id"], lv2=l2["name"]))
ROW = {o["id"]: o for o in ROWS}
by_c = collections.defaultdict(list)
for o in ROWS:
    by_c[o["code"]].append(o)

for e in INCIDENTS:
    e["date"] = str(e["date"])
    e["status"] = R.incident_status(e)
    e["recent"] = e["date"] >= R.RECENT_FROM
INC = {e["id"]: e for e in INCIDENTS}

DEF_RANK = {"행위·네트워크·아이덴티티": 3, "EDR·XDR": 2, "레거시": 1}

# 근거 집계
EV = {o["id"]: dict(real=set(), recent=set(), advisory=set(), intel=set(), research=set(), excluded=set(),
                    impact=collections.Counter(), actor=collections.Counter()) for o in ROWS}
for e in INCIDENTS:
    for k in e.get("ent") or []:
        if k not in EV:
            continue
        x = EV[k]
        st = e["status"]
        if st == "실제 사고":
            x["real"].add(e["id"])
            if e["recent"]:
                x["recent"].add(e["id"])
            x["impact"].update(i for i in e["impact"] if i not in ("영향 없음(사전 차단)", "영향 미상"))
            x["actor"].update(e["actor_type"])
        elif st == "정부 경보":
            x["advisory"].add(e["id"])
        elif st == "위협인텔":
            x["intel"].add(e["id"])
        elif st == "실증·연구":
            x["research"].add(e["id"])
        else:
            x["excluded"].add(e["id"])

for o in ROWS:
    subj = set()
    for t in o["attack"]:
        subj |= {p["subject"] for p in ATK["procedures"].get(t, [])}
    o["n_attack"] = len(subj)         # ATT&CK 주체(그룹·소프트웨어·캠페인) 수 — 빈발 신호
    o["n_edge_cve"] = 0               # 참고: KEV 플래그는 kev 필드


def likelihood(o, x):
    if len(x["real"]) >= R.LIKELY_HIGH_REAL:
        return "상"
    if x["real"] or x["advisory"] or x["intel"] or x["research"] or o.get("kev") or o["n_attack"] >= 1:
        return "중"
    return "하"


def evidence_level(o, x):
    if x["real"]:
        return "실제 사고 확인"
    if o["n_attack"] >= 10 or o.get("kev"):
        return "ATT&CK 사례·KEV 다수"
    if o["n_attack"] >= 1 or x["intel"] or x["advisory"]:
        return "ATT&CK 사례·KEV 일부"
    return "이론·연구"


for o in ROWS:
    x = EV[o["id"]]
    o["likelihood"] = likelihood(o, x)
    o["risk"] = R.RISK_MATRIX[(o["likelihood"], o["severity"])]
    o["level"] = evidence_level(o, x)
    o["n_real"], o["n_recent"] = len(x["real"]), len(x["recent"])
    o["n_adv"], o["n_intel"], o["n_research"] = len(x["advisory"]), len(x["intel"]), len(x["research"])
    o["lk_why"] = (f"실제 사고 {o['n_real']}건(최근 {o['n_recent']}) · ATT&CK 주체 {o['n_attack']} · "
                   f"KEV {'O' if o.get('kev') else '-'} · 정부 경보 {o['n_adv']} · 위협인텔 {o['n_intel']} · 실증 {o['n_research']}")
    txt = TEXT.get(o["id"]) or {}
    o["oneline"], o["summary"], o["reference"], o["detect"] = (txt.get(k) or "" for k in
                                                              ("oneline", "summary", "reference", "detect"))
    o["text_src"] = "분석가 작성" if txt else "문구 미작성"

rank_key = {r: i for i, r in enumerate(R.RISK_ORDER)}
for i, o in enumerate(sorted(ROWS, key=lambda o: (rank_key[o["risk"]], -o["n_real"], -DEF_RANK[o["defense_level"]],
                                                  -o["n_attack"], o["id"])), 1):
    o["priority"] = i


def nist_label(k):
    return f"{k} {NIST.get(k, '')}"


def atk_label(t):
    return f"{t} {ATK['techniques'][t]['name']}" if t in ATK["techniques"] else t


def link_label(k):
    return f"{k} {LINKS[k]['name']}" if k in LINKS else k


def ref_label(k):
    return REFS[k]["short"] if k in REFS else k


def inc_label(i):
    e = INC[i]
    return f"{i} {e['title']}({e['date']})"


def first_case(o):
    real = sorted(EV[o["id"]]["real"], key=lambda i: (INC[i]["date"], i), reverse=True)
    if real:
        e = INC[real[0]]
        return f"[실제 사고] {e['title']}({e['date']}) ({real[0]})"
    other = sorted(EV[o["id"]]["advisory"] | EV[o["id"]]["intel"] | EV[o["id"]]["research"],
                   key=lambda i: (INC[i]["date"], i), reverse=True)
    if other:
        e = INC[other[0]]
        lab = {"정부 경보": "[정부 경보]", "위협인텔": "[위협인텔]", "실증·연구": "[실증]"}[e["status"]]
        return f"{lab} {e['title']}({e['date']}) ({other[0]})"
    if o["n_attack"]:
        return f"[ATT&CK 주체 {o['n_attack']}]"
    return "공개 사고 미확인"


def first_control(o):
    b = [ln for ln in (o["detect"] or "").splitlines() if ln.startswith("- ")]
    if b:
        return b[0][2:].strip()
    return "대응 기준: NIST " + "·".join(o["nist"]) if o["nist"] else ""


def run_validate():
    import subprocess
    r = subprocess.run([sys.executable, str(Path(__file__).resolve().parent / "validate.py"), "--strict"],
                       capture_output=True, text=True, encoding="utf-8")
    lines = [ln for ln in r.stdout.splitlines() if ln.strip()]
    if r.returncode != 0:
        print("\n".join(lines))
        raise SystemExit("검증 오류·경고 — 고친 뒤 다시 빌드(--skip-validate로 무시 가능)")
    return next((ln for ln in reversed(lines) if ln.startswith("세부위협")), "")


VALIDATION = "검증 생략(--skip-validate)"

# 서식
BORDER = Border(*(Side(style="thin", color="D0D0D0"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "ATT&CK 사례·KEV 다수": "2E86C1", "ATT&CK 사례·KEV 일부": "B9770E", "이론·연구": "7F8C8D"}
HDR, NAVY, GRAY = "34495E", "1F3864", "8C8C8C"


def _w(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hdr(ws, row, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c)
        cell.fill = PatternFill("solid", fgColor=HDR)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BORDER


def _risk(cell, risk):
    bg, fg = RISK_COLOR[risk]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=risk in ("매우 높음", "높음"), color=fg)
    cell.alignment = CENTER


def _level(cell, level):
    cell.font = Font(bold=True, color=LEVEL_COLOR[level])
    cell.alignment = CENTER


def _row(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.alignment = WRAP
        c.border = BORDER


def _kv(ws, rows, widths):
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    _w(ws, list(widths))


def summ(o):
    return o["summary"] or f"(문구 미작성) {o['en']}"


def build(version):
    wb = openpyxl.Workbook()
    today = datetime.date.today().isoformat()
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    sc = collections.Counter(o["severity"] for o in ROWS)
    st = collections.Counter(e["status"] for e in INCIDENTS)
    gapc = collections.Counter(g for o in ROWS for g in o["legacy_gap"])
    defc = collections.Counter(o["defense_level"] for o in ROWS)
    years = sorted({e["date"][:4] for e in INCIDENTS})
    written = sum(1 for o in ROWS if o["text_src"] == "분석가 작성")
    n_lv2 = sum(len(d["lv2"]) for d in DOMAINS)
    top = sorted(ROWS, key=lambda o: o["priority"])[:10]
    scn = SCENARIOS.get("scenarios") or []
    kev_rows = [o for o in ROWS if o.get("kev")]
    lolbas_rows = [o for o in ROWS if o.get("lolbas")]
    drv_rows = [o for o in ROWS if o.get("loldriver")]

    # 개요
    ws = wb.active
    ws.title = "개요"
    meta = [
        (f"고도화·레거시 우회 엔터프라이즈 침해 보안위협 매트릭스 {version}", ""),
        ("관점", "ATT&CK Enterprise에서 고도화·최근빈발·레거시 보안체계(시그니처 AV·경계 방화벽·레거시 SIEM)로 막기 어려운 기법을 "
                "선별해 '레거시 방어 공백'·'요구 방어 수준'을 축으로 붙인 렌즈 매트릭스. 자산이 무엇이든 '레거시로 막기 어렵다'로 묶음"),
        ("생성일", today), ("", ""),
        ("■ 기준 데이터", ""),
        ("사고 DB", f"{len(INCIDENTS)}건({years[0]}~{years[-1]}) — 실제 사고 {st['실제 사고']} · 정부 경보 {st.get('정부 경보',0)} · "
                  f"위협인텔 {st.get('위협인텔',0)} · 실증·연구 {st.get('실증·연구',0)}. 실제 사고가 매핑된 세부위협 "
                  f"{sum(1 for o in ROWS if o['n_real'])}/{len(ROWS)}"),
        ("MITRE ATT&CK", f"Enterprise v19.2 — 기법 {len(ATK['techniques'])}개 연계, 주체(그룹·소프트웨어·캠페인) 절차 "
                         f"{sum(len(v) for v in ATK['procedures'].values())}건을 '최근빈발' 신호로 집계"),
        ("CISA KEV", f"{KEV['count']}건({KEV['source']}) — 연도별 등재 {KEV['by_year']}, 경계·엣지 장비 비중 {KEV['edge_share']}. "
                     f"KEV 악용 근거 행 {len(kev_rows)}개"),
        ("LOLBAS·LOLDrivers", f"LOLBin {len(lolbas_rows)}개 행·취약 드라이버(BYOVD) {len(drv_rows)}개 행에 '서명된 정상 도구·커널 하위 계층' 공백 연계"),
        ("대응 기준", "NIST SP 800-53 Rev.5 통제(탐지·대응). 참고 문헌 " + str(len(REFS)) + "종"),
        ("다른 매트릭스", "통합 v4.1 요약 ID(AI·CL·OT·SC·ID·PH) 연계 — 렌즈라 겹침이 많아 재수록 대신 연계만(D9)"),
        ("", ""),
        ("■ 분류 체계(D6)", ""),
        ("Lv1 전술(도메인)", f"ATT&CK Enterprise 침투~영향 12개 — " + " → ".join(f"[{d['code']}]{d['ko']}" for d in DOMAINS)),
        ("Lv2·Lv3", f"위협분류 {n_lv2} · 세부위협 {len(ROWS)}(ENT-<전술>-NN). 정찰·자원 개발은 영역 1·정찰이 다룸(D7)"),
        ("고유 축", "레거시 방어 공백 유형(8종) · 요구 방어 수준(레거시 / EDR·XDR / 행위·네트워크·아이덴티티)"),
        ("", ""),
        ("■ 시트", ""),
        ("보고서용 간략 매트릭스", "우선순위 순 세부위협·핵심 요약·위험도·요구 방어 수준·대표 사례 1건·핵심 대응 1줄"),
        ("매트릭스 뷰", "전술(열)별 세부위협을 위험도 색으로 배치"),
        ("통합 매트릭스", "분류·고유 축·교차 매핑·위험 평가·실제 근거·대응 기준 전체 열"),
        ("통합매트릭스_LITE", "핵심 열 발췌"),
        ("전술 요약", "전술별 경계·위협 수·위험도/근거 분포·대표 위협"),
        ("레거시 공백·방어 수준 요약", "고유 축(레거시 공백 유형 × 요구 방어 수준)별 세부위협·위험도"),
        ("공격 체인 시나리오", f"실제 사고 기반 공격 흐름 {len(scn)}개" if scn else "시트 양식만 — ③단계에서 작성"),
        ("역매핑_사고사례", f"사고 DB {len(INCIDENTS)}건과 매핑 세부위협·CVE·출처"),
        ("ATT&CK·KEV 교차", "기법별 연계 행·주체 수, KEV 악용·LOLBin·BYOVD 드라이버 표시"),
        ("대응 기준 연계", "NIST 통제별 연계 세부위협·위험도 분포"),
        ("참고 문헌", "ATT&CK·KEV·LOLBAS·벤더 보고 카탈로그"),
        ("다른 매트릭스 연계", "AI·클라우드·OT·공급망·신원·물리 매트릭스 항목별 연계 세부위협"),
        ("평가 기준", "발생가능성·심각도·위험도·근거수준·선정 기준·경계·가드레일"),
        ("변경이력", "버전별 변경 내역"), ("", ""),
        ("■ 결과 요약", ""),
        ("세부위협", f"{len(ROWS)}개 / 위협분류 {n_lv2} / 전술 {len(DOMAINS)}"),
        ("위험도", " · ".join(f"{k} {rc.get(k,0)}" for k in R.RISK_ORDER)),
        ("근거 수준", " · ".join(f"{k} {lc.get(k,0)}" for k in R.LEVELS)),
        ("심각도", " · ".join(f"{k} {sc.get(k,0)}" for k in ("상", "중", "하"))),
        ("요구 방어 수준", " · ".join(f"{k} {defc.get(k,0)}" for k in R.DEFENSE_LEVELS)),
        ("레거시 공백 유형(행 수)", " · ".join(f"{k} {gapc.get(k,0)}" for k in R.LEGACY_GAPS)),
        ("관점 문구", f"{written}/{len(ROWS)}개 작성" + (" — 전 항목 완료" if written == len(ROWS) else
                     " — 문구는 ③단계") + f" | validate.py {VALIDATION}"),
        ("우선 위협 Top 10", "\n".join(f"{o['priority']}. {o['id']} {o['name']} — {o['risk']}, 실제 사고 {o['n_real']}건"
                                    f"({o['defense_level']})" for o in top)),
        ("", ""), ("■ 주의", ""),
        ("발생가능성", "상 = 실제 사고 2건 이상. 선정은 고도화·최근빈발·레거시난제 셋 중 둘 이상으로 가리고, 우선순위에 요구 방어 수준(방어 난이도)을 보조(D5)"),
        ("중복·연계", "렌즈 매트릭스라 기존 6개와 겹침 — 재수록 대신 통합 v4.1 요약 ID로 연계(D9). 위험도·사고 수는 같은 매트릭스 안에서만 비교"),
        ("가드레일(D11)", "방어 자료 — 위험·탐지·대응 중심. 공세 제작법·악용 절차는 적지 않고 식별·탐지·완화 수준으로 서술"),
        ("출처 표기", "MITRE ATT&CK® © The MITRE Corporation · CISA KEV(공공 영역) · LOLBAS·LOLDrivers(MIT) · NIST SP 800-53(공공 영역)"),
    ]
    _kv(ws, meta, (28, 140))
    ws.cell(1, 1).font = Font(size=15, bold=True, color=NAVY)

    # 보고서용 간략 매트릭스
    ws = wb.create_sheet("보고서용 간략 매트릭스")
    cols = ["순위", "전술(Lv1)", "ENT-ID", "세부위협(Lv3)", "핵심 요약", "위험도", "발생가능성", "심각도", "요구 방어 수준",
            "레거시 공백", "실제 사고", "대표 사례", "핵심 대응"]
    ws.cell(1, 1, f"보고서용 간략 매트릭스 {version}").font = Font(size=13, bold=True, color=NAVY)
    ws.cell(2, 1, f"순위 = 위험도 → 실제 사고 → 방어 난이도 → ATT&CK 주체 수 | {today}").font = Font(size=9, color="595959")
    _hdr(ws, 4, cols)
    for r, o in enumerate(sorted(ROWS, key=lambda o: o["priority"]), 5):
        _row(ws, r, [o["priority"], o["domain"], o["id"], o["name"], o["oneline"] or summ(o), o["risk"], o["likelihood"],
                     o["severity"], o["defense_level"], " · ".join(o["legacy_gap"]), o["n_real"], first_case(o),
                     first_control(o)])
        ws.cell(r, 3).font = Font(bold=True)
        ws.cell(r, 4).font = Font(bold=True)
        if not o["oneline"]:
            ws.cell(r, 5).font = Font(color=GRAY)
        _risk(ws.cell(r, 6), o["risk"])
        for j in (1, 7, 8, 11):
            ws.cell(r, j).alignment = CENTER
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(ROWS)+4}"
    _w(ws, [6, 15, 13, 28, 42, 9, 7, 7, 14, 24, 7, 44, 48])

    # 매트릭스 뷰
    ws = wb.create_sheet("매트릭스 뷰")
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도, [n]=실제 사고 수, 괄호=세부위협 수. 열 = ATT&CK 전술").font = Font(bold=True, color=NAVY)
    for j, d in enumerate(DOMAINS, 1):
        c = ws.cell(2, j, f"[{d['code']}] {d['ko']}\n{d['en']}\n({len(by_c[d['code']])})")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        for i, o in enumerate(sorted(by_c[d["code"]], key=lambda o: (rank_key[o["risk"]], -o["n_real"], o["id"])), 3):
            cell = ws.cell(i, j, f"{o['id']} {o['name']}" + (f" [{o['n_real']}]" if o["n_real"] else ""))
            _risk(cell, o["risk"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9, bold=o["risk"] == "매우 높음", color=RISK_COLOR[o["risk"]][1])
            cell.border = BORDER
    ws.row_dimensions[2].height = 52
    ws.freeze_panes = "A3"
    _w(ws, [22] * len(DOMAINS))

    # 통합 매트릭스
    ws = wb.create_sheet("통합 매트릭스")
    g1 = ["전술(Lv1)", "ENT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "영문명", "설명 출처"]
    g2 = ["레거시 공백 유형", "요구 방어 수준", "위협 주체", "플랫폼"]
    g3 = ["ATT&CK ID", "ATT&CK 기법명", "ATT&CK 주체 수", "KEV", "LOLBin", "BYOVD 드라이버",
          "AI 연계", "클라우드 연계", "OT 연계", "공급망 연계", "신원 연계", "물리 연계"]
    g4 = ["발생가능성", "심각도", "위험도", "우선순위", "발생가능성 근거(자동)", "심각도 근거"]
    g5 = ["실제 사고 수", "최근(2025~)", "정부 경보", "위협인텔", "근거 수준", "관련 사례 ID", "주요 영향(실제 사고)"]
    g6 = ["NIST SP 800-53", "탐지·대응 포인트", "참고 문헌"]
    groups = [("분류 체계", len(g1), "2E5496"), ("고유 축", len(g2), "117A65"), ("교차 매핑", len(g3), "1F7A8C"),
              ("위험 평가", len(g4), "A04000"), ("실제 근거", len(g5), "1E8449"), ("대응 기준·참고 문헌", len(g6), "6C3483")]
    cols = g1 + g2 + g3 + g4 + g5 + g6
    t = ws.cell(1, 1, f"고도화 엔터프라이즈 침해 매트릭스 {version}")
    t.font = Font(size=13, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, f"사고 DB {len(INCIDENTS)}건 + ATT&CK v19.2 + CISA KEV + LOLBAS·LOLDrivers | {today}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    c0 = 1
    for g, span, color in groups:
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + span - 1)
        c = ws.cell(3, c0, g)
        c.fill, c.font, c.alignment = PatternFill("solid", fgColor=color), Font(bold=True, color="FFFFFF"), CENTER
        c0 += span
    _hdr(ws, 4, cols)
    ci = {c: i + 1 for i, c in enumerate(cols)}
    for r, o in enumerate(ROWS, 5):
        x = EV[o["id"]]
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summ(o), o["en"], o["text_src"],
                     " · ".join(o["legacy_gap"]), o["defense_level"], " · ".join(o["actors"]), " · ".join(o["platforms"]),
                     ", ".join(o["attack"]), "\n".join(atk_label(t) for t in o["attack"]), o["n_attack"],
                     "O" if o.get("kev") else "", "O" if o.get("lolbas") else "", "O" if o.get("loldriver") else "",
                     "\n".join(link_label(k) for k in o["ai"]), "\n".join(link_label(k) for k in o["cl"]),
                     "\n".join(link_label(k) for k in o["ot"]), "\n".join(link_label(k) for k in o["sc"]),
                     "\n".join(link_label(k) for k in o["idn"]), "\n".join(link_label(k) for k in o["ph"]),
                     o["likelihood"], o["severity"], o["risk"], o["priority"], o["lk_why"], o["severity_why"],
                     o["n_real"], o["n_recent"], o["n_adv"], o["n_intel"], o["level"],
                     ", ".join(sorted(x["real"]))
                     + (("\n정부 경보: " + ", ".join(sorted(x["advisory"]))) if x["advisory"] else "")
                     + (("\n위협인텔: " + ", ".join(sorted(x["intel"]))) if x["intel"] else "")
                     + (("\n실증: " + ", ".join(sorted(x["research"]))) if x["research"] else ""),
                     " · ".join(f"{k} {v}" for k, v in x["impact"].most_common(4)),
                     "\n".join(nist_label(k) for k in o["nist"]), o["detect"], "\n".join(ref_label(k) for k in o["refs"])])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, ci["요약설명"]).font = Font(color=GRAY)
        for c in ("요구 방어 수준", "ATT&CK 주체 수", "KEV", "LOLBin", "BYOVD 드라이버", "발생가능성", "심각도", "우선순위"):
            ws.cell(r, ci[c]).alignment = CENTER
        _risk(ws.cell(r, ci["위험도"]), o["risk"])
        _level(ws.cell(r, ci["근거 수준"]), o["level"])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(ROWS)+4}"
    widths = {"전술(Lv1)": 14, "ENT-ID": 13, "위협분류(Lv2)": 22, "세부위협(Lv3)": 28, "핵심 요약": 34, "요약설명": 56,
              "영문명": 28, "설명 출처": 9, "레거시 공백 유형": 26, "요구 방어 수준": 14, "위협 주체": 18, "플랫폼": 16,
              "ATT&CK ID": 16, "ATT&CK 기법명": 30, "ATT&CK 주체 수": 8, "KEV": 5, "LOLBin": 6, "BYOVD 드라이버": 8,
              "AI 연계": 18, "클라우드 연계": 20, "OT 연계": 18, "공급망 연계": 18, "신원 연계": 18, "물리 연계": 14,
              "발생가능성": 7, "심각도": 7, "위험도": 8, "우선순위": 7, "발생가능성 근거(자동)": 40, "심각도 근거": 46,
              "실제 사고 수": 7, "최근(2025~)": 7, "정부 경보": 7, "위협인텔": 7, "근거 수준": 14, "관련 사례 ID": 26,
              "주요 영향(실제 사고)": 24, "NIST SP 800-53": 26, "탐지·대응 포인트": 60, "참고 문헌": 22}
    _w(ws, [widths.get(c, 7) for c in cols])

    # LITE
    ws = wb.create_sheet("통합매트릭스_LITE")
    ws.cell(1, 1, f"고도화 엔터프라이즈 침해 매트릭스 {version} — LITE").font = Font(size=13, bold=True, color=NAVY)
    lcols = ["전술(Lv1)", "ENT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "레거시 공백", "요구 방어 수준",
             "플랫폼", "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고", "ATT&CK 주체", "KEV", "탐지·대응"]
    lc_ = {c: i + 1 for i, c in enumerate(lcols)}
    _hdr(ws, 2, lcols)
    for r, o in enumerate(ROWS, 3):
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summ(o), " · ".join(o["legacy_gap"]),
                     o["defense_level"], " · ".join(o["platforms"]), o["likelihood"], o["severity"], o["risk"],
                     o["priority"], o["level"], o["n_real"], o["n_attack"], "O" if o.get("kev") else "", o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, lc_["요약설명"]).font = Font(color=GRAY)
        _risk(ws.cell(r, lc_["위험도"]), o["risk"])
        _level(ws.cell(r, lc_["근거 수준"]), o["level"])
        for c in ("요구 방어 수준", "발생가능성", "심각도", "우선순위", "실제 사고", "ATT&CK 주체", "KEV"):
            ws.cell(r, lc_[c]).alignment = CENTER
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(lcols))}{len(ROWS)+2}"
    _w(ws, [14, 13, 22, 28, 34, 56, 26, 14, 16, 7, 7, 8, 7, 14, 7, 8, 5, 60])

    # 전술 요약
    ws = wb.create_sheet("전술 요약")
    dcols = ["전술(Lv1)", "설명", "넣지 않는 것(경계)", "위협분류", "세부위협"] + R.RISK_ORDER + R.LEVELS + \
            ["실제 사고(중복 제거)", "최고위험 세부위협(상위 3)"]
    _hdr(ws, 1, dcols)
    for i, d in enumerate(DOMAINS, 2):
        os_ = by_c[d["code"]]
        rr = collections.Counter(o["risk"] for o in os_)
        ll = collections.Counter(o["level"] for o in os_)
        incs = set().union(*(EV[o["id"]]["real"] for o in os_)) if os_ else set()
        top3 = sorted(os_, key=lambda o: o["priority"])[:3]
        _row(ws, i, [f"[{d['code']}] {d['ko']}", d["desc"], d["excludes"], len(d["lv2"]), len(os_)]
             + [rr.get(k, 0) for k in R.RISK_ORDER] + [ll.get(k, 0) for k in R.LEVELS]
             + [len(incs), "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in top3)])
    ws.freeze_panes = "B2"
    _w(ws, [16, 48, 40, 7, 7, 7, 7, 7, 7, 11, 13, 13, 9, 9, 52])

    # 레거시 공백·방어 수준 요약(고유 축)
    ws = wb.create_sheet("레거시 공백·방어 수준 요약")
    ws.cell(1, 1, "레거시 방어 공백 유형 × 요구 방어 수준 — 이 매트릭스의 차별적 축(D6)").font = Font(size=13, bold=True, color=NAVY)
    _hdr(ws, 3, ["레거시 공백 유형", "세부위협 수", "매우 높음·높음", "요구 방어 수준 분포", "대표 세부위협(위험 상위 3)"])
    r = 4
    for g in R.LEGACY_GAPS:
        os_ = [o for o in ROWS if g in o["legacy_gap"]]
        dl = collections.Counter(o["defense_level"] for o in os_)
        top3 = sorted(os_, key=lambda o: o["priority"])[:3]
        _row(ws, r, [g, len(os_), sum(1 for o in os_ if o["risk"] in ("매우 높음", "높음")),
                     " · ".join(f"{k} {dl[k]}" for k in R.DEFENSE_LEVELS if dl[k]),
                     "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in top3)])
        ws.cell(r, 2).alignment = CENTER
        ws.cell(r, 3).alignment = CENTER
        r += 1
    r += 1
    ws.cell(r, 1, "요구 방어 수준별").font = Font(bold=True, color=NAVY)
    r += 1
    _hdr(ws, r, ["요구 방어 수준", "세부위협 수", "매우 높음·높음", "대표 세부위협(위험 상위 5)"])
    r += 1
    for dl in R.DEFENSE_LEVELS:
        os_ = [o for o in ROWS if o["defense_level"] == dl]
        top5 = sorted(os_, key=lambda o: o["priority"])[:5]
        _row(ws, r, [dl, len(os_), sum(1 for o in os_ if o["risk"] in ("매우 높음", "높음")),
                     "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in top5)])
        ws.cell(r, 2).alignment = CENTER
        ws.cell(r, 3).alignment = CENTER
        r += 1
    _w(ws, [26, 10, 12, 30, 64])

    # 공격 체인 시나리오
    ws = wb.create_sheet("공격 체인 시나리오")
    if scn:
        ws.cell(1, 1, f"공격 체인 시나리오 {len(scn)}개 — 실제 사고 기반 공격 흐름(방어 관점)").font = Font(bold=True, color=NAVY)
        _hdr(ws, 3, ["시나리오", "제목", "전술", "세부위협", "행위(근거 사고)", "탐지·차단 포인트"])
        r = 4
        for s in scn:
            for step in s.get("steps") or []:
                _row(ws, r, [s["id"], s.get("title", ""), step.get("전술", ""), ", ".join(step.get("ent") or []),
                             step.get("행위", ""), step.get("탐지", "")])
                r += 1
        _w(ws, [10, 40, 12, 26, 58, 50])
    else:
        ws.cell(1, 1, "공격 체인 시나리오 — ③단계(data/scenarios.yaml)에서 작성").font = Font(bold=True, color=NAVY)
        _w(ws, [120])

    # 역매핑_사고사례
    ws = wb.create_sheet("역매핑_사고사례")
    _hdr(ws, 1, ["사고 ID", "사례명", "시점", "유형", "집계 상태", "지역", "업종", "행위자", "주체 유형", "영향",
                 "CVE", "매핑 세부위협", "ATT&CK", "출처"])
    for i, e in enumerate(sorted(INCIDENTS, key=lambda e: (e["date"], e["id"]), reverse=True), 2):
        _row(ws, i, [e["id"], e["title"], e["date"], e["kind"], e["status"], e.get("region", ""),
                     " · ".join(e.get("sector") or []), e.get("actor", ""), " · ".join(e.get("actor_type") or []),
                     " · ".join(e.get("impact") or []), ", ".join(e.get("cve") or []), ", ".join(e.get("ent") or []),
                     ", ".join(e.get("attack_ref") or []), "\n".join(str(s) for s in e.get("sources") or [])])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:N{len(INCIDENTS)+1}"
    _w(ws, [9, 42, 9, 10, 12, 9, 18, 24, 16, 22, 18, 26, 14, 58])

    # ATT&CK·KEV 교차
    ws = wb.create_sheet("ATT&CK·KEV 교차")
    _hdr(ws, 1, ["ATT&CK ID", "기법명", "전술", "주체 수", "연계 세부위협 수", "연계 세부위협", "KEV/LOLBin/BYOVD"])
    used = sorted({t for o in ROWS for t in o["attack"]})
    for i, t in enumerate(used, 2):
        rws = [o for o in ROWS if t in o["attack"]]
        flags = []
        if any(o.get("kev") for o in rws):
            flags.append("KEV")
        if any(o.get("lolbas") for o in rws):
            flags.append("LOLBin")
        if any(o.get("loldriver") for o in rws):
            flags.append("BYOVD")
        _row(ws, i, [t, ATK["techniques"][t]["name"], ATK["techniques"][t]["tactics"],
                     len({p["subject"] for p in ATK["procedures"].get(t, [])}), len(rws),
                     ", ".join(o["id"] for o in rws), " · ".join(flags)])
    ws.freeze_panes = "B2"
    _w(ws, [16, 34, 26, 8, 10, 30, 20])

    # 대응 기준 연계
    ws = wb.create_sheet("대응 기준 연계")
    _hdr(ws, 1, ["NIST SP 800-53", "통제명", "연계 세부위협 수", "매우 높음·높음", "연계 세부위협"])
    nr = collections.defaultdict(list)
    for o in ROWS:
        for k in o["nist"]:
            nr[k].append(o)
    for i, k in enumerate(sorted(nr), 2):
        os_ = nr[k]
        _row(ws, i, [k, NIST.get(k, ""), len(os_), sum(1 for o in os_ if o["risk"] in ("매우 높음", "높음")),
                     ", ".join(o["id"] for o in sorted(os_, key=lambda o: o["priority"]))])
    ws.freeze_panes = "B2"
    _w(ws, [16, 40, 10, 12, 50])

    # 참고 문헌
    ws = wb.create_sheet("참고 문헌")
    _hdr(ws, 1, ["키", "짧은 이름", "제목", "발행처", "연도", "쓰는 자리", "주 전술", "인용 세부위협", "이용 조건", "확인 경로", "확인일", "URL"])
    for i, (k, v) in enumerate(sorted(REFS.items()), 2):
        cited = [o["id"] for o in ROWS if k in o["refs"]]
        _row(ws, i, [k, v["short"], v["title"], v["publisher"], str(v["year"]), " · ".join(v["use"]),
                     " · ".join(v["domains"]), ", ".join(cited), v["terms"], v["verify"], str(v["checked"]),
                     "\n".join(v["urls"])])
    ws.freeze_panes = "B2"
    _w(ws, [14, 26, 44, 20, 6, 16, 22, 26, 28, 44, 11, 56])

    # 다른 매트릭스 연계
    ws = wb.create_sheet("다른 매트릭스 연계")
    _hdr(ws, 1, ["매트릭스", "연계 ID", "명칭", "통합 v4.1 위험도", "연계 세부위협"])
    lr = collections.defaultdict(list)
    for o in ROWS:
        for fld in ("ai", "cl", "ot", "sc", "idn", "ph"):
            for k in o[fld]:
                lr[k].append(o["id"])
    for i, k in enumerate(sorted(lr), 2):
        v = LINKS.get(k, {})
        _row(ws, i, [v.get("matrix", ""), k, v.get("name", ""), v.get("risk", ""), ", ".join(lr[k])])
    ws.freeze_panes = "B2"
    _w(ws, [24, 16, 40, 12, 40])

    # 평가 기준
    ws = wb.create_sheet("평가 기준")
    crit = [
        ("■ 발생가능성", ""),
        ("상", f"실제 사고 {R.LIKELY_HIGH_REAL}건 이상(다른 매트릭스와 같은 수식)"),
        ("중", "실제 사고 1건, 또는 KEV 악용·정부 경보·위협인텔·실증 중 하나, 또는 ATT&CK 주체(그룹·소프트웨어·캠페인) 사례 1개 이상"),
        ("하", "근거 없음"),
        ("선정 기준(D5)", "고도화·최근빈발·레거시난제 셋 중 둘 이상으로 포함 여부를 가림. 자산 특화 기법은 제외·연계(D7)"),
        ("우선순위", "위험도 → 실제 사고 수 → 요구 방어 수준(방어 난이도) → ATT&CK 주체 수 → ENT-ID"),
        ("", ""),
        ("■ 심각도", ""),
        ("상", TAX["assessment"]["severity_rule"]["상"]),
        ("중", TAX["assessment"]["severity_rule"]["중"]),
        ("하", TAX["assessment"]["severity_rule"]["하"]),
        ("", ""),
        ("■ 위험도 = 발생가능성 × 심각도", " / ".join(f"({a}·{b})→{R.RISK_MATRIX[(a,b)]}" for a in ("상", "중", "하") for b in ("상", "중", "하"))),
        ("근거 수준", " · ".join(R.LEVELS)),
        ("", ""),
        ("■ 고유 축(D6)", "레거시 방어 공백 유형(시그니처 회피·파일리스·서명된 정상 도구·암호화 채널·정상 인증 위장·경계 밖 장비·커널 펌웨어·신뢰 관계) · 요구 방어 수준(레거시 / EDR·XDR / 행위·네트워크·아이덴티티)"),
        ("■ 경계(D7)", "호스트·엔드포인트·네트워크·C2·방어 회피·온프레미스 AD 인접 기법만. 클라우드·신원·OT·AI·공급망 특화는 통합 v4.1 요약 ID로 연계만"),
        ("■ 가드레일(D11)", "방어 자료 — 위험·탐지·대응 중심. 공세 제작법·악용 절차는 적지 않고 식별·탐지·완화 수준으로 서술"),
    ]
    _kv(ws, crit, (28, 140))

    # 변경이력
    ws = wb.create_sheet("변경이력")
    _hdr(ws, 1, ["버전", "날짜", "구분", "대상", "변경 내용", "사유"])
    for i, e in enumerate(CHANGELOG.get("entries") or [], 2):
        _row(ws, i, list(e))
    _w(ws, [8, 12, 10, 40, 90, 50])
    return wb


def write_csv(path):
    cols = ["전술(Lv1)", "ENT-ID", "위협분류(Lv2)", "세부위협(Lv3)", "영문명", "레거시 공백 유형", "요구 방어 수준", "위협 주체",
            "플랫폼", "ATT&CK", "ATT&CK 주체 수", "KEV", "LOLBin", "BYOVD 드라이버", "AI 연계", "클라우드 연계", "OT 연계",
            "공급망 연계", "신원 연계", "물리 연계", "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수",
            "최근(2025~)", "발생가능성 근거", "심각도 근거", "관련 사례 ID", "NIST SP 800-53", "참고 문헌"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for o in sorted(ROWS, key=lambda o: o["priority"]):
            x = EV[o["id"]]
            w.writerow([o["domain"], o["id"], o["lv2"], o["name"], o["en"], " · ".join(o["legacy_gap"]), o["defense_level"],
                        " · ".join(o["actors"]), " · ".join(o["platforms"]), ", ".join(o["attack"]), o["n_attack"],
                        "O" if o.get("kev") else "", "O" if o.get("lolbas") else "", "O" if o.get("loldriver") else "",
                        ", ".join(o["ai"]), ", ".join(o["cl"]), ", ".join(o["ot"]), ", ".join(o["sc"]), ", ".join(o["idn"]),
                        ", ".join(o["ph"]), o["likelihood"], o["severity"], o["risk"], o["priority"], o["level"],
                        o["n_real"], o["n_recent"], o["lk_why"], o["severity_why"], ", ".join(sorted(x["real"])),
                        " · ".join(o["nist"]), " · ".join(o["refs"])])


def main():
    global VALIDATION
    ap = argparse.ArgumentParser()
    ap.add_argument("--version", default="v1")
    ap.add_argument("--skip-validate", action="store_true")
    a = ap.parse_args()
    if not a.skip_validate:
        VALIDATION = run_validate()
    xlsx = ROOT / "output" / f"{NAME}_{a.version}.xlsx"
    csvp = ROOT / "output" / f"{NAME}_{a.version}.csv"
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb = build(a.version)
    wb.save(xlsx)
    write_csv(csvp)
    print(f"빌드 완료 — 세부위협 {len(ROWS)} · 사고 {len(INCIDENTS)} · {VALIDATION}")
    print(f"  → {xlsx.relative_to(ROOT)} ({len(wb.sheetnames)}시트) · {csvp.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
