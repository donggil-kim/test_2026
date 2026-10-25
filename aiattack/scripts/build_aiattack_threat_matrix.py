# -*- coding: utf-8 -*-
"""
AI 활용 공격 보안위협 매트릭스 빌더
----------------------------------
분류 : 공격자 킬체인 7단계(RT·RD·VE·IA·AO·CC·IM) → 위협분류(Lv2) → 세부위협(Lv3) — data/taxonomy.yaml
고유 : AI 역할·자율성(보조·증강·반자율·완전자율) · AI 고유성(신규/증폭) · AI 역량 성숙도(실사용·실증·연구시연)
교차 : MITRE ATLAS 2026.09(공격자 AI 관점·성숙도·사례) · ATT&CK v19.2(가속하는 기존 기법) · OWASP · 통합 v4.1 연계 ID
근거 : AI 활용 공격 사고 DB(data/incidents.yaml) · AI 역량 성숙도 신호(D4) · ATLAS 사례 · 위협인텔 · 실증·연구
대응 : NIST SP 800-53 Rev.5 (탐지·대응 통제)
논리 : 발생가능성 수식은 다른 매트릭스와 같게 유지(상 = 실제 사고 2건 이상)하되, 사고가 적은 신흥 영역이므로
       'AI 역량 성숙도'를 보조 신호로 더해 실증된 역량이 과소평가되지 않게 함(D4). 위협인텔·실증·ATLAS 사례는 '중' 상한.

문구 : data/text/*.yaml(핵심 요약·요약설명·참조·탐지·대응) — ③단계. 없으면 '문구 미작성'으로 표시(v1 데이터 빌드)
시나리오 : data/scenarios.yaml — ③단계

실행 : python3 scripts/build_aiattack_threat_matrix.py [--version v1] [--skip-validate]
       → output/통합_AI활용공격_보안위협_매트릭스_<버전>.xlsx · .csv
"""
import argparse
import collections
import csv
import datetime
import json
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import taxonomy_rules as R  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
REF = ROOT / "reference"
NAME = "통합_AI활용공격_보안위협_매트릭스"


def load_yaml(p, default=None):
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else default


TAX = load_yaml(DATA / "taxonomy.yaml")
FW = load_yaml(DATA / "frameworks.yaml")
INCIDENTS = load_yaml(DATA / "incidents.yaml", []) or []
CHANGELOG = load_yaml(DATA / "changelog.yaml", {}) or {}
SCENARIOS = load_yaml(DATA / "scenarios.yaml", {}) or {}
ATLAS = json.loads((REF / "atlas" / "atlas-2026.09-subset.json").read_text(encoding="utf-8"))
ATK = json.loads((REF / "attack" / "enterprise-attack-v19.2-subset.json").read_text(encoding="utf-8"))
NIST = json.loads((REF / "nist" / "sp800-53r5-names.json").read_text(encoding="utf-8"))["names"]
LINKS = json.loads((REF / "links" / "other_matrices.json").read_text(encoding="utf-8"))["ids"]
REFS = FW["references"]
OWASP = FW.get("owasp") or {}

TEXT = {}
for f in sorted((DATA / "text").glob("*.yaml")) if (DATA / "text").exists() else []:
    TEXT.update(load_yaml(f, {}) or {})

DOMAINS, ROWS = [], []
for d in TAX["domains"]:
    DOMAINS.append(d)
    for l2 in d["lv2"]:
        for x in l2["lv3"]:
            ROWS.append(dict(x, code=d["code"], domain=f"[{d['code']}] {d['ko']}", lv2_id=l2["id"], lv2=l2["name"],
                             single=(x["id"] == l2["id"])))
ROW = {o["id"]: o for o in ROWS}
DOMAIN = {d["code"]: d for d in DOMAINS}
STAGE_OF = dict(zip(R.DOMAIN_ORDER, R.STAGES))

for e in INCIDENTS:
    e["date"] = str(e["date"])
    e["status"] = R.incident_status(e)
    e["recent"] = e["date"] >= R.RECENT_FROM
INC = {e["id"]: e for e in INCIDENTS}

# ===========================================================================
# 근거 집계
# ===========================================================================
EV = {o["id"]: dict(real=set(), recent=set(), confirmed=set(), presumed=set(), intel=set(), research=set(),
                    excluded=set(), atlas_cases=set(), impact=collections.Counter(), actor=collections.Counter())
      for o in ROWS}

for e in INCIDENTS:
    for k in e.get("aia") or []:
        if k not in EV:
            continue
        x = EV[k]
        if e["status"] == "실제 사고":
            x["real"].add(e["id"])
            if e["recent"]:
                x["recent"].add(e["id"])
            x["impact"].update(i for i in e["impact"] if i not in ("영향 없음(사전 차단)", "영향 미상"))
            x["actor"].update(e["actor_type"])
        elif e["status"] == "위협인텔":
            x["intel"].add(e["id"])
        elif e["status"] == "실증·연구":
            x["research"].add(e["id"])
        else:
            x["excluded"].add(e["id"])
        (x["confirmed"] if e["ai_use"] == "확인" else x["presumed"]).add(e["id"])

for o in ROWS:
    for t in o["atlas"]:
        EV[o["id"]]["atlas_cases"] |= set(ATLAS["techniques"].get(t, {}).get("cases") or [])


# ===========================================================================
# 위험 평가 (D4 — AI 역량 성숙도 보조 신호)
# ===========================================================================
def likelihood(o, x):
    if len(x["real"]) >= R.LIKELY_HIGH_REAL:
        return "상"
    if x["real"] or x["intel"] or x["research"] or x["atlas_cases"] or o["maturity"] in ("실사용", "실증"):
        return "중"
    return "하"


def evidence_level(o, x):
    if x["real"]:
        return "실제 사고 확인"
    if o["maturity"] == "실사용" or x["intel"]:
        return "실사용 역량(성숙도)"
    if o["maturity"] == "실증" or x["research"] or x["atlas_cases"]:
        return "실증·공개 PoC"
    return "이론·연구 시연"


for o in ROWS:
    x = EV[o["id"]]
    o["likelihood"] = likelihood(o, x)
    o["risk"] = R.RISK_MATRIX[(o["likelihood"], o["severity"])]
    o["level"] = evidence_level(o, x)
    o["n_real"], o["n_recent"] = len(x["real"]), len(x["recent"])
    o["n_conf"], o["n_pres"] = len(x["confirmed"]), len(x["presumed"])
    o["n_intel"], o["n_research"], o["n_atlas"] = len(x["intel"]), len(x["research"]), len(x["atlas_cases"])
    o["lk_why"] = (f"실제 사고 {o['n_real']}건(최근 {o['n_recent']}) · AI 활용 확인 {o['n_conf']}·추정 {o['n_pres']} · "
                   f"위협인텔 {o['n_intel']} · 실증·연구 {o['n_research']} · ATLAS 사례 {o['n_atlas']} · 역량 성숙도 {o['maturity']}")
    txt = TEXT.get(o["id"]) or {}
    o["oneline"], o["summary"], o["reference"], o["detect"] = (txt.get(k) or "" for k in
                                                              ("oneline", "summary", "reference", "detect"))
    o["text_src"] = "분석가 작성" if txt else "문구 미작성"

rank_key = {r: i for i, r in enumerate(R.RISK_ORDER)}
mat_key = {m: R.MATURITY_RANK[m] for m in R.MATURITY}
for i, o in enumerate(sorted(ROWS, key=lambda o: (rank_key[o["risk"]], -o["n_real"], -o["n_atlas"],
                                                  -mat_key[o["maturity"]], o["id"])), 1):
    o["priority"] = i


# ===========================================================================
# 라벨
# ===========================================================================
def nist_label(k):
    return f"{k} {NIST.get(k, '')}"


def atlas_label(t):
    return f"{t} {ATLAS['techniques'][t]['name']} [{ATLAS['techniques'][t]['maturity']}]"


def atk_label(t):
    return f"{t} {ATK['techniques'][t]['name']}" if t in ATK["techniques"] else t


def owasp_label(k):
    v = OWASP.get(k)
    return f"{k} {v[-1]}" if v else k


def link_label(k):
    return f"{k} {LINKS[k]['name']}" if k in LINKS else k


def ref_label(k):
    return REFS[k]["short"] if k in REFS else k


def inc_label(i):
    e = INC[i]
    return f"{i} {e['title']}({e['date']})"


def first_case(o):
    real = sorted(EV[o["id"]]["real"], key=lambda i: (INC[i]["date"], i), reverse=True)
    if real:
        e = INC[real[0]]
        return f"[실제 사고] {e['title']}({e['date']}) ({real[0]})"
    other = sorted(EV[o["id"]]["intel"] | EV[o["id"]]["research"], key=lambda i: (INC[i]["date"], i), reverse=True)
    if other:
        e = INC[other[0]]
        lab = {"위협인텔": "[위협인텔]", "실증·연구": "[실증]"}[e["status"]]
        return f"{lab} {e['title']}({e['date']}) ({other[0]})"
    if o["n_atlas"]:
        return f"[ATLAS 사례 {o['n_atlas']}건·성숙도 {o['maturity']}]"
    return "공개 사고 미확인"


def first_control(o):
    b = [ln for ln in (o["detect"] or "").splitlines() if ln.startswith("- ")]
    if b:
        return b[0][2:].strip()
    return "대응 기준: NIST " + "·".join(o["nist"]) if o["nist"] else ""


def run_validate():
    import subprocess
    cmd = [sys.executable, str(Path(__file__).resolve().parent / "validate.py"), "--strict"]
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    lines = [ln for ln in r.stdout.splitlines() if ln.strip()]
    if r.returncode != 0:
        print("\n".join(lines))
        raise SystemExit("검증 오류·경고 — 데이터를 고친 뒤 다시 빌드(--skip-validate로 무시 가능)")
    return next((ln for ln in reversed(lines) if ln.startswith("세부위협")), "")


VALIDATION = "검증 생략(--skip-validate)"

# ===========================================================================
# xlsx 서식
# ===========================================================================
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "실사용 역량(성숙도)": "2E86C1", "실증·공개 PoC": "B9770E", "이론·연구 시연": "7F8C8D"}
HDR, NAVY, GRAY = "34495E", "1F3864", "8C8C8C"


def _w(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hdr(ws, row, cols, start=1, color=HDR):
    for j, c in enumerate(cols, start):
        cell = ws.cell(row, j, c)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BORDER


def _risk(cell, risk):
    bg, fg = RISK_COLOR[risk]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=risk in ("매우 높음", "높음"), color=fg)
    cell.alignment = CENTER


def _level(cell, level):
    cell.font = Font(bold=True, color=LEVEL_COLOR[level])
    cell.alignment = CENTER


def _row(ws, r, vals, start=1):
    for j, v in enumerate(vals, start):
        c = ws.cell(r, j, v)
        c.alignment = WRAP
        c.border = BORDER


def _kv_sheet(ws, rows, widths=(24, 140)):
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    _w(ws, list(widths))


def summary_cell(o):
    return o["summary"] or f"(문구 미작성) {o['en']}"


def build(version):
    wb = openpyxl.Workbook()
    today = datetime.date.today().isoformat()
    rc = collections.Counter(o["risk"] for o in ROWS)
    lc = collections.Counter(o["level"] for o in ROWS)
    sc = collections.Counter(o["severity"] for o in ROWS)
    st = collections.Counter(e["status"] for e in INCIDENTS)
    rolec = collections.Counter(o["ai_role"] for o in ROWS)
    novc = collections.Counter(o["novelty"] for o in ROWS)
    matc = collections.Counter(o["maturity"] for o in ROWS)
    real = [e for e in INCIDENTS if e["status"] == "실제 사고"]
    reused = [e for e in INCIDENTS if e.get("origin")]
    years = sorted({e["date"][:4] for e in INCIDENTS})
    by_year = collections.Counter(e["date"][:4] for e in INCIDENTS)
    written = sum(1 for o in ROWS if o["text_src"] == "분석가 작성")
    n_lv2 = sum(len(d["lv2"]) for d in DOMAINS)
    by_c = collections.defaultdict(list)
    for o in ROWS:
        by_c[o["code"]].append(o)
    top = sorted(ROWS, key=lambda o: o["priority"])[:10]
    scn = SCENARIOS.get("scenarios") or []

    # ---------------- 개요 ----------------
    ws = wb.active
    ws.title = "개요"
    meta = [
        (f"AI 활용 공격 보안위협 매트릭스 {version}", ""),
        ("관점", "공격자가 AI(LLM·에이전트·생성형·적대적 ML)를 무기로 삼아 취약점을 발굴·악용하고 침투·사회공학·악성코드·C2를 "
                "가속·자동화·대규모화하는 위협을 공격자 킬체인 7단계로 정리한 렌즈 매트릭스. AI 매트릭스(v3.2)가 '우리 AI를 지키는' "
                "관점이라면 이 매트릭스는 '공격자의 AI가 (AI든 비-AI든) 대상을 치는' 관점"),
        ("생성일", today), ("", ""),
        ("■ 기준 데이터", ""),
        ("AI 활용 공격 사고 DB", f"{len(INCIDENTS)}건({years[0]}~{years[-1]}) — 실제 사고 {st['실제 사고']} · 위협인텔 "
                           f"{st.get('위협인텔', 0)} · 실증·연구 {st.get('실증·연구', 0)}, AI 활용 확인 "
                           f"{sum(1 for e in INCIDENTS if e['ai_use']=='확인')}·추정 {sum(1 for e in INCIDENTS if e['ai_use']=='추정')}. "
                           f"통합 v4.1 재인용 {len(reused)}건(origin) + 신규 {len(INCIDENTS)-len(reused)}건. "
                           f"연도별 " + "·".join(f"{y}:{by_year[y]}" for y in years) + " (급증)"),
        ("MITRE ATLAS", f"2026.09 — 공격자 AI 관점 기법 {len(ATLAS['techniques'])}개 연계(성숙도 Realized·Demonstrated·Feasible 포함). "
                        f"기법별 ATLAS 사례 수를 발생가능성 'AI 역량 성숙도' 신호로 참고(D4)"),
        ("MITRE ATT&CK", f"Enterprise v19.2 — AI가 가속하는 기존 기법 {len(ATK['techniques'])}개를 연계(정찰·익스플로잇 확보·피싱·실행·반출 등)"),
        ("OWASP", f"LLM Top 10(2025)·Agentic — 공격에 악용되는 제3자 AI 측 {len(OWASP)}개 연계(대부분 해당 없음)"),
        ("대응 기준", "NIST SP 800-53 Rev.5 통제(탐지·대응). NIST AI RMF는 거버넌스 맥락 명칭만"),
        ("참고 문헌", f"{len(REFS)}종 — 벤더·정부·연구 위협 인텔(분류 근거·위협 내용 인용, 발생가능성에는 넣지 않음)"),
        ("다른 매트릭스", "통합 v4.1 요약 ID(AI·CL·OT·SC·ID·PH) 연계 — 렌즈라 겹침이 많아 재수록 대신 연계만(작업방향 5장·D9)"),
        ("", ""),
        ("■ 분류 체계(D3)", ""),
        ("Lv1 도메인(킬체인 7단계)", " → ".join(f"[{d['code']}] {d['ko']}" for d in DOMAINS)),
        ("Lv2 위협분류", f"{n_lv2}개(AIA-<도메인>-NN). 하위 세부위협이 하나뿐이면 Lv2 = Lv3"),
        ("Lv3 세부위협", f"{len(ROWS)}개(AIA-<도메인>-NN.m). ID는 고정(재사용 안 함)"),
        ("고유 축", "AI 역할·자율성(보조·증강·반자율·완전자율) · AI 고유성(신규/증폭) · AI 역량 성숙도(실사용·실증·연구시연). "
                  "국가배후 여부는 행 속성이 아니라 사고 DB 행위자 필드로만 다룸"),
        ("", ""),
        ("■ 시트 구성", ""),
        ("보고서용 간략 매트릭스", "우선순위 순 세부위협 · 핵심 요약 · 위험도 · 실제 사고 수 · 대표 사례 1건 · 핵심 대응 1줄"),
        ("매트릭스 뷰", "킬체인 단계(열)별 세부위협을 위험도 색으로 배치"),
        ("통합 매트릭스", "분류·고유 축·교차 매핑·위험 평가·실제 근거·대응 기준 전체 열"),
        ("통합매트릭스_LITE", "핵심 열 발췌(필터·보고용)"),
        ("도메인 요약", "킬체인 단계별 경계·위협 수·위험도/근거 분포·대표 위협"),
        ("AI 역할·자율성 요약", "고유 축(AI 역할·자율성 × AI 고유성 × 성숙도)별 세부위협·위험도 — 이 매트릭스의 차별적 축"),
        ("공격 체인 시나리오", f"실제 사고 기반 공격 흐름 {len(scn)}개" if scn else "시트 양식만 — 시나리오(data/scenarios.yaml)는 ③단계"),
        ("역매핑_사고사례", f"사고 DB {len(INCIDENTS)}건과 매핑 세부위협·AI 활용·성숙도·출처(근거 추적)"),
        ("교차 매핑 연계", "ATLAS·ATT&CK·OWASP 항목별 연계 행"),
        ("대응 기준 연계", "NIST SP 800-53 통제별 연계 세부위협과 위험도 분포"),
        ("참고 문헌", "벤더·정부·연구 위협 인텔 카탈로그 — 쓰는 자리·주 도메인·인용 세부위협·URL·이용 조건·확인 경로"),
        ("다른 매트릭스 연계", "AI·클라우드·OT·공급망·신원·물리 매트릭스 항목별로 연결된 AI 활용 공격 세부위협"),
        ("평가 기준", "발생가능성·심각도·위험도·근거수준 산정 규칙(AI 역량 성숙도 신호 포함), 경계 규칙, 가드레일"),
        ("변경이력", "버전별 변경 내역"),
        ("", ""),
        ("■ 결과 요약", ""),
        ("세부위협(Lv3)", f"{len(ROWS)}개 / 위협분류 {n_lv2}개 / 도메인 {len(DOMAINS)}개"),
        ("위험도", " · ".join(f"{k} {rc.get(k, 0)}" for k in R.RISK_ORDER)),
        ("근거 수준", " · ".join(f"{k} {lc.get(k, 0)}" for k in R.LEVELS)),
        ("심각도", " · ".join(f"{k} {sc.get(k, 0)}" for k in ("상", "중", "하"))),
        ("AI 역할·자율성", " · ".join(f"{k} {rolec.get(k, 0)}" for k in R.AI_ROLES)),
        ("AI 고유성", " · ".join(f"{k} {novc.get(k, 0)}" for k in R.NOVELTY)),
        ("AI 역량 성숙도", " · ".join(f"{k} {matc.get(k, 0)}" for k in R.MATURITY)),
        ("관점 문구", f"{written}/{len(ROWS)}개 작성" + (" — 전 항목 완료" if written == len(ROWS) else
                     " — 문구는 ③단계(요약설명 자리에 영문명 회색 표시, 대표 사례는 사고 DB·ATLAS 근거, 핵심 대응은 NIST 통제)")
         + f" | scripts/validate.py {VALIDATION}"),
        ("우선 위협 Top 10", "\n".join(f"{o['priority']}. {o['id']} {o['name']} — {o['risk']}, 실제 사고 {o['n_real']}건"
                                    f"(성숙도 {o['maturity']})" for o in top)),
        ("", ""), ("■ 주의", ""),
        ("발생가능성(D4)", "수식은 다른 매트릭스와 같게 유지(상 = 실제 사고 2건 이상). 다만 AI 활용 공격은 사고가 적어도 급증·빠른 변화라 "
                       "'AI 역량 성숙도'(실사용/실증/연구시연)를 보조 신호로 더함 — 실증된 위험 역량이 사고 건수만으로 과소평가되지 않게 함. "
                       "위협인텔·실증·ATLAS 사례·성숙도는 '중' 상한"),
        ("근거 밀도 차이", f"자율 취약점 발굴·익스플로잇 생성(VE)·역량 확보(RD)·일부 정찰은 실제 사고보다 실증·위협인텔 근거가 큼. "
                     f"실제 사고가 매핑된 세부위협 {sum(1 for o in ROWS if o['n_real'])}/{len(ROWS)}개 — 나머지는 성숙도·실증·위협인텔로 평가"),
        ("AI 활용 추정", "벤더가 정황(코드 패턴·산출물·광고)으로 AI 활용을 판단한 사례는 ai_use '추정'으로 구분(hackerbot-claw·prt-scan·EvilTokens 등)"),
        ("중복·연계", "렌즈 매트릭스라 기존 6개와 겹침 — 재수록 대신 통합 v4.1 요약 ID로 연계(D9). 위험도·사고 수는 같은 매트릭스 안에서만 비교"),
        ("가드레일(D11)", "방어 자료 — 위험·탐지·대응 중심. 공세 AI 제작법·악용 절차는 적지 않고 역량의 존재·영향·탐지 단서까지만 식별 수준으로 서술"),
        ("출처 표기", "MITRE ATLAS®·ATT&CK® © The MITRE Corporation · NIST SP 800-53(공공 영역) · OWASP(CC BY-SA) · 벤더·정부·연구 보고서는 참고 문헌 시트"),
    ]
    _kv_sheet(ws, meta, (26, 140))
    ws.cell(1, 1).font = Font(size=15, bold=True, color=NAVY)

    # ---------------- 보고서용 간략 매트릭스 ----------------
    ws = wb.create_sheet("보고서용 간략 매트릭스")
    bcols = ["순위", "단계(Lv1)", "AIA-ID", "세부위협(Lv3)", "핵심 요약", "위험도", "발생가능성", "심각도", "AI 역할", "고유성",
             "성숙도", "실제 사고 수", "대표 사례", "핵심 대응"]
    ws.cell(1, 1, f"보고서용 간략 매트릭스 {version} — 본문 삽입용(상세는 '통합 매트릭스')").font = Font(size=13, bold=True, color=NAVY)
    ws.cell(2, 1, f"순위 = 위험도 → 실제 사고 수 → ATLAS 사례 → 성숙도 → AIA-ID | {today}").font = Font(size=9, color="595959")
    _hdr(ws, 4, bcols)
    for r, o in enumerate(sorted(ROWS, key=lambda o: o["priority"]), 5):
        _row(ws, r, [o["priority"], o["domain"], o["id"], o["name"], o["oneline"] or summary_cell(o), o["risk"],
                     o["likelihood"], o["severity"], o["ai_role"], o["novelty"], o["maturity"], o["n_real"],
                     first_case(o), first_control(o)])
        ws.cell(r, 3).font = Font(bold=True)
        ws.cell(r, 4).font = Font(bold=True)
        if not o["oneline"]:
            ws.cell(r, 5).font = Font(color=GRAY)
        _risk(ws.cell(r, 6), o["risk"])
        for j in (1, 7, 8, 9, 10, 11, 12):
            ws.cell(r, j).alignment = CENTER
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(bcols))}{len(ROWS) + 4}"
    _w(ws, [6, 18, 13, 30, 46, 9, 7, 7, 8, 7, 8, 8, 46, 50])
    ws.print_title_rows = "4:4"

    # ---------------- 매트릭스 뷰 ----------------
    ws = wb.create_sheet("매트릭스 뷰")
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도(빨강 매우 높음·주황 높음·노랑 보통·초록 낮음), [n]=실제 사고 수, "
                  "괄호=세부위협 수. 열 = 공격자 킬체인 단계").font = Font(bold=True, color=NAVY)
    for j, d in enumerate(DOMAINS, 1):
        c = ws.cell(2, j, f"[{d['code']}] {d['ko']}\n{d['en']}\n({len(by_c[d['code']])})")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        for i, o in enumerate(sorted(by_c[d["code"]], key=lambda o: (rank_key[o["risk"]], -o["n_real"], o["id"])), 3):
            cell = ws.cell(i, j, f"{o['id']} {o['name']}" + (f" [{o['n_real']}]" if o["n_real"] else ""))
            _risk(cell, o["risk"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9, bold=o["risk"] == "매우 높음", color=RISK_COLOR[o["risk"]][1])
            cell.border = BORDER
    ws.row_dimensions[2].height = 52
    ws.freeze_panes = "A3"
    _w(ws, [27] * len(DOMAINS))

    # ---------------- 통합 매트릭스 ----------------
    ws = wb.create_sheet("통합 매트릭스")
    g_class = ["단계(Lv1)", "AIA-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "영문명", "킬체인 단계", "설명 출처"]
    g_axis = ["AI 역할·자율성", "AI 고유성", "AI 역량 성숙도", "공격 대상", "위협 주체"]
    g_cross = ["ATLAS ID", "ATLAS 기법명(성숙도)", "ATT&CK ID", "ATT&CK 기법명", "OWASP",
               "AI 연계", "클라우드 연계", "OT 연계", "공급망 연계", "신원 연계", "물리 연계"]
    g_risk = ["발생가능성", "심각도", "위험도", "우선순위", "발생가능성 근거(자동 산정)", "심각도 근거"]
    g_ev = ["실제 사고 수", "최근(2025~)", "AI 활용 확인", "AI 활용 추정", "위협인텔", "실증·연구", "ATLAS 사례 수", "근거 수준",
            "관련 사례 ID", "주요 영향(실제 사고)"]
    g_ctl = ["NIST SP 800-53", "탐지·대응 포인트", "참고 문헌"]
    groups = [("분류 체계", len(g_class), "2E5496"), ("고유 축", len(g_axis), "117A65"),
              ("교차 매핑", len(g_cross), "1F7A8C"), ("위험 평가", len(g_risk), "A04000"),
              ("실제 근거", len(g_ev), "1E8449"), ("대응 기준·참고 문헌", len(g_ctl), "6C3483")]
    cols = g_class + g_axis + g_cross + g_risk + g_ev + g_ctl
    t = ws.cell(1, 1, f"AI 활용 공격 보안위협 매트릭스 {version} — 분류·고유 축·교차 매핑·위험 평가·실제 근거")
    t.font = Font(size=13, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, f"AI 활용 공격 사고 DB {len(INCIDENTS)}건 + MITRE ATLAS 2026.09 + ATT&CK v19.2 + OWASP | {today}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    c0 = 1
    for g, span, color in groups:
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + span - 1)
        c = ws.cell(3, c0, g)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        c0 += span
    _hdr(ws, 4, cols)
    ci = {c: i + 1 for i, c in enumerate(cols)}
    for r, o in enumerate(ROWS, 5):
        x = EV[o["id"]]
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summary_cell(o), o["en"],
                     ", ".join(o["stage"]), o["text_src"],
                     o["ai_role"], o["novelty"], o["maturity"], " · ".join(o["targets"]), " · ".join(o["actors"]),
                     ", ".join(o["atlas"]), "\n".join(atlas_label(t) for t in o["atlas"]),
                     ", ".join(o["attack"]), "\n".join(atk_label(t) for t in o["attack"]),
                     "\n".join(owasp_label(k) for k in o["owasp"]),
                     "\n".join(link_label(k) for k in o["ai"]), "\n".join(link_label(k) for k in o["cl"]),
                     "\n".join(link_label(k) for k in o["ot"]), "\n".join(link_label(k) for k in o["sc"]),
                     "\n".join(link_label(k) for k in o["idn"]), "\n".join(link_label(k) for k in o["ph"]),
                     o["likelihood"], o["severity"], o["risk"], o["priority"], o["lk_why"], o["severity_why"],
                     o["n_real"], o["n_recent"], o["n_conf"], o["n_pres"], o["n_intel"], o["n_research"], o["n_atlas"],
                     o["level"],
                     ", ".join(sorted(x["real"]))
                     + (("\n위협인텔: " + ", ".join(sorted(x["intel"]))) if x["intel"] else "")
                     + (("\n실증·연구: " + ", ".join(sorted(x["research"]))) if x["research"] else ""),
                     " · ".join(f"{k} {v}" for k, v in x["impact"].most_common(4)),
                     "\n".join(nist_label(k) for k in o["nist"]), o["detect"], "\n".join(ref_label(k) for k in o["refs"])])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, ci["요약설명"]).font = Font(color=GRAY)
        for c in ("AI 역할·자율성", "AI 고유성", "AI 역량 성숙도", "발생가능성", "심각도", "우선순위"):
            ws.cell(r, ci[c]).alignment = CENTER
        _risk(ws.cell(r, ci["위험도"]), o["risk"])
        _level(ws.cell(r, ci["근거 수준"]), o["level"])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(ROWS) + 4}"
    widths = {"단계(Lv1)": 16, "AIA-ID": 13, "위협분류(Lv2)": 22, "세부위협(Lv3)": 28, "핵심 요약": 34, "요약설명": 58,
              "영문명": 30, "킬체인 단계": 10, "설명 출처": 9, "AI 역할·자율성": 10, "AI 고유성": 8, "AI 역량 성숙도": 10,
              "공격 대상": 20, "위협 주체": 18, "ATLAS ID": 20, "ATLAS 기법명(성숙도)": 40, "ATT&CK ID": 14,
              "ATT&CK 기법명": 28, "OWASP": 22, "AI 연계": 22, "클라우드 연계": 20, "OT 연계": 16, "공급망 연계": 20,
              "신원 연계": 20, "물리 연계": 16, "발생가능성": 7, "심각도": 7, "위험도": 8, "우선순위": 7,
              "발생가능성 근거(자동 산정)": 40, "심각도 근거": 46, "실제 사고 수": 7, "최근(2025~)": 7, "AI 활용 확인": 7,
              "AI 활용 추정": 7, "위협인텔": 7, "실증·연구": 7, "ATLAS 사례 수": 7, "근거 수준": 12, "관련 사례 ID": 26,
              "주요 영향(실제 사고)": 24, "NIST SP 800-53": 26, "탐지·대응 포인트": 60, "참고 문헌": 22}
    _w(ws, [widths.get(c, 7) for c in cols])

    # ---------------- 통합매트릭스_LITE ----------------
    ws = wb.create_sheet("통합매트릭스_LITE")
    ws.cell(1, 1, f"AI 활용 공격 보안위협 매트릭스 {version} — LITE").font = Font(size=13, bold=True, color=NAVY)
    lcols = ["단계(Lv1)", "AIA-ID", "위협분류(Lv2)", "세부위협(Lv3)", "핵심 요약", "요약설명", "킬체인 단계", "AI 역할", "고유성",
             "성숙도", "공격 대상", "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수", "ATLAS 사례", "탐지·대응"]
    lc_ = {c: i + 1 for i, c in enumerate(lcols)}
    _hdr(ws, 2, lcols)
    for r, o in enumerate(ROWS, 3):
        _row(ws, r, [o["domain"], o["id"], o["lv2"], o["name"], o["oneline"], summary_cell(o), ", ".join(o["stage"]),
                     o["ai_role"], o["novelty"], o["maturity"], " · ".join(o["targets"]), o["likelihood"], o["severity"],
                     o["risk"], o["priority"], o["level"], o["n_real"], o["n_atlas"], o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, lc_["요약설명"]).font = Font(color=GRAY)
        _risk(ws.cell(r, lc_["위험도"]), o["risk"])
        _level(ws.cell(r, lc_["근거 수준"]), o["level"])
        for c in ("AI 역할", "고유성", "성숙도", "발생가능성", "심각도", "우선순위", "실제 사고 수", "ATLAS 사례"):
            ws.cell(r, lc_[c]).alignment = CENTER
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(lcols))}{len(ROWS) + 2}"
    _w(ws, [16, 13, 22, 28, 34, 58, 10, 8, 7, 8, 20, 7, 7, 8, 7, 12, 7, 7, 60])

    # ---------------- 도메인 요약 ----------------
    ws = wb.create_sheet("도메인 요약")
    dcols = ["단계(Lv1)", "설명", "넣지 않는 것(경계)", "위협분류(Lv2)", "세부위협(Lv3)"] + R.RISK_ORDER + R.LEVELS + \
            ["매핑 실제 사고(중복 제거)", "최근(2025~)", "최고위험 세부위협(상위 3)"]
    _hdr(ws, 1, dcols)
    for i, d in enumerate(DOMAINS, 2):
        os_ = by_c[d["code"]]
        rr = collections.Counter(o["risk"] for o in os_)
        ll = collections.Counter(o["level"] for o in os_)
        incs = set().union(*(EV[o["id"]]["real"] for o in os_)) if os_ else set()
        top3 = sorted(os_, key=lambda o: o["priority"])[:3]
        _row(ws, i, [f"[{d['code']}] {d['ko']}", d["desc"], d["excludes"], len(d["lv2"]), len(os_)]
             + [rr.get(k, 0) for k in R.RISK_ORDER] + [ll.get(k, 0) for k in R.LEVELS]
             + [len(incs), sum(1 for e in incs if INC[e]["recent"]),
                "\n".join(f"{o['id']} {o['name']}({o['risk']}, 사고 {o['n_real']})" for o in top3)])
    ws.freeze_panes = "B2"
    _w(ws, [18, 52, 42, 8, 8, 7, 7, 7, 7, 11, 12, 11, 10, 10, 9, 9, 52])

    # ---------------- AI 역할·자율성 요약(고유 축 시트) ----------------
    ws = wb.create_sheet("AI 역할·자율성 요약")
    ws.cell(1, 1, "AI 역할·자율성 × AI 고유성 × 역량 성숙도 — 이 매트릭스의 차별적 축(작업방향 3.3·D3)").font = \
        Font(size=13, bold=True, color=NAVY)
    ws.cell(2, 1, "AI 역할·자율성(보조→증강→반자율→완전자율)이 '사람 속도로 대응 가능한가'를, AI 고유성(신규/증폭)이 "
                  "'진짜 새로운 것/빨라진 것'을 가름").font = Font(size=9, color="595959")
    acols = ["AI 역할·자율성", "세부위협 수", "매우 높음·높음", "신규", "증폭", "실사용", "실증", "연구시연", "대표 세부위협(위험 상위 3)"]
    _hdr(ws, 4, acols)
    r = 5
    for role in R.AI_ROLES:
        os_ = [o for o in ROWS if o["ai_role"] == role]
        top3 = sorted(os_, key=lambda o: o["priority"])[:3]
        _row(ws, r, [role, len(os_), sum(1 for o in os_ if o["risk"] in ("매우 높음", "높음")),
                     sum(1 for o in os_ if o["novelty"] == "신규"), sum(1 for o in os_ if o["novelty"] == "증폭"),
                     sum(1 for o in os_ if o["maturity"] == "실사용"), sum(1 for o in os_ if o["maturity"] == "실증"),
                     sum(1 for o in os_ if o["maturity"] == "연구시연"),
                     "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in top3)])
        for j in range(2, 9):
            ws.cell(r, j).alignment = CENTER
        r += 1
    r += 1
    ws.cell(r, 1, "AI 고유성별").font = Font(bold=True, color=NAVY)
    r += 1
    _hdr(ws, r, ["AI 고유성", "세부위협 수", "매우 높음·높음", "완전자율+반자율", "대표 세부위협(위험 상위 5)"])
    r += 1
    for nov in R.NOVELTY:
        os_ = [o for o in ROWS if o["novelty"] == nov]
        top5 = sorted(os_, key=lambda o: o["priority"])[:5]
        _row(ws, r, [nov, len(os_), sum(1 for o in os_ if o["risk"] in ("매우 높음", "높음")),
                     sum(1 for o in os_ if o["ai_role"] in ("완전자율", "반자율")),
                     "\n".join(f"{o['id']} {o['name']}({o['risk']})" for o in top5)])
        for j in (2, 3, 4):
            ws.cell(r, j).alignment = CENTER
        r += 1
    _w(ws, [16, 10, 12, 8, 8, 8, 8, 10, 64])

    # ---------------- 공격 체인 시나리오 ----------------
    ws = wb.create_sheet("공격 체인 시나리오")
    if scn:
        ws.cell(1, 1, f"공격 체인 시나리오 {len(scn)}개 — 실제 사고 기반 공격 흐름(방어 관점)").font = Font(bold=True, color=NAVY)
        scols = ["시나리오", "제목", "단계", "세부위협", "행위(근거 사고)", "탐지·차단 포인트"]
        _hdr(ws, 3, scols)
        r = 4
        for sc_ in scn:
            for step in sc_.get("steps") or []:
                _row(ws, r, [sc_["id"], sc_.get("title", ""), step.get("단계", ""), ", ".join(step.get("aia") or []),
                             step.get("행위", ""), step.get("탐지", "")])
                r += 1
        _w(ws, [10, 40, 10, 24, 60, 50])
    else:
        ws.cell(1, 1, "공격 체인 시나리오 — ③단계(data/scenarios.yaml)에서 작성").font = Font(bold=True, color=NAVY)
        ws.cell(3, 1, "실제 사고를 공격 흐름으로 재구성해 한 행=한 행위로 나눈 세부위협을 흐름 안에서 보고 "
                      "초크 포인트(가장 싸고 효과 큰 차단 지점)를 드러낼 예정").alignment = WRAP
        _w(ws, [120])

    # ---------------- 역매핑_사고사례 ----------------
    ws = wb.create_sheet("역매핑_사고사례")
    _hdr(ws, 1, ["사고 ID", "사례명", "시점", "유형", "집계 상태", "AI 활용", "성숙도", "지역", "업종", "행위자", "주체 유형",
                 "영향", "매핑 세부위협", "origin(재인용)", "출처"])
    for i, e in enumerate(sorted(INCIDENTS, key=lambda e: (e["date"], e["id"]), reverse=True), 2):
        _row(ws, i, [e["id"], e["title"], e["date"], e["kind"], e["status"], e["ai_use"], e["maturity"],
                     e.get("region", ""), " · ".join(e.get("sector") or []), e.get("actor", ""),
                     " · ".join(e.get("actor_type") or []), " · ".join(e.get("impact") or []),
                     ", ".join(e.get("aia") or []), ", ".join(str(o) for o in e.get("origin") or []),
                     "\n".join(str(s) for s in e.get("sources") or [])])
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:O{len(INCIDENTS) + 1}"
    _w(ws, [9, 44, 9, 10, 11, 8, 9, 9, 18, 24, 16, 24, 24, 16, 60])

    # ---------------- 교차 매핑 연계 ----------------
    ws = wb.create_sheet("교차 매핑 연계")
    _hdr(ws, 1, ["체계", "ID", "명칭", "성숙도/비고", "연계 세부위협 수", "연계 세부위협"])
    r = 2
    used_atlas = sorted({t for o in ROWS for t in o["atlas"]})
    for t in used_atlas:
        rws = [o["id"] for o in ROWS if t in o["atlas"]]
        _row(ws, r, ["ATLAS", t, ATLAS["techniques"][t]["name"], ATLAS["techniques"][t]["maturity"], len(rws),
                     ", ".join(rws)])
        r += 1
    for t in sorted({t for o in ROWS for t in o["attack"]}):
        rws = [o["id"] for o in ROWS if t in o["attack"]]
        _row(ws, r, ["ATT&CK", t, ATK["techniques"].get(t, {}).get("name", ""), "가속하는 기존 기법", len(rws), ", ".join(rws)])
        r += 1
    for k in sorted({k for o in ROWS for k in o["owasp"]}):
        rws = [o["id"] for o in ROWS if k in o["owasp"]]
        _row(ws, r, ["OWASP", k, (OWASP.get(k) or ["", ""])[-1], "악용되는 제3자 AI 측", len(rws), ", ".join(rws)])
        r += 1
    ws.freeze_panes = "B2"
    _w(ws, [10, 16, 40, 18, 10, 50])

    # ---------------- 대응 기준 연계 ----------------
    ws = wb.create_sheet("대응 기준 연계")
    _hdr(ws, 1, ["NIST SP 800-53", "통제명", "연계 세부위협 수", "매우 높음·높음", "연계 세부위협"])
    nist_rows = collections.defaultdict(list)
    for o in ROWS:
        for k in o["nist"]:
            nist_rows[k].append(o)
    for i, k in enumerate(sorted(nist_rows), 2):
        os_ = nist_rows[k]
        _row(ws, i, [k, NIST.get(k, ""), len(os_), sum(1 for o in os_ if o["risk"] in ("매우 높음", "높음")),
                     ", ".join(o["id"] for o in sorted(os_, key=lambda o: o["priority"]))])
    ws.freeze_panes = "B2"
    _w(ws, [16, 40, 10, 12, 50])

    # ---------------- 참고 문헌 ----------------
    ws = wb.create_sheet("참고 문헌")
    _hdr(ws, 1, ["키", "짧은 이름", "제목", "발행처", "연도", "쓰는 자리", "주 도메인", "인용 세부위협", "이용 조건", "확인 경로", "확인일", "URL"])
    for i, (k, v) in enumerate(sorted(REFS.items()), 2):
        cited = [o["id"] for o in ROWS if k in o["refs"]]
        _row(ws, i, [k, v["short"], v["title"], v["publisher"], str(v["year"]), " · ".join(v["use"]),
                     " · ".join(v["domains"]), ", ".join(cited), v["terms"], v["verify"], str(v["checked"]),
                     "\n".join(v["urls"])])
    ws.freeze_panes = "B2"
    _w(ws, [16, 26, 46, 22, 6, 18, 12, 24, 30, 46, 11, 60])

    # ---------------- 다른 매트릭스 연계 ----------------
    ws = wb.create_sheet("다른 매트릭스 연계")
    _hdr(ws, 1, ["매트릭스", "연계 ID", "명칭", "통합 v4.1 위험도", "연계 AI 활용 세부위협"])
    link_rows = collections.defaultdict(list)
    for o in ROWS:
        for fld in ("ai", "cl", "ot", "sc", "idn", "ph"):
            for k in o[fld]:
                link_rows[k].append(o["id"])
    for i, k in enumerate(sorted(link_rows), 2):
        v = LINKS.get(k, {})
        _row(ws, i, [v.get("matrix", ""), k, v.get("name", ""), v.get("risk", ""), ", ".join(link_rows[k])])
    ws.freeze_panes = "B2"
    _w(ws, [24, 16, 40, 12, 40])

    # ---------------- 평가 기준 ----------------
    ws = wb.create_sheet("평가 기준")
    crit = [
        ("■ 발생가능성(likelihood) — D4", ""),
        ("상", f"실제 사고 {R.LIKELY_HIGH_REAL}건 이상(다른 매트릭스와 같은 수식)"),
        ("중", "실제 사고 1건, 또는 AI 역량 성숙도가 '실사용/실증', 또는 위협인텔·실증·연구·ATLAS 사례 중 하나 이상"),
        ("하", "근거 없음(성숙도 '연구시연'이고 사고·ATLAS 사례 없음)"),
        ("AI 역량 성숙도 신호", "AI 활용 공격은 사고가 적어도 급증·빠른 변화라, 발생가능성에 성숙도(실사용/실증/연구시연)를 보조 신호로 더해 "
                          "실증된 위험 역량이 사고 건수만으로 과소평가되지 않게 함. ATLAS maturity(Realized·Demonstrated·Feasible)와 정합"),
        ("상한", "위협인텔·실증·연구·ATLAS 사례·성숙도는 발생가능성 '중'까지만 반영(실제 사고 2건 이상일 때만 '상')"),
        ("", ""),
        ("■ 심각도(severity)", ""),
        ("상", TAX["assessment"]["severity_rule"]["상"]),
        ("중", TAX["assessment"]["severity_rule"]["중"]),
        ("하", TAX["assessment"]["severity_rule"]["하"]),
        ("", ""),
        ("■ 위험도(risk) = 발생가능성 × 심각도", " / ".join(f"({a}·{b})→{R.RISK_MATRIX[(a, b)]}"
                                                    for a in ("상", "중", "하") for b in ("상", "중", "하"))),
        ("우선순위", "위험도 → 실제 사고 수 → ATLAS 사례 수 → 성숙도 → AIA-ID 순"),
        ("", ""),
        ("■ 근거 수준(evidence level)", " · ".join(R.LEVELS)),
        ("실제 사고 확인", "실제 사고가 매핑된 행"),
        ("실사용 역량(성숙도)", "성숙도 '실사용' 또는 위협인텔"),
        ("실증·공개 PoC", "성숙도 '실증' 또는 실증·연구·ATLAS 사례"),
        ("이론·연구 시연", "그 외"),
        ("", ""),
        ("■ 사고 집계", "실제 사고 = kind ∈ {사고·캠페인·공시·사례연구(익명)}이면서 verification ∉ {행위자 주장·논란·원인 번복}. "
                     "위협인텔·실증·연구는 '중' 상한. AI 활용 '확인/추정'을 보존하되 둘 다 그 행의 사례로는 집계(문구에서 '추정' 구분)"),
        ("■ 고유 축(D3)", "AI 역할·자율성(보조·증강·반자율·완전자율) · AI 고유성(신규/증폭) · AI 역량 성숙도(실사용·실증·연구시연)"),
        ("■ 경계 규칙(D9)", "렌즈 매트릭스라 기존 6개와 겹침 — 재수록 대신 통합 v4.1 요약 ID(AI·CL·OT·SC·ID·PH)로 연계. "
                       "AI 도구 자체가 오염·위장되는 공급망 위협은 AI 매트릭스 D08·공급망. 위험도·사고 수는 같은 매트릭스 안에서만 비교"),
        ("■ 가드레일(D11)", "방어 자료 — 위험·탐지·대응 중심. 공세 AI 제작법·악용 절차는 적지 않고 역량의 존재·영향·탐지 단서까지만 식별 수준으로 서술"),
    ]
    _kv_sheet(ws, crit, (26, 140))

    # ---------------- 변경이력 ----------------
    ws = wb.create_sheet("변경이력")
    _hdr(ws, 1, ["버전", "날짜", "구분", "대상", "변경 내용", "사유"])
    for i, e in enumerate(CHANGELOG.get("entries") or [], 2):
        _row(ws, i, list(e))
    _w(ws, [8, 12, 10, 40, 90, 50])

    return wb


def write_csv(path):
    cols = ["단계(Lv1)", "AIA-ID", "위협분류(Lv2)", "세부위협(Lv3)", "영문명", "킬체인 단계", "AI 역할·자율성", "AI 고유성",
            "AI 역량 성숙도", "공격 대상", "위협 주체", "ATLAS", "ATT&CK", "OWASP", "AI 연계", "클라우드 연계", "OT 연계",
            "공급망 연계", "신원 연계", "물리 연계", "발생가능성", "심각도", "위험도", "우선순위", "근거 수준", "실제 사고 수",
            "최근(2025~)", "AI 활용 확인", "AI 활용 추정", "위협인텔", "실증·연구", "ATLAS 사례 수", "발생가능성 근거",
            "심각도 근거", "관련 사례 ID", "NIST SP 800-53", "참고 문헌"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for o in sorted(ROWS, key=lambda o: o["priority"]):
            x = EV[o["id"]]
            w.writerow([o["domain"], o["id"], o["lv2"], o["name"], o["en"], ", ".join(o["stage"]), o["ai_role"],
                        o["novelty"], o["maturity"], " · ".join(o["targets"]), " · ".join(o["actors"]),
                        ", ".join(o["atlas"]), ", ".join(o["attack"]), ", ".join(o["owasp"]),
                        ", ".join(o["ai"]), ", ".join(o["cl"]), ", ".join(o["ot"]), ", ".join(o["sc"]),
                        ", ".join(o["idn"]), ", ".join(o["ph"]), o["likelihood"], o["severity"], o["risk"],
                        o["priority"], o["level"], o["n_real"], o["n_recent"], o["n_conf"], o["n_pres"], o["n_intel"],
                        o["n_research"], o["n_atlas"], o["lk_why"], o["severity_why"], ", ".join(sorted(x["real"])),
                        " · ".join(o["nist"]), " · ".join(o["refs"])])


def main():
    global VALIDATION
    ap = argparse.ArgumentParser()
    ap.add_argument("--version", default="v1")
    ap.add_argument("--skip-validate", action="store_true")
    a = ap.parse_args()
    if not a.skip_validate:
        VALIDATION = run_validate()
    xlsx = ROOT / "output" / f"{NAME}_{a.version}.xlsx"
    csvp = ROOT / "output" / f"{NAME}_{a.version}.csv"
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb = build(a.version)
    wb.save(xlsx)
    write_csv(csvp)
    print(f"빌드 완료 — 세부위협 {len(ROWS)} · 사고 {len(INCIDENTS)} · {VALIDATION}")
    print(f"  → {xlsx.relative_to(ROOT)} ({len(wb.sheetnames)}시트) · {csvp.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
