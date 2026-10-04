# -*- coding: utf-8 -*-
"""ATT&CK 원천 데이터 로더 — ICS v19.2(STIX) · Enterprise v19.2(xlsx)"""
import collections
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
ICS_JSON = ROOT / "reference/attack-ics/ics-attack-19.2.json"
ENT_XLSX = ROOT / "reference/attack-enterprise/enterprise-attack-v19.2.xlsx"


def clean(text):
    """ATT&CK 설명의 인용 표기·마크다운 링크·HTML 태그 제거."""
    t = re.sub(r"\(Citation:[^)]*\)", "", text or "")
    t = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", t)
    t = re.sub(r"</?[a-z]+>", "", t).replace("`", "")
    return re.sub(r"\s+", " ", t).strip()


def first_sentence(text, limit=260):
    t = clean(text)
    m = re.match(r"(.+?\.)(\s|$)", t)
    s = m.group(1) if m else t
    return s if len(s) <= limit else s[:limit - 1] + "…"


def _eid(o):
    for r in o.get("external_references", []):
        if r.get("source_name") == "mitre-attack":
            return r.get("external_id")
    return None


class ICS:
    """ICS 도메인: 전술·기법·자산·완화책·탐지 전략·절차(캠페인·소프트웨어·그룹)."""

    def __init__(self, path=ICS_JSON):
        bundle = json.loads(Path(path).read_text(encoding="utf-8"))
        objs = {o["id"]: o for o in bundle["objects"]}
        live = [o for o in bundle["objects"] if not o.get("revoked") and not o.get("x_mitre_deprecated")]
        self.version = "19.2"

        matrix = next(o for o in live if o["type"] == "x-mitre-matrix")
        self.tactics = [(objs[t]["x_mitre_shortname"], _eid(objs[t]), objs[t]["name"]) for t in matrix["tactic_refs"]]

        self.tech = {}
        for o in live:
            if o["type"] != "attack-pattern":
                continue
            e = _eid(o)
            self.tech[e] = dict(
                eid=e, name=o["name"], desc=clean(o.get("description")),
                tactics=[k["phase_name"] for k in o.get("kill_chain_phases", [])],
                is_sub=bool(o.get("x_mitre_is_subtechnique")), parent=e.split(".")[0],
                url=f"https://attack.mitre.org/techniques/{e.replace('.', '/')}")

        # 폐기(revoked) ID → 대체 ID (v19 ICS 하위기법 개편)
        self.revoked = {}
        for r in bundle["objects"]:
            if r["type"] == "relationship" and r["relationship_type"] == "revoked-by":
                s, t = objs.get(r["source_ref"]), objs.get(r["target_ref"])
                if s and t and s["type"] == "attack-pattern":
                    self.revoked[_eid(s)] = _eid(t)

        self.assets = {_eid(o): o["name"] for o in live if o["type"] == "x-mitre-asset"}
        self.mitigations = {_eid(o): (o["name"], clean(o.get("description"))) for o in live if o["type"] == "course-of-action"}
        self.subjects = {}
        for o in live:
            if o["type"] in ("campaign", "malware", "tool", "intrusion-set"):
                kind = {"campaign": "캠페인", "malware": "소프트웨어", "tool": "소프트웨어", "intrusion-set": "그룹"}[o["type"]]
                self.subjects[_eid(o)] = dict(name=o["name"], kind=kind, desc=clean(o.get("description")),
                                              refs=[(r.get("source_name"), r.get("url")) for r in o.get("external_references", [])
                                                    if r.get("url") and r.get("source_name") != "mitre-attack"])
        analytics = {o["id"]: clean(o.get("description")) for o in live if o["type"] == "x-mitre-analytic"}
        self.det_strategies = {}
        for o in live:
            if o["type"] == "x-mitre-detection-strategy":
                self.det_strategies[o["id"]] = (_eid(o), o["name"], [analytics[a] for a in o.get("x_mitre_analytic_refs", []) if a in analytics])

        self.targets = collections.defaultdict(list)     # 기법 → 자산 ID
        self.mitig = collections.defaultdict(list)       # 기법 → 완화책 ID
        self.detect = collections.defaultdict(list)      # 기법 → (DET ID, 이름, 분석 설명)
        self.procs = collections.defaultdict(list)       # 기법 → (주체 ID, 절차 설명)
        self.subject_uses = collections.defaultdict(list)  # 주체 → 기법
        for r in live:
            if r["type"] != "relationship":
                continue
            s, t = objs.get(r["source_ref"]), objs.get(r["target_ref"])
            if not s or not t or s.get("revoked") or t.get("revoked"):
                continue
            rt = r["relationship_type"]
            if rt == "targets" and s["type"] == "attack-pattern" and t["type"] == "x-mitre-asset":
                self.targets[_eid(s)].append(_eid(t))
            elif rt == "mitigates" and t["type"] == "attack-pattern":
                self.mitig[_eid(t)].append(_eid(s))
            elif rt == "detects" and t["type"] == "attack-pattern" and s["id"] in self.det_strategies:
                self.detect[_eid(t)].append(self.det_strategies[s["id"]])
            elif rt == "uses" and t["type"] == "attack-pattern" and _eid(s) in self.subjects:
                self.procs[_eid(t)].append((_eid(s), clean(r.get("description"))))
                self.subject_uses[_eid(s)].append(_eid(t))

    def ot_subjects(self):
        """Enterprise 절차를 OT 근거로 인정하는 주체: ICS 캠페인·소프트웨어(그룹은 IT 활동이 섞여 제외)."""
        return {k for k, v in self.subjects.items() if v["kind"] in ("캠페인", "소프트웨어")}


class Enterprise:
    """Enterprise 도메인: 기법 정보, OT 주체의 절차, 완화책."""

    def __init__(self, ot_subjects, path=ENT_XLSX):
        wb = openpyxl.load_workbook(path, read_only=True)
        rows = wb["techniques"].iter_rows(values_only=True)
        h = {k: i for i, k in enumerate(next(rows))}
        self.tech = {}
        for r in rows:
            e = r[h["ID"]]
            self.tech[e] = dict(eid=e, name=r[h["name"]], desc=clean(r[h["description"]]),
                                tactics=[x.strip() for x in str(r[h["tactics"]] or "").split(",") if x.strip()],
                                is_sub=str(r[h["is sub-technique"]]) == "True", parent=e.split(".")[0], url=r[h["url"]])
        rows = wb["mitigations"].iter_rows(values_only=True)
        h = {k: i for i, k in enumerate(next(rows))}
        self.mitigations = {r[h["ID"]]: (r[h["name"]], clean(r[h["description"]])) for r in rows}

        self.procs = collections.defaultdict(list)
        self.subject_uses = collections.defaultdict(list)
        self.mitig = collections.defaultdict(list)
        rows = wb["relationships"].iter_rows(values_only=True)
        h = {k: i for i, k in enumerate(next(rows))}
        for r in rows:
            src, rt, tgt, ttype = r[h["source ID"]], r[h["mapping type"]], r[h["target ID"]], str(r[h["target type"]])
            if not ttype.startswith("technique"):
                continue
            if rt == "uses" and src in ot_subjects:
                self.procs[tgt].append((src, clean(r[h["mapping description"]])))
                self.subject_uses[src].append(tgt)
            elif rt == "mitigates":
                self.mitig[tgt].append(src)
