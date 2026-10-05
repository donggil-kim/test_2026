# -*- coding: utf-8 -*-
"""
통합 OT/ICS/IoT 보안위협 매트릭스 빌더
-------------------------------------
뼈대 : MITRE ATT&CK for ICS v19.2 킬체인(전술 12) + 사전 단계 [RD]
보강 : Enterprise v19.2 기법 중 ICS로 표현되지 않는 IT/OT 경계 기법(같은 캠페인·소프트웨어의 Enterprise 절차 근거)
근거 : OT 사고 DB(data/incidents.yaml) · ATT&CK ICS/Enterprise 절차 · CISA ICS 권고(CSAF)·KEV
논리 : 통합 AI 매트릭스 v3.2 · 클라우드 매트릭스 v5와 같은 발생가능성 수식, 심각도는 OT 결과 기준

실행 : python3 scripts/build_ot_threat_matrix.py [--worksheet DIR]
       --worksheet 를 주면 문구 작성용 근거 정리본(전술별 .md)을 DIR에 함께 만든다.
"""
import argparse
import collections
import csv
import datetime
import re
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import taxonomy_rules as R  # noqa: E402
from attack_data import ICS, Enterprise, ROOT, first_sentence  # noqa: E402

VERSION = "v3"
DATA = ROOT / "data"
OUT = ROOT / "output" / f"통합_OT보안위협_매트릭스_{VERSION}.xlsx"
REGISTRY = DATA / "id_registry.yaml"
TEXT_DIR = DATA / "text"
CHANGELOG = DATA / "changelog.yaml"
VULN_CSV = DATA / "ics_advisory_cves.csv"
INCIDENTS = DATA / "incidents.yaml"
EMB3D_FILE = DATA / "emb3d.yaml"           # scripts/prepare_emb3d.py로 만든 MITRE EMB3D 최소 추출본
EMB3D_URL = "https://emb3d.mitre.org/threats/{}.html"
SCENARIOS = DATA / "scenarios.yaml"        # 공격 체인 시나리오(실제 사고 기반)
STANDARDS = DATA / "standards.yaml"        # 표준 요구사항 명칭(NIST 800-53·IEC 62443)

RISK_MATRIX = {
    ("상", "상"): "매우 높음", ("상", "중"): "높음", ("상", "하"): "보통",
    ("중", "상"): "높음", ("중", "중"): "보통", ("중", "하"): "낮음",
    ("하", "상"): "보통", ("하", "중"): "낮음", ("하", "하"): "낮음",
}
LIKELY_HIGH_REAL = 2   # AI 매트릭스 v3.2 · 클라우드 v5와 동일


# ===========================================================================
# 1) 원천 데이터
# ===========================================================================
ics = ICS()
ent = Enterprise(ics.ot_subjects())
incidents = yaml.safe_load(INCIDENTS.read_text(encoding="utf-8"))
emb3d = yaml.safe_load(EMB3D_FILE.read_text(encoding="utf-8")) if EMB3D_FILE.exists() else dict(threats=[], mitigations=[])
std_names = yaml.safe_load(STANDARDS.read_text(encoding="utf-8")) if STANDARDS.exists() else {}
E3T = {t["id"]: t for t in emb3d["threats"]}        # EMB3D 위협(TID)
E3M = {m["id"]: m for m in emb3d["mitigations"]}    # EMB3D 완화책(MID)
EMB3D_KEYS = set(R.EMB3D_NEW)                        # EMB3D 신설 세부위협 행 키
WARN = collections.defaultdict(set)   # 판정 규칙이 없는 Enterprise 기법 → 등장 위치


def norm(tid, where=""):
    """기법 ID → 매트릭스 행 키. 폐기 ID 변환, Enterprise 통합 규칙 적용. 제외·미판정은 None."""
    if tid in EMB3D_KEYS:
        return tid
    t = ics.revoked.get(tid, tid)
    if t in ics.tech:
        return t
    seen = set()
    while t in R.ENT_MERGE and t not in seen:
        seen.add(t)
        t = ics.revoked.get(R.ENT_MERGE[t], R.ENT_MERGE[t])
    if t in ics.tech or t in R.ENT_INCLUDE:
        return t
    if t in R.ENT_DROP:
        return None
    p = t.split(".")[0]
    if p != t and p in R.ENT_INCLUDE:
        return p
    WARN[t].add(where)
    return None


def tech_info(key):
    if key in EMB3D_KEYS or key in R.EMB3D_GROUP_NAME:
        tids = R.EMB3D_NEW[key][4] if key in EMB3D_KEYS else []
        return dict(eid=key, name=" / ".join(E3T[t]["name"] for t in tids if t in E3T) or key, desc="", tactics=[],
                    is_sub="." in key, parent=key.split(".")[0], url="\n".join(EMB3D_URL.format(t) for t in tids))
    return ics.tech.get(key) or ent.tech.get(key) or dict(eid=key, name=key, desc="", tactics=[], is_sub="." in key,
                                                           parent=key.split(".")[0], url="")


def name_ko(key):
    if key in EMB3D_KEYS:
        return R.EMB3D_NEW[key][3]
    return R.NAME_KO.get(key) or R.EMB3D_GROUP_NAME.get(key) or tech_info(key)["name"]


# ===========================================================================
# 2) 근거 집계
# ===========================================================================
def empty_ev():
    return dict(real=set(), recent=set(), research=set(), intel=set(), excluded=set(), subj=set(),
                basis=collections.defaultdict(set), cves=set(), advs=set(), recent_cves=set(), kev=set(),
                iot_inc=set(), emb3d={})


ev = collections.defaultdict(empty_ev)

# ATT&CK 절차 — ICS는 모든 주체, Enterprise는 ICS 캠페인·소프트웨어의 절차만
for tid, lst in ics.procs.items():
    k = norm(tid, "ICS 절차")
    for subj, _ in lst:
        ev[k]["subj"].add(subj)
ent_seen = collections.defaultdict(set)   # Enterprise 기법 → 근거 주체(판정표 출력용)
for tid, lst in ent.procs.items():
    k = norm(tid, "Enterprise 절차(" + ",".join(sorted({s for s, _ in lst})) + ")")
    for subj, _ in lst:
        ent_seen[tid].add(subj)
        if k:
            ev[k]["subj"].add(subj)


def incident_status(e):
    if e["kind"] in R.RESEARCH_KINDS:
        return "실증·연구"
    if e["kind"] in R.INTEL_KINDS:
        return "위협인텔"
    if e["verification"] in R.EXCLUDED_VERIFICATION:
        return "제외(검증)"
    if e["relevance"] in R.EXCLUDED_RELEVANCE:
        return "제외(IT 한정)"
    return "실제 사고"


for e in incidents:
    links = {}
    for ref in e.get("attack_ref") or []:
        rid, _, dom = ref.partition("@")   # 'S0608@ICS' = 소프트웨어의 ICS 절차만(일반 기능인 Enterprise 절차 제외)
        for t in ics.subject_uses.get(rid, []) + ([] if dom == "ICS" else ent.subject_uses.get(rid, [])):
            k = norm(t, e["id"])
            if k:
                links.setdefault(k, "ATT&CK 공식")
            elif t not in R.ENT_DROP and t not in ics.tech:
                ent_seen[t].add(rid)
    for t in (e.get("ics") or []) + (e.get("ent") or []):
        k = norm(t, e["id"])
        if k and k not in links:
            links[k] = "분석"
        if t.startswith("T1") and not (t.startswith("T169")):
            ent_seen[t].add(e["id"])
    e["links"] = links
    e["status"] = incident_status(e)
    e["recent"] = str(e["date"]) >= R.RECENT_FROM
    for k, why in links.items():
        x = ev[k]
        x["basis"][why].add(e["id"])
        if e["status"] == "실제 사고":
            x["real"].add(e["id"])
            if e["recent"]:
                x["recent"].add(e["id"])
            if e["relevance"] == "IoT 기기":
                x["iot_inc"].add(e["id"])
        elif e["status"] == "실증·연구":
            x["research"].add(e["id"])
        elif e["status"] == "위협인텔":
            x["intel"].add(e["id"])
        else:
            x["excluded"].add(e["id"])


# 공개 취약점(CWE 규칙) · KEV
def cwe_rule(cwe, av):
    for tid, cwes, avs, _ in R.CWE_RULES:
        if cwe in cwes and (avs is None or av in avs):
            return tid
    return None


def kev_class(r):
    cve, kv, prod = r["CVE"], r["KEV 벤더"], r["KEV 제품"] or ""
    if cve in R.KEV_OT:
        return R.KEV_OT[cve]
    if kv in R.NETWORK_OS_KEV:
        t = R.NETWORK_OS_KEV[kv]
        if any(x in prod for x in R.NETWORK_OS_EXPOSED):
            t = "T0819"
        return ("OT 네트워크 장비 내장 OS", t, f"{kv} {prod} — {r['권고명']}")
    if kv in R.KEV_IT_VENDORS:
        return ("범용 IT 구성요소", None, f"{kv} {prod}")
    WARN[f"KEV 미판정 {cve}"].add(kv)
    return ("미판정", None, f"{kv} {prod}")


vuln_rows = [r for r in csv.DictReader(open(VULN_CSV, encoding="utf-8-sig"))]
vstat = dict(adv_all=len({r["권고 ID"] for r in vuln_rows}),
             adv_icsma=len({r["권고 ID"] for r in vuln_rows if r["권고 ID"].startswith("ICSMA")}))
ics_rows = [r for r in vuln_rows if r["권고 ID"].startswith("ICSA")]
vstat.update(adv=len({r["권고 ID"] for r in ics_rows}), cve=len({r["CVE"] for r in ics_rows if r["CVE"]}))
cve_cwe_tech = {}
kev_table = {}
for r in ics_rows:
    cve = r["CVE"]
    if not cve:
        continue
    t = cwe_rule(r["CWE"], r["공격 경로(AV)"])
    if t and cve not in cve_cwe_tech:
        cve_cwe_tech[cve] = t
    if t and cve_cwe_tech.get(cve) == t:
        x = ev[norm(t)]
        x["cves"].add(cve)
        x["advs"].add(r["권고 ID"])
        if r["공개일"] >= R.RECENT_FROM:
            x["recent_cves"].add(cve)
    if r["KEV"] == "Y" and cve not in kev_table:
        cat, t2, why = kev_class(r)
        kev_table[cve] = dict(cve=cve, vendor=r["KEV 벤더"], product=r["KEV 제품"], adv=r["권고 ID"], adv_title=r["권고명"],
                              cat=cat, tech=norm(t2) if t2 else None, why=why, cwe=r["CWE"], added=r["KEV 등재일"])
        if t2:
            ev[norm(t2)]["kev"].add(cve)
vstat["mapped_cve"] = len(cve_cwe_tech)
vstat["kev_all"] = len(kev_table)
vstat["kev_ot"] = sum(1 for v in kev_table.values() if v["tech"])


# MITRE EMB3D — 장치 위협(TID) ↔ 행: 신설 > EMB3D 인용(근거 문단의 ATT&CK ID) > 분석(EMB3D_MAP)
def resolve_cited(t):
    """EMB3D가 인용한 ATT&CK ID → 행 키(폐기 ID·Enterprise 통합 규칙 적용). 매트릭스 밖 ID는 None(경고 없음)"""
    t = ics.revoked.get(t, t)
    seen = set()
    while t in R.ENT_MERGE and t not in seen:
        seen.add(t)
        t = ics.revoked.get(R.ENT_MERGE[t], R.ENT_MERGE[t])
    return t if (t in ics.tech or t in R.ENT_INCLUDE) else None


# OWASP IoT Top 10 → 행 (분석 교차 매핑, 근거 집계에는 넣지 않음)
owasp_of = collections.defaultdict(list)       # 행 키 → [OWASP ID]
for oid, en, ko, desc, keys in R.OWASP_IOT:
    for k in keys:
        owasp_of[k].append(oid)

emb3d_links = collections.defaultdict(dict)    # TID → {행 키: 연계 근거}
emb3d_outside = collections.defaultdict(list)  # TID → 매트릭스 밖 인용 ID
for tid, (keys, _) in R.EMB3D_MAP.items():
    for k in keys:
        emb3d_links[tid][k] = "분석"
for tid, t in E3T.items():
    for c in t["atk_cited"]:
        k = resolve_cited(c)
        if k:
            emb3d_links[tid][k] = "EMB3D 인용"
        else:
            emb3d_outside[tid].append(c)
for k, (_, _, _, _, tids) in R.EMB3D_NEW.items():
    for tid in tids:
        emb3d_links[tid][k] = "신설"
for tid, links in emb3d_links.items():
    if tid not in E3T:
        WARN[f"EMB3D 추출본에 없는 위협 {tid}"].add("taxonomy_rules")
        continue
    for k, basis in links.items():
        ev[k]["emb3d"][tid] = basis


def emb3d_obs(x):
    return [t for t in x["emb3d"] if E3T[t]["maturity"] == "관찰"]


def emb3d_judgement(tid):
    """EMB3D 위협 판정 (판정, 사유) — 신설 > 공식 연계 > 분석 연계 > 미연계"""
    bases = set(emb3d_links.get(tid, {}).values())
    if "신설" in bases:
        return "신설", "ATT&CK에 대응 기법이 없는 장치 하드웨어 위협 — 효과가 나타나는 전술 아래 세부위협 신설"
    if "EMB3D 인용" in bases:
        why = "EMB3D 근거 문단이 인용한 ATT&CK 기법"
        return "공식 연계", why + (" + 분석: " + R.EMB3D_MAP[tid][1] if tid in R.EMB3D_MAP else "")
    if "분석" in bases:
        return "분석 연계", R.EMB3D_MAP[tid][1]
    return "미연계", ""


# ===========================================================================
# 3) 골격 조립 (전술 → 기법(Lv2) → 하위기법(Lv3))
# ===========================================================================
def has_ev(k):
    x = ev.get(k)
    return bool(x and (x["real"] or x["research"] or x["intel"] or x["subj"] or x["cves"] or x["kev"] or x["excluded"]
                       or x["emb3d"]))


def tactic_items(code, short):
    items = []
    if code != "RD":
        parents = sorted(e for e, t in ics.tech.items() if not t["is_sub"] and short in t["tactics"])
        for p in parents:
            subs = sorted(e for e, t in ics.tech.items() if t["is_sub"] and t["parent"] == p and short in t["tactics"])
            leaves = ([p] if has_ev(p) else []) + subs if subs else [p]
            items.append(("ICS", p, leaves, bool(subs)))
    by_parent = collections.defaultdict(list)
    for k, tacs in R.ENT_INCLUDE.items():
        if code in tacs:
            by_parent[k.split(".")[0]].append(k)
    for p in sorted(by_parent):
        keys = by_parent[p]
        subs = sorted(k for k in keys if k != p)
        if subs:
            leaves = ([p] if p in keys and has_ev(p) else []) + subs
        else:
            leaves = [p]
        items.append(("Enterprise 편입", p, leaves, bool(subs)))
    groups = collections.defaultdict(list)
    for k, (c, g, *_) in R.EMB3D_NEW.items():
        if c == code:
            groups[g].append(k)
    for g in sorted(groups):
        ks = sorted(groups[g])
        items.append(("EMB3D 신설", g, ks, ks != [g]))
    return items


registry = yaml.safe_load(REGISTRY.read_text(encoding="utf-8")) if REGISTRY.exists() else {}
registry.setdefault("lv2", {})
registry.setdefault("lv3", {})


def lv2_number(code, p):
    key = f"{code}|{p}"
    if key not in registry["lv2"]:
        used = [v for k, v in registry["lv2"].items() if k.startswith(code + "|")]
        registry["lv2"][key] = max(used, default=0) + 1
    return registry["lv2"][key]


def lv3_number(code, p, leaf):
    if leaf == p:
        return 0
    key = f"{code}|{leaf}"
    if key not in registry["lv3"]:
        used = [v for k, v in registry["lv3"].items() if k.startswith(code + "|" + p + ".")]
        registry["lv3"][key] = max(used, default=0) + 1
    return registry["lv3"][key]


TEXT = {}
for f in sorted(TEXT_DIR.glob("*.yaml")) if TEXT_DIR.exists() else []:
    TEXT.update(yaml.safe_load(f.read_text(encoding="utf-8")) or {})


def text_for(key, code):
    return TEXT.get(f"{key}@{code}") or TEXT.get(key)


merged_into = collections.defaultdict(list)
for e_id in sorted(R.ENT_MERGE):
    k = norm(e_id, "통합 규칙")
    if k:
        merged_into[k].append(e_id)


def severity(key, code):
    for k in (key, key.split(".")[0]):
        if k in R.SEVERITY:
            return R.SEVERITY[k]
    return R.TACTIC_SEVERITY[code], f"전술 기준값({R.TACTIC_BY_CODE[code][2]}) — {R.TACTIC_SEVERITY_REASON[code]}"


def likelihood(x):
    real = len(x["real"])
    if real >= LIKELY_HIGH_REAL:
        return "상"
    if real or x["subj"] or x["kev"] or x["intel"] or x["cves"] or x["research"] or x["emb3d"]:
        return "중"
    return "하"


def evidence_level(x):
    if x["real"]:
        return "실제 사고 확인"
    if x["subj"] or x["kev"] or x["intel"] or emb3d_obs(x):
        return "실사용 기법 포함"
    if x["research"] or x["cves"] or x["emb3d"]:
        return "실증·공개 취약점"
    return "이론·시나리오"


def assets_of(key, src):
    if src == "ICS":
        a = ics.targets.get(key) or ics.targets.get(key.split(".")[0]) or []
    elif src == "EMB3D 신설":
        a = R.EMB3D_ASSETS
    else:
        a = R.ENT_ASSETS.get(key.split(".")[0], [])
    return sorted(set(a))


def emb3d_mids(tids):
    return sorted({m for t in tids if t in E3T for m in E3T[t]["mids"]})


def _iec_sort(x):      # 'CR 3.14 RE 1' → (3, 14, 1) 정렬용
    n = [int(v) for v in re.findall(r"\d+", x)]
    return (n + [0, 0, 0])[:3]


def standards_of(key, src, mitig_ids):
    """행의 완화책에서 IEC 62443(3-3·4-2)·NIST SP 800-53 요구사항을 모음.
    ATT&CK 완화책은 STIX labels, EMB3D 신설 행은 EMB3D 완화책의 62443-4-2 매핑."""
    iec33, iec42, nist = set(), set(), set()
    if src == "EMB3D 신설":
        for m in emb3d_mids(R.EMB3D_NEW[key][4]):
            for rid, _ in (E3M.get(m, {}).get("iec62443_4_2") or []):
                iec42.add(rid)
    else:
        src_tbl = ics.mitig if src == "ICS" else ent.mitig
        for m in (src_tbl.get(key) or src_tbl.get(key.split(".")[0]) or []):
            std = ics.mitig_std.get(m)      # Enterprise 완화책은 표준 라벨이 없어 ICS 표에서만
            if std:
                iec33.update(std["iec33"])
                iec42.update(std["iec42"])
                nist.update(std["nist"])
    return (sorted(iec33, key=_iec_sort), sorted(iec42, key=_iec_sort), sorted(nist))


def mitigations_of(key, src):
    if src == "EMB3D 신설":
        return [f"EMB3D {m} {E3M[m]['name']}" for m in emb3d_mids(R.EMB3D_NEW[key][4]) if m in E3M]
    if src == "ICS":
        ids = ics.mitig.get(key) or ics.mitig.get(key.split(".")[0]) or []
        return [f"{m} {ics.mitigations[m][0]}" for m in sorted(set(ids)) if m in ics.mitigations]
    ids = ent.mitig.get(key) or ent.mitig.get(key.split(".")[0]) or []
    return [f"{m} {ent.mitigations[m][0]}" for m in sorted(set(ids)) if m in ent.mitigations]


def subj_label(s):
    if s in ics.subjects:
        return f"{s} {ics.subjects[s]['name']}"
    return s


rows, domains = [], []
for code, short, ko, en, tac_id, tac_desc in R.TACTICS:
    items = tactic_items(code, short)
    lv2_count = 0
    for src, p, leaves, has_subs in items:
        n = lv2_number(code, p)
        lv2_count += 1
        for leaf in leaves:
            m = lv3_number(code, p, leaf)
            otc = f"OTC-{code}-{n:02d}" + (f".{m}" if has_subs else "")
            general = has_subs and leaf == p
            x = ev[leaf]
            sev, sev_why = severity(leaf, code)
            lk = likelihood(x)
            assets = assets_of(leaf, src)
            purdue = sorted({lv for a in assets for lv in R.ASSET_PURDUE.get(a, [])}, key=R.PURDUE_ORDER.index)
            prof = {}
            for pname, _, pas in R.PROFILES:
                if pname == "IoT·임베디드":
                    on = leaf in R.IOT_APPLICABLE or p in R.IOT_APPLICABLE or bool(x["iot_inc"]) or bool(x["emb3d"])
                else:
                    on = any(a in pas for a in assets)
                prof[pname] = "●" if on else ""
            txt = text_for(leaf, code)
            info = tech_info(leaf)
            lv3_name = (name_ko(p) + " (일반·상위기법)") if general else name_ko(leaf)
            if txt and txt.get("name"):
                lv3_name = txt["name"]
            basis_n = {b: len(v & (x["real"] | x["research"] | x["intel"])) for b, v in x["basis"].items()}
            inc_ids = sorted(x["real"])
            extra = []
            if x["research"]:
                extra.append("실증: " + ", ".join(sorted(x["research"])))
            if x["intel"]:
                extra.append("위협인텔: " + ", ".join(sorted(x["intel"])))
            if x["excluded"]:
                extra.append("집계 제외: " + ", ".join(sorted(x["excluded"])))
            rows.append(dict(
                code=code, tactic=en, domain=f"[{code}] {ko}", otc=otc, lv2=name_ko(p), lv3=lv3_name, key=leaf, parent=p,
                general=general, src=src, en_name=info["name"] + (" (일반)" if general else ""),
                merged=", ".join(merged_into.get(leaf, [])),
                summary=(txt or {}).get("summary") or "", reference=(txt or {}).get("reference") or "",
                detect=(txt or {}).get("detect") or "", text_src="분석가 작성" if txt else "작성 예정",
                placeholder=first_sentence(info.get("desc", "")),
                assets=assets, purdue=purdue, prof=prof, mitig=mitigations_of(leaf, src),
                std=standards_of(leaf, src, None),
                cloud=R.CLOUD_LINK.get(leaf) or R.CLOUD_LINK.get(p, ""), ai=R.AI_LINK.get(leaf) or R.AI_LINK.get(p, ""),
                owasp=", ".join(owasp_of.get(leaf, [])),
                likelihood=lk, severity=sev, sev_why=sev_why, risk=RISK_MATRIX[(lk, sev)],
                real=len(x["real"]), recent=len(x["recent"]), atk=len(x["subj"]), cves=len(x["cves"]),
                recent_cves=len(x["recent_cves"]), advs=len(x["advs"]), kev=len(x["kev"]), research=len(x["research"]),
                intel=len(x["intel"]), level=evidence_level(x),
                emb3d="\n".join(f"{t} {R.EMB3D_NAME_KO.get(t, E3T[t]['name'])} ({E3T[t]['maturity']}·{b})"
                                for t, b in sorted(x["emb3d"].items())),
                n_emb3d=len(x["emb3d"]),
                lk_why=(f"실제 사고 {len(x['real'])}건(최근 {len(x['recent'])}) · ATT&CK 사례 {len(x['subj'])}건 · "
                        f"공개 취약점 {len(x['cves'])}건 · KEV {len(x['kev'])}건 · 실증·연구 {len(x['research'])}건"
                        + (f" · 위협인텔 {len(x['intel'])}건" if x["intel"] else "")
                        + (f" · EMB3D 위협 {len(x['emb3d'])}건(관찰 {len(emb3d_obs(x))})" if x["emb3d"] else "")),
                basis=" · ".join(f"{b} {c}" for b, c in sorted(basis_n.items()) if c),
                inc_ids=", ".join(inc_ids) + (("\n" + "\n".join(extra)) if extra else ""),
                subjects=", ".join(subj_label(s) for s in sorted(x["subj"])),
                kev_ids=", ".join(sorted(x["kev"])),
                link=info.get("url") or ""))
    domains.append(dict(code=code, ko=ko, en=en, tac_id=tac_id, desc=tac_desc, lv2=lv2_count))

REGISTRY.write_text(yaml.safe_dump(dict(lv2=dict(sorted(registry["lv2"].items())), lv3=dict(sorted(registry["lv3"].items()))),
                                   allow_unicode=True, sort_keys=False), encoding="utf-8")


# ===========================================================================
# 4) xlsx 출력 — 클라우드 v5 서식 준용
# ===========================================================================
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "실사용 기법 포함": "2E86C1",
               "실증·공개 취약점": "B9770E", "이론·시나리오": "7F8C8D"}
SRC_COLOR = {"ICS": "D6EAF8", "Enterprise 편입": "FEF9E7", "EMB3D 신설": "E8F8F5"}
SRC_MARK = {"ICS": "", "Enterprise 편입": "◆ ", "EMB3D 신설": "◇ "}
HDR, NAVY, GRAY = "34495E", "1F3864", "8C8C8C"


def _w(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hdr(ws, row, cols, start=1):
    for j, c in enumerate(cols, start):
        cell = ws.cell(row, j, c)
        cell.fill = PatternFill("solid", fgColor=HDR)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BORDER


def _risk(cell, risk):
    bg, fg = RISK_COLOR[risk]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=risk in ("매우 높음", "높음"), color=fg)
    cell.alignment = CENTER


def _level(cell, level):
    cell.font = Font(bold=True, color=LEVEL_COLOR[level])
    cell.alignment = CENTER


def _row(ws, r, vals, start=1):
    for j, v in enumerate(vals, start):
        c = ws.cell(r, j, v)
        c.alignment = WRAP
        c.border = BORDER


def _title(ws, text, sub=None, span=10):
    c = ws.cell(1, 1, text)
    c.font = Font(size=14, bold=True, color=NAVY)
    if sub:
        ws.cell(2, 1, sub).font = Font(size=9, color="595959")


def summary_cell(o):
    if o["summary"]:
        return o["summary"]
    return f"(작성 예정) {o['placeholder']}"


def section(ws, r, text):
    c = ws.cell(r, 1, text)
    c.font = Font(bold=True, size=11, color=NAVY)
    return r + 1


def write_xlsx(path):
    wb = openpyxl.Workbook()
    today = datetime.date.today().isoformat()
    log = yaml.safe_load(CHANGELOG.read_text(encoding="utf-8")) if CHANGELOG.exists() else {}
    rc = collections.Counter(o["risk"] for o in rows)
    lc = collections.Counter(o["level"] for o in rows)
    keys = {o["key"] for o in rows}
    written = {o["key"] for o in rows if o["text_src"] == "분석가 작성"}
    counted = [e for e in incidents if e["status"] == "실제 사고"]
    mapped = [e for e in counted if e["links"]]
    # 근거 편중: ATT&CK 공식 절차가 촘촘한 대형 사고 몇 건이 여러 기법의 '실제 사고' 근거를 떠받치는 정도
    top = sorted((e for e in incidents if e["status"] == "실제 사고"), key=lambda e: -len(e["links"]))[:5]
    top_ids = {e["id"] for e in top}
    real_keys = [k for k in keys if ev[k]["real"]]
    only_top = sum(1 for k in real_keys if ev[k]["real"] <= top_ids)
    n_high = sum(1 for o in rows if o["likelihood"] == "상")
    high_dep = sum(1 for o in rows if o["likelihood"] == "상" and len(ev[o["key"]]["real"] - top_ids) < LIKELY_HIGH_REAL)
    n_ent_rows = sum(1 for o in rows if o["src"] == "Enterprise 편입")
    ent_keys = sorted({o["parent"] for o in rows if o["src"] == "Enterprise 편입"})
    n_e3_rows = sum(1 for o in rows if o["src"] == "EMB3D 신설")
    e3_judge = collections.Counter(emb3d_judgement(t)[0] for t in E3T)
    e3_rows_linked = sum(1 for o in rows if o["n_emb3d"])
    std_req = {i: {req for o in rows for req in o["std"][i]} for i in range(3)}
    std_rows_n = sum(1 for o in rows if any(o["std"]))

    # ---------------- 개요 ----------------
    ws = wb.active
    ws.title = "개요"
    meta = [
        (f"통합 OT/ICS/IoT 보안위협 매트릭스 {VERSION}", ""),
        ("기준", "MITRE ATT&CK for ICS v19.2 킬체인(전술→기법→하위기법) 뼈대 · 통합 AI 매트릭스 v3.2 · 클라우드 매트릭스 v5의 위험평가·근거 로직 이식"),
        ("생성일", today), ("", ""),
        ("■ 기준 데이터", ""),
        ("MITRE ATT&CK for ICS", f"v{ics.version} — 전술 12 · 기법 {sum(1 for t in ics.tech.values() if not t['is_sub'])} · "
                                f"하위기법 {sum(1 for t in ics.tech.values() if t['is_sub'])} · 자산 {len(ics.assets)} · "
                                f"캠페인·소프트웨어·그룹 {len(ics.subjects)} · 완화책 {len(ics.mitigations)}"),
        ("MITRE ATT&CK Enterprise", f"v19.2 — ICS 캠페인·소프트웨어의 Enterprise 절차로 IT/OT 경계 기법 판정(편입 {len(ent_keys)}개 기법 · "
                                   f"{n_ent_rows}행, 통합 {len(R.ENT_MERGE)}개, 제외 {len(R.ENT_DROP)}개)"),
        ("OT 보안사고 DB", f"{len(incidents)}건(2000~2025, 공개 출처) — 실제 사고 집계 {len(counted)}건, 기법 매핑 {len(mapped)}건 · "
                          f"실증·연구 {sum(1 for e in incidents if e['status'] == '실증·연구')}건 · 위협인텔 {sum(1 for e in incidents if e['status'] == '위협인텔')}건 · "
                          f"집계 제외 {sum(1 for e in incidents if e['status'].startswith('제외'))}건"),
        ("CISA ICS 권고(CSAF)", f"ICS 권고 {vstat['adv']}건 · CVE {vstat['cve']}개(의료기기 권고 {vstat['adv_icsma']}건 제외) — CWE 규칙으로 기법 매핑 {vstat['mapped_cve']}개"),
        ("CISA KEV", f"ICS 권고와 교차 {vstat['kev_all']}개 — OT·IoT 제품·임베디드 구성요소·OT 네트워크 장비 OS {vstat['kev_ot']}개만 실사용 근거로 반영(범용 IT 구성요소 제외)"),
        ("MITRE EMB3D", f"{emb3d.get('source', {}).get('file', '')} — 임베디드 장치 위협 {len(E3T)}개 · 완화책 {len(E3M)}개: "
                        f"공식 연계 {e3_judge.get('공식 연계', 0)} · 분석 연계 {e3_judge.get('분석 연계', 0)} · 신설 {e3_judge.get('신설', 0)}"
                        f"(세부위협 {n_e3_rows}행) · 미연계 {e3_judge.get('미연계', 0)}"),
        ("다른 매트릭스", "통합 클라우드 매트릭스 v5(CTC) · 통합 AI 매트릭스 v3.2(UT) 연계 열"), ("", ""),
        ("■ 시트 구성", ""),
        ("매트릭스 뷰", "전술(열)별 세부위협을 위험도 색으로 배치한 한눈 보기(ATT&CK for ICS Navigator와 같은 배치)"),
        ("통합 매트릭스", "분류체계 · 교차매핑(자산·Purdue·적용 프로파일·완화책·타 매트릭스) · 위험평가 · 실제근거 · 탐지·대응 전체 열"),
        ("통합매트릭스_LITE", "핵심 열 발췌(필터·보고용) + 탐지·대응 포인트"),
        ("도메인 요약", "전술별 위협 수 · 위험도/근거수준 분포 · 매핑 사고 수 · 최고위험 항목"),
        ("자산·계층 요약", "ATT&CK ICS 자산 18종 · Purdue 계층 · 적용 프로파일별 위험 분포"),
        ("업종 요약", "업종(전력·수처리·제조·석유가스·교통·식품·IoT·빌딩)별 실제 사고·영향·대표 세부위협"),
        ("역매핑_사고사례", f"사고 DB {len(incidents)}건과 매핑 기법·매핑 근거(근거 추적용)"),
        ("취약점 근거", "기법별 공개 취약점(CWE 규칙)·KEV 집계, KEV 판정 내역, CWE 규칙"),
        ("Enterprise 판정", "ICS 캠페인·소프트웨어·사고 DB에 등장한 Enterprise 기법의 편입·통합·제외 판정과 근거"),
        ("EMB3D 판정", "EMB3D 장치 위협 전체의 연계(공식 인용·분석)·신설 판정, 성숙도, CWE·완화책"),
        ("표준 연계", "IEC 62443-3-3·4-2·NIST SP 800-53 요구사항 ↔ 세부위협(완화책의 표준 매핑 집계, 명칭 포함, 원문 미수록)"),
        ("OWASP IoT Top10", "OWASP IoT Top 10(2018) 위험 범주 ↔ 세부위협(분석 교차 매핑)"),
        ("공격 체인 시나리오", "실제 사고 기반 IT 침투→OT 영향 흐름 — 단계별 OTC-ID 연결·탐지 포인트·초크 포인트"),
        ("평가 기준", "발생가능성·심각도·위험도·근거수준 산정 규칙과 AI·클라우드 매트릭스 근거 대응"),
        ("변경이력", "버전별 변경 내역"), ("", ""),
        ("■ 분류 체계", ""),
        ("Lv1 도메인", "ATT&CK for ICS 전술 12개 + 사전 단계 [RD] 자원 개발. 예: [IR] 대응 기능 억제, [IP] 공정 제어 훼손"),
        ("Lv2 위협분류", "ATT&CK 기법(ICS T08xx·T169x, 편입 Enterprise T1xxx). 예: OTC-IP-02 = T0836 운전 파라미터 변조"),
        ("Lv3 세부위협", "하위기법 또는 기법 자체 — 요약설명 · 참조(OT 관점 · 실제 사례) · 탐지·대응 포인트. "
                       "'.0 (일반·상위기법)'은 하위기법으로 특정되지 않은 상위기법 근거를 보존한 행. ID는 data/id_registry.yaml로 고정(폐지 ID 재사용 안 함)"),
        ("OT 특화 검토", "Enterprise 기법은 ICS 기법으로 표현되면 통합(근거 합산), 표현되지 않는 IT/OT 경계 메커니즘만 원래 전술에 대응하는 ICS 전술 아래 편입. "
                       "자격증명 접근은 [LM], 반출은 운영 정보 탈취(T0882)에 통합 — 판정은 'Enterprise 판정' 시트"),
        ("적용 프로파일", "제어계통 · 안전계통 · 원격 필드 · 감시·운영 · IT/OT 경계 = 기법의 대상 자산(ATT&CK 'targets')에서 자동 부여, "
                        "IoT·임베디드 = EMB3D 연계 행 + 임베디드 기기 일반에 성립하는 기법(분석자 지정) + IoT 사고 매핑"),
        ("EMB3D 연계", "임베디드 장치 위협(TID)을 EMB3D가 인용한 ATT&CK 기법(공식)이나 정의가 같은 기법(분석)에 연계하고, "
                     "ATT&CK에 대응 기법이 없는 하드웨어 위협(디버그 포트·결함 주입·물리 추출·부채널·신뢰 루트)은 세부위협으로 신설(◇) — 판정은 'EMB3D 판정' 시트"),
        ("", ""), ("■ 결과 요약", ""),
        ("세부위협(Lv3)", f"{len(rows)}행 / 고유 기법 {len(keys)}개 / 도메인 {len(domains)}개 "
                       f"(ICS {len(rows) - n_ent_rows - n_e3_rows}행 · Enterprise 편입 {n_ent_rows}행 · EMB3D 신설 {n_e3_rows}행)"),
        ("EMB3D 연계 행", f"{e3_rows_linked}행에 EMB3D 위협이 연계됨(IoT·임베디드 프로파일 부여)"),
        ("표준 연계", f"{std_rows_n}행에 표준 요구사항 매핑 — IEC 62443-3-3 {len(std_req[0])}개 · 62443-4-2 {len(std_req[1])}개 · NIST SP 800-53 {len(std_req[2])}개"),
        ("위험도", " · ".join(f"{k} {rc.get(k, 0)}" for k in RISK_COLOR)),
        ("근거 수준", " · ".join(f"{k} {lc.get(k, 0)}" for k in LEVEL_COLOR)),
        ("OT 관점 문구", f"고유 기법 {len(written & keys)}/{len(keys)}개 작성" + (" — 전 기법 완료" if written >= keys else
                                                                       " — 나머지는 ATT&CK 원문 첫 문장(회색, '작성 예정')")),
        ("", ""), ("■ 주의", ""),
        ("발생가능성", f"공개 OT 사고는 적고 과소보고되므로 '중' 이하에 몰리는 경향 — '이론·시나리오'는 '발생하지 않음'이 아니라 '공개 근거 없음'. "
                     f"수식은 AI·클라우드 매트릭스와 같게 유지(상=실제 사고 {LIKELY_HIGH_REAL}건 이상)"),
        ("근거 해석", "Program Download·Modify Parameter처럼 설계상 인증이 없는 기능의 악용은 CVE가 없어 '공개 취약점' 경로로 근거가 쌓이지 않음. "
                    "IT 랜섬웨어로 인한 예방적 OT 중단이 가장 잦은 실제 영향이라 [IM] 생산·매출 손실 등에 사고가 몰림('OT 관련도' 열로 구분)"),
        ("근거 편중", f"ATT&CK 공식 절차가 촘촘한 대형 사고 5건({'·'.join(e['id'] for e in top)})이 각각 "
                    f"{min(len(e['links']) for e in top)}~{max(len(e['links']) for e in top)}개 기법에 매핑됨 — 실제 사고 근거가 있는 기법 "
                    f"{len(real_keys)}개 중 {only_top}개는 이 5건에만 의존하고, 발생가능성 '상' {n_high}행 중 {high_dep}행은 이 5건이 없으면 '상' 기준에 못 미침. "
                    "'사고 매핑 근거' 열(ATT&CK 공식·분석)과 실제 사고 수로 근거의 폭을 함께 볼 것"),
        ("사고 집계", "행위자 주장만 있거나 원인이 번복된 사례는 실제 사고에서 제외(역매핑 시트에 보존). 같은 사건은 출처가 여럿이어도 1건"),
        ("위험평가", "심각도는 OT 결과 기준의 기법별 기준값 — 업종·공정(화학·전력 vs 빌딩 공조)과 자산 중요도에 맞게 조정 권장"),
        ("중복 표시", "한 기법이 여러 전술에 속하면 전술마다 반복 표시(ATT&CK 원칙). 근거 수는 같은 기법 기준"),
        ("출처 표기", "MITRE ATT&CK® © The MITRE Corporation · CISA 권고·KEV(미국 정부 저작물) · 사고 출처는 '역매핑_사고사례' 시트. "
                    "IEC 62443 등 유료 표준은 원문 미수록"),
        ("EMB3D 표기", "©2026 The MITRE Corporation. This work is reproduced and distributed with the permission of The MITRE Corporation. "
                     "— EMB3D™ 이용 조건(내부 업무·상업적 이용, 사본에 저작권 표시·라이선스 문구 포함, EMB3D@mitre.org에 이용 통지)은 "
                     "data/emb3d.yaml 머리말 참고"),
    ]
    for i, (a, b) in enumerate(meta, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    ws.cell(1, 1).font = Font(size=15, bold=True, color=NAVY)
    _w(ws, [24, 130])

    # ---------------- 매트릭스 뷰 ----------------
    ws = wb.create_sheet("매트릭스 뷰")
    by_c = collections.defaultdict(list)
    for o in rows:
        by_c[o["code"]].append(o)
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도 (빨강 매우 높음 · 주황 높음 · 노랑 보통 · 초록 낮음), [n] = 실제 사고 수, "
                  "◆ = Enterprise 편입 기법, ◇ = EMB3D 신설(장치 하드웨어 위협)")
    ws.cell(1, 1).font = Font(bold=True, color=NAVY)
    rank = {"매우 높음": 0, "높음": 1, "보통": 2, "낮음": 3}
    for j, d in enumerate(domains, 1):
        c = ws.cell(2, j, f"[{d['code']}] {d['ko']}\n{d['en']}\n({len(by_c[d['code']])})")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        for i, o in enumerate(sorted(by_c[d["code"]], key=lambda o: (rank[o["risk"]], -o["real"], o["otc"])), 3):
            cell = ws.cell(i, j, SRC_MARK[o["src"]] + f"{o['otc']} {o['lv3']}" + (f" [{o['real']}]" if o["real"] else ""))
            _risk(cell, o["risk"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(size=9, bold=o["risk"] == "매우 높음", color=RISK_COLOR[o["risk"]][1])
            cell.border = BORDER
    ws.row_dimensions[2].height = 48
    ws.freeze_panes = "A3"
    _w(ws, [20] * len(domains))

    # ---------------- 통합 매트릭스 ----------------
    ws = wb.create_sheet("통합 매트릭스")
    prof_names = [p[0] for p in R.PROFILES]
    groups = [("분류 체계", 11, "2E5496"), ("교차 매핑", 10 + len(prof_names), "1F7A8C"),
              ("위험 평가", 5, "A04000"), ("실제 근거", 11, "1E8449"), ("탐지·대응", 1, "6C3483")]
    cols = (["도메인(Lv1)", "OTC-ID", "위협분류(Lv2)", "세부위협(Lv3)", "요약설명", "참조",
             "ATT&CK·EMB3D ID", "ATT&CK·EMB3D 위협명", "구분", "통합된 Enterprise 기법", "설명 출처",
             "대상 자산", "Purdue 계층"] + prof_names + ["완화책(ATT&CK·EMB3D)", "클라우드 매트릭스 연계", "AI 매트릭스 연계(UT)",
             "OWASP IoT Top10", "EMB3D 연계(성숙도·근거)", "IEC 62443-3-3(SR)", "IEC 62443-4-2(CR 등)", "NIST SP 800-53",
             "발생가능성", "심각도", "위험도", "발생가능성 근거 (자동 산정)", "심각도 근거",
             "실제 사고 수", "최근 사고(2025~)", "ATT&CK 사례 수", "공개 취약점(CVE)", "KEV(OT)", "실증·연구 수", "근거 수준",
             "사고 매핑 근거", "관련 사례 ID", "ATT&CK 사례 주체", "ATT&CK 링크", "탐지·대응 포인트"])
    t = ws.cell(1, 1, f"통합 OT/ICS/IoT 보안위협 매트릭스 {VERSION} — 분류체계 · 위험평가 · 실제근거")
    t.font = Font(size=13, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(2, 1, f"ATT&CK for ICS v{ics.version} + Enterprise 경계 기법 + OT 사고 DB {len(incidents)}건 + CISA ICS 권고 {vstat['adv']}건·KEV | {today}")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    c0 = 1
    for g, span, color in groups:
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + span - 1)
        c = ws.cell(3, c0, g)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = CENTER
        c0 += span
    _hdr(ws, 4, cols)
    risk_col = cols.index("위험도") + 1
    level_col = cols.index("근거 수준") + 1
    for r, o in enumerate(rows, 5):
        _row(ws, r, [o["domain"], o["otc"], o["lv2"], o["lv3"], summary_cell(o), o["reference"],
                     o["key"], o["en_name"], o["src"], o["merged"], o["text_src"],
                     "\n".join(f"{a} {R.ASSET_KO.get(a, a)}" for a in o["assets"]), ", ".join(o["purdue"])]
             + [o["prof"][p] for p in prof_names]
             + ["\n".join(o["mitig"]), o["cloud"], o["ai"], o["owasp"], o["emb3d"],
                ", ".join(o["std"][0]), ", ".join(o["std"][1]), ", ".join(o["std"][2]),
                o["likelihood"], o["severity"], o["risk"], o["lk_why"], o["sev_why"],
                o["real"], o["recent"], o["atk"], o["cves"], o["kev"], o["research"], o["level"],
                o["basis"], o["inc_ids"], o["subjects"], o["link"], o["detect"]])
        ws.cell(r, 9).fill = PatternFill("solid", fgColor=SRC_COLOR[o["src"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, 5).font = Font(color=GRAY)
        for j in range(14, 14 + len(prof_names)):
            ws.cell(r, j).alignment = CENTER
        for j in (risk_col - 2, risk_col - 1):
            ws.cell(r, j).alignment = CENTER
        _risk(ws.cell(r, risk_col), o["risk"])
        _level(ws.cell(r, level_col), o["level"])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(rows) + 4}"
    _w(ws, [14, 12, 16, 18, 60, 80, 10, 22, 10, 16, 9, 20, 10] + [6] * len(prof_names)
       + [26, 20, 18, 12, 30, 16, 20, 16, 7, 7, 8, 30, 26, 7, 7, 7, 7, 7, 7, 11, 14, 20, 26, 22, 70])

    # ---------------- 통합매트릭스_LITE ----------------
    ws = wb.create_sheet("통합매트릭스_LITE")
    ws.cell(1, 1, f"통합 OT/ICS/IoT 보안위협 매트릭스 {VERSION} — LITE").font = Font(size=13, bold=True, color=NAVY)
    lcols = ["도메인(Lv1)", "OTC-ID", "위협분류(Lv2)", "세부위협(Lv3)", "요약설명", "ATT&CK·EMB3D ID", "Purdue 계층", "적용 프로파일",
             "발생가능성", "심각도", "위험도", "근거 수준", "실제 사고 수", "ATT&CK 사례 수", "탐지·대응 포인트"]
    _hdr(ws, 2, lcols)
    for r, o in enumerate(rows, 3):
        _row(ws, r, [o["domain"], o["otc"], o["lv2"], o["lv3"], summary_cell(o), o["key"], ", ".join(o["purdue"]),
                     " · ".join(p for p in prof_names if o["prof"][p]), o["likelihood"], o["severity"], o["risk"],
                     o["level"], o["real"], o["atk"], o["detect"]])
        if o["text_src"] != "분석가 작성":
            ws.cell(r, 5).font = Font(color=GRAY)
        _risk(ws.cell(r, 11), o["risk"])
        _level(ws.cell(r, 12), o["level"])
        for j in (9, 10, 13, 14):
            ws.cell(r, j).alignment = CENTER
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(lcols))}{len(rows) + 2}"
    _w(ws, [14, 12, 16, 18, 60, 10, 12, 22, 7, 7, 8, 11, 7, 7, 70])

    # ---------------- 도메인 요약 ----------------
    ws = wb.create_sheet("도메인 요약")
    dcols = ["도메인", "ATT&CK 전술", "전술 설명", "위협분류(Lv2)", "세부위협(Lv3)", "매우 높음", "높음", "보통", "낮음",
             "실제 사고 확인", "실사용 기법", "실증·공개", "이론", "매핑 사고 수(중복 제거)", "최고위험 세부위협(상위 3)"]
    _hdr(ws, 1, dcols)
    for i, d in enumerate(domains, 2):
        os_ = by_c[d["code"]]
        rr = collections.Counter(o["risk"] for o in os_)
        ll = collections.Counter(o["level"] for o in os_)
        incs = set()
        for o in os_:
            incs |= ev[o["key"]]["real"]
        top = sorted(os_, key=lambda o: (rank[o["risk"]], -o["real"], -o["atk"], o["otc"]))[:3]
        _row(ws, i, [f"[{d['code']}] {d['ko']}", f"{d['tac_id']} {d['en']}", d["desc"], d["lv2"], len(os_),
                     rr.get("매우 높음", 0), rr.get("높음", 0), rr.get("보통", 0), rr.get("낮음", 0),
                     ll.get("실제 사고 확인", 0), ll.get("실사용 기법 포함", 0), ll.get("실증·공개 취약점", 0), ll.get("이론·시나리오", 0),
                     len(incs), "\n".join(f"{o['otc']} {o['lv3']} ({o['risk']})" for o in top)])
    _w(ws, [18, 30, 46, 9, 9, 7, 7, 7, 7, 9, 9, 9, 7, 12, 50])

    # ---------------- 자산·계층 요약 ----------------
    ws = wb.create_sheet("자산·계층 요약")
    r = section(ws, 1, "1. ATT&CK for ICS 자산별 위협 분포")
    acols = ["자산 ID", "자산", "Purdue 계층", "적용 프로파일", "연결 세부위협 수", "매우 높음", "높음", "실제 사고 연결 세부위협", "최고위험 세부위협(상위 3)"]
    _hdr(ws, r, acols)
    r += 1
    for a in sorted(ics.assets):
        os_ = [o for o in rows if a in o["assets"]]
        rr = collections.Counter(o["risk"] for o in os_)
        top = sorted(os_, key=lambda o: (rank[o["risk"]], -o["real"], o["otc"]))[:3]
        prof = [p[0] for p in R.PROFILES if a in p[2]]
        _row(ws, r, [a, f"{ics.assets[a]} ({R.ASSET_KO.get(a, '')})", ", ".join(R.ASSET_PURDUE.get(a, [])), ", ".join(prof),
                     len(os_), rr.get("매우 높음", 0), rr.get("높음", 0), sum(1 for o in os_ if o["real"]),
                     "\n".join(f"{o['otc']} {o['lv3']} ({o['risk']})" for o in top)])
        r += 1
    r = section(ws, r + 1, "2. Purdue 계층 × 전술 (해당 계층 자산을 대상으로 하는 세부위협 수)")
    _hdr(ws, r, ["계층"] + [f"[{d['code']}]" for d in domains] + ["합계"])
    r += 1
    for lv in R.PURDUE_ORDER:
        cnt = [sum(1 for o in by_c[d["code"]] if lv in o["purdue"]) for d in domains]
        _row(ws, r, [lv] + cnt + [sum(cnt)])
        r += 1
    _row(ws, r, ["계층 미지정"] + [sum(1 for o in by_c[d["code"]] if not o["purdue"]) for d in domains]
         + [sum(1 for o in rows if not o["purdue"])])
    r = section(ws, r + 2, "3. 적용 프로파일별 위험 분포")
    _hdr(ws, r, ["프로파일", "정의", "부여 기준", "세부위협 수", "매우 높음", "높음", "실제 사고 확인", "최고위험 세부위협(상위 3)"])
    r += 1
    for pname, pdesc, pas in R.PROFILES:
        os_ = [o for o in rows if o["prof"][pname]]
        rr = collections.Counter(o["risk"] for o in os_)
        top = sorted(os_, key=lambda o: (rank[o["risk"]], -o["real"], o["otc"]))[:3]
        rule = ("대상 자산: " + ", ".join(f"{a} {R.ASSET_KO[a]}" for a in pas)) if pas else \
            "EMB3D 연계 행 + 임베디드 기기 일반에 성립하는 기법(분석자 지정) + IoT 사고 매핑"
        _row(ws, r, [pname, pdesc, rule, len(os_), rr.get("매우 높음", 0), rr.get("높음", 0),
                     sum(1 for o in os_ if o["level"] == "실제 사고 확인"),
                     "\n".join(f"{o['otc']} {o['lv3']} ({o['risk']})" for o in top)])
        r += 1
    _w(ws, [12, 30, 22, 30, 10, 9, 9, 12, 48] + [8] * 6)

    # ---------------- 업종 요약 ----------------
    ws = wb.create_sheet("업종 요약")
    ws.cell(1, 1, "업종 프로파일별 위험 요약 — 사고 DB의 업종(sector)을 업종군으로 묶어 실제 사고·영향·대표 세부위협을 집계").font = Font(bold=True, color=NAVY)
    key_row = {}        # 기법 키 → 대표 행(위험도 높은 쪽)
    for o in sorted(rows, key=lambda o: rank[o["risk"]]):
        key_row.setdefault(o["key"], o)
    sec_map = {}
    for name, secs, _ in R.SECTOR_PROFILES:
        for s in secs:
            sec_map[s] = name
    _hdr(ws, 2, ["업종", "실제 사고 수", "주요 영향 유형", "대표 세부위협(사고 수 상위 5)", "특성"])
    r = 3
    for name, secs, desc in R.SECTOR_PROFILES:
        incs = [e for e in incidents if e["status"] == "실제 사고" and e.get("sector") in secs]
        impact = collections.Counter(im for e in incs for im in e.get("impact", []))
        tech_cnt = collections.Counter()
        for e in incs:
            for k in {kk.split(".")[0] if kk.split(".")[0] in key_row else kk for kk in e["links"]}:
                tech_cnt[k] += 1
        top = [f"{key_row[k]['otc']} {key_row[k]['lv3']} ({n}건)" for k, n in tech_cnt.most_common(5) if k in key_row]
        _row(ws, r, [name, len(incs), ", ".join(f"{im}({n})" for im, n in impact.most_common(4)),
                     "\n".join(top), desc])
        r += 1
    # 기타(프로파일에 안 잡힌 sector)
    other = [e for e in incidents if e["status"] == "실제 사고" and e.get("sector") not in sec_map]
    if other:
        _row(ws, r, ["(기타)", len(other), "", "", "업종군 미분류: " + ", ".join(sorted({e.get("sector", "") for e in other}))])
    _w(ws, [18, 11, 40, 52, 60])

    # ---------------- 역매핑_사고사례 ----------------
    ws = wb.create_sheet("역매핑_사고사례")
    icols = ["사건ID", "사건명", "기준일", "사례유형", "검증 수준", "OT 관련도", "집계", "국가", "업종", "행위자",
             "물리적 영향", "매핑 기법(근거)", "사건 요약", "출처"]
    _hdr(ws, 1, icols)
    for i, e in enumerate(sorted(incidents, key=lambda e: e["id"]), 2):
        mp = "\n".join(f"{k} {name_ko(k)} ({why})" for k, why in sorted(e["links"].items()))
        _row(ws, i, [e["id"], e["title"], str(e["date"]), e["kind"], e["verification"], e["relevance"], e["status"],
                     e.get("country", ""), e.get("sector", ""), e.get("actor", ""), ", ".join(e.get("impact", [])),
                     mp, e["summary"], "\n".join(e.get("sources", []))])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(icols))}{len(incidents) + 1}"
    _w(ws, [9, 34, 9, 12, 14, 13, 10, 12, 12, 20, 18, 34, 70, 50])

    # ---------------- 취약점 근거 ----------------
    ws = wb.create_sheet("취약점 근거")
    r = section(ws, 1, f"1. 기법별 공개 취약점·KEV 근거 — CISA ICS 권고 {vstat['adv']}건 · CVE {vstat['cve']}개 중 CWE 규칙 매핑 {vstat['mapped_cve']}개")
    _hdr(ws, r, ["ATT&CK ID", "세부위협", "CWE 규칙", "권고 수", "CVE 수", "최근 CVE(2025~)", "KEV(OT) 수", "KEV(OT) CVE"])
    r += 1
    rule_of = {t: (", ".join(sorted(c)), why) for t, c, _, why in R.CWE_RULES}
    vkeys = sorted({k for k, x in ev.items() if k and (x["cves"] or x["kev"])})
    for k in vkeys:
        x = ev[k]
        rule = rule_of.get(k)
        _row(ws, r, [k, name_ko(k), (rule[1] + " — " + rule[0]) if rule else "(KEV 판정만)", len(x["advs"]), len(x["cves"]),
                     len(x["recent_cves"]), len(x["kev"]), ", ".join(sorted(x["kev"]))])
        r += 1
    r = section(ws, r + 1, f"2. ICS 권고와 교차된 KEV {vstat['kev_all']}개 판정 — 근거 반영 {vstat['kev_ot']}개")
    _hdr(ws, r, ["CVE", "KEV 벤더", "KEV 제품", "ICS 권고", "권고명", "구분", "근거 반영", "매핑 기법", "판정 근거", "CWE", "KEV 등재일"])
    r += 1
    corder = {"OT·IoT 제품": 0, "임베디드 구성요소": 1, "OT 네트워크 장비 내장 OS": 2, "범용 IT 구성요소": 3, "미판정": 4}
    for v in sorted(kev_table.values(), key=lambda v: (corder.get(v["cat"], 9), v["cve"])):
        _row(ws, r, [v["cve"], v["vendor"], v["product"], v["adv"], v["adv_title"], v["cat"], "Y" if v["tech"] else "N",
                     (f"{v['tech']} {name_ko(v['tech'])}" if v["tech"] else ""), v["why"], v["cwe"], v["added"]])
        r += 1
    r = section(ws, r + 1, "3. CWE → 기법 규칙 (한 CVE는 한 기법에만, 위에서부터 우선 적용)")
    _hdr(ws, r, ["순위", "기법", "CWE", "공격 경로 조건", "사유"])
    r += 1
    for i, (t, cwes, avs, why) in enumerate(R.CWE_RULES, 1):
        _row(ws, r, [i, f"{t} {name_ko(t)}", ", ".join(sorted(cwes)), ", ".join(sorted(avs)) if avs else "-", why])
        r += 1
    _w(ws, [16, 26, 40, 9, 9, 10, 9, 30, 50, 10, 11])

    # ---------------- Enterprise 판정 ----------------
    ws = wb.create_sheet("Enterprise 판정")
    ws.cell(1, 1, "Enterprise 기법 판정 — ICS 캠페인·소프트웨어의 Enterprise 절차와 사고 DB 분석 매핑에 등장한 기법").font = Font(bold=True, color=NAVY)
    _hdr(ws, 2, ["Enterprise ID", "기법명", "원래 전술", "판정", "대상(배치 전술 또는 통합 대상)", "근거 주체·사고", "사유"])
    rr_ = 3
    seen_ids = set(ent_seen) | set(R.ENT_INCLUDE) | set(R.ENT_MERGE) | set(R.ENT_DROP)
    seen_ids = {s for s in seen_ids if not s.startswith("T08") and not s.startswith("T169")}
    for e_id in sorted(seen_ids):
        info = ent.tech.get(e_id, {})
        if e_id in R.ENT_INCLUDE:
            judge, target, why = "편입", ", ".join(f"[{c}]" for c in R.ENT_INCLUDE[e_id]), "ICS 기법으로 표현되지 않는 IT/OT 경계 메커니즘"
        elif e_id in R.ENT_MERGE:
            k = norm(e_id, "판정표")
            judge, target, why = "통합", f"{k} {name_ko(k)}" if k else "", "ICS 기법이 같은 행위를 OT 자산에서 표현"
        elif e_id in R.ENT_DROP:
            judge, target, why = "제외", "", R.ENT_DROP[e_id]
        else:
            judge, target, why = "미판정", "", "판정 규칙 없음 — 근거 미반영"
        _row(ws, rr_, [e_id, info.get("name", ""), ", ".join(info.get("tactics", [])), judge, target,
                       ", ".join(sorted(ent_seen.get(e_id, []))), why])
        rr_ += 1
    _w(ws, [12, 40, 30, 8, 34, 40, 50])

    # ---------------- EMB3D 판정 ----------------
    ws = wb.create_sheet("EMB3D 판정")
    ws.cell(1, 1, "MITRE EMB3D 장치 위협 판정 — 신설(◇ 세부위협) · 공식 연계(EMB3D 근거 문단이 인용한 ATT&CK 기법) · "
                  "분석 연계(위협 정의가 같은 행위를 가리키는 기법)").font = Font(bold=True, color=NAVY)
    ws.cell(2, 1, "©2026 The MITRE Corporation. This work is reproduced and distributed with the permission of The MITRE Corporation."
            ).font = Font(size=9, color="595959")
    _hdr(ws, 3, ["TID", "위협(EMB3D)", "한글명", "분류", "성숙도", "판정", "연계 세부위협(근거)", "EMB3D 인용 ATT&CK ID",
                 "판정 사유", "CWE", "CVE 예시 수", "완화책(MID)"])
    otc_by_key = collections.defaultdict(list)
    for o in rows:
        otc_by_key[o["key"]].append(o["otc"])
    for i, (tid, t) in enumerate(sorted(E3T.items()), 4):
        judge, why = emb3d_judgement(tid)
        linked = "\n".join(f"{'/'.join(otc_by_key.get(k, ['-']))} {name_ko(k)} ({b})"
                           for k, b in sorted(emb3d_links.get(tid, {}).items()))
        cited = ", ".join(c + (" (매트릭스 밖)" if c in emb3d_outside.get(tid, []) else "") for c in t["atk_cited"])
        _row(ws, i, [tid, t["name"], R.EMB3D_NAME_KO.get(tid, ""), R.EMB3D_CATEGORY_KO.get(t["category"], t["category"]),
                     t["maturity"], judge, linked, cited, why, ", ".join(t["cwes"]), len(t["cves"]), ", ".join(t["mids"])])
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:L{len(E3T) + 3}"
    _w(ws, [9, 34, 26, 14, 11, 10, 44, 24, 50, 22, 8, 30])

    otc_row = {o["otc"]: o for o in rows}      # OTC-ID → 행(표준·OWASP·시나리오 시트 공용)
    otc_by_key = collections.defaultdict(list)  # 기법 키 → [OTC-ID] (한 기법이 여러 전술에 반복)
    for o in rows:
        otc_by_key[o["key"]].append(o["otc"])

    # ---------------- 표준 연계 ----------------
    ws = wb.create_sheet("표준 연계")
    r = section(ws, 1, "1. IEC 62443·NIST SP 800-53 요구사항 → 세부위협 (ATT&CK 완화책의 표준 매핑·EMB3D 완화책의 62443-4-2 매핑 집계)")
    ws.cell(r, 1, "출처: ATT&CK for ICS 완화책의 STIX labels(IEC 62443-3-3:2013 · 62443-4-2:2019 · NIST SP 800-53 Rev.5), "
                  "EMB3D 완화책의 IEC 62443-4-2 매핑. 요구사항 원문은 유료 표준이라 수록하지 않고 ID만 연계.").font = Font(size=9, color="595959")
    r += 1
    std_titles = {"iec33": "IEC 62443-3-3 (SR)", "iec42": "IEC 62443-4-2 (CR·EDR·HDR·NDR)", "nist": "NIST SP 800-53 Rev.5"}
    std_src = {"iec33": "iec_62443_3_3", "iec42": "iec_62443_4_2", "nist": "nist_800_53"}
    std_idx = {"iec33": 0, "iec42": 1, "nist": 2}

    def std_name(fam, req):
        tbl = std_names.get(std_src[fam], {})
        if req in tbl:
            return tbl[req]
        base = re.sub(r"\s*RE\(\d+\)$", "", req)      # 'CR 1.5 RE(1)' → 'CR 1.5' + 강화요건
        if base != req and base in tbl:
            return tbl[base] + f" (강화요건 {req[len(base):].strip()})"
        WARN[f"표준 명칭 없음 {fam} {req}"].add("standards.yaml")
        return ""

    std_rows = {"iec33": collections.defaultdict(list), "iec42": collections.defaultdict(list),
                "nist": collections.defaultdict(list)}
    for o in rows:
        for fam, i in std_idx.items():
            for req in o["std"][i]:
                std_rows[fam][req].append(o)
    _hdr(ws, r, ["표준", "요구사항 ID", "요구사항 명칭(제목)", "연계 세부위협 수", "연계 세부위협(OTC-ID)"])
    r += 1
    for fam in ("iec42", "iec33", "nist"):
        for req in sorted(std_rows[fam], key=_iec_sort if fam != "nist" else None):
            os_ = std_rows[fam][req]
            _row(ws, r, [std_titles[fam], req, std_name(fam, req), len(os_),
                         ", ".join(sorted(o["otc"] for o in os_))])
            r += 1
    r = section(ws, r + 1, "2. 세부위협별 표준 요구사항 (요구사항이 매핑된 행만)")
    _hdr(ws, r, ["OTC-ID", "세부위협", "IEC 62443-3-3", "IEC 62443-4-2", "NIST SP 800-53"])
    r += 1
    for o in rows:
        if any(o["std"]):
            _row(ws, r, [o["otc"], o["lv3"], ", ".join(o["std"][0]), ", ".join(o["std"][1]), ", ".join(o["std"][2])])
            r += 1
    _w(ws, [28, 16, 52, 12, 70])

    # ---------------- OWASP IoT Top 10 ----------------
    ws = wb.create_sheet("OWASP IoT Top10")
    ws.cell(1, 1, "OWASP IoT Top 10 (2018) 연계 — IoT 위험 범주를 매트릭스 세부위협에 분석 매핑(교차 참조용, 근거 집계 아님). "
                  "출처: OWASP/www-project-internet-of-things").font = Font(bold=True, color=NAVY)
    _hdr(ws, 2, ["ID", "범주(EN)", "범주(KO)", "요약", "연계 세부위협(OTC-ID)"])
    r = 3
    for oid, en, ko, desc, keys in R.OWASP_IOT:
        otcs = []
        for k in keys:
            if k not in otc_by_key:
                WARN[f"OWASP {oid} 잘못된 행 키 {k}"].add("taxonomy_rules")
            else:
                otcs += otc_by_key[k]
        _row(ws, r, [oid, en, ko, desc, ", ".join(sorted(otcs))])
        r += 1
    _w(ws, [6, 40, 28, 70, 40])

    # ---------------- 공격 체인 시나리오 ----------------
    scn_data = yaml.safe_load(SCENARIOS.read_text(encoding="utf-8")) if SCENARIOS.exists() else {"scenarios": []}
    otc_set = set(otc_row)
    ws = wb.create_sheet("공격 체인 시나리오")
    ws.cell(1, 1, "공격 체인 시나리오 — 실제 사고 기반 IT 침투 → OT 영향 흐름. 각 단계는 매트릭스 OTC-ID에 연결(위험도 색). "
                  "굵은 단계 = 흐름을 끊기 좋은 초크 포인트(진입·실행·영향 분기)").font = Font(bold=True, color=NAVY)
    r = 3
    for scn in scn_data["scenarios"]:
        c = ws.cell(r, 1, f"{scn['id']}. {scn['title']}")
        c.font = Font(size=12, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
        ws.cell(r, 1, f"근거: {scn['ref']}  ·  사고: {', '.join(scn['based_on'])}  ·  {scn.get('note', '')}").font = Font(size=9, color="595959")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
        _hdr(ws, r, ["#", "전술", "OTC-ID", "세부위협", "행위", "탐지·대응 포인트"])
        r += 1
        for i, st in enumerate(scn["steps"], 1):
            otc = st["otc"]
            if otc not in otc_set:
                WARN[f"시나리오 {scn['id']} 잘못된 OTC {otc}"].add("scenarios")
            o = otc_row.get(otc)
            _row(ws, r, [i, st["전술"], otc, o["lv3"] if o else "(없음)", st["행위"], st["탐지"]])
            if o:
                _risk(ws.cell(r, 3), o["risk"])
            r += 1
        r += 1
    _w(ws, [4, 16, 12, 22, 50, 60])

    # ---------------- 평가 기준 ----------------
    ws = wb.create_sheet("평가 기준")
    crit = [
        ("■ 발생가능성 (근거에서 자동 산정 — AI 매트릭스 v3.2 · 클라우드 v5와 같은 수식)", ""),
        ("상", f"실제 사고 {LIKELY_HIGH_REAL}건 이상"),
        ("중", "실제 사고 1건, 또는 ATT&CK 사례(ICS 절차·OT 주체의 Enterprise 절차)·KEV(OT)·위협인텔·공개 취약점·실증 연구·"
              "EMB3D 위협(관찰·알려진 취약점·개념 증명) 중 하나 이상"),
        ("하", "근거 없음(이론·시나리오)"),
        ("근거 대응", "AI '실제 사고(ATLAS Incident+OWASP 인용)' ↔ 클라우드 '사고 DB 실제 사고' ↔ OT '사고 DB 실제 사고(ICS 캠페인 포함)' / "
                    "AI 'Realized 기법' ↔ 클라우드 'ATT&CK 클라우드 사례' ↔ OT 'ATT&CK ICS 절차 + KEV(OT) + 위협인텔(능력 발견)' / "
                    "AI '실증·공개 취약점' ↔ 클라우드 '연구·노출 + 벤더 매트릭스' ↔ OT '연구·시연 + CISA ICS 권고 CVE'. 실제 사고 외 근거는 건수와 무관하게 '중'까지"),
        ("", ""),
        ("■ 심각도 (OT 결과 기준 — 기법별 기준값, 없으면 전술 기준값)", ""),
        ("상", "인명·환경·설비 피해 가능, 제어 상실·조작, 안전·보호·대응 기능 무력화, 제어기 로직·펌웨어·운전 모드 변조, 광역·장기 운영 중단"),
        ("중", "감시 기능 상실·조작, 운영 정보 탈취, OT 진입·실행·지속성·확산·은닉 등 후속 공격 기반"),
        ("하", "탐색·수집 위주로 직접 피해 없음, 공격자 측 준비(자원 개발)"),
        ("전술 기준값", " · ".join(f"[{c}] {R.TACTIC_SEVERITY[c]}" for c, *_ in R.TACTICS)),
        ("", ""),
        ("■ 위험도 (발생가능성 × 심각도)", "심각도 상 / 중 / 하"),
        ("발생가능성 상", "매우 높음 / 높음 / 보통"),
        ("발생가능성 중", "높음 / 보통 / 낮음"),
        ("발생가능성 하", "보통 / 낮음 / 낮음"),
        ("", ""),
        ("■ 근거 수준 (높은 순)", ""),
        ("실제 사고 확인", "사고 DB의 실제 사고(사고·캠페인·사례연구·정부 경보·공시, 검증된 건)와 연결"),
        ("실사용 기법 포함", "실제 사고 연결은 없으나 ATT&CK 절차(ICS 전 주체, Enterprise는 ICS 캠페인·소프트웨어), KEV(OT), 위협인텔(능력 발견), "
                         "EMB3D '관찰' 위협 존재"),
        ("실증·공개 취약점", "연구·시연 사례, CISA ICS 권고의 공개 취약점(CWE 규칙 매핑), EMB3D '알려진 취약점'·'개념 증명' 위협만 존재"),
        ("이론·시나리오", "공개 근거 없음 — '발생하지 않음'이 아님"),
        ("", ""),
        ("■ 사고 DB 집계", ""),
        ("실제 사고", "사례유형 사고·캠페인·사례연구(익명)·정부 경보·공시 중 검증 수준이 '행위자 주장'·'논란·원인 번복'이 아니고 OT 관련도가 'IT 한정'이 아닌 건"),
        ("매핑 근거", "ATT&CK 공식 = 사고가 ATT&CK 캠페인·소프트웨어로 문서화돼 그 공식 절차를 그대로 매핑 / 분석 = 사고 출처를 읽고 분석자가 매핑"),
        ("폐기 ID 변환", "v19 ICS 개편으로 폐기된 ID는 대체 하위기법으로 변환: " + ", ".join(f"{a}→{b}" for a, b in sorted(ics.revoked.items()) if a and b)),
        ("최근 사고", f"기준일 {R.RECENT_FROM}년 이후 실제 사고"),
        ("", ""),
        ("■ 공개 취약점·KEV", ""),
        ("공개 취약점", "CISA ICS 권고(ICSA, 의료기기 ICSMA 제외)의 CVE를 CWE·공격 경로 규칙으로 1개 기법에 매핑 — 노출 여부가 환경에 달린 공개 애플리케이션 악용(T0819)은 CWE로 매핑하지 않음"),
        ("KEV(OT)", "ICS 권고와 교차된 KEV 중 OT·IoT 제품, 임베디드 구성요소, OT 네트워크 장비 내장 OS만 반영(범용 IT 구성요소 제외) — 실사용 근거(중 상한)"),
        ("", ""),
        ("■ EMB3D 연계", ""),
        ("공식 연계", "EMB3D 근거 문단이 인용한 ATT&CK 기법(폐기 ID 변환·Enterprise 통합 규칙 적용) — Mobile·매트릭스 밖 Enterprise ID는 연계하지 않고 판정 시트에 표시"),
        ("분석 연계", "위협 정의가 같은 행위를 가리키는 기법에 분석자가 연계(사유는 'EMB3D 판정' 시트)"),
        ("신설", "ATT&CK에 대응 기법이 없는 장치 하드웨어 위협을 효과가 나타나는 전술 아래 세부위협으로 추가(◇) — 신설 행의 완화책은 EMB3D 완화책(MID)"),
        ("근거 반영", "성숙도 '관찰' = 실사용 기법 포함, '알려진 취약점'·'개념 증명' = 실증·공개 취약점 — 모두 발생가능성 '중' 상한"),
        ("", ""),
        ("■ 참조 열 사례 라벨", "'■ 실제 사례' 줄 앞에 붙는 라벨 — 사례명(YYYY-MM): 경위·결과 (사고 ID, ATT&CK ID)"),
        ("[실제 사고]", "사고 DB에서 실제 사고로 집계된 건"),
        ("[위협인텔]", "배포 전에 발견된 공격 도구의 능력(예: PIPEDREAM) — 실제 사용은 확인되지 않음"),
        ("[공개 취약점]", "CISA ICS 권고·KEV의 취약점"),
        ("[실증]", "연구·시연(예: PLC-Blaster, Aurora 시험)"),
        ("[시나리오]", "공개 사례가 없는 기법의 가정 사례 — '공개 사고 미확인 — 사유' 줄과 함께 표기"),
        ("검증", "scripts/validate.py가 라벨과 사고 DB 집계 상태, 인용한 사고와 기법 매핑, 인용한 ATT&CK 주체의 절차 보유, OTC·INC·ATT&CK ID의 존재를 확인"),
    ]
    for i, (a, b) in enumerate(crit, 1):
        ws.cell(i, 1, a).font = Font(bold=a.startswith("■"), color=NAVY if a.startswith("■") else "000000")
        ws.cell(i, 2, b).alignment = WRAP
    _w(ws, [40, 130])

    # ---------------- 변경이력 ----------------
    ws = wb.create_sheet("변경이력")
    ws.cell(1, 1, "변경이력 — 개정은 새 버전 파일로 저장하고 이 시트 상단에 기록").font = Font(bold=True, color=NAVY)
    _hdr(ws, 2, ["버전", "날짜", "구분", "대상", "변경 내용", "사유"])
    for i, e in enumerate(log.get("entries", []), 3):
        _row(ws, i, e)
    _w(ws, [8, 11, 14, 20, 80, 50])

    path.parent.mkdir(exist_ok=True)
    wb.save(path)


# ===========================================================================
# 5) 문구 작성용 근거 정리본 (선택)
# ===========================================================================
def write_worksheet(outdir):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    inc_by = {e["id"]: e for e in incidents}
    done = set()
    for d in domains:
        lines = [f"# [{d['code']}] {d['ko']} — 근거 정리본\n"]
        for o in [o for o in rows if o["code"] == d["code"]]:
            k = o["key"]
            lines.append(f"## {o['otc']} {k} {o['lv3']} / {o['en_name']}  (위험 {o['risk']}: 발생 {o['likelihood']} × 심각 {o['severity']}, {o['level']})")
            if k in done:
                lines.append(f"(앞 전술에서 정리 — 같은 기법)\n")
                continue
            done.add(k)
            info = tech_info(k)
            lines.append(f"- 설명: {info.get('desc', '')[:1200]}")
            lines.append(f"- 자산: {', '.join(R.ASSET_KO.get(a, a) for a in o['assets'])} | Purdue: {', '.join(o['purdue'])} | 통합: {o['merged']}")
            procs = (ics.procs.get(k, []) if k in ics.tech else []) + ent.procs.get(k, [])
            for m in merged_into.get(k, []):
                procs += [(s, f"[{m}] {t}") for s, t in ent.procs.get(m, [])]
            for s, t in procs:
                lines.append(f"  - 절차 {subj_label(s)}: {t[:400]}")
            for i in sorted(ev[k]["real"] | ev[k]["research"] | ev[k]["intel"] | ev[k]["excluded"]):
                e = inc_by[i]
                lines.append(f"  - 사고 {i} [{e['status']}·{e['links'].get(k)}] {e['title']}({e['date']}): {e['summary'][:300]}")
            if ev[k]["cves"] or ev[k]["kev"]:
                lines.append(f"  - 취약점: CVE {len(ev[k]['cves'])}개(최근 {len(ev[k]['recent_cves'])}) · KEV {', '.join(sorted(ev[k]['kev']))}")
            lines.append(f"- 완화책: {'; '.join(o['mitig'])}")
            dets = ics.detect.get(k, []) if k in ics.tech else []
            for det_id, det_name, analytics in dets:
                for a in analytics[:3]:
                    lines.append(f"  - 탐지 {det_id}: {a[:300]}")
            lines.append("")
        (outdir / f"{d['code']}.md").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--worksheet", help="문구 작성용 근거 정리본 출력 폴더")
    a = ap.parse_args()
    write_xlsx(OUT)
    if a.worksheet:
        write_worksheet(a.worksheet)
    rc = collections.Counter(o["risk"] for o in rows)
    lc = collections.Counter(o["level"] for o in rows)
    print(f"{OUT.relative_to(ROOT)} — 세부위협 {len(rows)}행 · 고유 기법 {len({o['key'] for o in rows})}개")
    print("  위험도:", dict(rc), "| 근거 수준:", dict(lc))
    print("  문구 작성:", sum(1 for o in rows if o["text_src"] == "분석가 작성"), "/", len(rows), "행")
    if WARN:
        print("  [경고] 판정 규칙 없음:")
        for t, w in sorted(WARN.items()):
            print(f"    {t}: {', '.join(sorted(x for x in w if x))[:160]}")
