"""통합 보안위협 매트릭스 빌드 — 요약 문안(YAML)과 원본 6종을 보고서용 워크북 하나로 묶는다.

입력
  - 요약(재구성 Lv3) 문안: integrated/data/<Lv0>/<도메인·전술>.yaml (ai · cloud · ot · supplychain · identity · physical)
  - 원본: sources/ai_v3.2 (AI v3.2 LITE, Lv3 124), output/통합_클라우드보안위협_매트릭스_v5.xlsx (154행),
          ot/output/통합_OT보안위협_매트릭스_v5.xlsx (141행), supplychain/output/통합_공급망보안위협_매트릭스_v2.xlsx (60),
          identity/output/통합_신원보안위협_매트릭스_v2.xlsx (65), physical/output/통합_물리인적보안위협_매트릭스_v2.xlsx (45)
출력
  - integrated/output/통합_보안위협_매트릭스_v3.xlsx
  - integrated/output/통합_요약매트릭스_v3.csv (요약 매트릭스 검토·diff용)
  - integrated/output/통합_요약목록_v3.md (도메인별 요약 위협 목록)
  - --release: integrated/output/통합_보안위협_매트릭스_v3_배포본.xlsx
    (고객 배포본 — 클라우드 상세 시트의 CSA Top Threats 열, 신원 상세 시트의 CSA CCM·CIS Controls 열을 ID만 남김)

사용: python3 integrated/scripts/build_integrated_matrix.py [--release]
"""
import argparse
import collections
import csv
import datetime
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common  # noqa: E402

VERSION = "v3"
TITLE = "통합 보안위협 매트릭스"
OUT_DIR = os.path.join(common.ROOT, "integrated", "output")
XLSX = os.path.join(OUT_DIR, f"통합_보안위협_매트릭스_{VERSION}.xlsx")
XLSX_RELEASE = os.path.join(OUT_DIR, f"통합_보안위협_매트릭스_{VERSION}_배포본.xlsx")
CSV = os.path.join(OUT_DIR, f"통합_요약매트릭스_{VERSION}.csv")
MD = os.path.join(OUT_DIR, f"통합_요약목록_{VERSION}.md")
LV0 = {"ai": "AI 보안위협", "cloud": "클라우드 보안위협", "ot": "OT 보안위협",  # 순서 = common.KINDS
       "supplychain": "공급망 보안위협", "identity": "신원 보안위협", "physical": "물리·인적 보안위협"}
KIND_OF = {v: k for k, v in LV0.items()}
# 요약표 '공격 단계'의 AI(ATLAS) 용어를 클라우드(ATT&CK v19.2 국문) 용어로 통일. 대응 용어가 없는 단계는 원문 유지
STAGE_NORM = {"초기 접근": "초기 침투", "내부 확산": "횡적 이동", "유출": "반출"}
# 공급망·신원·물리 원본의 '다른 매트릭스 연계' 열 — 요약 ID(AI-·CL-·OT-)와 상세 ID(OTC-·SCT-·IDT-·PHT-)가 섞여 있음
TAXO_LINK_COLS = {
    "supplychain": ["클라우드 매트릭스 연계", "AI 매트릭스 연계", "OT 매트릭스 연계"],
    "identity": ["클라우드 매트릭스 연계", "AI 매트릭스 연계", "OT 매트릭스 연계", "공급망 매트릭스 연계"],
    "physical": ["아이덴티티 연계", "공급망 연계", "OT 연계", "클라우드 연계", "AI 연계"],
}
LINK_ID_RE = re.compile(r"\b(?:AI-\d{2}-\d{2}|CL-[A-Z0-9]{2}-\d{2}|OT-[A-Z0-9]{2}-\d{2}|OTC-[A-Z0-9]{2}-\d{2}(?:\.\d+)?"
                        r"|SCT-[A-Z]{2}-\d{2}(?:\.\d+)?|IDT-[A-Z]{2}-\d{2}(?:\.\d+)?|PHT-[A-Z]{2}-\d{2}(?:\.\d+)?)\b")
DETAIL_SHEET = {"ai": "AI 위협 상세", "cloud": "클라우드 위협 상세", "ot": "OT 위협 상세",
                "supplychain": "공급망 위협 상세", "identity": "신원 위협 상세", "physical": "물리·인적 위협 상세"}

# ---------------------------------------------------------------- 서식 (클라우드 v5 빌드 스크립트와 같은 색 체계)
FONT = "맑은 고딕"
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RISK_COLOR = {"매우 높음": ("C0392B", "FFFFFF"), "높음": ("E67E22", "FFFFFF"),
              "보통": ("F7DC6F", "000000"), "낮음": ("A9DFBF", "000000")}
LEVEL_COLOR = {"실제 사고 확인": "1E8449", "실사용 기법 포함": "2E86C1",
               "실증·공개 취약점": "B9770E", "이론·시나리오": "7F8C8D"}
LV0_FILL = {"ai": "EBDEF0", "cloud": "D6EAF8", "ot": "FDEBD0",
            "supplychain": "D5F5E3", "identity": "FCF3CF", "physical": "F2D7D5"}
LV0_BAND = {"ai": "6C3483", "cloud": "1F618D", "ot": "AF601A",
            "supplychain": "1E8449", "identity": "9A7D0A", "physical": "922B21"}
NAVY, HDR = "1F3864", "34495E"
RISK_RANK = {r: i for i, r in enumerate(common.RISK_ORDER)}


def font(size=10, bold=False, color="000000"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(color):
    return PatternFill("solid", fgColor=color)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub, ncols):
    c = ws.cell(1, 1, text)
    c.font, c.fill, c.alignment = font(13, True, "FFFFFF"), fill(NAVY), Alignment(vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 24
    c = ws.cell(2, 1, sub)
    c.font = font(9, color="404040")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)


def groups(ws, row, spec):
    c0 = 1
    for name, span, color in spec:
        if span > 1:
            ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c0 + span - 1)
        c = ws.cell(row, c0, name)
        c.fill, c.font, c.alignment = fill(color), font(10, True, "FFFFFF"), CENTER
        c0 += span


def header(ws, row, cols, color=HDR):
    for j, name in enumerate(cols, 1):
        c = ws.cell(row, j, name)
        c.fill, c.font, c.alignment, c.border = fill(color), font(10, True, "FFFFFF"), CENTER, BORDER
    ws.row_dimensions[row].height = 30


def put_row(ws, r, vals, size=10):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.font, c.alignment, c.border = font(size), WRAP, BORDER


def print_setup(ws, title_rows=None, paper="A3"):
    """인쇄: 가로·폭 맞춤(높이 자동), 표 머리글 반복."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if paper == "A3" else ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if title_rows:
        ws.print_title_rows = title_rows


def paint_risk(c, risk, size=10):
    bg, fg = RISK_COLOR[risk]
    c.fill, c.font, c.alignment = fill(bg), font(size, risk in ("매우 높음", "높음"), fg), CENTER


def paint_level(c, level, size=10):
    c.font, c.alignment = font(size, True, LEVEL_COLOR[level]), CENTER


# ---------------------------------------------------------------- 요약 행 구성
def compress_ids(ids):
    """ATT&CK ID를 상위기법별로 묶어 짧게 표기: T1552, T1552.001·.008 / 비 ATT&CK(AZT…)는 그대로."""
    grouped = collections.OrderedDict()
    for t in ids:
        parent, _, sub = t.partition(".")
        grouped.setdefault(parent, []).append(sub)
    out = []
    for parent, subs in grouped.items():
        if "" in subs:
            out.append(parent)
        subs = [s for s in subs if s]
        if subs:
            out.append(f"{parent}.{subs[0]}" + "".join(f"·.{s}" for s in subs[1:]))
    return ", ".join(out)


def compress_sr(srs):
    """IEC 62443-3-3 SR 목록을 번호순으로 묶어 짧게 표기: ['SR 2.1', 'SR 1.1'] → 'SR 1.1·2.1'."""
    nums = sorted({x.replace("SR", "").strip() for x in srs}, key=lambda v: [int(p) for p in v.split(".")])
    return "SR " + "·".join(nums)


def ot_refs(entry, ot_src):
    """OT 관련 기준: ATT&CK ID(ICS·Enterprise 편입, 상위기법별 축약) + EMB3D TID(EMB3D 신설 행)
    + IEC 62443-3-3 SR(구성 원본 완화책의 요구사항 ID 합집합)."""
    rows = [common.member_row(ot_src, m, entry["domain"]) for m in entry["members"]]
    attack = [m for m in entry["members"] if m.startswith("T")]
    tids = sorted({t for m, r in zip(entry["members"], rows) if m.startswith("EMB3D-")
                   for t in re.findall(r"TID-(\d+)", r["EMB3D 연계(성숙도·근거)"] or "")}, key=int)
    srs = {x.strip() for r in rows for x in re.split(r",\s*", r["IEC 62443-3-3(SR)"] or "") if x.strip()}
    return ([f"ATT&CK {compress_ids(attack)}"] if attack else []) + \
           (["EMB3D TID-" + "·".join(tids)] if tids else []) + \
           ([f"IEC 62443-3-3 {compress_sr(srs)}"] if srs else [])


def uniq(seq):
    return list(collections.OrderedDict.fromkeys(x for x in seq if x))


def lines_of(rows, col):
    """원본 행들의 여러 줄 셀 값 → 줄 목록(중복 제거, 원본 순서)."""
    return uniq(l.strip() for r in rows for l in str(r.get(col) or "").split("\n"))


def num_key(v):
    return [int(p) if p.isdigit() else p for p in re.split(r"[.\-]", v)]


def taxo_refs(kind, rows):
    """공급망·신원·물리 관련 기준: ATT&CK ID(상위기법별 축약) + Lv0별 교차 매핑 ID.
    공급망: SLSA v1.2 위협 · OWASP CI/CD · OWASP OSS / 신원: Browser & Identity Attacks(SAT) · OWASP NHI /
    물리: ITKB 내부자 관측 기법 · ISO/IEC 27001:2022 부속서 A · ISMS-P (모두 원본 행 값의 ID만)."""
    attack = uniq(t.strip() for r in rows for t in str(r["ATT&CK ID"] or "").replace("\n", ",").split(",")
                  if t.strip().startswith("T"))
    out = [f"ATT&CK {compress_ids(attack)}"] if attack else []

    def ids(label, col, pat, sort=False):
        found = uniq(m for l in lines_of(rows, col) for m in [re.match(pat, l)] if m for m in [m.group(1)])
        if sort:
            found = sorted(found, key=num_key)
        if found:
            out.append(f"{label} {', '.join(found)}")

    if kind == "supplychain":
        ids("SLSA", "SLSA v1.2 위협", r"^([A-Z]\d?)\s")
        ids("OWASP", "OWASP CI/CD Top10", r"^(CICD-SEC-\d+)")
        ids("OWASP", "OWASP OSS Top10", r"^(OSS-RISK-\d+)")
    elif kind == "identity":
        ids("SAT", "Browser & Identity Attacks(SAT)", r"^(SAT\d+)")
        ids("OWASP", "OWASP NHI Top10", r"^(NHI\d+)")
    elif kind == "physical":
        itkb = uniq(m for l in lines_of(rows, "ITKB 내부자 관측 기법") for m in re.findall(r"^(T\d{4}(?:\.\d{3})?)", l))
        if itkb:
            out.append(f"ITKB {compress_ids(itkb)}")
        ids("ISO 27001", "ISO/IEC 27001:2022", r"^(\d+\.\d+)", sort=True)
        ids("ISMS-P", "ISMS-P", r"^(\d+\.\d+\.\d+)", sort=True)
    return out


def build_rows(src, summ):
    """요약 행 구성. 반환: (행 목록, 요약 항목으로 옮기지 못한 OT 원본 연계 참조 {'원본키→참조': (요약 ID, 대상 Lv0)})."""
    ai_src, cl_src, ot_src = src["ai"], src["cloud"], src["ot"]
    ai_sum, cl_sum, ot_sum = summ["ai"], summ["cloud"], summ["ot"]
    every = ai_sum + cl_sum + ot_sum
    tactic_ko = {t: label.split("] ", 1)[1] for t, label in cl_src["tactics"].items()}
    ot_ko = {t: label.split("] ", 1)[1] for t, label in ot_src["tactics"].items()}
    to_summary = {kind: {m: e["id"] for e in summ[kind] for m in e["members"]} for kind in common.KINDS}
    names = {e["id"]: e["name"] for e in every}
    order = {e["id"]: i for i, e in enumerate(every)}

    # 연계 위협: 문안의 수동 연계 + 클라우드 원본 'AI 매트릭스 연계(UT)' 열
    #           + OT 원본 '클라우드 매트릭스 연계'(ATT&CK ID)·'AI 매트릭스 연계(UT)'(Lv3 ID) 열 → 양방향
    links = collections.defaultdict(set)

    def link(a, b):
        links[a].add(b)
        links[b].add(a)

    for e in every:
        for target in e.get("links") or []:
            link(e["id"], target)
    for e in cl_sum:
        for m in e["members"]:
            for r in cl_src["tech"][m]["rows"]:
                for ut in re.findall(r"UT-\d+\.\d+", r["AI 매트릭스 연계(UT)"] or ""):
                    link(e["id"], to_summary["ai"][ut])
    # OT 원본 연계 중 요약 항목으로 바로 옮길 수 없는 참조(클라우드 v5에 없는 기법, UT Lv2 ID)는 건너뛰고 문안 links로 지정
    skipped = collections.OrderedDict()
    for e in ot_sum:
        for m in e["members"]:
            for r in ot_src["tech"][m]["rows"]:
                refs = [("cloud", t) for t in re.findall(r"\bT\d{4}(?:\.\d{3})?\b", r["클라우드 매트릭스 연계"] or "")] + \
                       [("ai", u) for u in re.findall(r"UT-\d+(?:\.\d+)?", r["AI 매트릭스 연계(UT)"] or "")]
                for kind, ref in refs:
                    if ref in to_summary[kind]:
                        link(e["id"], to_summary[kind][ref])
                    else:
                        skipped.setdefault(f"{m}→{ref}", (e["id"], kind))
    # 공급망·신원·물리 원본 '다른 매트릭스 연계' 열: 요약 ID(AI-·CL-·OT-)는 그대로, 상세 ID(OTC-·SCT-·IDT-·PHT-)는
    # 그 원본이 묶인 요약 항목 ID로 바꿈. 바꿀 수 없는 참조는 건너뛰고 경고(문안 links로 대체)
    otc = {c: to_summary["ot"][k] for k, t in ot_src["tech"].items() for c in t["otc"]}
    detail = {"OTC": otc, "SCT": to_summary["supplychain"], "IDT": to_summary["identity"], "PHT": to_summary["physical"]}
    for kind in common.TAXO:
        for e in summ[kind]:
            for m in e["members"]:
                r = src[kind]["tech"][m]["row"]
                for col in TAXO_LINK_COLS[kind]:
                    for ref in LINK_ID_RE.findall(str(r[col] or "")):
                        head = ref.split("-")[0]
                        target = detail[head].get(ref) if head in detail else (ref if ref in names else None)
                        if target and target != e["id"]:
                            link(e["id"], target)
                        elif not target:
                            skipped.setdefault(f"{m}→{ref}", (e["id"], head))
    unknown = [t for k in links for t in links[k] if t not in names]
    if unknown:
        raise ValueError(f"연계 위협 ID 없음: {unknown}")

    rows = []
    for e in ai_sum:
        ev = common.evaluate_ai(e, ai_src)
        mem = [ai_src["lv3"][m] for m in e["members"]]
        owasp = uniq(x for d in mem for x in re.split(r",\s*", d["OWASP 매핑 (주)"] or ""))
        atlas = uniq(x.split(".")[0] for d in mem for x in re.split(r",\s*", d["ATLAS 기법 (주)"] or ""))
        refs = ([f"OWASP {', '.join(owasp)}"] if owasp else []) + ([f"ATLAS {', '.join(atlas)}"] if atlas else [])
        stages = uniq(STAGE_NORM.get(x, x) for x in
                      (re.sub(r"\([A-Za-z][^()]*\)$", "", d["주 공격 단계"]).strip() for d in mem))
        rows.append(dict(e, kind="ai", lv0=LV0["ai"], lv1=ai_src["domains"][e["domain"]], ev=ev,
                         stage=", ".join(stages), refs="\n".join(refs), origin=", ".join(e["members"])))
    for e in cl_sum:
        ev = common.evaluate_cloud(e, cl_src)
        tech = [cl_src["tech"][m] for m in e["members"]]
        others = sorted({t for x in tech for t in x["tactics"]} - {e["domain"]}, key=common.TACTICS.index)
        stage = tactic_ko[e["domain"]] + (f" (연관: {', '.join(tactic_ko[t] for t in others)})" if others else "")
        attack = [m for m in e["members"] if m.startswith("T")]
        azure = [m for m in e["members"] if not m.startswith("T")]
        csa = sorted({s for x in tech for s in re.findall(r"SI-\d+", x["row"]["CSA Top Threats 2026"] or "")})
        refs = ([f"ATT&CK {compress_ids(attack)}"] if attack else []) + \
               ([f"Azure ATRM {', '.join(azure)}"] if azure else []) + ([f"CSA {', '.join(csa)}"] if csa else [])
        rows.append(dict(e, kind="cloud", lv0=LV0["cloud"], lv1=cl_src["tactics"][e["domain"]], ev=ev,
                         stage=stage, refs="\n".join(refs), origin=", ".join(c for x in tech for c in x["ctc"])))
    for e in ot_sum:
        ev = common.evaluate_attack(e, ot_src)
        tech = [ot_src["tech"][m] for m in e["members"]]
        others = sorted({t for x in tech for t in x["tactics"]} - {e["domain"]}, key=common.OT_TACTICS.index)
        stage = ot_ko[e["domain"]] + (f" (연관: {', '.join(ot_ko[t] for t in others)})" if others else "")
        rows.append(dict(e, kind="ot", lv0=LV0["ot"], lv1=ot_src["tactics"][e["domain"]], ev=ev, stage=stage,
                         refs="\n".join(ot_refs(e, ot_src)), origin=", ".join(c for x in tech for c in x["otc"])))
    for kind, cfg in common.TAXO.items():  # 공급망·신원·물리: 공격 단계 = 구성 원본 '공격 단계' 합집합(원본 용어)
        sk = src[kind]
        for e in summ[kind]:
            ev = common.evaluate_attack(e, sk)
            mem = [sk["tech"][m]["row"] for m in e["members"]]
            stages = uniq(x.strip() for r in mem for x in str(r[cfg["stage"]] or "").split(","))
            rows.append(dict(e, kind=kind, lv0=LV0[kind], lv1=sk["tactics"][e["domain"]], ev=ev, stage=", ".join(stages),
                             refs="\n".join(taxo_refs(kind, mem)), origin=", ".join(e["members"])))
    for r in rows:
        r["links_text"] = "\n".join(f"{t} {names[t]}" for t in sorted(links[r["id"]], key=order.get))
        r["case1"] = first_case(r["cases"])
        r["control1"] = r["controls"].split("\n")[0][2:].strip()
    # 우선순위: 같은 Lv0 안에서 위험도 → 실제 사고 수 → 최근 사고(2025~) → 위협 ID 순
    for kind in LV0:
        ranked = sorted((r for r in rows if r["kind"] == kind), key=priority_key)
        for i, r in enumerate(ranked, 1):
            r["rank"] = i
    return rows, skipped


def priority_key(r):
    return (RISK_RANK[r["ev"]["risk"]], -r["ev"]["incidents"], -r["ev"]["recent"], r["id"])


def first_case(cases):
    """대표 사례 첫 줄 → '[유형] 사례명(시점)' (보고서용 간략 표기)."""
    line = cases.split("\n")[0].strip()
    if line.startswith("- 공개 사고 미확인"):
        return "공개 사고 미확인"
    m = common.CASE_RE.match(line)
    n = common.NAMED_RE.match(m.group("body")) if m.group("tag") != "시나리오" else None
    if not n:
        body = m.group("body")
        return f"[{m.group('tag')}] " + (body if len(body) <= 40 else body[:39] + "…")
    name = n.group("name").strip()
    return f"[{m.group('tag')}] {name}({n.group('date')})" if n.group("date") else f"[{m.group('tag')}] {name}"


# ---------------------------------------------------------------- 재구성 매핑 (원본 Lv3 → 요약 Lv3)
def mapping_rows(src, summ, rows):
    ai_src, ai_sum = src["ai"], summ["ai"]
    rep = {r["id"]: r["ev"]["rep"] for r in rows}
    out = []
    for e in ai_sum:
        n = len(e["members"])
        for m in e["members"]:
            d = ai_src["lv3"][m]
            kind = "단독 유지" if n == 1 else "유사 위협 통합"
            if d["domain"] != e["domain"]:
                kind = f"도메인 이동({d['domain']}→{e['domain']})" + ("" if n == 1 else " · 유사 위협 통합")
            out.append([LV0["ai"], d["도메인 (Lv1)"], m, "-", d["세부 위협 (Lv3)"], d["위험도"],
                        e["id"], e["name"], ai_src["domains"][e["domain"]], kind,
                        "●" if rep[e["id"]] == m else "", e["basis"]])
    for lv0, id_col in [("cloud", "CTC-ID"), ("ot", "OTC-ID")]:  # ATT&CK 기반 원본: 원본 행(전술 반복 행 포함) 단위
        sk = src[lv0]
        for e in summ[lv0]:
            n = len(e["members"])
            for m in e["members"]:
                tech = sk["tech"][m]
                for r in tech["rows"]:
                    t = r["tactic"]
                    if t == e["domain"]:
                        kind = "단독 유지" if n == 1 else "유사 위협 통합"
                    elif e["domain"] in tech["tactics"]:
                        kind = f"전술 중복 통합({t}→{e['domain']})"
                    else:
                        kind = f"전술 이동({t}→{e['domain']})" + ("" if n == 1 else " · 유사 위협 통합")
                    out.append([LV0[lv0], r["도메인(Lv1)"], r[id_col], m, r["세부위협(Lv3)"], r["위험도"],
                                e["id"], e["name"], sk["tactics"][e["domain"]], kind,
                                "●" if rep[e["id"]] == m else "", e["basis"]])
    for lv0 in common.TAXO:  # 자체 분류 원본: 세부위협 1행 단위, ATT&CK ID 열은 원본 'ATT&CK ID'(없으면 '-')
        sk = src[lv0]
        for e in summ[lv0]:
            n = len(e["members"])
            for m in e["members"]:
                r = sk["tech"][m]["row"]
                kind = "단독 유지" if n == 1 else "유사 위협 통합"
                if r["tactic"] != e["domain"]:
                    kind = f"도메인 이동({r['tactic']}→{e['domain']})" + ("" if n == 1 else " · 유사 위협 통합")
                attack = ", ".join(t.strip() for t in str(r["ATT&CK ID"] or "").replace("\n", ",").split(",") if t.strip())
                out.append([LV0[lv0], r["도메인(Lv1)"], m, attack or "-", r["세부위협(Lv3)"], r["위험도"],
                            e["id"], e["name"], sk["tactics"][e["domain"]], kind,
                            "●" if rep[e["id"]] == m else "", e["basis"]])
    return out


# ---------------------------------------------------------------- 시트 작성
# 시트 부제에 쓰는 원본 목록
SOURCE_NOTE = ("AI v3.2 LITE(Lv3 124) · 클라우드 v5(154행) · OT v5(141행) · 공급망 v2(세부위협 60) · 신원 v2(65) · "
               "물리·인적 v2(45)")
SUMMARY_COLS = ["구분(Lv0)", "도메인(Lv1)", "위협 ID", "세부 위협(Lv3)", "핵심 요약", "위협 설명", "공격 시나리오", "대표 사례",
                "발생가능성", "심각도", "위험도", "근거 수준", "실제 사고 수", "최근 사고(2025~)", "우선순위", "대응 방안",
                "공격 단계", "관련 기준", "연계 위협", "원본 Lv3 ID"]


def summary_values(r):
    ev = r["ev"]
    return [r["lv0"], r["lv1"], r["id"], r["name"], r["summary"], r["description"], r["scenario"], r["cases"],
            ev["likelihood"], ev["severity"], ev["risk"], ev["evidence"], ev["incidents"], ev["recent"], r["rank"],
            r["controls"], r["stage"], r["refs"], r["links_text"], r["origin"]]


def sheet_summary(wb, rows, today):
    ws = wb.create_sheet("통합 요약 매트릭스")
    title(ws, f"{TITLE} {VERSION} — 보고서용 요약 (Lv0 구분 → Lv1 도메인 → Lv3 세부 위협)",
          " · ".join(f"{common.LABEL[k]} {sum(r['kind'] == k for r in rows)}개" for k in common.KINDS) + " | "
          f"원본: {SOURCE_NOTE} | {today}", len(SUMMARY_COLS))
    groups(ws, 3, [("분류 체계", 4, "2E5496"), ("위협 내용", 4, "117A65"), ("위험 평가", 7, "A04000"),
                   ("대응", 1, "6C3483"), ("교차 매핑·추적", 4, "1F7A8C")])
    header(ws, 4, SUMMARY_COLS)
    for i, r in enumerate(rows, 5):
        put_row(ws, i, summary_values(r))
        ws.cell(i, 1).fill = fill(LV0_FILL[r["kind"]])
        ws.cell(i, 3).font = font(10, True)
        ws.cell(i, 4).font = font(10, True)
        for col in (9, 10, 13, 14, 15):
            ws.cell(i, col).alignment = CENTER
        paint_risk(ws.cell(i, 11), r["ev"]["risk"])
        paint_level(ws.cell(i, 12), r["ev"]["evidence"])
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(SUMMARY_COLS))}{len(rows) + 4}"
    widths(ws, [11, 18, 10, 22, 30, 58, 52, 56, 7, 7, 9, 11, 7, 7, 7, 44, 18, 24, 30, 22])
    print_setup(ws, "3:4")


BRIEF_COLS = ["구분(Lv0)", "도메인(Lv1)", "위협 ID", "세부 위협(Lv3)", "핵심 요약", "위험도", "우선순위", "실제 사고 수",
              "대표 사례", "핵심 대응"]


def sheet_brief(wb, rows, today):
    ws = wb.create_sheet("보고서용 간략 매트릭스")
    title(ws, f"보고서용 간략 매트릭스 {VERSION} — 본문 삽입용 (상세 문안은 '통합 요약 매트릭스')",
          "우선순위 = 같은 구분(Lv0) 안에서 위험도 → 실제 사고 수 → 최근 사고(2025~) 순 | 대표 사례 = 첫 번째 대표 사례 | "
          f"핵심 대응 = 첫 번째 대응 방안 | {today}", len(BRIEF_COLS))
    groups(ws, 3, [("분류 체계", 4, "2E5496"), ("요약", 1, "117A65"), ("위험 평가", 3, "A04000"),
                   ("사례·대응", 2, "6C3483")])
    header(ws, 4, BRIEF_COLS)
    for i, r in enumerate(rows, 5):
        put_row(ws, i, [r["lv0"], r["lv1"], r["id"], r["name"], r["summary"], r["ev"]["risk"], r["rank"],
                        r["ev"]["incidents"], r["case1"], r["control1"]])
        ws.cell(i, 1).fill = fill(LV0_FILL[r["kind"]])
        ws.cell(i, 3).font = font(10, True)
        ws.cell(i, 4).font = font(10, True)
        paint_risk(ws.cell(i, 6), r["ev"]["risk"])
        for col in (7, 8):
            ws.cell(i, col).alignment = CENTER
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(BRIEF_COLS))}{len(rows) + 4}"
    widths(ws, [11, 18, 10, 24, 44, 9, 8, 8, 36, 46])
    print_setup(ws, "3:4")


def sheet_matrix_view(wb, rows, src):
    ws = wb.create_sheet("매트릭스 뷰")
    ws.cell(1, 1, "매트릭스 뷰 — 셀 색 = 위험도 (빨강 매우 높음 · 주황 높음 · 노랑 보통 · 초록 낮음), [n] = 실제 사고 수, "
                  "같은 도메인 안에서 우선순위 순 정렬").font = font(11, True, NAVY)
    r0 = 3
    for kind in common.KINDS:
        doms = common.domains_of(kind, src[kind])
        items = [r for r in rows if r["kind"] == kind]
        band = ws.cell(r0, 1, f"{LV0[kind]} — {len(items)}개 위협 · {len(doms)}개 도메인")
        band.fill, band.font, band.alignment = fill(LV0_BAND[kind]), font(11, True, "FFFFFF"), Alignment(vertical="center")
        ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=len(doms))
        ws.row_dimensions[r0].height = 20
        depth = 0
        for j, (code, label) in enumerate(doms.items(), 1):
            col = sorted((r for r in items if r["domain"] == code), key=priority_key)
            h = ws.cell(r0 + 1, j, f"{label}\n({len(col)})")
            h.fill, h.font, h.alignment, h.border = fill(NAVY), font(9, True, "FFFFFF"), CENTER, BORDER
            for i, r in enumerate(col, r0 + 2):
                n = r["ev"]["incidents"]
                c = ws.cell(i, j, f"{r['id']} {r['name']}" + (f" [{n}]" if n else ""))
                paint_risk(c, r["ev"]["risk"], 9)
                c.alignment, c.border = Alignment(wrap_text=True, vertical="top"), BORDER
            depth = max(depth, len(col))
        ws.row_dimensions[r0 + 1].height = 44
        r0 += depth + 4
    widths(ws, [20] * max(len(common.domains_of(k, src[k])) for k in common.KINDS))
    print_setup(ws)


def sheet_domain_summary(wb, rows, src):
    ws = wb.create_sheet("도메인 요약")
    cols = ["구분(Lv0)", "도메인(Lv1)", "원본 Lv3 수", "요약 위협 수", "매우 높음", "높음", "보통", "낮음",
            "실제 사고 확인", "주요 고위험 위협"]
    title(ws, "도메인 요약 — Lv1 도메인별 요약 위협 수·위험도 분포", "원본 Lv3 수: AI는 Lv3 항목 수, 클라우드·OT는 v5 행 수"
          "(전술 중복 행 포함, 전술 이동한 원본은 이동한 도메인에 셈), 공급망·신원·물리는 원본 세부위협 수. "
          "위험도·우선순위는 같은 Lv0 안에서 비교", len(cols))
    header(ws, 4, cols)
    r = 5
    for kind in common.KINDS:
        doms = common.domains_of(kind, src[kind])
        tot = collections.Counter()
        for code, label in doms.items():
            items = [x for x in rows if x["kind"] == kind and x["domain"] == code]
            if kind == "ai":
                n_src = sum(len(x["members"]) for x in items)
            else:
                n_src = sum(len(src[kind]["tech"][m]["rows"]) for x in items for m in x["members"])
            dist = collections.Counter(x["ev"]["risk"] for x in items)
            real = sum(x["ev"]["evidence"] == "실제 사고 확인" for x in items)
            top = sorted(items, key=priority_key)[:2]
            put_row(ws, r, [LV0[kind], label, n_src, len(items)] + [dist[k] for k in common.RISK_ORDER] +
                    [real, "\n".join(f"{x['id']} {x['name']} ({x['ev']['risk']}, 우선순위 {x['rank']})" for x in top)])
            ws.cell(r, 1).fill = fill(LV0_FILL[kind])
            for col in range(3, 10):
                ws.cell(r, col).alignment = CENTER
            tot.update({"src": n_src, "n": len(items), "real": real, **dist})
            r += 1
        put_row(ws, r, [LV0[kind], "합계", tot["src"], tot["n"]] + [tot[k] for k in common.RISK_ORDER] + [tot["real"], ""])
        for col in range(1, 11):
            ws.cell(r, col).font = font(10, True)
            ws.cell(r, col).fill = fill("F2F3F4")
        for col in range(3, 10):
            ws.cell(r, col).alignment = CENTER
        r += 2
    for k, col in zip(common.RISK_ORDER, range(5, 9)):
        ws.cell(4, col).fill = fill(RISK_COLOR[k][0])
        ws.cell(4, col).font = font(10, True, RISK_COLOR[k][1])
    ws.freeze_panes = "C5"
    widths(ws, [16, 30, 10, 10, 9, 9, 9, 9, 11, 60])
    print_setup(ws, "4:4", "A4")


AI_DETAIL_COLS = ["통합 위협 ID", "구분(Lv0)", "도메인(Lv1)", "Lv2 ID", "위협 분류(Lv2)", "원본 ID(Lv3)", "세부 위협(Lv3)",
                  "요약설명", "참조", "발생가능성", "심각도", "위험도", "근거 수준", "실제 사고 수", "심각도 근거",
                  "주 공격 단계", "ATLAS 기법 (주)", "OWASP 매핑 (주)", "핵심 통제 요약(Lv2)",
                  "ATLAS 사고 사례 ID", "OWASP 인용 사고 수"]
CLOUD_SRC_TAIL = ["ATT&CK ID", "통합된 기법", "구분", "Cloud 범위", "Cloud 플랫폼", "기타 플랫폼", "설명 출처",
                  "CSA Top Threats 2026", "AI 매트릭스 연계(UT)", "AWS (TTC)", "Azure (ATRM)", "Kubernetes",
                  "발생가능성 근거 (자동 산정)", "최근 사고(2025~)", "실증·연구 수", "ATT&CK 사례 수",
                  "사고 매핑 근거", "관련 사례 ID", "ATT&CK 링크", "탐지·대응 포인트"]
CLOUD_DETAIL_COLS = ["통합 위협 ID", "구분(Lv0)", "도메인(Lv1)", "Lv2 ID", "위협 분류(Lv2)", "원본 ID(Lv3)", "세부 위협(Lv3)",
                     "요약설명", "참조", "발생가능성", "심각도", "위험도", "근거 수준", "실제 사고 수", "심각도 근거"] + CLOUD_SRC_TAIL


def detail_sheet(wb, name, sub, cols, data, spec, col_widths):
    ws = wb.create_sheet(name)
    title(ws, name, sub, len(cols))
    groups(ws, 3, spec)
    header(ws, 4, cols)
    for i, (kind, vals, risk, level) in enumerate(data, 5):
        put_row(ws, i, vals, 9)
        ws.cell(i, 1).font = font(9, True)
        ws.cell(i, 2).fill = fill(LV0_FILL[kind])
        for col in (10, 11, 14):
            ws.cell(i, col).alignment = CENTER
        paint_risk(ws.cell(i, 12), risk, 9)
        paint_level(ws.cell(i, 13), level, 9)
    ws.freeze_panes = "H5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(data) + 4}"
    widths(ws, col_widths)
    print_setup(ws, "3:4")


def sheet_ai_detail(wb, ai_src, ai_sum):
    back = {m: e["id"] for e in ai_sum for m in e["members"]}
    data = []
    for k, d in ai_src["lv3"].items():
        vals = [back[k], LV0["ai"], d["도메인 (Lv1)"], d["UT-ID"], d["위협 분류 (Lv2)"], k, d["세부 위협 (Lv3)"],
                d["요약설명"], d["참조"], d["발생가능성"], d["심각도"], d["위험도"], d["근거 수준"], d["실제 사고 수"],
                d["심각도 근거"], d["주 공격 단계"], d["ATLAS 기법 (주)"], d["OWASP 매핑 (주)"], d["핵심 통제 요약"],
                ", ".join(d["atlas_incidents"]), d["owasp_incidents"]]
        data.append(("ai", vals, d["위험도"], d["근거 수준"]))
    detail_sheet(wb, "AI 위협 상세", f"원본 그대로: 통합 AI 보안위협 매트릭스 v3.2 LITE 통합매트릭스 Lv3 {len(data)}개 "
                 "(위험도는 원본 수식 값을 정적 값으로 기록) + 통합 위협 ID 역참조 · 실제 사고 수 산정 근거(ATLAS 사고 사례·OWASP 인용)",
                 AI_DETAIL_COLS, data,
                 [("통합 추적", 2, "7B241C"), ("분류 체계 (원본)", 7, "2E5496"), ("위험 평가 (원본)", 6, "A04000"),
                  ("교차 매핑·통제 (원본)", 4, "1F7A8C"), ("실제 사고 근거", 2, "1E8449")],
                 [10, 11, 18, 7, 16, 8, 18, 50, 80, 7, 7, 9, 11, 7, 22, 16, 22, 13, 30, 16, 8])


def csa_ids(value):
    """CSA Top Threats 연계 표기에서 이슈 ID만 남김: 'SI-01 부적절한 IAM · SI-07 APT' → 'SI-01 · SI-07'."""
    return " · ".join(re.findall(r"SI-\d+", value or "")) or None


def sheet_cloud_detail(wb, cl_src, cl_sum, release):
    back = {m: e["id"] for e in cl_sum for m in e["members"]}
    data = []
    for r in cl_src["rows"]:
        tid = str(r["ATT&CK ID"])
        tail = [csa_ids(r[c]) if (release and c == "CSA Top Threats 2026") else r[c] for c in CLOUD_SRC_TAIL]
        vals = [back[tid], LV0["cloud"], r["도메인(Lv1)"], tid.split(".")[0], r["위협분류(Lv2)"], r["CTC-ID"],
                r["세부위협(Lv3)"], r["요약설명"], r["참조"], r["발생가능성"], r["심각도"], r["위험도"], r["근거 수준"],
                r["실제 사고 수"], r["심각도 근거"]] + tail
        data.append(("cloud", vals, r["위험도"], r["근거 수준"]))
    csa_note = " · CSA 열은 이슈 ID만 표기(배포본)" if release else ""
    detail_sheet(wb, "클라우드 위협 상세", f"원본 그대로: 통합 클라우드 보안위협 매트릭스 v5 통합 매트릭스 {len(data)}행 "
                 f"(전술 중복 행 포함) + 통합 위협 ID 역참조. Lv2 ID는 ATT&CK 상위기법 ID{csa_note}",
                 CLOUD_DETAIL_COLS, data,
                 [("통합 추적", 2, "7B241C"), ("분류 체계 (원본)", 7, "2E5496"), ("위험 평가 (원본)", 6, "A04000"),
                  ("기법·교차 매핑 (원본)", 12, "1F7A8C"), ("실제 근거 (원본)", 7, "1E8449"), ("탐지·대응", 1, "6C3483")],
                 [10, 11, 15, 8, 20, 12, 22, 48, 60, 7, 7, 9, 11, 7, 24,
                  10, 12, 12, 11, 14, 14, 10, 24, 18, 20, 20, 16, 28, 8, 7, 7, 14, 22, 26, 48])


OT_PROFILES = ["제어계통", "안전계통", "원격 필드", "감시·운영", "IT/OT 경계", "IoT·임베디드"]
OT_SRC_TAIL = ["ATT&CK·EMB3D ID", "ATT&CK·EMB3D 위협명", "구분", "통합된 Enterprise 기법", "설명 출처", "대상 자산", "Purdue 계층",
               "적용 프로파일", "완화책(ATT&CK·EMB3D)", "클라우드 매트릭스 연계", "AI 매트릭스 연계(UT)", "OWASP IoT Top10",
               "EMB3D 연계(성숙도·근거)", "IEC 62443-3-3(SR)", "IEC 62443-4-2(CR 등)", "NIST SP 800-53",
               "발생가능성 근거 (자동 산정)", "최근 사고(2025~)", "ATT&CK 사례 수", "공개 취약점(CVE)", "KEV(OT)", "실증·연구 수",
               "사고 매핑 근거", "관련 사례 ID", "ATT&CK 사례 주체", "ATT&CK 링크", "탐지·대응 포인트"]
OT_DETAIL_COLS = CLOUD_DETAIL_COLS[:15] + OT_SRC_TAIL


def sheet_ot_detail(wb, ot_src, ot_sum):
    back = {m: e["id"] for e in ot_sum for m in e["members"]}
    data = []
    for r in ot_src["rows"]:
        key = str(r["ATT&CK·EMB3D ID"])
        profile = " · ".join(p for p in OT_PROFILES if r.get(p)) or None  # 원본 ● 6열 → 한 열
        tail = [profile if c == "적용 프로파일" else r[c] for c in OT_SRC_TAIL]
        vals = [back.get(key), LV0["ot"], r["도메인(Lv1)"], r["OTC-ID"].split(".")[0], r["위협분류(Lv2)"], r["OTC-ID"],
                r["세부위협(Lv3)"], r["요약설명"], r["참조"], r["발생가능성"], r["심각도"], r["위험도"], r["근거 수준"],
                r["실제 사고 수"], r["심각도 근거"]] + tail
        data.append(("ot", vals, r["위험도"], r["근거 수준"]))
    detail_sheet(wb, "OT 위협 상세", f"원본 그대로: 통합 OT 보안위협 매트릭스 v5 통합 매트릭스 {len(data)}행 "
                 "(전술 반복 행 포함) + 통합 위협 ID 역참조. Lv2 ID는 원본 OTC Lv2, 적용 프로파일은 원본 ● 6열을 한 열로 표기",
                 OT_DETAIL_COLS, data,
                 [("통합 추적", 2, "7B241C"), ("분류 체계 (원본)", 7, "2E5496"), ("위험 평가 (원본)", 6, "A04000"),
                  ("기법·교차 매핑 (원본)", 16, "1F7A8C"), ("실제 근거 (원본)", 10, "1E8449"), ("탐지·대응", 1, "6C3483")],
                 [10, 11, 15, 10, 20, 12, 22, 48, 60, 7, 7, 9, 11, 7, 24,
                  11, 20, 11, 14, 9, 26, 10, 18, 30, 18, 18, 8, 30, 16, 20, 16, 28, 7, 7, 7, 7, 7, 14, 18, 26, 22, 48])


# 공급망·신원·물리 상세 시트: 앞 15열은 클라우드·OT 상세와 같고, 나머지는 원본 열 순서 그대로.
# 원본 ● 열(프로파일:·신원:·주체:)은 한 열로 합치고, 원본 '우선순위'는 요약 우선순위와 구분해 '원본 우선순위'로 표기
TAXO_HEAD_USED = {"도메인(Lv1)", "위협분류(Lv2)", "세부위협(Lv3)", "요약설명", "참조", "발생가능성", "심각도", "위험도",
                  "근거 수준", "실제 사고 수", "심각도 근거"}
TAXO_DOT_GROUPS = {"프로파일:": "적용 프로파일", "신원:": "신원 유형", "주체:": "위협 주체"}
# 배포본에서 ID만 남기는 열 — CSA CCM(CSA 이용 조건)·CIS Controls(CC BY-NC-ND 4.0)는 통제명(국문 번역)을 빼고 ID만
TAXO_RELEASE_IDS = {("identity", "CSA CCM v4.1"): r"[A-Z]{3}-\d{2}", ("identity", "CIS Controls v8.1"): r"^\d+\.\d+"}


def taxo_tail(header, id_col):
    """원본 열 → 상세 시트 뒤쪽 열 [(표시명, 원본 열 목록)] (● 열 묶음은 원본 열 여러 개)."""
    out, seen = [], set()
    for c in header:
        if c in TAXO_HEAD_USED or c == id_col:
            continue
        group = next((g for g in TAXO_DOT_GROUPS if str(c).startswith(g)), None)
        if group:
            if group not in seen:
                seen.add(group)
                out.append((TAXO_DOT_GROUPS[group], [x for x in header if str(x).startswith(group)]))
            continue
        out.append(("원본 우선순위" if c == "우선순위" else c, [c]))
    return out


def release_ids(value, pat):
    """여러 줄 기준 표기에서 ID만 남김: 'IAM-13 강한 인증\nIAM-02 …' → 'IAM-13 · IAM-02'."""
    found = [m for l in str(value or "").split("\n") for m in re.findall(pat, l.strip())]
    return " · ".join(uniq(found)) or None


def sheet_taxo_detail(wb, kind, sk, summ, release):
    back = {m: e["id"] for e in summ for m in e["members"]}
    tail = taxo_tail(sk["header"], sk["id_col"])
    cols = CLOUD_DETAIL_COLS[:15] + [name for name, _ in tail]
    data = []
    for r in sk["rows"]:
        key = r[sk["id_col"]]
        vals = []
        for name, src_cols in tail:
            if len(src_cols) > 1:  # ● 열 묶음 → '오픈소스 · CI/CD'
                vals.append(" · ".join(c.split(":", 1)[1] for c in src_cols if r.get(c)) or None)
            elif release and (kind, src_cols[0]) in TAXO_RELEASE_IDS:
                vals.append(release_ids(r[src_cols[0]], TAXO_RELEASE_IDS[(kind, src_cols[0])]))
            else:
                vals.append(r[src_cols[0]])
        vals = [back[key], LV0[kind], r["도메인(Lv1)"], key.rsplit(".", 1)[0] if "." in key else key, r["위협분류(Lv2)"], key,
                r["세부위협(Lv3)"], r["요약설명"], r["참조"], r["발생가능성"], r["심각도"], r["위험도"], r["근거 수준"],
                r["실제 사고 수"], r["심각도 근거"]] + vals
        data.append((kind, vals, r["위험도"], r["근거 수준"]))
    note = " · CSA CCM·CIS Controls 열은 ID만 표기(배포본)" if release and kind == "identity" else ""
    wide = {"핵심 요약": 30, "탐지·대응 포인트": 48, "관련 사례 ID": 22, "발생가능성 근거 (자동 산정)": 26, "ATT&CK 기법명": 22}
    detail_sheet(wb, DETAIL_SHEET[kind], f"원본 그대로: {common.TAXO[kind]['name']} 통합 매트릭스 {len(data)}행 "
                 f"+ 통합 위협 ID 역참조. Lv2 ID는 원본 세부위협 ID의 상위 번호, 원본 ● 열(프로파일·신원 유형·위협 주체)은 한 열로 표기{note}",
                 cols, data,
                 [("통합 추적", 2, "7B241C"), ("분류 체계 (원본)", 7, "2E5496"), ("위험 평가 (원본)", 6, "A04000"),
                  ("교차 매핑·근거·대응 기준 (원본)", len(tail), "1F7A8C")],
                 [10, 11, 15, 10, 20, 12, 22, 48, 60, 7, 7, 9, 11, 7, 24] + [wide.get(name, 16) for name, _ in tail])


def sheet_mapping(wb, maps):
    ws = wb.create_sheet("부록-재구성 매핑")
    cols = ["구분(Lv0)", "원본 도메인(Lv1)", "원본 ID(Lv3)", "ATT&CK ID", "원본 세부 위협(Lv3)", "원본 위험도",
            "통합 위협 ID", "통합 세부 위협(Lv3)", "통합 도메인(Lv1)", "처리 유형", "위험 대표", "통합 근거"]
    kinds = collections.Counter(m[9].split(" · ")[0].split("(")[0] for m in maps)
    title(ws, "부록 — 재구성 매핑 (원본 Lv3 → 통합 요약 Lv3)",
          f"AI {sum(m[0] == LV0['ai'] for m in maps)}개 · 클라우드 {sum(m[0] == LV0['cloud'] for m in maps)}행 · "
          f"OT {sum(m[0] == LV0['ot'] for m in maps)}행 · "
          + " · ".join(f"{common.LABEL[k]} {sum(m[0] == LV0[k] for m in maps)}개" for k in common.TAXO) + " 전수 | "
          + " · ".join(f"{k} {v}" for k, v in kinds.most_common())
          + " | 위험 대표 ● = 요약 위험도(발생가능성·심각도)를 결정한 원본", len(cols))
    groups(ws, 3, [("원본", 6, "2E5496"), ("통합 요약", 3, "7B241C"), ("재구성", 3, "1F7A8C")])
    header(ws, 4, cols)
    for i, m in enumerate(maps, 5):
        put_row(ws, i, m, 9)
        ws.cell(i, 1).fill = fill(LV0_FILL[KIND_OF[m[0]]])
        paint_risk(ws.cell(i, 6), m[5], 9)
        ws.cell(i, 7).font = font(9, True)
        if not m[9].startswith(("단독", "유사")):
            ws.cell(i, 10).font = font(9, True, "A04000")
        ws.cell(i, 11).alignment = CENTER
        ws.cell(i, 11).font = font(9, True, "C0392B")
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{len(maps) + 4}"
    widths(ws, [11, 20, 12, 10, 28, 9, 10, 26, 22, 22, 7, 70])
    print_setup(ws, "3:4")


CRITERIA = [
    ("■ 1. 재구성 원칙", None),
    ("단위", "원본 Lv3(AI 124개, 클라우드 고유 기법 126개·154행, OT 고유 기법·항목 118개·141행, 공급망 세부위협 60개, 신원 65개, "
             "물리·인적 45개)를 공격 메커니즘·대상·통제가 같은 것끼리 묶어 요약 Lv3로 재구성 (AI 61개, 클라우드 52개, OT 50개, "
             "공급망 33개, 신원 35개, 물리·인적 27개). 원본 Lv2는 요약표에서 생략하고 원본 상세 시트로 추적"),
    ("통합 기준", "같은 공격 흐름의 연속 단계이거나, 대상만 다른 같은 기법이거나, 같은 통제로 막히는 위협을 통합. 피해 규모·근거가 뚜렷이 "
                "다른 위협(예: 코드 저장소·CRM·DB 수집)은 분리 유지"),
    ("클라우드 전술 중복", "ATT&CK에서 여러 전술에 걸친 기법(원본의 중복 행)은 주 전술 1곳에만 배치하고 나머지는 '공격 단계'의 연관 전술로 표기: "
                       "T1078→초기 침투, T1133→초기 침투, T1053→지속성, T1072→실행, T1098.x→지속성(T1098.003만 권한 상승), "
                       "T1543.005→지속성, T1556.x→지속성, T1484.x→권한 상승, T1557→자격증명 접근"),
    ("OT 전술 중복", "ATT&CK for ICS에서 여러 전술에 걸친 19개 기법(원본 반복 행 23)도 주 전술 1곳에 배치: T0866·T0886·T0859·T1694.x·"
                   "T1056→횡적 이동, T1053·T1543→지속성, T1484.001→권한 상승, T0858·T0874→실행, T0851→회피, T1692.x→공정 제어 훼손, "
                   "T1693.x→대응 기능 억제, T0887→탐색. 전술마다 위험 값이 다른 반복 기법은 주 전술 행 값으로 평가"),
    ("경계 이동", "도메인·전술 경계를 넘긴 10건 — 클라우드·AI 6건: T1684.x 은닉→초기 침투, T1546 지속성·권한 상승→실행, T1525 "
                "지속성→실행, T1550.x 횡적 이동→자격증명 접근, AZT704 영향→수집, AI UT-14.2 D04→D06 / OT 4건(EMB3D 신설 행): "
                "EMB3D-CO.2 수집→초기 침투, EMB3D-LM.1·LM.2 횡적 이동→수집, EMB3D-EV 회피→수집. 공급망·신원·물리는 원본 도메인 "
                "경계를 그대로 유지(이동 0건)"),
    ("OT 묶음 기준", "영향은 ATT&CK ICS 영향 분류 단위(가용성·생산 / 제어 / 감시 / 안전·보호·설비 / 운영 정보), 대응 기능 억제는 "
                    "운영자 대응을 막는 수단별(파괴·정지·경보 억제·통신 차단·펌웨어·자격증명 변경), 제어기 로직 기법(태스크 변조·"
                    "프로그램 변조·프로그램 다운로드)은 킬체인 위치가 달라 전술별 분리, Enterprise 편입 기법은 메커니즘이 같은 ICS "
                    "기법과 묶고 아니면 Enterprise끼리 묶음, EMB3D 신설 7행은 하드웨어 공격 2개 항목으로 묶음"),
    ("공급망·신원·물리 묶음 기준", "원본 위협분류(Lv2)를 출발점으로, 같은 Lv2 안에서 통제가 같은 세부위협을 묶음. 피해 규모·근거 밀도가 "
                              "두드러진 세부위협은 단독 유지(예: 정상 패키지 악성 버전 게시 · MFA 미적용 · 재직자 기밀 유출). Lv2가 "
                              "세부위협 1개뿐이면 같은 도메인의 인접 Lv2와 묶을지 검토(예: 의존성 혼동·재점유 + 버전 미고정, "
                              "헬프데스크 사칭 + 계정 복구 절차). 영향(IM) 도메인은 피해 유형 단위로 유지"),
    ("단독 유지", "다른 항목과 메커니즘·통제가 달라 묶을 대상이 없는 항목. '통합 근거'에 '단독 유지 — 사유'로 기록"),
    ("■ 2. 위험 평가 (요약 Lv3 재산정)", None),
    ("위험도", "구성 원본별 위험도(발생가능성 × 심각도)의 최댓값. 발생가능성과 심각도를 서로 다른 원본에서 가져와 조합하지 않음"
             "(교차 결합 금지). 3×3 표: 상×상 매우 높음 / 상×중·중×상 높음 / 중×중·상×하·하×상 보통 / 그 외 낮음 (원본 여섯 매트릭스와 같은 표)"),
    ("발생가능성 보정", "구성 원본의 실제 사고 합집합이 2건 이상이면, 사고가 1건 이상 확인된 원본만 발생가능성을 '상'으로 보정"),
    ("발생가능성·심각도", "위험도 최댓값을 낸 원본(대표 원본)의 값을 표시. 대표 원본은 '부록-재구성 매핑'의 '위험 대표' 열에 ● 표시"),
    ("실제 사고 수", "AI: ATLAS Incident 유형 사례 ID(중복 제거) + OWASP 인용 사고 수(원본 건수 합에서, 참조에 같은 사례명(시점)으로 "
                  "실린 중복분을 뺌). 클라우드: 사고 DB 680건 중 집계 대상 유형(사고·캠페인·사례연구(익명)·규제공시(8-K)·"
                  "위협인텔 보고서, 집계 포함=Y)의 사고 ID(중복 제거). OT: OT 보안사고 DB 85건 중 '집계 = 실제 사고'(76건)인 사고 ID"
                  "(중복 제거) — 실증·연구, 위협인텔(배포 전 발견된 공격 도구), 원인 번복 사례는 제외. 공급망·신원·물리: 각 사고 DB"
                  "(167건·126건·100건) 중 '집계 상태 = 실제 사고'(152건·116건·89건)인 사고 ID(중복 제거) — 실증·연구, 위협인텔, "
                  "물리의 정부 경보·제외(검증)는 세지 않음"),
    ("최근 사고(2025~)", "실제 사고 중 2025년 이후 사고 수. 클라우드·OT·공급망·신원·물리: 사고 DB 기준일(시점). AI: ATLAS 사례 일자 + "
                       "참조에 시점이 적힌 OWASP 인용 사고(시점이 없는 인용 사고는 제외)"),
    ("우선순위", "같은 Lv0 안에서 위험도 → 실제 사고 수 → 최근 사고(2025~) → 위협 ID 순으로 매긴 순위 (1 = 최우선). 영향(IM) 도메인도 "
              "같은 수식으로 매기고, 공급망·신원·물리는 개요에 영향 도메인을 뺀 '단계 위협 기준' 상위 항목을 함께 표기"),
    ("근거 수준", "구성 원본 중 가장 강한 수준: 실제 사고 확인 > 실사용 기법 포함 > 실증·공개 취약점 > 이론·시나리오"),
    ("Lv0 간 비교 유의", "클라우드는 사고 DB(680건), OT·공급망·신원·물리는 각 사고 DB와 ATT&CK 공식 절차 매핑을 근거로 해 실제 사고 근거 "
                       "밀도가 AI(ATLAS 사례·OWASP 인용)와 다르고, 심각도 기준도 원본마다 다름(OT 물리적 결과, 공급망 '상' 포화, "
                       "신원·물리 '상' 제한). 위험도·사고 수·우선순위는 같은 Lv0 안에서 비교하고, Lv0 간 등급을 직접 비교하지 않음"),
    ("Lv0 간 중복", "신원은 클라우드 계정·자격증명 위협과, 공급망은 AI 공급망·클라우드 공급망 침해와, 물리는 신원 내부자 위협과 겹치는 "
                  "부분이 있음. 원본 경계 규칙에 따라 관점별로 각 Lv0에 두고 '연계 위협'으로 연결 — Lv0별 위협 수의 합은 고유 "
                  "위협 수가 아님"),
    ("공급망 위험도 포화", "공급망 원본은 심각도 '상'이 60개 중 42개라 요약도 대부분 '매우 높음' — 공급망 항목 간 비교는 실제 사고 수·"
                         "최근 사고·우선순위로 함"),
    ("OT 근거 편중", "ATT&CK 공식 절차가 촘촘한 대형 사고 5건(Stuxnet·2015/2016 우크라이나·Triton·2025 폴란드)이 원본 36~46행에 매핑돼 "
                   "OT 요약 대부분이 '실제 사고 확인'임. OT 항목 간 비교는 실제 사고 수·우선순위로 하고, 개요에 이 5건에서만 실제 사고가 "
                   "확인되는 항목을 표기"),
    ("OT 우선순위 해석", "OT 심각도는 물리적 결과 기준이라 대응 기능 억제·공정 제어 훼손·영향 단계가 상위권에 몰림. 진입·확산 차단 "
                       "우선순위는 매트릭스 뷰·도메인 요약의 초기 침투·횡적 이동 열에서 확인"),
    ("물리 공개 사고 편향", "물리·인적 원본은 공개 사고 기준이라 기소·판결로 드러나는 기술 유출·내부자 사건이 높게 나오고, 공개·귀속되지 "
                          "않는 물리 침입·은닉 감시는 과소평가될 수 있음"),
    ("신규 공격면 유의", "발생가능성은 확인된 실제 사고 기준이라, 공개 취약점은 많으나 사고 보고가 적은 신규 공격면"
                      "(예: 간접 프롬프트 인젝션·AI 패키지 환각)은 과소평가될 수 있음. OT는 사고 공개가 적고(과소보고) 설계상 인증이 "
                      "없는 기능의 악용은 CVE가 없어 '이론·시나리오'는 '공개 근거 없음'으로 읽음"),
    ("■ 3. 문안 작성 기준", None),
    ("핵심 요약", "한 줄 50자 이내. '수단 + 대상 + 결과'를 명사형으로 — 보고서 본문·슬라이드용"),
    ("위협 설명", "'- '로 시작하는 2줄. 1줄: 공격 방법·대상·결과를 담아 '~위협'으로 종결. 2줄: 영향·발생 조건·탐지 곤란 사유"),
    ("공격 시나리오", "1~2줄, 'A → B → C' 단계 흐름. 구성 원본의 참조·사례에 있는 경로만 사용"),
    ("대표 사례", "1~2줄, '[유형] 사례명(시점): 요지'. 구성 원본 참조에 있는 사례만 쓰며(빌드 검증기가 사례명·시점·유형 대조), "
              "실제 사고를 우선하고 같은 사고가 여러 행에 반복되지 않게 배치(OT·공급망·신원·물리는 원본 참조의 사고 ID 기준으로 "
              "빌드 검증이 반복을 경고하고, 불가피한 반복만 허용 항목·사유를 등록). 원본에 사례가 없으면 '공개 사고 미확인 — 사유'"),
    ("대응 방안", "2~3줄 명사형. 예방(권한·구성) → 차단·보호 → 탐지 순. OT는 계정 잠금·패치·재부팅이 운전에 주는 영향을 고려해 "
              "보상 통제(망 분리·허용 목록·OT 네트워크 감시)를 함께 적음"),
    ("물리·인적 작성 제약", "방어 자료로 한정 — 수법은 탐지·식별에 필요한 징후 수준까지만 적고 침입·도청·촬영·포섭의 실행 방법은 쓰지 않음. "
                          "드론은 탐지·식별·신고·차폐만 권고(기체 무력화·전파 교란은 권한 있는 기관 영역). 인적 징후는 객관적 행동·"
                          "직무 사실만 쓰고 보호 특성(출신·종교·성별 등)은 쓰지 않음. 개인은 실명 대신 직책·관계로 적고, 임직원 "
                          "모니터링은 개인정보보호법·노사 협의 범위로 한정(빌드 검증기가 금지 표현 검사)"),
    ("관련 기준", "AI: OWASP 주 매핑 + ATLAS 기법(상위 ID). 클라우드: ATT&CK ID(상위기법별 축약) + CSA Top Threats 2026 이슈 ID. "
              "OT: ATT&CK ID(ICS·Enterprise 편입, 상위기법별 축약) + EMB3D TID(EMB3D 신설 행) + IEC 62443-3-3 SR(구성 원본 "
              "완화책의 요구사항 ID). 공급망: ATT&CK + SLSA v1.2 위협 + OWASP CI/CD·OSS Top 10 ID. 신원: ATT&CK + Browser & "
              "Identity Attacks(SAT) + OWASP NHI Top 10 ID. 물리·인적: ATT&CK(있는 행만) + ITKB 내부자 관측 기법 + ISO/IEC "
              "27001:2022 부속서 A + ISMS-P 항목 번호. 기준 문서 내용은 싣지 않고 연계 ID만 표기"),
    ("연계 위협", "여섯 Lv0 교차 연계. 클라우드 원본 'AI 매트릭스 연계(UT)' 열, OT 원본 '클라우드 매트릭스 연계'·'AI 매트릭스 연계(UT)' 열, "
              "공급망·신원·물리 원본의 '다른 매트릭스 연계' 열(상세 ID는 요약 ID로 변환), 요약 작성 시 검토한 연계를 양방향으로 표기"),
    ("공격 단계", "AI: 구성 원본 '주 공격 단계'(ATLAS 전술)를 클라우드(ATT&CK v19.2 국문) 용어로 통일 — 초기 접근→초기 침투, "
              "내부 확산→횡적 이동, 유출→반출 (방어 회피처럼 1:1 대응 용어가 없는 단계는 원문 유지). 클라우드·OT: 주 전술 + "
              "연관 전술(원본 중복·이동 전 전술). 공급망·신원·물리: 구성 원본 '공격 단계' 합집합(원본 용어 — 물리는 접근·내부자화 등 "
              "물리·인적 단계 포함). 원본 상세 시트는 원문 그대로"),
    ("배포본", "CSA Top Threats 2026·CCM은 'CSA 이용 조건(개인·비상업적, 재배포 금지)', CIS Controls v8.1은 CC BY-NC-ND 4.0이므로, "
             "고객 배포본(--release)은 클라우드 상세 시트의 CSA 열과 신원 상세 시트의 CSA CCM·CIS Controls 열을 ID만 남김. "
             "외부 배포 전 각 출처 이용약관 확인"),
    ("■ 4. 사례 유형", None),
    ("[실제 사고]", "공개 보도·공시·위협 인텔로 확인된 실제 침해·악용 (AI: ATLAS Incident·OWASP 인용 사고, 클라우드·OT·공급망·신원·물리: "
                  "사고 DB 집계 대상)"),
    ("[공개 취약점]", "CVE·공개 보안 권고로 확인된 취약점 (OT: CISA ICS 권고·KEV, 신원·공급망: KEV 판정 포함)"),
    ("[실증]", "연구·레드팀·벤더가 실제 제품에서 재현한 공격 (원본 '실증·연구' 포함)"),
    ("[ATT&CK 사례]", "MITRE ATT&CK에 기록된 그룹·소프트웨어·캠페인의 해당 기법 사용 사례"),
    ("[시나리오]", "OWASP 등 기준 문서의 가상 공격 시나리오"),
    ("[위협인텔]", "OT·신원·물리. 배포 전에 발견된 공격 도구의 능력 분석(OT: PIPEDREAM·COSMICENERGY) 또는 위협인텔 보고 — 실제 사고로 세지 않음"),
    ("[EMB3D]", "OT 전용. MITRE EMB3D 장치 위협(TID) 인용 — 시연·분석 근거로, 구성 원본에 실제 사고·실증 사례가 없을 때만 "
                "사용(빌드 검증)"),
    ("[시연]", "신원 전용. Browser & Identity Attacks Matrix(SAT) 등 공개 시연 — 실제 사고로 세지 않음"),
    ("[정부 경보]", "물리 전용. 정부·수사기관의 공개 경보(실제 사례 요지 포함) — 원본 규칙에 따라 실제 사고로 세지 않음"),
    ("■ 5. 용어", None),
    ("자격증명", "credential. 비밀번호·액세스 키·토큰·인증서 등 인증 수단 전체 (인증정보·크리덴셜 대신 사용)"),
    ("반출 / 유출", "반출: 공격자가 데이터를 밖으로 빼내는 행위(Exfiltration) / 유출: 데이터가 외부로 새어 나간 결과"),
    ("에이전트", "LLM이 도구 호출·코드 실행·외부 시스템 연동으로 작업을 자율 수행하는 AI 시스템"),
    ("제어 평면 / 데이터 평면", "클라우드 자원을 만들고 설정하는 관리 API 계층 / 워크로드·데이터에 직접 접근하는 계층"),
    ("제어기", "PLC·RTU·DCS 제어기·안전 제어기처럼 현장 공정을 직접 제어하는 장치"),
    ("EWS / HMI", "엔지니어링 워크스테이션(제어 로직 작성·다운로드) / 운영자가 공정을 감시·조작하는 화면"),
    ("안전계통(SIS)", "위험 상황에서 공정을 안전하게 정지시키는 독립 계통"),
    ("Purdue 계층", "OT 자산을 L0 현장 장치부터 L3.5 IT/OT 경계(DMZ)까지 나눈 참조 모델"),
    ("게시 자격 / 레지스트리", "패키지를 레지스트리(npm·PyPI 등 패키지 저장소)에 올릴 수 있는 토큰·OIDC 신뢰 설정 / 패키지·이미지를 "
                           "배포하는 공개·사설 저장소"),
    ("IdP / 비인간 신원(NHI)", "사용자 인증·SSO를 제공하는 ID 공급자(Entra ID·Okta·AD FS 등) / 서비스 계정·API 키·워크로드 ID·OAuth 앱·"
                             "AI 에이전트처럼 사람이 아닌 신원"),
    ("내부자 / 비의도 내부자", "조직의 정상 접근 권한을 가진 재직자·협력사·파견 인력 / 악의 없이 분실·부주의로 정보를 노출한 내부자"),
]

def sheet_criteria(wb):
    ws = wb.create_sheet("부록-작성·평가 기준")
    title(ws, "부록 — 작성·평가 기준", "요약 Lv3 재구성·위험 재산정·문안 작성 규칙 (원본 산정 규칙은 각 원본 워크북의 평가 기준 시트 참조)", 2)
    r = 4
    for k, v in CRITERIA:
        if v is None:
            c = ws.cell(r, 1, k)
            c.font = font(11, True, NAVY)
            r += 1
            continue
        put_row(ws, r, [k, v])
        ws.cell(r, 1).font = font(10, True)
        ws.cell(r, 1).fill = fill("F2F3F4")
        r += 1
    widths(ws, [20, 120])
    print_setup(ws, paper="A4")


def sheet_overview(wb, rows, src, maps, today, release):
    ai_src, cl_src, ot_src = src["ai"], src["cloud"], src["ot"]
    sc_src, id_src, ph_src = src["supplychain"], src["identity"], src["physical"]
    ws = wb.create_sheet("개요", 0)
    ws.cell(1, 1, f"{TITLE} {VERSION}").font = font(16, True, NAVY)
    variant = "고객 배포본" if release else "내부본"
    ws.cell(2, 1, f"컨설팅 보고서용 요약 위협 매트릭스 — AI·클라우드·OT·공급망·신원·물리·인적 | {variant} | 작성 {today}").font = \
        font(10, color="404040")
    merge_rows = []  # 서술형 줄: B~K 병합
    r = 4

    def section(text):
        nonlocal r
        ws.cell(r, 1, text).font = font(12, True, NAVY)
        r += 1

    def line(k, v=""):
        nonlocal r
        ws.cell(r, 1, k).font = font(10, True)
        ws.cell(r, 1).alignment = Alignment(vertical="top")
        c = ws.cell(r, 2, v)
        c.font, c.alignment = font(10), Alignment(wrap_text=True, vertical="top")
        merge_rows.append(r)
        r += 1

    def table_header(cells):
        """cells: [(열 시작, 열 끝, 제목)]"""
        nonlocal r
        for c0, c1, text in cells:
            for col in range(c0, c1 + 1):
                ws.cell(r, col).fill, ws.cell(r, col).border = fill(HDR), BORDER
            cell = ws.cell(r, c0, text)
            cell.font, cell.alignment = font(10, True, "FFFFFF"), CENTER
            if c1 > c0:
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        r += 1

    n = {k: sum(x["kind"] == k for x in rows) for k in common.KINDS}
    section("■ 목적")
    line("", "AI·클라우드·OT·소프트웨어 공급망·신원·물리·인적 보안위협을 하나의 체계(Lv0 구분 → Lv1 도메인 → Lv3 세부 위협)로 "
             "요약·통합해, 보고서에 바로 쓸 수 있는 위협 목록·위험 평가·대표 사례·대응 방안을 제공. 원본 상세 내용은 별도 시트로 "
             "함께 수록해 근거를 추적")
    r += 1
    section("■ 원본 (각 원본 데이터를 그대로 사용)")
    line("AI 보안위협", f"통합 AI 보안위협 매트릭스 v3.2 LITE — Lv1 도메인 10 · Lv2 36 · Lv3 {len(ai_src['lv3'])} | MITRE ATLAS "
                    "v2026.09 · OWASP LLM Top 10 2026 · Agentic Top 10 2026 · GenAI Data Security 2026 v1.0 · MCP Top 10(2025 beta)")
    line("클라우드 보안위협", f"통합 클라우드 보안위협 매트릭스 v5 — ATT&CK 전술 14 · Lv3 {len(cl_src['rows'])}행(고유 기법 "
                       f"{len(cl_src['tech'])}개) | ATT&CK Cloud v19.2 · AWS TTC · Azure ATRM · Kubernetes · CSA Top Threats 2026 · "
                       "클라우드 보안사고 DB 680건")
    line("OT 보안위협", f"통합 OT/ICS/IoT 보안위협 매트릭스 v5 — ATT&CK for ICS 전술 12 + 자원 개발 · Lv3 {len(ot_src['rows'])}행"
                     f"(고유 기법·항목 {len(ot_src['tech'])}개: ICS 기법 · Enterprise 편입 · EMB3D 신설) | ATT&CK for ICS·Enterprise "
                     "v19.2 · MITRE EMB3D · CISA ICS 권고·KEV · IEC 62443·NIST SP 800-53 연계 · OWASP IoT Top 10 · "
                     f"OT 보안사고 DB {ot_src['n_incidents']}건")
    line("공급망 보안위협", f"통합 소프트웨어 공급망 보안위협 매트릭스 v2 — 공급망 단계 8 · 위협분류 32 · 세부위협 {len(sc_src['rows'])} | "
                      "ATT&CK v19.2 · SLSA v1.2 · OWASP CI/CD·OSS Top 10 · SAP Risk Explorer · CNCF 침해 유형 · CISA KEV · OSV 악성 "
                      f"패키지 · S2C2F·SSDF·Scorecard·NIST SP 800-53 | 공급망 보안사고 DB {sc_src['n_incidents']}건")
    line("신원 보안위협", f"통합 신원·계정 보안위협 매트릭스 v2 — 신원 공격면 10 · 위협분류 36 · 세부위협 {len(id_src['rows'])} | "
                     "ATT&CK v19.2 · Browser & Identity Attacks Matrix · OWASP NHI Top 10 · NIST SP 800-63-4 · OWASP ASVS 5.0 · "
                     f"CISA KEV · CSA CCM v4.1·CIS v8.1·NIST SP 800-53·ISMS-P | 신원 보안사고 DB {id_src['n_incidents']}건")
    domestic = sum(1 for x in ph_src["incidents"] if x.get("지역") == "한국")
    line("물리·인적 보안위협", f"물리·인적 보안위협 매트릭스 v2 — 물리·인적 공격면 7 · 위협분류 31 · 세부위협 {len(ph_src['rows'])} | "
                        "ATT&CK v19.2 물리·인적 접점 · MITRE CTID 내부자 TTP 지식베이스(ITKB) · NIST SP 800-53 · ISO/IEC 27001:2022 "
                        "부속서 A · ISMS-P · CERT 내부자 가이드 · NPSA·CISA 참고 문헌 | 물리·인적 보안사고 DB "
                        f"{ph_src['n_incidents']}건(국내 {domestic}건)")
    r += 1
    section("■ 구조")
    line("Lv0 구분", "AI 보안위협 / 클라우드 보안위협 / OT 보안위협(ICS·IoT 포함) / 공급망 보안위협(소프트웨어 공급망) / 신원 보안위협"
                   "(계정·IdP·비인간 신원·고객 계정) / 물리·인적 보안위협(물리 침입·감시·포섭·내부자)")
    line("Lv1 도메인", "AI: 원본 10개 도메인(D01~D10) / 클라우드: 원본 ATT&CK 14개 전술(RD~IM) / OT: 원본 ATT&CK for ICS 12개 전술 + "
                    "자원 개발(RD~IM) / 공급망: 원본 8개 단계(DE~IM) / 신원: 원본 10개 공격면(AU~IM) / 물리·인적: 원본 7개 공격면"
                    "(RC~IM) 유지")
    line("Lv3 세부 위협", f"원본 Lv3를 내용 분석 후 재구성·재작성 — AI {len(ai_src['lv3'])} → {n['ai']}개, 클라우드 "
                       f"{len(cl_src['rows'])}행({len(cl_src['tech'])}기법) → {n['cloud']}개, OT {len(ot_src['rows'])}행"
                       f"({len(ot_src['tech'])}기법·항목) → {n['ot']}개, 공급망 {len(sc_src['rows'])} → {n['supplychain']}개, "
                       f"신원 {len(id_src['rows'])} → {n['identity']}개, 물리·인적 {len(ph_src['rows'])} → {n['physical']}개, "
                       f"합계 {sum(n.values())}개")
    line("요약 열", "핵심 요약 · 위협 설명 · 공격 시나리오 · 대표 사례 · 위험 평가(발생가능성·심각도·위험도·근거 수준·실제 사고 수·"
                  "최근 사고·우선순위) · 대응 방안 · 공격 단계 · 관련 기준 · 연계 위협 · 원본 Lv3 ID")
    r += 1
    section("■ 결과 요약")
    cols = ["구분(Lv0)", "원본 Lv3", "요약 Lv3", "매우 높음", "높음", "보통", "낮음", "실제 사고 확인", "실사용 기법 포함",
            "실증·공개 취약점", "이론·시나리오"]
    table_header([(j, j, c) for j, c in enumerate(cols, 1)])
    for k, col in zip(common.RISK_ORDER, range(4, 8)):
        ws.cell(r - 1, col).fill, ws.cell(r - 1, col).font = fill(RISK_COLOR[k][0]), font(10, True, RISK_COLOR[k][1])
    n_src = {"ai": len(ai_src["lv3"]), "cloud": f"{len(cl_src['rows'])}행 / {len(cl_src['tech'])}기법",
             "ot": f"{len(ot_src['rows'])}행 / {len(ot_src['tech'])}기법"}
    n_src.update({k: len(src[k]["rows"]) for k in common.TAXO})
    for kind in common.KINDS:
        items = [x for x in rows if x["kind"] == kind]
        risk = collections.Counter(x["ev"]["risk"] for x in items)
        lvl = collections.Counter(x["ev"]["evidence"] for x in items)
        vals = [LV0[kind], n_src[kind], len(items)] + [risk[k] for k in common.RISK_ORDER] + \
               [lvl[k] for k in common.EVIDENCE_ORDER]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.font, c.alignment, c.border = font(10, j == 1), CENTER, BORDER
        ws.cell(r, 1).fill = fill(LV0_FILL[kind])
        r += 1
    kinds = collections.Counter(m[9].split(" · ")[0].split("(")[0] for m in maps)
    line("재구성 유형", " · ".join(f"{k} {v}" for k, v in kinds.most_common()) + " (원본 Lv3 단위, '부록-재구성 매핑' 시트)")
    ot_items = [x for x in rows if x["kind"] == "ot"]
    major = [x["id"] for x in ot_items if x["ev"]["incident_ids"] and set(x["ev"]["incident_ids"]) <= common.OT_MAJOR_INCIDENTS]
    line("OT 근거 편중", f"OT 요약 {len(ot_items)}개 중 {sum(x['ev']['evidence'] == '실제 사고 확인' for x in ot_items)}개가 '실제 사고 "
                      "확인' — ATT&CK 공식 절차가 촘촘한 대형 사고 5건(Stuxnet·2015/2016 우크라이나·Triton·2025 폴란드)이 원본 "
                      f"36~46행에 매핑된 영향. 실제 사고가 이 5건에서만 확인되는 항목 {len(major)}개"
                      + (f"({', '.join(major)})" if major else "") + ". OT 항목 간 비교는 실제 사고 수·우선순위로")
    sc_items = [x for x in rows if x["kind"] == "supplychain"]
    line("공급망 위험도 포화", f"공급망 요약 {len(sc_items)}개 중 {sum(x['ev']['risk'] == '매우 높음' for x in sc_items)}개가 '매우 "
                         "높음' — 원본 심각도 '상'이 60개 중 42개(신원·물리처럼 '상'을 제한하는 기준을 쓰지 않음). 공급망 항목 간 "
                         "비교는 실제 사고 수·최근 사고·우선순위로")
    line("Lv0 간 중복", "신원↔클라우드(계정·자격증명), 공급망↔AI·클라우드(공급망 침해), 물리↔신원(내부자) 위협은 관점별로 각 Lv0에 두고 "
                      "'연계 위협'으로 연결 — Lv0별 위협 수의 합은 고유 위협 수가 아님")
    r += 1
    for kind in LV0:
        if kind != "ai":  # 인쇄 시 표 제목과 본문이 쪽 경계에서 갈라지지 않도록
            ws.row_breaks.append(Break(id=r - 1))
        section(f"■ 우선 위협 Top 10 — {LV0[kind]} (위험도 → 실제 사고 수 → 최근 사고 순)")
        table_header([(1, 1, "우선순위"), (2, 2, "위협 ID"), (3, 6, "세부 위협(Lv3)"), (7, 7, "위험도"),
                      (8, 8, "실제 사고 수"), (9, 9, "최근 사고(2025~)"), (10, 11, "근거 수준")])
        ranked = sorted((x for x in rows if x["kind"] == kind), key=lambda x: x["rank"])
        for x in ranked[:10]:
            vals = {1: x["rank"], 2: x["id"], 3: x["name"], 7: x["ev"]["risk"], 8: x["ev"]["incidents"],
                    9: x["ev"]["recent"], 10: x["ev"]["evidence"]}
            for col in range(1, 12):
                ws.cell(r, col).border = BORDER
            for col, v in vals.items():
                c = ws.cell(r, col, v)
                c.font, c.alignment = font(10), CENTER
            ws.cell(r, 2).font = font(10, True)
            ws.cell(r, 3).alignment = Alignment(vertical="center")
            ws.cell(r, 1).fill = fill(LV0_FILL[kind])
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)
            paint_risk(ws.cell(r, 7), x["ev"]["risk"])
            paint_level(ws.cell(r, 10), x["ev"]["evidence"])
            r += 1
        if kind in common.TAXO:  # 결과 유형인 영향(IM) 도메인을 뺀 단계 위협 기준 상위 — 원본 매트릭스의 우선 위협 표기와 같은 방식
            stage = [x for x in ranked if x["domain"] != "IM"][:5]
            line("단계 위협 기준", "영향(IM) 도메인 제외 상위 5: " + " → ".join(f"{x['id']} {x['name']}" for x in stage))
        r += 1
    section("■ 위험도 산정 (발생가능성 × 심각도, 원본 여섯 매트릭스와 같은 표)")
    table_header([(1, 1, "발생가능성 \\ 심각도"), (2, 2, "상"), (3, 3, "중"), (4, 4, "하")])
    for lik in ["상", "중", "하"]:
        c = ws.cell(r, 1, lik)
        c.font, c.fill, c.alignment, c.border = font(10, True), fill("F2F3F4"), CENTER, BORDER
        for j, sev in enumerate(["상", "중", "하"], 2):
            c = ws.cell(r, j, common.risk_of(lik, sev))
            paint_risk(c, common.risk_of(lik, sev))
            c.border = BORDER
        r += 1
    line("", "요약 위험도 = 구성 원본별 위험도의 최댓값(발생가능성·심각도를 서로 다른 원본에서 조합하지 않음) | 구성 원본의 실제 사고 "
             "합집합이 2건 이상이면 사고가 확인된 원본만 발생가능성 '상'으로 보정 | 전술마다 값이 다른 반복 기법은 요약 항목의 주 전술 "
             "행 값 | 근거 수준 = 구성 원본 중 가장 강한 수준")
    r += 1
    section("■ 시트 안내")
    for k, v in [("보고서용 간략 매트릭스", "본문 삽입용 — 위협·핵심 요약·위험도·우선순위·대표 사례 1건·핵심 대응 1줄"),
                 ("통합 요약 매트릭스", "상세 본표(부록용) — 설명·시나리오·사례·위험 평가·대응 방안·교차 매핑 전체"),
                 ("매트릭스 뷰", "Lv0별로 도메인(열)마다 요약 위협을 위험도 색으로 배치한 한눈 보기"),
                 ("도메인 요약", "Lv1 도메인별 원본·요약 항목 수, 위험도 분포, 주요 위협"),
                 ("AI 위협 상세", f"AI 원본 Lv3 {len(ai_src['lv3'])}개 원문 그대로 + 통합 위협 ID 역참조·실제 사고 산정 근거"),
                 ("클라우드 위협 상세", f"클라우드 원본 {len(cl_src['rows'])}행 원문 그대로 + 통합 위협 ID 역참조"
                                   + (" (CSA 열은 이슈 ID만)" if release else "")),
                 ("OT 위협 상세", f"OT 원본 {len(ot_src['rows'])}행 원문 그대로 + 통합 위협 ID 역참조 (적용 프로파일 6열은 한 열로)"),
                 ("공급망 위협 상세", f"공급망 원본 세부위협 {len(sc_src['rows'])}개 원문 그대로 + 통합 위협 ID 역참조 (프로파일 ● 열은 한 열로)"),
                 ("신원 위협 상세", f"신원 원본 세부위협 {len(id_src['rows'])}개 원문 그대로 + 통합 위협 ID 역참조 (신원 유형 ● 열은 한 열로)"
                                 + (" (CSA CCM·CIS Controls 열은 ID만)" if release else "")),
                 ("물리·인적 위협 상세", f"물리·인적 원본 세부위협 {len(ph_src['rows'])}개 원문 그대로 + 통합 위협 ID 역참조 "
                                    "(위협 주체 ● 열은 한 열로)"),
                 ("부록-재구성 매핑", "원본 Lv3 → 요약 Lv3 대응, 처리 유형, 위험 대표 원본, 통합 근거"),
                 ("부록-작성·평가 기준", "재구성 원칙, 위험 재산정 규칙, 문안·사례 표기 기준, 용어")]:
        line(k, v)
    r += 1
    section("■ 출처·이용조건")
    csa = ("원문은 싣지 않고 이슈·통제 ID만 연계(이 배포본은 상세 시트 포함 ID만 표기)" if release else
           "요약표는 이슈 ID만 연계. 내부본의 클라우드·신원 상세 시트는 원본 이슈·통제명을 유지 — 고객 배포는 배포본(--release) 사용")
    for k, v in [("MITRE ATT&CK®", "Enterprise·ICS v19.2 © The MITRE Corporation — ATT&CK 이용약관에 따른 출처 표기"),
                 ("MITRE ATLAS™", "v2026.09 © The MITRE Corporation"),
                 ("MITRE EMB3D™", "©2026 The MITRE Corporation. This work is reproduced and distributed with the permission of "
                                  "The MITRE Corporation. — 위협 ID·명칭·성숙도만 연계(설명문 미수록). 아래 라이선스 적용"),
                 ("EMB3D 라이선스", "LICENSE (EMB3D Terms of Use, https://emb3d.mitre.org/subtabs/terms-of-use.html) The MITRE "
                                   "Corporation (MITRE) hereby grants you a non-exclusive, royalty-free license to use, copy, and create "
                                   "derivative works of EMB3D™ for internal business purposes or commercial use. Any copy you make for "
                                   "such purposes is authorized provided that you reproduce MITRE's copyright designation and this "
                                   "license in any such copy and provide MITRE notice of such use at EMB3D@mitre.org. For all other uses "
                                   "of EMB3D™, contact MITRE at techtransfer@mitre.org. All rights not expressly granted are hereby "
                                   "reserved."),
                 ("MITRE CTID ITKB", "Insider Threat TTP Knowledge Base — Apache-2.0 (물리·인적 상세 시트에 관측 기법 ID·명칭)"),
                 ("OWASP", "GenAI Security Project — LLM Top 10 2026 · Agentic Top 10 2026 · GenAI Data Security 2026 · "
                           "MCP Top 10, IoT Top 10(2018, 범주 ID만), Top 10 CI/CD Security Risks · Top 10 OSS Risks · Non-Human "
                           "Identities Top 10 · ASVS 5.0 · Automated Threats(ID·명칭만). 각 문서의 라이선스 조건(CC BY-SA 등) 확인"),
                 ("공급망 기준", "SLSA v1.2·OpenSSF S2C2F(Community Specification License 1.0) · OpenSSF Scorecard(Apache-2.0) · SAP "
                              "Risk Explorer(Apache-2.0) · CNCF 공급망 침해 목록(CC BY 4.0) · NIST SSDF(공공 영역) — ID·명칭만"),
                 ("Push Security", "Browser & Identity Attacks Matrix(SaaS attacks) — CC BY 4.0, 기법 ID·명칭만"),
                 ("AWS · Microsoft", "AWS Threat Technique Catalog · Azure Threat Research Matrix · Threat Matrix for "
                                     "Kubernetes — 각 원저작권자"),
                 ("CSA", "Top Threats to Cloud Computing 2026 · CCM v4.1 © Cloud Security Alliance — 개인·비상업적 용도, "
                         "재배포 금지. " + csa),
                 ("CIS Controls", "v8.1 © Center for Internet Security — CC BY-NC-ND 4.0, 세이프가드 ID·명칭만"
                                  + (" (이 배포본은 ID만)" if release else " (배포본은 ID만)")),
                 ("CISA", "ICS 권고(CSAF 2.0)·KEV 카탈로그·내부자 위협·무인기 지침 — 미국 정부 저작물(공공 영역)"),
                 ("표준·인증기준", "IEC 62443-3-3·4-2, ISO/IEC 27001:2022, ISO 22341은 유료 표준이라 요구사항 번호·제목만 연계(원문 미수록) / "
                                "NIST SP 800-53 Rev.5·800-63-4 — 공공 영역 / ISMS-P 인증기준 — 항목 번호·명칭"),
                 ("NPSA", "영국 NPSA 지침(물리·인적 참고 문헌) — 서지·URL만 수록하고 요지를 의역, 원문 미수록"),
                 ("사고 DB", "클라우드 보안사고 DB 680건 — Wiz, ramimac, SEC, GTI 등 공개 출처 종합 / OT 보안사고 DB "
                            f"{ot_src['n_incidents']}건 · 공급망 {sc_src['n_incidents']}건 · 신원 {id_src['n_incidents']}건 · 물리·인적 "
                            f"{ph_src['n_incidents']}건 — 정부 발표·수사·판결·보안업체·연구기관 보고서·언론 등 공개 출처 종합"),
                 ("·", "외부 배포 전 각 출처의 이용약관(특히 상업적 이용·변형물 조건)을 확인. EMB3D는 사용 시 EMB3D@mitre.org에 "
                       "통지하는 조건이 있음")]:
        line(k, v)
    r += 1
    section("■ 활용 시 유의사항")
    for v in ["위험도·실제 사고 수·우선순위는 같은 Lv0 안에서 비교 — 원본마다 사고 DB 규모와 심각도 기준이 달라(OT 물리적 결과, "
              "공급망 '상' 포화, 신원·물리 '상' 제한) Lv0 간 등급을 직접 비교하지 않음",
              "발생가능성은 확인된 실제 사고 기준 — 공개 취약점은 많으나 사고 보고가 적은 신규 공격면(예: 간접 프롬프트 인젝션·AI "
              "패키지 환각)은 과소평가될 수 있음. OT는 사고 공개가 적고 설계상 인증이 없는 기능의 악용은 CVE가 없어 "
              "'이론·시나리오'는 '공개 근거 없음'으로 읽음",
              "OT 우선순위는 심각도 '상' 전술(대응 기능 억제·공정 제어 훼손·영향)이 상위권에 몰림 — 진입·확산 차단 우선순위는 "
              "매트릭스 뷰·도메인 요약의 초기 침투·횡적 이동 열에서 확인. 공급망·신원·물리는 개요의 '단계 위협 기준'(영향 도메인 제외)을 함께 봄",
              "물리·인적은 공개 사고 기준이라 기소·판결로 드러나는 기술 유출·내부자 사건이 높게 나오고 은닉 감시·물리 침입은 "
              "과소평가될 수 있음. 방어 자료로 한정해 실행 방법은 싣지 않고, 드론은 탐지·식별·신고·차폐만 권고",
              "요약 위험도는 구성 원본에서 재산정한 값(규칙: 부록-작성·평가 기준). 원본 개별 값은 원본 상세 시트에 그대로 유지",
              "AI 실제 사고 수 = ATLAS 사례 ID(중복 제거) + OWASP 인용 사고 — OWASP 인용 사고는 ID가 없어 참조의 사례명(시점)으로 "
              "원본 간 중복을 걸러냄",
              "대표 사례는 구성 원본 참조에 수록된 사례만 사용(빌드 시 자동 대조) — 새로운 사실을 추가하지 않음. OT는 근거가 "
              "대형 사고에 몰려 같은 사고(Stuxnet·Triton 등)가 여러 항목에 나올 수 있으며, 각 항목은 그 사고의 다른 단계·기능을 보여 줌",
              "모든 값은 정적 값(수식 없음). 재생성: python3 integrated/scripts/build_integrated_matrix.py [--release]"]:
        line("·", v)
    widths(ws, [20, 14, 10, 10, 9, 9, 9, 13, 13, 13, 12])
    print_setup(ws, paper="A4")
    for row in merge_rows:  # 서술형 줄은 B~K 병합으로 폭 확보
        text = ws.cell(row, 2).value or ""
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=11)
        ws.row_dimensions[row].height = 15 * (1 + len(text) // 95)

# ---------------------------------------------------------------- 요약 목록 (Markdown, 저장소에서 바로 보기용)
def write_markdown(path, rows, src):
    out = [f"# 통합 보안위협 요약 목록 {VERSION} — AI·클라우드·OT·공급망·신원·물리·인적", "",
           "> 빌드 산출물(`build_integrated_matrix.py`가 생성) — 직접 고치지 말고 `integrated/data/` 문안을 고친 뒤 다시 빌드",
           "", "| 구분 | 위협 수 | 매우 높음 | 높음 | 보통 | 낮음 |", "|---|---:|---:|---:|---:|---:|"]
    for kind in common.KINDS:
        items = [x for x in rows if x["kind"] == kind]
        dist = collections.Counter(x["ev"]["risk"] for x in items)
        out.append(f"| {LV0[kind]} | {len(items)} | " + " | ".join(str(dist[k]) for k in common.RISK_ORDER) + " |")
    for kind in LV0:
        out += ["", f"## 우선 위협 Top 10 — {LV0[kind]}", "",
                "| 순위 | 위협 ID | 세부 위협(Lv3) | 핵심 요약 | 위험도 | 실제 사고 | 최근 사고(2025~) |",
                "|---:|---|---|---|---|---:|---:|"]
        for x in sorted((x for x in rows if x["kind"] == kind), key=lambda x: x["rank"])[:10]:
            out.append(f"| {x['rank']} | {x['id']} | {x['name']} | {x['summary']} | {x['ev']['risk']} | "
                       f"{x['ev']['incidents']} | {x['ev']['recent']} |")
    for kind in common.KINDS:
        out += ["", f"## {LV0[kind]}"]
        for code, label in common.domains_of(kind, src[kind]).items():
            items = [x for x in rows if x["kind"] == kind and x["domain"] == code]
            out += ["", f"### {label} ({len(items)})", "",
                    "| 위협 ID | 세부 위협(Lv3) | 핵심 요약 | 위험도 | 우선순위 | 실제 사고 | 근거 수준 | 원본 Lv3 |",
                    "|---|---|---|---|---:|---:|---|---|"]
            for x in items:
                origin = ", ".join(x["members"])
                out.append(f"| {x['id']} | {x['name']} | {x['summary']} | {x['ev']['risk']} | {x['rank']} | "
                           f"{x['ev']['incidents']} | {x['ev']['evidence']} | {origin} |")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")


# ---------------------------------------------------------------- 산출물 재검증
def verify(path, rows, n_detail, n_maps, release):
    """저장한 워크북을 다시 열어 시트·행 수·역참조·값 유효성을 확인."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    expect = ["개요", "보고서용 간략 매트릭스", "통합 요약 매트릭스", "매트릭스 뷰", "도메인 요약"] + \
        [DETAIL_SHEET[k] for k in common.KINDS] + ["부록-재구성 매핑", "부록-작성·평가 기준"]
    assert wb.sheetnames == expect, wb.sheetnames
    ids = {r["id"] for r in rows}

    def body(name, ncol):
        ws = wb[name]
        return [[ws.cell(i, j).value for j in range(1, ncol + 1)] for i in range(5, ws.max_row + 1)
                if ws.cell(i, 1).value is not None]

    summ = body("통합 요약 매트릭스", len(SUMMARY_COLS))
    assert len(summ) == len(rows) and {v[2] for v in summ} == ids
    assert all(v[10] in RISK_COLOR and v[11] in LEVEL_COLOR for v in summ)
    assert all(v[4] and len(v[4]) <= common.LIMITS["summary"] for v in summ)
    for lv0 in LV0.values():  # 우선순위는 Lv0마다 1..N 중복 없이
        ranks = sorted(v[14] for v in summ if v[0] == lv0)
        assert ranks == list(range(1, len(ranks) + 1)), lv0
    brief = body("보고서용 간략 매트릭스", len(BRIEF_COLS))
    assert [v[2] for v in brief] == [v[2] for v in summ] and all(v[5] in RISK_COLOR and v[8] for v in brief)
    for name, n in [(DETAIL_SHEET[k], n_detail[k]) for k in common.KINDS]:
        det = body(name, 14)
        assert len(det) == n and all(v[0] in ids for v in det), name
        assert all(v[11] in RISK_COLOR for v in det), name
    maps = body("부록-재구성 매핑", 12)
    assert len(maps) == n_maps and all(v[6] in ids for v in maps)
    assert sorted({v[6] for v in maps if v[10] == "●"}) == sorted(ids)  # 요약 항목마다 위험 대표 원본 표시
    csa_col = 15 + CLOUD_SRC_TAIL.index("CSA Top Threats 2026")
    csa = [v for v in (wb["클라우드 위협 상세"].cell(i, csa_col + 1).value for i in range(5, n_detail["cloud"] + 5)) if v]
    assert csa and all(re.fullmatch(r"SI-\d+( · SI-\d+)*", v) for v in csa) == release, "CSA 열 표기"
    ws = wb[DETAIL_SHEET["identity"]]  # 신원 상세의 CSA CCM 열: 배포본은 통제 ID만
    head = [ws.cell(4, j).value for j in range(1, ws.max_column + 1)]
    ccm = [v for v in (ws.cell(i, head.index("CSA CCM v4.1") + 1).value for i in range(5, n_detail["identity"] + 5)) if v]
    assert ccm and all(re.fullmatch(r"[A-Z]{3}-\d{2}( · [A-Z]{3}-\d{2})*", v) for v in ccm) == release, "CCM 열 표기"
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    raise AssertionError(f"수식 셀: {ws.title}!{c.coordinate}")
    return len(summ)


# ---------------------------------------------------------------- 실행
def main():
    ap = argparse.ArgumentParser(description="통합 보안위협 매트릭스 빌드")
    ap.add_argument("--release", action="store_true",
                    help="고객 배포본: 클라우드 상세의 CSA 열, 신원 상세의 CSA CCM·CIS Controls 열을 ID만 남김")
    args = ap.parse_args()
    src, summ = common.load_sources(), common.load_summary()
    errs, warns = common.check_all(src, summ)
    for w in warns:
        print("  [경고]", w)
    if errs:
        for e in errs:
            print("  [오류]", e)
        sys.exit(f"검증 오류 {len(errs)}건 — 빌드 중단")
    rows, skipped = build_rows(src, summ)
    maps = mapping_rows(src, summ, rows)
    n_detail = {k: len(src[k]["lv3"]) if k == "ai" else len(src[k]["rows"]) for k in common.KINDS}
    for kind in common.KINDS:
        assert len([m for m in maps if m[0] == LV0[kind]]) == n_detail[kind], f"재구성 매핑 행 수: {kind}"

    today = datetime.date.today().isoformat()
    wb = Workbook()
    wb.remove(wb.active)
    sheet_brief(wb, rows, today)
    sheet_summary(wb, rows, today)
    sheet_matrix_view(wb, rows, src)
    sheet_domain_summary(wb, rows, src)
    sheet_ai_detail(wb, src["ai"], summ["ai"])
    sheet_cloud_detail(wb, src["cloud"], summ["cloud"], args.release)
    sheet_ot_detail(wb, src["ot"], summ["ot"])
    for kind in common.TAXO:
        sheet_taxo_detail(wb, kind, src[kind], summ[kind], args.release)
    sheet_mapping(wb, maps)
    sheet_criteria(wb)
    sheet_overview(wb, rows, src, maps, today, args.release)
    tabs = ["1F3864", "117A65", "C0392B", "E67E22", "2E5496"] + [LV0_BAND[k] for k in common.KINDS] + ["7F8C8D", "7F8C8D"]
    for ws, color in zip(wb.worksheets, tabs):
        ws.sheet_properties.tabColor = color
    os.makedirs(OUT_DIR, exist_ok=True)
    path = XLSX_RELEASE if args.release else XLSX
    wb.save(path)
    if not args.release:  # CSV·목록은 두 판이 같으므로 내부본 빌드 때만 갱신
        with open(CSV, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(SUMMARY_COLS)
            for r in rows:
                w.writerow(summary_values(r))
        write_markdown(MD, rows, src)
    n = verify(path, rows, n_detail, len(maps), args.release)
    print(f"저장: {os.path.relpath(path, common.ROOT)} (재검증 통과: 요약 {n}행)")
    for kind in LV0:
        items = [x for x in rows if x["kind"] == kind]
        dist = collections.Counter(x["ev"]["risk"] for x in items)
        lvl = collections.Counter(x["ev"]["evidence"] for x in items)
        print(f"  {LV0[kind]} {len(items)}개 | " + " · ".join(f"{k} {dist[k]}" for k in common.RISK_ORDER)
              + " | " + " · ".join(f"{k} {lvl[k]}" for k in common.EVIDENCE_ORDER))
    print(f"  재구성 매핑 {len(maps)}행 | 연계 위협이 있는 항목 {sum(bool(x['links_text']) for x in rows)}개")
    if skipped:  # 요약 ID로 바로 옮길 수 없는 원본 연계 참조 → 문안 links로 대체했는지 확인
        manual = {e["id"]: list(e.get("links") or []) for k in common.KINDS for e in summ[k]}
        prefix = {"cloud": "CL-", "ai": "AI-", "AI": "AI-", "CL": "CL-", "OT": "OT-", "OTC": "OT-",
                  "SCT": "SC-", "IDT": "ID-", "PHT": "PH-"}
        covered, missing = [], []
        for k, (sid, kind) in skipped.items():
            subs = [t for t in manual.get(sid, []) if t.startswith(prefix[kind])]
            (covered if subs else missing).append(f"{k}({sid}→{'·'.join(subs)})" if subs else f"{k}({sid})")
        if covered:
            print(f"  원본 연계 중 요약 ID로 바로 옮길 수 없는 참조 {len(covered)}건 — 문안 links로 대체: " + ", ".join(covered))
        for x in missing:
            print(f"  [경고] 원본 연계 누락: {x} — 해당 항목 문안에 대상 Lv0의 links 지정 필요")


if __name__ == "__main__":
    main()
