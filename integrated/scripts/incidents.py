"""통합 사고사례 목록 — 요약 위협이 참조하는 사고·사례를 여섯 원본의 사고 DB에서 모아 한 목록으로 묶는다.

수록 범위
  1) 위험 산정 사고: 요약 항목의 구성 원본(members)이 실제 사고로 센 사고
     (클라우드·OT·공급망·신원·물리 사고 DB ID, AI는 ATLAS Incident 사례 ID와 참조에 사례명이 실린 OWASP 인용 사고)
  2) 참고 사고: Lv0 간 통합으로 흡수한 다른 Lv0 원본(xmembers)이 센 사고 — 위험 산정에는 넣지 않음(결정 V3)
  3) 대표 사례 인용: 요약 항목 '대표 사례'에 쓴 사례. 사고 DB에 있으면 그 기록, 없으면([공개 취약점]·[ATT&CK 사례]·
     [EMB3D] 등) 대표 사례 줄과 원본 참조 줄로 '인용 사례' 기록을 만듦
같은 사고가 여러 DB에 실린 경우(교차 DB 동일 사고)는 한 행으로 묶는다
  - 신원·물리 사고 DB의 '다른 DB 사고 ID' 열(자동)과 data/incident_aliases.yaml(검토 확정)
  - 묶음 안 첫 ID(별칭 파일 순서, 없으면 DB 우선순위)가 사례명·시점·요약 등을 가져올 대표 기록

사례 ID 표기: 클라우드 DB `CL/INC-0001` · OT DB `OT/INC-001`(두 DB가 INC- 접두사를 같이 씀) · SCI-·IDI-·PHI- · ATLAS
`AML.CS0000` · OWASP 인용 사고 `OWASP:사례명(시점)` · 사고 DB 밖 인용 사례 `CITE:[유형] 사례명(시점)`
"""
import collections
import os
import re

import openpyxl
import yaml

import common

ALIAS_YAML = os.path.join(common.DATA_DIR, "incident_aliases.yaml")
ATLAS_RE = re.compile(r"\bAML\.CS\d{4}\b")
DB_LABEL = {"cloud": "클라우드 DB", "ot": "OT DB", "supplychain": "공급망 DB", "identity": "신원 DB",
            "physical": "물리·인적 DB", "ai": "ATLAS", "owasp": "OWASP 인용", "cite": "대표 사례 인용"}
# 묶음에 별칭 파일 순서가 없을 때 대표 기록을 고르는 순서 — 국문 사례명·지역·행위자 열이 있는 DB 우선
DB_PRIORITY = ["supplychain", "identity", "physical", "ot", "ai", "owasp", "cloud", "cite"]
# 사례 구분 — 실제 사고로 센 기록은 '실제 사고', 나머지는 원본 집계 값 또는 대표 사례 유형
CATEGORY_ORDER = ["실제 사고", "실제 사고(사고 DB 밖 인용)", "정부 경보", "위협인텔", "실증·연구", "시연", "공개 취약점",
                  "ATT&CK 사례", "EMB3D", "집계 제외"]
TAG_CATEGORY = {"실제 사고": "실제 사고(사고 DB 밖 인용)", "실증": "실증·연구", "공개 취약점": "공개 취약점",
                "ATT&CK 사례": "ATT&CK 사례", "위협인텔": "위협인텔", "시연": "시연", "정부 경보": "정부 경보",
                "EMB3D": "EMB3D"}
SOURCE_NOTE_RE = re.compile(r"\(([^()]*(?:\([^()]*\)[^()]*)*)\)\s*$")


def qid(kind, raw):
    """원본 사례 ID → 목록 안에서 유일한 ID (클라우드·OT DB만 접두사)."""
    return {"cloud": f"CL/{raw}", "ot": f"OT/{raw}"}.get(kind, raw)


def display_id(rec):
    """원본 사례 ID 표시: 'INC-0001(클라우드 DB)' · 'SCI-001' · 'OWASP 인용: 사례명(시점)'."""
    if rec["db"] in ("cloud", "ot"):
        return f"{rec['raw']}({DB_LABEL[rec['db']]})"
    if rec["db"] == "owasp":
        return f"OWASP 인용: {rec['raw']}"
    if rec["db"] == "cite":
        return "사고 DB 외(대표 사례 인용)"
    return rec["raw"]


def _rec(db, raw, name, date, category, raw_type, counted, region=None, actor=None, impact=None, summary=None, source=None):
    return dict(qid=qid(db, raw) if db in ("cloud", "ot") else raw, db=db, raw=raw, name=str(name or "").strip(),
                date=str(date or "").strip(), category=category, raw_type=raw_type, counted=counted, region=region,
                actor=actor, impact=impact, summary=summary, source=source)


# ---------------------------------------------------------------- 사고 DB 로더
def load_records(src):
    """여섯 원본의 사고·사례 기록 {qid: 기록} — 클라우드·OT는 워크북 '역매핑_사고사례' 시트를 다시 읽음."""
    recs = {}
    wb = openpyxl.load_workbook(common.CLOUD_XLSX, read_only=True, data_only=True)
    rows = list(wb["역매핑_사고사례"].iter_rows(values_only=True))
    for x in rows[1:]:
        d = dict(zip(rows[0], x))
        if not d["사건ID"]:
            continue
        counted = d["사례유형"] in common.CLOUD_REAL_KINDS and d["집계 포함"] == "Y"
        cat = "실제 사고" if counted else ("실증·연구" if d["사례유형"] == "연구·노출" else "집계 제외")
        actor = re.search(r"행위자: ([^)]+)\)", d["사건 요약"] or "")
        r = _rec("cloud", d["사건ID"], d["대표 제목"], str(d["기준일"] or "")[:10], cat,
                 f"{d['사례유형']} · 집계 {d['집계 포함']}", counted, None,
                 actor.group(1) if actor and actor.group(1) != "미상" else None, d["영향유형"], d["사건 요약"], d["출처"])
        recs[r["qid"]] = r
    wb = openpyxl.load_workbook(common.OT_XLSX, read_only=True, data_only=True)
    rows = list(wb["역매핑_사고사례"].iter_rows(values_only=True))
    for x in rows[1:]:
        d = dict(zip(rows[0], x))
        if not d["사건ID"]:
            continue
        counted = d["집계"] == common.OT_REAL_STATUS
        cat = "실제 사고" if counted else {"제외(검증)": "집계 제외"}.get(d["집계"], d["집계"])
        r = _rec("ot", d["사건ID"], d["사건명"], d["기준일"], cat, f"{d['사례유형']} · {d['집계']}", counted, d["국가"],
                 d["행위자"], d["물리적 영향"], d["사건 요약"], d["출처"])
        recs[r["qid"]] = r
    for kind in common.TAXO:
        for d in src[kind]["incidents"]:
            counted = d["집계 상태"] == common.TAXO_REAL_STATUS
            cat = "실제 사고" if counted else {"제외(검증)": "집계 제외"}.get(d["집계 상태"], d["집계 상태"])
            r = _rec(kind, d["사례 ID"], d["사례명"], d["시점"], cat, f"{d['사례유형']} · {d['집계 상태']}", counted,
                     d.get("지역"), d.get("행위자"), d.get("영향"), d.get("요약"), d.get("출처"))
            recs[r["qid"]] = r
    for c in src["ai"]["cases"]:
        counted = c["type"] == "Incident"
        r = _rec("ai", c["id"], c["name"], c.get("date"), "실제 사고" if counted else "실증·연구",
                 f"ATLAS {c['type']}", counted, None, c.get("actor"), c.get("target"), c.get("summary"), c.get("url"))
        recs[r["qid"]] = r
    for d in src["ai"]["lv3"].values():  # OWASP 인용 사고(ID 없음) — 참조 줄의 사례명(시점)·요지·출처
        for label, date in d["owasp_named"]:
            key = f"OWASP:{label}"
            if key in recs:
                continue
            line = next(l for l in (d["참조"] or "").split("\n") if label.split("(")[0] in l and "OWASP" in l)
            n = common.NAMED_RE.match(common.CASE_RE.match(line.strip()).group("body"))
            gist, note = _split_note(n.group("gist") if n else line)
            r = _rec("owasp", label, label.rsplit("(", 1)[0] if date else label, date, "실제 사고", "OWASP 인용 사고", True,
                     summary=gist, source=note)
            r["qid"] = key
            recs[key] = r
    return recs


def _split_note(text):
    """'요지 (출처 메모)' → (요지, 출처 메모)."""
    m = SOURCE_NOTE_RE.search(text or "")
    if not m:
        return (text or "").strip(), None
    return text[:m.start()].strip(), m.group(1).strip()


def load_aliases(src):
    """동일 사고 묶음 [[qid…]] — 별칭 파일(순서 = 대표 우선) + 신원·물리 DB '다른 DB 사고 ID' 열."""
    groups = [list(g["ids"]) for g in yaml.safe_load(open(ALIAS_YAML, encoding="utf-8")) or []]
    for kind in ("identity", "physical"):
        for d in src[kind]["incidents"]:
            for other in re.findall(r"\b(?:SCI|IDI|PHI)-\d+\b", str(d.get("다른 DB 사고 ID") or "")):
                groups.append([d["사례 ID"], other])
    return groups


# ---------------------------------------------------------------- 요약 항목 → 사고 참조
def _member_incidents(kind, sk, key, domain):
    """구성 원본 하나가 실제 사고로 센 사고 qid 목록."""
    if kind == "ai":
        d = sk["lv3"][key]
        return list(d["atlas_incidents"]) + [f"OWASP:{label}" for label, _ in d["owasp_named"]]
    return [qid(kind, i) for i in common.member_row(sk, key, domain)["real_incidents"]] if kind in ("cloud", "ot") \
        else list(common.member_row(sk, key, domain)["real_incidents"])


def _case_sources(entry, src):
    """대표 사례 대조용 (Lv0, 참조 문자열) 목록 — 구성 원본 + Lv0 간 통합 원본."""
    return [(k, common.source_refs(k, src[k], m)) for k, m, _ in common.entry_sources(entry)]


def case_refs(entry, src, recs):
    """대표 사례 줄 → [(qid 목록, 인용 기록 또는 None)] — 사례명(시점)이 실린 원본 참조 줄의 사례 ID를 찾음."""
    out = []
    sources = [(k, common.normalize_terms(r)) for k, r in _case_sources(entry, src)]
    for line in common._lines(entry["cases"]):
        m = common.CASE_RE.match(line.strip())
        if not m or m.group("tag") == "시나리오":
            continue
        tag, body = m.group("tag"), m.group("body")
        n = common.NAMED_RE.match(body)
        if not n:
            continue
        name = n.group("name").strip()
        label = f"{name}({n.group('date')})" if n.group("date") else name
        key = label if n.group("date") else f"] {name}:"
        hit_kind, hit = None, ""
        for k, refs in sources:
            hit = next((l for l in refs.split("\n") if key in l), "")
            if hit:
                hit_kind = k
                break
        ids = []
        if hit:
            ids = [qid(hit_kind, i) for i in dict.fromkeys(common.INCIDENT_RE.findall(hit))] + \
                  list(dict.fromkeys(ATLAS_RE.findall(hit)))
            if hit_kind == "ai" and not ids and "OWASP" in hit and f"OWASP:{label}" in recs:
                ids = [f"OWASP:{label}"]
        ids = [i for i in ids if i in recs]
        if ids:
            out.append((ids, None))
            continue
        _, note = _split_note(hit.strip()) if hit else (None, None)
        cite = _rec("cite", f"[{tag}] {label}", name, n.group("date"), TAG_CATEGORY.get(tag, tag), f"대표 사례 [{tag}]",
                    False, summary=n.group("gist").strip(), source=note)
        cite["qid"] = f"CITE:[{tag}] {label}"
        out.append(([cite["qid"]], cite))
    return out


# ---------------------------------------------------------------- 목록 구성
def build_catalog(entries, src):
    """entries: [{id, kind, domain, members, absorbs, xmembers, cases}] (common.load_summary 항목) → 목록 행(최근 순).
    행: id · 사례명 · 시점 · 사례 구분 · 원본 사례 ID(표시) · 참조 Lv0 · 위험 산정/참고/대표 사례 위협 ID · 최근 여부 …"""
    recs = load_records(src)
    counted, ref, cited = (collections.defaultdict(set) for _ in range(3))
    for e in entries:
        sk = src[e["kind"]]
        for m in e["members"]:
            for q in _member_incidents(e["kind"], sk, m, e["domain"]):
                counted[q].add(e["id"])
        for a in e.get("absorbs") or []:  # 흡수 원본은 흡수 전 위협의 도메인 기준(참고)
            for m in a["members"]:
                for q in _member_incidents(a["kind"], src[a["kind"]], m, a["domain"]):
                    ref[q].add(e["id"])
        for ids, cite in case_refs(e, src, recs):
            if cite and cite["qid"] not in recs:
                recs[cite["qid"]] = cite
            for q in ids:
                cited[q].add(e["id"])
    used = set(counted) | set(ref) | set(cited)
    missing = sorted(q for q in used if q not in recs)
    if missing:
        raise ValueError(f"사고 기록 없음: {missing[:10]}")

    # 동일 사고 묶음 — union-find, 대표는 묶음 목록에서 먼저 나온 ID
    parent = {}

    def find(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    order = {}
    for g in load_aliases(src):
        unknown = [q for q in g if q not in recs and not q.startswith("CITE:")]
        if unknown:
            raise ValueError(f"별칭 파일의 없는 사례 ID: {unknown}")
        g = [q for q in g if q in recs]  # 문안에서 인용하지 않게 된 CITE: 키는 건너뜀
        for q in g:
            order.setdefault(q, len(order))
        for q in g[1:]:
            a, b = find(g[0]), find(q)
            if a != b:
                parent[b] = a
    groups = collections.defaultdict(list)
    for q in used:
        groups[find(q)].append(q)
    # 묶음의 비참조 구성원도 원본 ID로 함께 표시(예: 같은 사고의 다른 DB 기록)
    for q in list(parent):
        root = find(q)
        if root in groups and q not in groups[root]:
            groups[root].append(q)

    rows = []
    for root, qs in groups.items():
        qs = sorted(qs, key=lambda q: (order.get(q, 10 ** 6), DB_PRIORITY.index(recs[q]["db"]), q))
        rep = recs[qs[0]]
        cats = [recs[q]["category"] for q in qs]
        category = min(cats, key=lambda c: CATEGORY_ORDER.index(c) if c in CATEGORY_ORDER else len(CATEGORY_ORDER))
        t_counted = set().union(*[counted[q] for q in qs])
        t_ref = set().union(*[ref[q] for q in qs]) - t_counted
        t_cited = set().union(*[cited[q] for q in qs])
        dbs = [recs[q]["db"] for q in qs]
        rows.append(dict(
            name=rep["name"], date=rep["date"] or next((recs[q]["date"] for q in qs if recs[q]["date"]), ""),
            category=category, region=_first(recs, qs, "region"), actor=_first(recs, qs, "actor"),
            impact=_first(recs, qs, "impact"), summary=rep["summary"] or _first(recs, qs, "summary"),
            source="\n".join(dict.fromkeys(str(recs[q]["source"]) for q in qs if recs[q]["source"])),
            origin="\n".join(f"{display_id(recs[q])} — {recs[q]['raw_type']}" for q in qs),
            qids=qs, dbs=dbs, counted=t_counted, ref=t_ref, cited=t_cited, rep=rep["qid"],
            multi_db=len({d for d in dbs if d != "cite"}) > 1))
    rows.sort(key=lambda r: (_date_key(r["date"]), r["name"]), reverse=True)
    for i, r in enumerate(rows, 1):
        r["uid"] = f"UC-{i:04d}"
        r["recent"] = r["category"].startswith("실제 사고") and _date_key(r["date"])[0] >= int(common.RECENT_FROM)
    return rows


def _first(recs, qs, field):
    return next((recs[q][field] for q in qs if recs[q][field]), None)


def _date_key(date):
    """시점 정렬 키: '2025-02~04' · '2023-09~' · '2020' · '2024-05-06' → (연, 월, 일)."""
    m = re.match(r"(\d{4})(?:-(\d{2}))?(?:-(\d{2}))?", str(date or ""))
    if not m:
        return (0, 0, 0)
    return tuple(int(x) if x else 0 for x in m.groups())
