"""통합 보안위협 매트릭스 — 원본 로더·근거 재산정·문안 검증 공통 모듈.

원본
  - AI: sources/ai_v3.2 (통합 AI 보안위협 매트릭스 v3.2, Lv3 124개)
  - 클라우드: output/통합_클라우드보안위협_매트릭스_v5.xlsx (Lv3 154행, 고유 기법 126개)
  - OT: ot/output/통합_OT보안위협_매트릭스_v5.xlsx (Lv3 141행, 고유 기법·항목 118개)
  - 공급망: supplychain/output/통합_공급망보안위협_매트릭스_v2.xlsx (세부위협 60개)
  - 신원: identity/output/통합_신원보안위협_매트릭스_v2.xlsx (세부위협 65개)
  - 물리·인적: physical/output/통합_물리인적보안위협_매트릭스_v2.xlsx (세부위협 45개)
요약(재구성 Lv3) 문안: integrated/data/<Lv0>/<도메인·전술>.yaml (ai · cloud · ot · supplychain · identity · physical)
"""
import collections
import glob
import os
import re

import openpyxl
import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
AI_DIR = os.path.join(ROOT, "sources", "ai_v3.2")
AI_XLSX = os.path.join(AI_DIR, "output", "통합_AI보안위협_매트릭스_v3.2_LITE.xlsx")
CLOUD_XLSX = os.path.join(ROOT, "output", "통합_클라우드보안위협_매트릭스_v5.xlsx")
OT_XLSX = os.path.join(ROOT, "ot", "output", "통합_OT보안위협_매트릭스_v5.xlsx")
SC_XLSX = os.path.join(ROOT, "supplychain", "output", "통합_공급망보안위협_매트릭스_v2.xlsx")
ID_XLSX = os.path.join(ROOT, "identity", "output", "통합_신원보안위협_매트릭스_v2.xlsx")
PH_XLSX = os.path.join(ROOT, "physical", "output", "통합_물리인적보안위협_매트릭스_v2.xlsx")
DATA_DIR = os.path.join(ROOT, "integrated", "data")

LEVEL = {"상": 3, "중": 2, "하": 1}
LEVEL_INV = {3: "상", 2: "중", 1: "하"}
RISK = {("상", "상"): "매우 높음", ("상", "중"): "높음", ("중", "상"): "높음",
        ("중", "중"): "보통", ("상", "하"): "보통", ("하", "상"): "보통"}
RISK_ORDER = ["매우 높음", "높음", "보통", "낮음"]
EVIDENCE_ORDER = ["실제 사고 확인", "실사용 기법 포함", "실증·공개 취약점", "이론·시나리오"]
CLOUD_REAL_KINDS = {"사고", "캠페인", "사례연구(익명)", "규제공시(8-K)", "위협인텔 보고서"}
OT_REAL_STATUS = "실제 사고"  # OT '역매핑_사고사례' 시트 '집계' 열에서 실제 사고로 세는 값
# OT 근거 편중 표기용 — ATT&CK 공식 절차가 촘촘해 원본 36~46행에 매핑된 대형 사고 5건(OT 방법론 §12)
OT_MAJOR_INCIDENTS = {"INC-002", "INC-003", "INC-004", "INC-008", "INC-009"}

AI_DOMAINS = ["D01", "D02", "D03", "D04", "D05", "D06", "D07", "D08", "D09", "D10"]
TACTICS = ["RD", "IA", "EX", "PE", "PV", "ST", "DI", "CA", "DS", "LM", "CO", "C2", "EF", "IM"]
OT_TACTICS = ["RD", "IA", "EX", "PE", "PV", "EV", "DS", "LM", "CO", "C2", "IR", "IP", "IM"]
SC_DOMAINS = ["DE", "SR", "DP", "BD", "SC", "PB", "CS", "IM"]
ID_DOMAINS = ["AU", "SE", "CR", "TS", "DI", "PE", "NH", "LC", "CI", "IM"]
PH_DOMAINS = ["RC", "PA", "SV", "HU", "IN", "EX", "IM"]
KINDS = ["ai", "cloud", "ot", "supplychain", "identity", "physical"]  # Lv0 순서
LABEL = {"ai": "AI", "cloud": "클라우드", "ot": "OT", "supplychain": "공급망", "identity": "신원", "physical": "물리·인적"}
DOMAINS = {"ai": AI_DOMAINS, "cloud": TACTICS, "ot": OT_TACTICS,
           "supplychain": SC_DOMAINS, "identity": ID_DOMAINS, "physical": PH_DOMAINS}
# 자체 분류 체계 원본(공급망·신원·물리) — 같은 빌드 틀(분류 체계 → 사고 DB → 관점 문구)로 만든 워크북이라 로더 하나로 읽음
#   id_col: 원본 세부위협 ID 열 · inc: 사고 DB ID 접두사 · prefix: 요약 ID 접두사 · stage: 공격 단계 열
TAXO = {
    "supplychain": dict(xlsx=SC_XLSX, id_col="SCT-ID", inc="SCI", prefix="SC", stage="공격 단계(ATT&CK 전술)",
                        name="통합 소프트웨어 공급망 보안위협 매트릭스 v2"),
    "identity": dict(xlsx=ID_XLSX, id_col="IDT-ID", inc="IDI", prefix="ID", stage="공격 단계(ATT&CK 전술)",
                     name="통합 신원·계정 보안위협 매트릭스 v2"),
    "physical": dict(xlsx=PH_XLSX, id_col="PHT-ID", inc="PHI", prefix="PH", stage="공격 단계",
                     name="물리·인적 보안위협 매트릭스 v2"),
}
TAXO_REAL_STATUS = "실제 사고"  # 세 원본 '역매핑_사고사례' 시트 '집계 상태' 열에서 실제 사고로 세는 값
INCIDENT_RE = re.compile(r"\b(?:INC|SCI|IDI|PHI)-\d+\b")  # 원본 참조 줄의 사고 ID (OT · 공급망 · 신원 · 물리)
CASE_TAGS = ["실제 사고", "공개 취약점", "실증", "ATT&CK 사례", "시나리오"]
# Lv0별 허용 사례 유형 — OT는 [위협인텔](배포 전에 발견된 공격 도구)·[EMB3D](EMB3D 장치 위협 인용)를 더 씀
# 신원은 [위협인텔]·[시연](SAT 시연), 물리는 [정부 경보]·[위협인텔]을 더 씀 — 모두 실제 사고로 세지 않음(원본 '집계 상태' 기준)
CASE_TAGS_BY_KIND = {"ai": CASE_TAGS, "cloud": CASE_TAGS, "ot": CASE_TAGS + ["위협인텔", "EMB3D"],
                     "supplychain": CASE_TAGS, "identity": CASE_TAGS + ["위협인텔", "시연"],
                     "physical": CASE_TAGS + ["정부 경보", "위협인텔"]}
# 같은 Lv0 안 대표 사례 반복 사용 경고 — 대형 사고 5건에 근거가 몰린 OT와 사고 DB 기반 공급망·신원·물리에 적용
# (AI·클라우드는 v1에서 수동 검토). 사례는 원본 참조 줄의 사고 ID(INC-·SCI-·IDI-·PHI-…)로 묶어,
# 같은 사고의 다른 표기('2016 우크라이나'·'2016 우크라이나 송전 변전소')도 반복으로 셈
CASE_REUSE_CHECK = {"ot", "supplychain", "identity", "physical"}
# 반복이 불가피한 사례 — {(Lv0, 사고 ID 또는 '사례명(시점)'): (허용 항목, 사유)}. 허용 항목 밖에서 또 쓰면 다시 경고
# OT: 대형 5건은 가장 특징적인 항목(IR-01 2025 폴란드 · IP-02 2016 우크라이나 · IM-05 Triton · PE-01 Stuxnet ·
#     EX-03 2015 우크라이나)에 두고, 구성 원본의 실제 사고가 다른 항목에 이미 배치된 사고뿐인 항목에서만 함께 씀
CASE_REUSE_ALLOWED = {
    ("ot", "INC-009"): ({"OT-PE-01", "OT-EX-02", "OT-EV-02"},
                        "OT-EX-02·EV-02는 구성 원본의 실제 사고가 대형 5건뿐 — 주기 실행 블록 감염·PLC 루트킷은 Stuxnet의 다른 단계"),
    ("ot", "INC-004"): ({"OT-IM-05", "OT-PV-01", "OT-LM-01"},
                        "OT-PV-01·LM-01은 구성 원본의 실제 사고가 Triton·Stuxnet뿐 — 펌웨어 0-day 권한 상승·SIS 프로그램 추가는 Triton의 다른 단계"),
    ("ot", "INC-008"): ({"OT-IR-01", "OT-RD-02"},
                        "OT-RD-02는 구성 원본의 실제 사고가 Triton·2025 폴란드뿐 — LLM 생성 스크립트·페이로드 재작성은 준비 단계"),
    ("ot", "INC-002"): ({"OT-EX-03", "OT-IR-05"},
                        "OT-IR-05의 실제 사고 줄이 모두 다른 항목에 배치된 사고 — 직렬-이더넷 변환기 펌웨어 덮어쓰기는 펌웨어 변조의 대표 사례"),
    ("ot", "INC-003"): ({"OT-IP-02", "OT-EV-01"},
                        "OT-EV-01의 대형 5건 밖 사례(2022 우크라이나·NotPetya)는 OT-EX-01·PV-02·PE-02에 배치 — 전력 프로토콜 이름 위장은 OT 특화 사례"),
    ("ot", "INC-006"): ({"OT-EX-01", "OT-PV-02"},
                        "OT-EX-01·PV-02 모두 대형 5건 밖 실제 사고가 2022 우크라이나뿐 — SCADA 명령 도구 실행·GPO 와이퍼 배포는 다른 단계"),
    ("ot", "INC-007"): ({"OT-IP-01", "OT-PE-03"},
                        "OT-PE-03의 대형 5건 밖 사례는 2022 우크라이나·리비우뿐 — 웹셸 지속성·Modbus 파라미터 변조는 다른 단계"),
    ("ot", "INC-013"): ({"OT-IA-06", "OT-EX-04"},
                        "OT-EX-04의 대형 5건 밖 사례는 Dragonfly·2022 우크라이나뿐 — 벤더 설치 파일 변조·메일 첨부 실행은 캠페인의 다른 경로"),
    ("ot", "INC-018"): ({"OT-DS-02", "OT-CO-03"},
                        "OT-DS-02·CO-03 모두 대형 5건 밖 실제 사고가 VPNFilter뿐 — 패킷 스니퍼·ssler 중간자 모듈은 다른 기능"),
}
SOURCE_TAG_ALIAS = {"실증·연구": "실증"}
# 물리·인적 작성 제약(결정 D9) — 방어 자료로 한정: 드론 기체 무력화·전파 교란은 권고하지 않고(탐지·식별·신고·차폐),
# 인적 징후에 보호 특성을 쓰지 않음. 대응 방안·전 필드 금지 표현
PHYSICAL_BANNED_CONTROLS = ["재밍", "전파 교란", "전파 차단", "격추", "기체 무력화", "드론 무력화"]
PHYSICAL_BANNED_ANY = ["출신", "종교", "인종", "성별", "성적 지향"]
RECENT_FROM = "2025"  # 최근 사고 기준 연도(이 해 1월 1일 이후)


def risk_of(likelihood, severity):
    return RISK.get((likelihood, severity), "낮음")


# ---------------------------------------------------------------- 원본 로더
def load_ai():
    """AI v3.2 원본: Lv3 행(dict), 도메인명, Lv2 통제, 사례(Incident) 매핑."""
    wb = openpyxl.load_workbook(AI_XLSX, read_only=True, data_only=True)
    rows = list(wb["통합매트릭스"].iter_rows(values_only=True))
    header = rows[3]
    lv3 = collections.OrderedDict()
    for r in rows[4:]:
        if r[3] and str(r[3]).startswith("UT-"):
            d = dict(zip(header, r))
            d["위험도"] = risk_of(d["발생가능성"], d["심각도"])  # 원본 셀은 수식 → 값으로 재계산
            d["domain"] = d["도메인 (Lv1)"][:3]
            lv3[r[3]] = d
    base = {b["id"]: b for b in yaml.safe_load(open(os.path.join(AI_DIR, "data", "v3.2", "base.yaml")))}
    lv2 = collections.OrderedDict((x["ut"], x) for x in yaml.safe_load(open(os.path.join(AI_DIR, "data", "v3.2", "lv2.yaml"))))
    cases = yaml.safe_load(open(os.path.join(AI_DIR, "data", "v3.2", "cases.yaml")))
    incidents = collections.defaultdict(set)
    for c in cases:
        if c["type"] == "Incident":
            for l in c["lv3"]:
                incidents[l].add(c["id"])
    atlas_dates = {c["id"]: str(c.get("date") or "") for c in cases if c["type"] == "Incident"}
    for k, d in lv3.items():
        d["atlas_incidents"] = sorted(incidents[k])
        d["owasp_incidents"] = base[k].get("owasp_incidents") or 0
        d["owasp_vulns"] = base[k].get("owasp_vulns") or 0
        if len(d["atlas_incidents"]) + d["owasp_incidents"] != d["실제 사고 수"]:
            raise ValueError(f"AI 실제 사고 수 재현 불일치: {k}")
        d["owasp_named"] = owasp_incident_lines(d["참조"])
        if len(d["owasp_named"]) > d["owasp_incidents"]:
            raise ValueError(f"AI OWASP 인용 사고 줄이 건수보다 많음: {k}")
    domains = collections.OrderedDict()
    for d in lv3.values():
        domains.setdefault(d["domain"], d["도메인 (Lv1)"])
    return {"lv3": lv3, "header": header, "lv2": lv2, "domains": domains, "cases": cases, "atlas_dates": atlas_dates}


SOURCE_RE = re.compile(r"\(([^()]*(?:ATLAS|OWASP)[^()]*)\)\s*$")


def owasp_incident_lines(ref):
    """참조 열에서 OWASP 출처로 집계된 [실제 사고] 줄 → [(사례명(시점), 시점)] (ATLAS 사례 ID가 붙은 줄은 제외)."""
    out = []
    for line in (ref or "").split("\n"):
        m = CASE_RE.match(line.strip())
        if not m or m.group("tag") != "실제 사고":
            continue
        src = SOURCE_RE.search(line.strip())
        if not src or "OWASP" not in src.group(1) or "AML.CS" in src.group(1):
            continue
        n = NAMED_RE.match(m.group("body"))
        name, date = (n.group("name").strip(), n.group("date") or "") if n else (m.group("body")[:30], "")
        out.append((f"{name}({date})" if date else name, date))
    return out


def load_cloud():
    """클라우드 v5 원본: 행(154), 기법별 대표 행·전술·실제 사고 ID 집합."""
    wb = openpyxl.load_workbook(CLOUD_XLSX, read_only=True, data_only=True)
    m = list(wb["통합 매트릭스"].iter_rows(values_only=True))
    header = m[3]
    rows = [dict(zip(header, r)) for r in m[4:] if r[1]]
    inc = list(wb["역매핑_사고사례"].iter_rows(values_only=True))
    ih = {k: i for i, k in enumerate(inc[0])}
    real = {x[ih["사건ID"]] for x in inc[1:] if x[ih["사례유형"]] in CLOUD_REAL_KINDS and x[ih["집계 포함"]] == "Y"}
    inc_dates = {x[ih["사건ID"]]: str(x[ih["기준일"]] or "") for x in inc[1:]}
    tech = collections.OrderedDict()
    for r in rows:
        t = str(r["ATT&CK ID"])
        r["tactic"] = r["도메인(Lv1)"][1:3]
        ids = {i.strip() for i in str(r["관련 사례 ID"] or "").replace("\n", ",").split(",") if i.strip()}
        r["real_incidents"] = sorted(ids & real)
        if len(r["real_incidents"]) != r["실제 사고 수"]:
            raise ValueError(f"클라우드 실제 사고 수 재현 불일치: {r['CTC-ID']}")
        if sum(inc_dates[i][:4] >= RECENT_FROM for i in r["real_incidents"]) != (r["최근 사고(2025~)"] or 0):
            raise ValueError(f"클라우드 최근 사고 수 재현 불일치: {r['CTC-ID']}")
        if t not in tech:
            tech[t] = {"row": r, "rows": [], "ctc": [], "tactics": []}
        tech[t]["rows"].append(r)
        tech[t]["ctc"].append(r["CTC-ID"])
        tech[t]["tactics"].append(r["tactic"])
    tactics = collections.OrderedDict()
    for r in rows:
        tactics.setdefault(r["tactic"], r["도메인(Lv1)"])
    for x in tech.values():
        x["ids"] = x["ctc"]
    return {"rows": rows, "header": header, "tech": tech, "tactics": tactics, "wb": CLOUD_XLSX, "inc_dates": inc_dates}


def load_ot():
    """OT v5 원본: 행(141), 키(ATT&CK·EMB3D ID)별 행 목록·전술·실제 사고 ID 집합."""
    wb = openpyxl.load_workbook(OT_XLSX, read_only=True, data_only=True)
    m = list(wb["통합 매트릭스"].iter_rows(values_only=True))
    header = m[3]
    rows = [dict(zip(header, r)) for r in m[4:] if r[1]]
    inc = list(wb["역매핑_사고사례"].iter_rows(values_only=True))
    ih = {k: i for i, k in enumerate(inc[0])}
    real = {x[ih["사건ID"]] for x in inc[1:] if x[ih["집계"]] == OT_REAL_STATUS}
    inc_dates = {x[ih["사건ID"]]: str(x[ih["기준일"]] or "") for x in inc[1:]}
    tech = collections.OrderedDict()
    for r in rows:
        k = str(r["ATT&CK·EMB3D ID"])
        r["tactic"] = r["도메인(Lv1)"][1:3]
        ids = {i.strip() for i in re.split(r"[,\n]", str(r["관련 사례 ID"] or "")) if i.strip()}
        r["real_incidents"] = sorted(ids & real)
        if len(r["real_incidents"]) != r["실제 사고 수"]:
            raise ValueError(f"OT 실제 사고 수 재현 불일치: {r['OTC-ID']}")
        if sum(inc_dates[i][:4] >= RECENT_FROM for i in r["real_incidents"]) != (r["최근 사고(2025~)"] or 0):
            raise ValueError(f"OT 최근 사고 수 재현 불일치: {r['OTC-ID']}")
        if k not in tech:
            tech[k] = {"row": r, "rows": [], "otc": [], "tactics": []}
        tech[k]["rows"].append(r)
        tech[k]["otc"].append(r["OTC-ID"])
        tech[k]["tactics"].append(r["tactic"])
    tactics = collections.OrderedDict()
    for r in rows:
        tactics.setdefault(r["tactic"], r["도메인(Lv1)"])
    for x in tech.values():
        x["ids"] = x["otc"]
    return {"rows": rows, "header": header, "tech": tech, "tactics": tactics, "wb": OT_XLSX, "inc_dates": inc_dates,
            "n_incidents": len(inc) - 1, "n_real": len(real)}


def load_taxo(kind):
    """자체 분류 체계 원본(공급망·신원·물리): 세부위협 행, 세부위협 ID별 행(전술 반복 없음 — 1행)·도메인·실제 사고 ID 집합.
    클라우드·OT와 같은 구조(tech·tactics)로 돌려줘 위험 재산정(evaluate_attack)·출처 대조를 그대로 쓴다."""
    cfg = TAXO[kind]
    wb = openpyxl.load_workbook(cfg["xlsx"], read_only=True, data_only=True)
    m = list(wb["통합 매트릭스"].iter_rows(values_only=True))
    header = m[3]
    rows = [dict(zip(header, r)) for r in m[4:] if r[1]]
    inc = list(wb["역매핑_사고사례"].iter_rows(values_only=True))
    incidents = [dict(zip(inc[0], x)) for x in inc[1:] if x[0]]
    real = {x["사례 ID"] for x in incidents if x["집계 상태"] == TAXO_REAL_STATUS}
    inc_dates = {x["사례 ID"]: str(x["시점"] or "") for x in incidents}
    pat = re.compile(rf"\b{cfg['inc']}-\d+\b")
    tech = collections.OrderedDict()
    for r in rows:
        k = r[cfg["id_col"]]
        r["tactic"] = r["도메인(Lv1)"][1:3]
        r["real_incidents"] = sorted(set(pat.findall(str(r["관련 사례 ID"] or ""))) & real)
        if len(r["real_incidents"]) != r["실제 사고 수"]:
            raise ValueError(f"{LABEL[kind]} 실제 사고 수 재현 불일치: {k}")
        if sum(inc_dates[i][:4] >= RECENT_FROM for i in r["real_incidents"]) != (r["최근 사고(2025~)"] or 0):
            raise ValueError(f"{LABEL[kind]} 최근 사고 수 재현 불일치: {k}")
        tech[k] = {"row": r, "rows": [r], "ids": [k], "tactics": [r["tactic"]]}
    tactics = collections.OrderedDict()
    for r in rows:
        tactics.setdefault(r["tactic"], r["도메인(Lv1)"])
    if list(tactics) != DOMAINS[kind]:
        raise ValueError(f"{LABEL[kind]} 도메인 순서 불일치: {list(tactics)}")
    return {"rows": rows, "header": header, "tech": tech, "tactics": tactics, "wb": cfg["xlsx"], "inc_dates": inc_dates,
            "incidents": incidents, "n_incidents": len(incidents), "n_real": len(real), "id_col": cfg["id_col"]}


def load_sources():
    src = {"ai": load_ai(), "cloud": load_cloud(), "ot": load_ot()}
    for kind in TAXO:
        src[kind] = load_taxo(kind)
    return src


def source_keys(kind, src):
    """Lv0별 구성 원본 키 → 원본 (AI: UT Lv3 ID, 클라우드: ATT&CK ID, OT: ATT&CK·EMB3D ID, 공급망·신원·물리: 세부위협 ID)."""
    return src["lv3"] if kind == "ai" else src["tech"]


def domains_of(kind, src):
    """Lv0별 Lv1 도메인 {코드: 표시명} (원본 순서)."""
    return src["domains"] if kind == "ai" else src["tactics"]


def source_refs(kind, src, key):
    """구성 원본의 참조 문자열 (전술 반복 행은 모든 행의 참조를 합침)."""
    if kind == "ai":
        return src["lv3"][key]["참조"] or ""
    return "\n".join(r["참조"] or "" for r in src["tech"][key]["rows"])


def member_row(src, key, domain):
    """구성 원본 키의 대표 행 — 요약 항목의 주 전술에 행이 있으면 그 행, 없으면(전술 이동) 원본 첫 행.
    전술마다 심각도가 다른 반복 기법(예: OT T0851 회피 '중'·대응 기능 억제 '상')을 주 전술 기준으로 평가한다."""
    t = src["tech"][key]
    for r in t["rows"]:
        if r["tactic"] == domain:
            return r
    return t["row"]


# ---------------------------------------------------------------- 요약 문안 로더
def load_summary():
    """요약 문안 {Lv0: [항목]} — data/<Lv0>/<도메인>.yaml을 도메인 순서로 읽고 항목에 domain 키를 붙인다."""
    out = {}
    for kind in KINDS:
        out[kind] = []
        for dom in DOMAINS[kind]:
            p = os.path.join(DATA_DIR, kind, f"{dom}.yaml")
            if os.path.exists(p):
                for e in yaml.safe_load(open(p, encoding="utf-8")) or []:
                    e["domain"] = dom
                    out[kind].append(e)
    return out


# ---------------------------------------------------------------- 근거 재산정
# 요약 위험도 = 구성 원본별 위험도(발생가능성 × 심각도)의 최댓값.
# 발생가능성·심각도를 서로 다른 원본에서 가져와 조합하지 않는다(교차 결합 금지).
# 실제 사고 합집합이 2건 이상이면, 사고가 1건 이상 확인된 원본만 발생가능성을 '상'으로 보정한다.
# 표시하는 발생가능성·심각도는 최댓값을 낸 원본(대표 원본)의 값이다.
def _representative(members, union_n):
    best = None
    for i, m in enumerate(members):
        lik = "상" if (union_n >= 2 and m["own"] >= 1) else m["lik"]
        risk = risk_of(lik, m["sev"])
        key = (RISK_ORDER.index(risk), -m["own"], -LEVEL[m["sev"]], -LEVEL[lik], i)
        if best is None or key < best[0]:
            best = (key, m["id"], lik, m["sev"])
    return best[1:]


def _result(entry, rep, lik, sev, evidence, n, recent, **extra):
    if entry.get("severity_override"):
        sev = entry["severity_override"]
    out = {"likelihood": lik, "severity": sev, "risk": risk_of(lik, sev), "evidence": evidence,
           "incidents": n, "recent": recent, "rep": rep}
    out.update(extra)
    return out


def evaluate_ai(entry, src):
    mem = [(m, src["lv3"][m]) for m in entry["members"]]
    atlas = sorted(set().union(*[set(d["atlas_incidents"]) for _, d in mem]))
    # OWASP 인용 사고는 ID 없이 건수만 있으므로, 참조에 같은 사례명(시점)으로 실린 중복분을 원본 간에 뺀다
    named = collections.Counter(k for _, d in mem for k in {k for k, _ in d["owasp_named"]})
    owasp = sum(d["owasp_incidents"] for _, d in mem) - sum(v - 1 for v in named.values())
    n = len(atlas) + owasp
    dates = dict(x for _, d in mem for x in d["owasp_named"])
    recent = sum(src["atlas_dates"].get(c, "")[:4] >= RECENT_FROM for c in atlas) + \
        sum(dt[:4] >= RECENT_FROM for dt in dates.values())
    members = [dict(id=m, lik=d["발생가능성"], sev=d["심각도"], own=len(d["atlas_incidents"]) + d["owasp_incidents"])
               for m, d in mem]
    rep, lik, sev = _representative(members, n)
    evidence = min((d["근거 수준"] for _, d in mem), key=EVIDENCE_ORDER.index)
    return _result(entry, rep, lik, sev, evidence, n, recent, incident_ids=atlas, owasp_incidents=owasp)


def evaluate_attack(entry, src):
    """ATT&CK 기반 원본(클라우드·OT) 공용 재산정."""
    mem = [(t, member_row(src, t, entry["domain"])) for t in entry["members"]]
    ids = sorted(set().union(*[set(r["real_incidents"]) for _, r in mem]))
    recent = sum(src["inc_dates"].get(i, "")[:4] >= RECENT_FROM for i in ids)
    members = [dict(id=t, lik=r["발생가능성"], sev=r["심각도"], own=len(r["real_incidents"])) for t, r in mem]
    rep, lik, sev = _representative(members, len(ids))
    evidence = min((r["근거 수준"] for _, r in mem), key=EVIDENCE_ORDER.index)
    return _result(entry, rep, lik, sev, evidence, len(ids), recent, incident_ids=ids)


evaluate_cloud = evaluate_ot = evaluate_attack
EVALUATE = {"ai": evaluate_ai, "cloud": evaluate_attack, "ot": evaluate_attack,
            "supplychain": evaluate_attack, "identity": evaluate_attack, "physical": evaluate_attack}


# ---------------------------------------------------------------- 문안 검증
CASE_RE = re.compile(r"^- \[(?P<tag>[^\]]+)\] (?P<body>.+)$")
NAMED_RE = re.compile(r"^(?P<name>[^:]+?)(?:\((?P<date>[^()]*(?:\([^()]*\))?[^()]*)\))?: (?P<gist>.+)$")
LIMITS = {"name": 30, "summary": 50, "description": 130, "scenario": 175, "cases": 130, "controls": 80}
BANNED = {"인증정보": "자격증명", "크리덴셜": "자격증명", "엑스필": "반출"}


def _lines(text):
    return [l for l in (text or "").split("\n") if l.strip()]


def normalize_terms(text):
    """원본 참조의 금지 용어를 통합본 용어로 바꾼 사본 — 출처 대조용(원본 사례명 '크리덴셜 스터핑' ↔ 요약 '자격증명 스터핑')."""
    for bad, good in BANNED.items():
        text = text.replace(bad, good)
    return text


def check_entry(entry, refs, kind):
    """문안 형식·길이·사례 출처 대조. refs: 구성 원본의 참조 문자열 목록."""
    errs, warns = [], []
    eid = entry.get("id", "?")
    for f in ["id", "name", "summary", "members", "description", "scenario", "cases", "controls", "basis"]:
        if not entry.get(f):
            errs.append(f"{eid}: '{f}' 누락")
    if errs:
        return errs, warns
    if len(entry["name"]) > LIMITS["name"]:
        warns.append(f"{eid}: 위협명 {len(entry['name'])}자 (>{LIMITS['name']})")
    summary = str(entry["summary"])
    if "\n" in summary.strip() or summary.startswith("- ") or len(summary) > LIMITS["summary"]:
        errs.append(f"{eid}: 핵심 요약은 '- ' 없는 한 줄 {LIMITS['summary']}자 이내 ({len(summary)}자)")
    desc = _lines(entry["description"])
    if len(desc) != 2 or not all(l.startswith("- ") for l in desc):
        errs.append(f"{eid}: 위협 설명은 '- '로 시작하는 2줄이어야 함")
    elif not desc[0].rstrip().endswith("위협"):
        errs.append(f"{eid}: 위협 설명 1줄은 '~위협'으로 끝나야 함")
    scen = _lines(entry["scenario"])
    if not 1 <= len(scen) <= 2 or not all(l.startswith("- ") and "→" in l for l in scen):
        errs.append(f"{eid}: 공격 시나리오는 '→'를 포함한 1~2줄")
    ctrl = _lines(entry["controls"])
    if not 2 <= len(ctrl) <= 3 or not all(l.startswith("- ") for l in ctrl):
        errs.append(f"{eid}: 대응 방안은 '- '로 시작하는 2~3줄")
    for field in ["description", "scenario", "cases", "controls"]:
        for l in _lines(entry[field]):
            if len(l) > LIMITS[field]:
                warns.append(f"{eid}: {field} 한 줄 {len(l)}자 (>{LIMITS[field]})")
    for bad, good in BANNED.items():
        for field in ["name", "summary", "description", "scenario", "cases", "controls"]:
            if bad in (entry.get(field) or ""):
                errs.append(f"{eid}: 용어 '{bad}' → '{good}'")
    cases = _lines(entry["cases"])
    if not 1 <= len(cases) <= 2:
        errs.append(f"{eid}: 대표 사례는 1~2줄")
    joined = normalize_terms("\n".join(refs))
    for c in cases:
        if c.startswith("- 공개 사고 미확인"):  # 사례가 없는 기법: 원본의 '공개 사고 미확인' 근거가 있을 때만 허용
            if "공개 사고 미확인" not in joined:
                errs.append(f"{eid}: 원본에 '공개 사고 미확인' 근거 없음")
            continue
        m = CASE_RE.match(c)
        if not m or m.group("tag") not in CASE_TAGS_BY_KIND[kind]:
            errs.append(f"{eid}: 사례 형식 오류: {c[:40]}")
            continue
        tag, body = m.group("tag"), m.group("body")
        if tag == "시나리오":
            if "[시나리오]" not in joined:
                errs.append(f"{eid}: 원본에 [시나리오] 항목 없음")
            continue
        n = NAMED_RE.match(body)
        if not n:
            errs.append(f"{eid}: 사례는 '사례명(시점): 요지' 형식: {body[:40]}")
            continue
        key = f"{n.group('name')}({n.group('date')})" if n.group("date") else f"] {n.group('name')}:"
        hit = None
        for line in joined.split("\n"):
            if key in line:
                hit = line
                break
        if not hit:
            errs.append(f"{eid}: 사례 '{key}'가 구성 원본 참조에 없음")
            continue
        st = re.match(r"^- \[([^\]]+)\]", hit.strip())
        stag = SOURCE_TAG_ALIAS.get(st.group(1), st.group(1)) if st else None
        if stag != tag:
            errs.append(f"{eid}: 사례 '{key}' 유형 불일치 (원본 [{stag}] / 요약 [{tag}])")
    return errs, warns


def entry_prefix(kind, domain):
    """요약 항목 ID 접두사 — AI-<도메인번호>- / CL-<전술>- / OT-<전술>- / SC-·ID-·PH-<도메인>-."""
    if kind in TAXO:
        return f"{TAXO[kind]['prefix']}-{domain}-"
    return {"ai": f"AI-{domain[1:]}-", "cloud": f"CL-{domain}-", "ot": f"OT-{domain}-"}[kind]


def case_incidents(cases, refs):
    """대표 사례 줄 → [('사례명(시점)', 사고 ID 목록)] ([시나리오]·'공개 사고 미확인'은 제외).
    사고 ID는 사례명(시점)이 실린 원본 참조 줄의 사고 ID(INC-·SCI-·IDI-·PHI-…)에서 가져오며, 없으면(공개 취약점·EMB3D 등) 빈 목록."""
    out = []
    for line in _lines(cases):
        m = CASE_RE.match(line.strip())
        if not m or m.group("tag") == "시나리오":
            continue
        n = NAMED_RE.match(m.group("body"))
        if not n:
            continue
        name = n.group("name").strip()
        label = f"{name}({n.group('date')})" if n.group("date") else name
        key = label if n.group("date") else f"] {name}:"  # check_entry의 출처 대조 키와 같음
        hit = next((l for l in normalize_terms(refs or "").split("\n") if key in l), "")
        out.append((label, list(dict.fromkeys(INCIDENT_RE.findall(hit)))))
    return out


def check_all(src, summ):
    """src·summ: {Lv0: 원본} · {Lv0: [요약 항목]}."""
    errs, warns = [], []
    # 1) 원본 전수 배정 (누락·중복 0)
    for kind in KINDS:
        keys = source_keys(kind, src[kind])
        used = collections.Counter(m for e in summ[kind] for m in e["members"])
        missing = [k for k in keys if k not in used]
        dup = [k for k, v in used.items() if v > 1]
        unknown = [k for k in used if k not in keys]
        for label, lst in [(f"{LABEL[kind]} 미배정", missing), (f"{LABEL[kind]} 중복", dup),
                           (f"{LABEL[kind]} 없는 ID", unknown)]:
            if lst:
                (warns if "미배정" in label else errs).append(f"{label}: {lst}")
    ids = [e["id"] for kind in KINDS for e in summ[kind]]
    for k, v in collections.Counter(ids).items():
        if v > 1:
            errs.append(f"위협 ID 중복: {k}")
    # 2) 항목별 문안·사례 검증
    valid = {kind: [e for e in summ[kind] if all(m in source_keys(kind, src[kind]) for m in e.get("members", []))]
             for kind in KINDS}
    for kind in KINDS:
        for e in valid[kind]:
            refs = [source_refs(kind, src[kind], m) for m in e["members"]]
            a, b = check_entry(e, refs, kind)
            errs += a
            warns += b
            if not e["id"].startswith(entry_prefix(kind, e["domain"])):
                errs.append(f"{e['id']}: ID와 {'도메인' if kind == 'ai' else '전술'}({e['domain']}) 불일치")
    # 3) 근거 수준과 사례 유형 정합성: '실제 사고 확인'이면 [실제 사고] 1건 이상,
    #    [EMB3D]는 구성 원본에 실제 사고·실증 사례가 없을 때만 (OT 결정 D7)
    for kind in KINDS:
        for e in valid[kind]:
            if not e.get("members"):
                continue
            cases = e.get("cases") or ""
            if EVALUATE[kind](e, src[kind])["evidence"] == "실제 사고 확인" and "[실제 사고]" not in cases:
                warns.append(f"{e['id']}: 근거 수준 '실제 사고 확인'이나 대표 사례에 [실제 사고] 없음")
            if "[EMB3D]" in cases:
                refs = "\n".join(source_refs(kind, src[kind], m) for m in e["members"])
                if re.search(r"^\s*- \[(실제 사고|실증|실증·연구)\]", refs, re.M):
                    warns.append(f"{e['id']}: [EMB3D] 사례는 구성 원본에 실제 사고·실증 사례가 없을 때만 사용")
    # 4) 물리·인적 작성 제약(결정 D9): 금지 표현
    for e in valid.get("physical", []):
        for bad in PHYSICAL_BANNED_CONTROLS:
            if bad in str(e.get("controls") or ""):
                errs.append(f"{e['id']}: 대응 방안에 '{bad}' — 드론·전파는 탐지·식별·신고·차폐만 권고")
        for bad in PHYSICAL_BANNED_ANY:
            for field in ["name", "summary", "description", "scenario", "cases", "controls"]:
                if bad in str(e.get(field) or ""):
                    errs.append(f"{e['id']}: '{bad}' — 인적 징후·사례에 보호 특성을 쓰지 않음")
    # 5) 대표 사례 분산: 같은 Lv0 안에서 같은 사고를 여러 항목에 쓰면 경고 (CASE_REUSE_CHECK 대상 Lv0)
    #    사고 ID가 없는 사례(공개 취약점·EMB3D 등)는 '사례명(시점)'으로 셈. 허용 목록의 항목 안에서만 반복을 허용
    for kind in KINDS:
        if kind not in CASE_REUSE_CHECK:
            continue
        seen, labels = collections.defaultdict(list), collections.defaultdict(list)
        for e in valid[kind]:
            refs = "\n".join(source_refs(kind, src[kind], m) for m in e["members"])
            for label, incs in case_incidents(e.get("cases"), refs):
                for key in incs or [label]:
                    if e["id"] not in seen[key]:
                        seen[key].append(e["id"])
                    if label not in labels[key]:
                        labels[key].append(label)
        for key, where in seen.items():
            allowed = CASE_REUSE_ALLOWED.get((kind, key), (set(), ""))[0]
            if len(where) > 1 and not set(where) <= allowed:
                name = key if labels[key] == [key] else f"{key} {'·'.join(labels[key])}"
                warns.append(f"대표 사례 '{name}' 반복 사용: {', '.join(where)}")
    return errs, warns


if __name__ == "__main__":
    src, summ = load_sources(), load_summary()
    errs, warns = check_all(src, summ)
    print(" / ".join(f"{LABEL[k]} 요약 {len(summ[k])}개" for k in KINDS))
    for w in warns:
        print("  [경고]", w)
    for e in errs:
        print("  [오류]", e)
    print("오류", len(errs), "경고", len(warns))
